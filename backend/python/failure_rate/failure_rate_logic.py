#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Failure Rate Integration Tool - Core Logic
"""
import math
import re
from pathlib import Path
from typing import Any, Callable, Dict, Optional, Union

import pandas as pd


def normalize_refdes_for_lookup(value) -> str:
    """
    Normalize RefDes for lookup, stripping leading zeros from numeric part.

    This ensures "U01" matches "U1" and "R001" matches "R1" across different
    data sources that may use inconsistent zero-padding.

    Args:
        value: Raw RefDes string (or NaN/None which returns empty string)

    Returns:
        Normalized RefDes with leading zeros stripped (e.g., "U01" -> "U1")
    """
    # Guard against NaN/None values to prevent "NAN" pollution in lookups
    # pd.isna() handles None, float NaN, pd.NA, and numpy NaN types
    if pd.isna(value):
        return ""
    s = str(value).strip().upper()
    # Match prefix (letters) + leading zeros + number + optional suffix (multi-letter)
    match = re.match(r'^([A-Z]+)0*(\d+)([A-Z]*)$', s)
    if match:
        prefix, num, suffix = match.groups()
        return f"{prefix}{int(num)}{suffix}"
    return s

# Tauri backend copy — source: src/apps/failure_rate/failure_rate_logic.py
from common import (
    try_read_table,
    ensure_columns_exist,
    sha256_of_path,
    write_snapshot,
    get_tool_logger,
    style_worksheet,
    write_df_to_sheet,
    CancellationToken,
    FMR_TOLERANCE,
    FileAccessError,
    ProcessingError,
    to_user_facing_text,
)
from common.refdes_utils import (
    extract_base_refdes,
    extract_instance_refdes,
    split_refdes_list,
)
from common.fmea_utils import classify_fmea_rows

# Initialize module logger
_logger = get_tool_logger("failure_rate")


# ==================== LOGIC CLASS ====================

class FMEALinkerLogic:
    """Core logic for linking FMEA data with failure rate predictions."""

    def __init__(
        self,
        log_func: Optional[Callable[[str], None]] = None,
        progress_func: Optional[Callable[[float], None]] = None,  # L7: Progress callback (0.0-1.0)
    ) -> None:
        self._log_func: Callable[[str], None] = log_func or print
        self._progress_func: Optional[Callable[[float], None]] = progress_func  # L7
        self.prediction_df: Optional[pd.DataFrame] = None
        self.fmea_df: Optional[pd.DataFrame] = None
        self.merged_df: Optional[pd.DataFrame] = None
        # Cancellation token for user-initiated stop requests
        self.cancel: CancellationToken = CancellationToken()

    def log(self, msg: str) -> None:
        """Log a message to both the GUI callback and the module logger."""
        self._log_func(msg)
        _logger.info(msg)

    def _report_progress(self, fraction: float) -> None:
        """L7: Report progress to GUI callback (0.0 to 1.0)."""
        if self._progress_func:
            self._progress_func(min(1.0, max(0.0, fraction)))

    def load_prediction(self, path: Union[str, Path], sheet_name=None) -> None:
        """Load prediction file containing failure rate data.

        Args:
            path: Path to the prediction file.
            sheet_name: Excel sheet name or index (None = first sheet).
        """
        self.log(f"Loading Prediction: {Path(path).name}")
        # Pass cancel_check so Cancel interrupts slow OneDrive reads /
        # retry loops instead of blocking until the read completes.
        self.prediction_df = try_read_table(
            path,
            sheet_name=sheet_name,
            log_func=self.log,
            cancel_check=self.cancel.is_cancelled,
        )
        if self.prediction_df is None or self.prediction_df.empty:
            raise FileAccessError(
                f"Failed to load prediction data from '{path}' or file is empty",
                file_path=str(path),
                operation="read",
            )
        self.log(f"  Loaded {len(self.prediction_df)} rows.")

    def load_fmea(self, path: Union[str, Path], sheet_name=None) -> None:
        """Load FMEA file containing failure mode data.

        Args:
            path: Path to the FMEA file.
            sheet_name: Excel sheet name or index (None = first sheet).
        """
        self.log(f"Loading FMEA: {Path(path).name}")
        self.fmea_df = try_read_table(
            path,
            sheet_name=sheet_name,
            log_func=self.log,
            cancel_check=self.cancel.is_cancelled,
        )
        if self.fmea_df is None or self.fmea_df.empty:
            raise FileAccessError(
                f"Failed to load FMEA data from '{path}' or file is empty",
                file_path=str(path),
                operation="read",
            )
        self.log(f"  Loaded {len(self.fmea_df)} rows.")

    def process(
        self,
        col_map: Dict[str, str],
        check_fmr: bool = False,
    ) -> pd.DataFrame:
        """Process FMEA and prediction data. Call self.cancel.cancel() to stop."""
        # Precondition check - ensure data is loaded
        if self.prediction_df is None or self.fmea_df is None:
            raise ProcessingError(
                "Data not loaded. Call load_prediction() and load_fmea() before process()."
            )

        self.log("Starting Processing...")
        pred = self.prediction_df.copy()
        fmea = self.fmea_df.copy()

        pred_ref_col = col_map['pred_ref']
        pred_fr_col = col_map['pred_fr']
        ensure_columns_exist(pred, [pred_ref_col, pred_fr_col], "Prediction file")
        
        # Normalize RefDes with zero-stripping (U01 -> U1, R001 -> R1)
        pred['RefDes_Norm'] = pred[pred_ref_col].apply(normalize_refdes_for_lookup)
        # Coerce the prediction failure rate to numeric. A cell that held a
        # non-empty but unparseable value (a formula string with no cached
        # result, '1.2 FIT', 'TBD', free text) becomes NaN here; fillna(0.0)
        # would then silently treat it as a real 0.0 failure rate and the linked
        # row would look identical to a genuinely-zero part. Track those RefDes
        # so linked rows get a visible Validation_Notes flag instead of an
        # invisible zero (Tier-1 fix: silent prediction-FR coercion).
        _pred_fr_numeric = pd.to_numeric(pred[pred_fr_col], errors='coerce')
        _pred_fr_raw = pred[pred_fr_col]
        # A cell is unparseable if it was non-empty text that coerced to NaN, OR a
        # non-finite number ('inf' / '1e999' coerce to inf, NOT NaN). Both must be
        # flagged AND zeroed so they never propagate as inf/garbage into Mode_FR
        # and Function_FR (mirrors the math.isfinite guard the BOM FMR check uses).
        _pred_fr_finite = _pred_fr_numeric.apply(lambda v: pd.notna(v) and math.isfinite(v))
        _unparseable_fr_mask = (
            (_pred_fr_numeric.isna()
             & _pred_fr_raw.notna()
             & (_pred_fr_raw.astype(str).str.strip() != ""))
            | (_pred_fr_numeric.notna() & ~_pred_fr_finite)
        )
        pred['FR_Clean'] = _pred_fr_numeric.where(_pred_fr_finite, 0.0)
        unparseable_fr_refs = set(pred.loc[_unparseable_fr_mask, 'RefDes_Norm'])
        unparseable_fr_refs.discard("")
        if unparseable_fr_refs:
            _ufr_sample = ', '.join(sorted(unparseable_fr_refs)[:5])
            _ufr_remaining = len(unparseable_fr_refs) - 5
            _ufr_suffix = f" (and {_ufr_remaining} more)" if _ufr_remaining > 0 else ""
            self.log(
                f"  WARNING: {len(unparseable_fr_refs)} Prediction RefDes have a "
                f"non-numeric Failure Rate (treated as 0.0): {_ufr_sample}{_ufr_suffix}"
            )

        # Validate and apply unit conversion
        valid_unit_modes = {'per_hour', 'per_million_hours', 'per_billion_hours'}
        unit_mode = col_map.get('unit_mode', 'per_hour')
        if unit_mode not in valid_unit_modes:
            self.log(f"  WARNING: Invalid unit_mode '{unit_mode}', defaulting to 'per_hour'")
            unit_mode = 'per_hour'

        if unit_mode == "per_million_hours":
            pred['FR_Clean'] = pred['FR_Clean'] / 1_000_000
        elif unit_mode == "per_billion_hours":
            pred['FR_Clean'] = pred['FR_Clean'] / 1_000_000_000
        # per_hour: no conversion needed
        
        # Check for duplicate RefDes in Prediction file - keep first occurrence
        pred_dupes = pred[pred.duplicated(subset=['RefDes_Norm'], keep=False)]
        if not pred_dupes.empty:
            dupe_refs = pred_dupes['RefDes_Norm'].unique()
            # L5: Show clear total count and sample of duplicates
            sample = ', '.join(dupe_refs[:5])
            remaining = len(dupe_refs) - 5
            suffix = f" (and {remaining} more)" if remaining > 0 else ""
            self.log(f"  WARNING: Found {len(dupe_refs)} duplicate RefDes entries in Prediction: {sample}{suffix}")

        # Drop duplicates keeping first occurrence, then create lookup
        pred_deduped = pred.drop_duplicates(subset=['RefDes_Norm'], keep='first')
        if len(pred_deduped) < len(pred):
            dropped_count = len(pred) - len(pred_deduped)
            self.log(f"  Note: Dropped {dropped_count} duplicate rows (keeping first occurrence)")
        fr_lookup = pred_deduped.set_index('RefDes_Norm')['FR_Clean'].to_dict()
        self.log(f"  Mapped {len(fr_lookup)} failure rates from Prediction.")

        fmea_cause_col = col_map['fmea_cause']
        fmea_ratio_col = col_map['fmea_ratio']
        fmea_usage_col = col_map['fmea_usage']
        fmea_func_col = col_map['fmea_func']
        # M3: Normalize fmea_func_col early - GUI may send "" or "None" for unselected
        if not fmea_func_col or fmea_func_col == "None":
            fmea_func_col = ""
        fmea_required = [fmea_cause_col, fmea_ratio_col, fmea_usage_col]
        if fmea_func_col:
            fmea_required.append(fmea_func_col)
        ensure_columns_exist(fmea, fmea_required, "FMEA file")

        # Check for duplicate RefDes in FMEA file (for user awareness)
        fmea_dupes = fmea[fmea.duplicated(subset=[fmea_cause_col], keep=False)]
        if not fmea_dupes.empty:
            self.log(f"  Note: Found {len(fmea_dupes)} rows with duplicate Failure Cause values (expected for multiple failure modes per component).")

        linked_count = 0
        missing_refs = set()

        # Initialize accumulators for vectorized column assignment
        # Using list accumulation pattern for 5-10x performance improvement over iterrows()
        extracted_refdes_list = []
        validation_refdes_list = []
        part_fr_list = []
        mode_fr_list = []
        corrected_ratio_list = []
        validation_notes_list = []

        # Build column index map for position-based access
        # This handles column names with spaces/special characters that break namedtuple attributes
        col_list = list(fmea.columns)
        cause_idx = col_list.index(fmea_cause_col)
        ratio_idx = col_list.index(fmea_ratio_col)
        usage_idx = col_list.index(fmea_usage_col)

        # Use itertuples(name=None, index=False) for 5-10x performance over iterrows()
        # M6 fix: index=False removes DataFrame index from tuple, simplifying access
        total_rows = len(fmea)
        for row_idx, row in enumerate(fmea.itertuples(name=None, index=False)):
            self.cancel.check("Processing cancelled by user.")
            # L7: Report progress every 100 rows to avoid callback overhead
            if row_idx % 100 == 0:
                self._report_progress(row_idx / total_rows * 0.8)  # 0-80% for main loop
            # Direct column access without index offset
            cause_text = str(row[cause_idx]) if pd.notna(row[cause_idx]) else ""
            notes = []

            # 1. Extract Base RefDes for Linking (e.g. U200 from U200-1)
            # Using centralized refdes_utils for consistent behavior across tools
            refdes_raw = extract_base_refdes(cause_text)
            # Normalize to match prediction lookup (strip leading zeros: U01 -> U1)
            refdes = normalize_refdes_for_lookup(refdes_raw) if refdes_raw else None
            extracted_refdes_list.append(refdes or "")

            # Lookup failure rate
            if refdes and refdes in fr_lookup:
                part_fr = fr_lookup[refdes]
                linked_count += 1
                # The RefDes matched, but its Prediction FR cell was non-numeric
                # and got coerced to 0.0. Flag it so a silent zero isn't mistaken
                # for a real link (Tier-1 fix: silent prediction-FR coercion).
                if refdes in unparseable_fr_refs:
                    notes.append("Prediction FR unparseable (treated as 0.0)")
            elif refdes:
                part_fr = 0.0
                missing_refs.add(refdes)
                notes.append("RefDes not in Prediction")
            else:
                part_fr = 0.0
            part_fr_list.append(part_fr)

            # 2. Extract Instance RefDes for Validation (e.g. U200-1).
            # FMR (Failure Mode Ratio) validation deliberately groups by the
            # INSTANCE / pin designator, NOT the base component: each instance
            # (U200-1, U200-2) carries its own failure modes (open/short/...)
            # whose ratios must sum to 1.0. Component-level failure-rate roll-up
            # is handled separately by the Part Usage column, which fractionally
            # allocates (e.g. 1/2 each) so Mode_FR sums back to the base part's
            # rate. Collapsing instances to a base RefDes HERE would wrongly sum
            # multiple instances' ratios to > 1.0 and emit false FMR warnings.
            # Do NOT "fix" this to group by base RefDes.
            if check_fmr:
                inst_ref = extract_instance_refdes(cause_text)
                validation_refdes_list.append(inst_ref if inst_ref else (refdes or ""))
            else:
                validation_refdes_list.append("")

            # Validate and convert usage value (C10 fix: check for inf/huge values)
            usage_val = pd.to_numeric(row[usage_idx], errors='coerce')
            if pd.isna(usage_val):
                usage = 1.0
                notes.append("Invalid Usage (NaN, defaulted 1.0)")
            else:
                usage = float(usage_val)
                # C10: Validate finite and reasonable range
                if not math.isfinite(usage) or usage <= 0 or usage > 1e6:
                    notes.append(f"Invalid Usage ({usage}, defaulted 1.0)")
                    usage = 1.0

            # Validate and convert ratio value (C10 fix: check for inf/huge values)
            ratio_val = pd.to_numeric(row[ratio_idx], errors='coerce')
            if pd.isna(ratio_val):
                ratio = 1.0
                notes.append("Invalid Ratio (NaN, defaulted 1.0)")
            else:
                ratio = float(ratio_val)
                # C10: Validate finite and reasonable range (0-1 expected for FMR)
                if not math.isfinite(ratio) or ratio < 0 or ratio > 10:
                    notes.append(f"Invalid Ratio ({ratio}, defaulted 1.0)")
                    ratio = 1.0

            corrected_ratio_list.append(ratio)

            # Calculate mode failure rate
            mode_fr = part_fr * usage * ratio
            mode_fr_list.append(mode_fr)

            # Format validation notes with trailing semicolon/space for consistency
            validation_notes_list.append("; ".join(notes) + ("; " if notes else ""))

        # Assign all columns at once (vectorized assignment)
        fmea['Extracted_RefDes'] = extracted_refdes_list
        fmea['Validation_RefDes'] = validation_refdes_list
        fmea['Part_FR'] = part_fr_list
        fmea['Mode_FR'] = mode_fr_list
        fmea['Corrected_Ratio'] = corrected_ratio_list
        fmea['Validation_Notes'] = validation_notes_list

        self.log(f"  Linked {linked_count} rows to Prediction data.")
        if missing_refs: self.log(f"  Warning: {len(missing_refs)} RefDes found in FMEA but missing in Prediction.")
        self._report_progress(0.85)  # L7: 85% after main linking

        if check_fmr:
            self.log("  Validating Failure Mode Ratios (Sum to 1.0)...")
            # Group by Validation_RefDes
            # Filter out empty validation refdes
            valid_rows = fmea[fmea['Validation_RefDes'] != ""]
            # Use Corrected_Ratio column for consistent validation with Mode_FR calculations
            ratio_sums = valid_rows.groupby('Validation_RefDes')['Corrected_Ratio'].sum()
            
            for ref, total in ratio_sums.items():
                self.cancel.check("Processing cancelled by user.")
                # Check if sum is 1.0 - use epsilon comparison for float reliability
                # M4 fix: Remove redundant round() - compare directly against tolerance
                if abs(total - 1.0) > FMR_TOLERANCE:
                    mask = fmea['Validation_RefDes'] == ref
                    fmea.loc[mask, 'Validation_Notes'] += f"FMR Sum {total:.2f} != 1.0; "
            self._report_progress(0.95)  # L7: 95% after FMR validation

        # Circuit/function-block roll-up (Tier-1 fix): a block row (FMEA Level =
        # Circuit/Function Block, or a cause cell listing multiple RefDes) is an
        # aggregate, not an independent part. Replace its phantom first-RefDes
        # rate with the SUM of its piece-part children and exclude it from the
        # function total so part rates are not double-counted. Returns False for
        # plain piece-part files, in which case the legacy Function-Description
        # roll-up runs unchanged.
        applied_block_rollup = self._apply_circuit_block_rollup(
            fmea, fmea_cause_col, fmea_func_col, fr_lookup
        )
        if not applied_block_rollup and fmea_func_col and fmea_func_col != "None":
            self.log("  Calculating Function Failure Rates...")
            func_fr = fmea.groupby(fmea_func_col)['Mode_FR'].transform('sum')
            fmea['Function_FR'] = func_fr

        self.merged_df = fmea
        self._report_progress(1.0)  # L7: 100% complete
        return fmea

    # FMEA-ID column synonyms used to tie a piece-part row to its circuit/
    # function-block parent by shared id prefix (block "CPU-001" owns
    # "CPU-001-R201-A"). Kept narrow to avoid grabbing a generic "ID Number".
    _FMEA_ID_SYNONYMS = ("fmea id", "fmea-id", "fmea_id", "fmeaid")

    def _detect_fmea_id_col(self, fmea: "pd.DataFrame") -> Optional[str]:
        """Return the FMEA-ID column name if present, else None."""
        for col in fmea.columns:
            col_text = str(col).lower().strip()
            if any(syn in col_text for syn in self._FMEA_ID_SYNONYMS):
                return col
        return None

    def _apply_circuit_block_rollup(
        self,
        fmea: "pd.DataFrame",
        fmea_cause_col: str,
        fmea_func_col: str,
        fr_lookup: dict,
    ) -> bool:
        """Roll circuit/function-block rows up from their piece-part children.

        A block row that HAS piece-part children (associated by shared FMEA-ID
        prefix, or by the block's listed RefDes when no FMEA-ID column exists)
        gets failure rate = SUM of its children's Mode_FR and is EXCLUDED from
        the function total (its children are counted instead).

        A CHILDLESS block is a LEAF: if its cause lists >1 RefDes its rate
        becomes the sum of those components' prediction lambdas (killing the
        phantom first-RefDes rate); a single-RefDes leaf keeps its computed rate.
        Leaf blocks are COUNTED in the function total. This is what prevents the
        regression where a block-only functional FMEA had every rate wiped to 0.

        Function_FR = the per-function sum of every COUNTED (non-rolled) row's
        Mode_FR, broadcast to all rows of that function (so it stays consistent
        when a function mixes block/child rows with ungrouped piece-parts).

        Returns True when a block structure was found and applied; False for a
        plain piece-part file (legacy Function-Description roll-up then runs).
        """
        token = self.cancel

        class _ClassifyCancel:
            def check(_self):
                if token is not None:
                    token.check("Processing cancelled by user.")

        # Classify on the ORIGINAL columns only — the result columns we just
        # added (Extracted_RefDes, etc.) would otherwise confuse the classifier.
        _added = {
            "Extracted_RefDes", "Validation_RefDes", "Part_FR", "Mode_FR",
            "Corrected_Ratio", "Validation_Notes", "Function_FR",
        }
        orig_cols = [c for c in fmea.columns if c not in _added]
        try:
            classifications, _level_col, _validated = classify_fmea_rows(
                fmea[orig_cols], log_func=None, cancel_token=_ClassifyCancel()
            )
        except InterruptedError:
            raise
        except Exception as exc:  # classification is best-effort; never break a run
            self.log(f"  Note: circuit-block classification skipped ({exc}).")
            return False

        if len(classifications) != len(fmea):
            return False
        row_types = [c.row_type for c in classifications]
        block_positions = [i for i, t in enumerate(row_types) if t == "circuit_block"]
        if not block_positions:
            return False  # plain piece-part file -> legacy function roll-up

        mode_fr = list(fmea["Mode_FR"])
        part_fr = list(fmea["Part_FR"])
        notes = list(fmea["Validation_Notes"])
        cause_vals = (
            list(fmea[fmea_cause_col]) if fmea_cause_col in fmea.columns
            else [""] * len(fmea)
        )

        # --- associate piece-part children to their block ---
        children: Dict[int, list] = {bp: [] for bp in block_positions}
        fmea_id_col = self._detect_fmea_id_col(fmea[orig_cols])
        if fmea_id_col is not None:
            ids = [str(v).strip() if pd.notna(v) else "" for v in fmea[fmea_id_col]]
            block_id_by_pos = {bp: ids[bp] for bp in block_positions if ids[bp]}
            for j, t in enumerate(row_types):
                if t == "circuit_block":
                    continue
                pp_id = ids[j]
                if not pp_id:
                    continue
                best_bp, best_len = None, -1
                for bp, bid in block_id_by_pos.items():
                    if (pp_id == bid or pp_id.startswith(bid + "-")) and len(bid) > best_len:
                        best_bp, best_len = bp, len(bid)
                if best_bp is not None:
                    children[best_bp].append(j)
            assoc = f"FMEA-ID column '{fmea_id_col}'"
        else:
            extracted = [str(v).strip().upper() for v in fmea["Extracted_RefDes"]]
            block_refsets = {}
            for bp in block_positions:
                raw = cause_vals[bp] if pd.notna(cause_vals[bp]) else ""
                refset = {normalize_refdes_for_lookup(tok) for tok in split_refdes_list(str(raw))}
                refset.discard("")
                block_refsets[bp] = refset
            for j, t in enumerate(row_types):
                if t == "circuit_block":
                    continue
                child_ref = normalize_refdes_for_lookup(extracted[j])
                if not child_ref:
                    continue
                for bp in block_positions:
                    if child_ref in block_refsets[bp]:
                        children[bp].append(j)
                        break
            assoc = "listed RefDes (no FMEA-ID column)"

        # --- set each block's rate; track which blocks are "rolled" (replaced by
        #     their children in the function total) vs "leaf" (counted directly) ---
        rolled = set()
        leaf_blocks = 0
        for bp in block_positions:
            kids = children[bp]
            if kids:
                r = sum(mode_fr[j] for j in kids)
                part_fr[bp] = r
                mode_fr[bp] = r
                notes[bp] = (notes[bp] or "") + (
                    f"Circuit-block roll-up of {len(kids)} piece-part row(s); "
                )
                rolled.add(bp)
            else:
                leaf_blocks += 1
                raw = cause_vals[bp] if (bp < len(cause_vals) and pd.notna(cause_vals[bp])) else ""
                listed = [normalize_refdes_for_lookup(t) for t in split_refdes_list(str(raw))]
                listed = [x for x in listed if x]
                if len(listed) > 1:
                    # Leaf block with multiple components but no enumerated child
                    # rows (block-only functional FMEA): rate = sum of the listed
                    # components' predicted lambdas, NOT the phantom first RefDes.
                    r = sum(fr_lookup.get(x, 0.0) for x in listed)
                    part_fr[bp] = r
                    mode_fr[bp] = r
                    notes[bp] = (notes[bp] or "") + (
                        f"Block roll-up of {len(listed)} listed component(s); "
                    )
                # else: single/zero-RefDes leaf -> keep its computed rate (do NOT zero).

        # --- Function_FR: per-function sum of every COUNTED (non-rolled) row's
        #     Mode_FR, broadcast to all rows of that function (consistent). ---
        function_fr: list = [None] * len(fmea)
        if fmea_func_col and fmea_func_col != "None" and fmea_func_col in fmea.columns:
            func_vals = list(fmea[fmea_func_col])
            eff = [(0.0 if i in rolled else mode_fr[i]) for i in range(len(fmea))]
            tmp = pd.DataFrame({"_func": func_vals, "_eff": eff})
            func_sum = tmp.groupby("_func")["_eff"].sum().to_dict()
            function_fr = [func_sum.get(func_vals[j]) for j in range(len(fmea))]
        else:
            # No function column: a block defines its own group (block + children).
            for bp in block_positions:
                function_fr[bp] = mode_fr[bp]
                for j in children[bp]:
                    function_fr[j] = mode_fr[bp]

        fmea["Part_FR"] = part_fr
        fmea["Mode_FR"] = mode_fr
        fmea["Validation_Notes"] = notes
        fmea["Function_FR"] = function_fr

        # --- structure log + self-check (sum of rolled blocks == sum of their
        #     children; flag piece-parts not attached to any rolled block) ---
        attached = set(j for v in children.values() for j in v)
        unattached_pp = sum(
            1 for j, t in enumerate(row_types)
            if t != "circuit_block" and j not in attached
        )
        self.log(
            f"  Circuit-block roll-up via {assoc}: {len(rolled)} block(s) rolled "
            f"from {len(attached)} child row(s), {leaf_blocks} leaf block(s)."
        )
        if unattached_pp and rolled:
            self.log(
                f"  Note: {unattached_pp} piece-part row(s) are not attached to any "
                f"block; counted directly in their function total."
            )
        return True

    def save_results(self, output_path: Union[str, Path]) -> None:
        """Save merged results to an Excel file."""
        if self.merged_df is None: return
        self.log(f"Saving report to: {output_path}")
        try:
            from openpyxl import Workbook

            export_df = self.merged_df.copy()
            for col in export_df.columns:
                export_df[col] = export_df[col].apply(to_user_facing_text)

            wb = Workbook()
            ws_main = wb.active
            ws_main.title = "Main"

            # Row styling based on validation status
            # M7 fix: Check column exists before accessing
            has_validation_col = 'Validation_Notes' in export_df.columns

            def main_row_style(row, idx):
                if not has_validation_col:
                    return 'default'
                notes = str(row.get('Validation_Notes', ''))
                has_ratio_sum_error = (
                    ('FMR Sum' in notes and '!= 1.0' in notes)
                    or ('Failure Mode Ratio Sum' in notes and '!= 1.0' in notes)
                )
                if has_ratio_sum_error:
                    return 'error'
                elif notes.strip():
                    return 'warning'
                return 'default'

            # Write and style Main sheet (with zebra striping for readability)
            write_df_to_sheet(ws_main, export_df)
            style_worksheet(ws_main, export_df, row_style_func=main_row_style, max_width=40, alternate_rows=True)

            # Apply scientific notation to failure rate columns
            from common.excel_styles import NUMBER_FORMATS
            from openpyxl.utils import get_column_letter
            fr_columns = ['Mode_FR', 'Function_FR', 'Failure_Rate', 'FR']
            for col_idx, col_name in enumerate(export_df.columns, start=1):
                if isinstance(col_name, str) and (col_name in fr_columns or col_name.endswith('_FR')):
                    col_letter = get_column_letter(col_idx)
                    for row_idx in range(2, ws_main.max_row + 1):
                        cell = ws_main[f"{col_letter}{row_idx}"]
                        if cell.value is not None and cell.value != '':
                            cell.number_format = NUMBER_FORMATS["scientific"]

            # Create Validation Warnings sheet if there are warnings
            # M7 fix: Only create sheet if column exists
            if has_validation_col and 'Extracted_RefDes' in export_df.columns:
                val_df = export_df[export_df['Validation_Notes'] != ""][['Extracted_RefDes', 'Validation_Notes']]
            else:
                val_df = pd.DataFrame()  # Empty DataFrame
            if not val_df.empty:
                ws_val = wb.create_sheet("Validation Warnings")
                write_df_to_sheet(ws_val, val_df)
                # All validation rows get warning styling
                style_worksheet(ws_val, val_df, row_style_func=lambda r, i: 'warning', max_width=60)

            wb.save(output_path)
            self.log("  Save complete.")
        except (IOError, OSError, PermissionError) as e:
            self.log(f"  Failed to save: {e}")
            raise
