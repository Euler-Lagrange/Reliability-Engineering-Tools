#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BOM Compare - Excel Export Module

Excel report writing for custom BOM-vs-BOM comparison results,
including FMEA coverage sheets and styled summary/detail sheets.
"""
from typing import Dict, List, Optional, Any

import pandas as pd

from common import (
    get_tool_logger,
    style_worksheet,
    write_df_to_sheet,
    to_reason_code_label,
    to_user_facing_text,
)

from .bom_compare_logic import BomCompareResult
from .fmea_coverage import FmeaCoverageResult

_logger = get_tool_logger("bom_compare")

# Excel sheet name constraints
EXCEL_INVALID_SHEET_CHARS = r'\/*?:[]'
EXCEL_MAX_SHEET_NAME_LEN = 31

def sanitize_sheet_name(name: str, max_len: int = EXCEL_MAX_SHEET_NAME_LEN) -> str:
    """Sanitize string for use as Excel sheet name.

    Excel sheet names have the following restrictions:
    - Cannot exceed 31 characters
    - Cannot contain: \\ / * ? : [ ]
    - Cannot be empty or consist only of whitespace

    Args:
        name: The proposed sheet name
        max_len: Maximum length (default 31 for Excel)

    Returns:
        Sanitized sheet name safe for Excel use
    """
    if not name or not name.strip():
        return "Sheet"

    # Replace invalid characters with underscore
    for char in EXCEL_INVALID_SHEET_CHARS:
        name = name.replace(char, '_')

    # Spaces are kept — the whole tool uses the group path's human naming
    # scheme ("Only In <file>", "Part Usage"). Collapse runs and trim so a
    # filename full of replaced characters still reads cleanly.
    name = ' '.join(name.split())

    # Remove consecutive underscores
    while '__' in name:
        name = name.replace('__', '_')

    # Strip leading/trailing underscores
    name = name.strip('_')

    # Truncate to max length (leave room for suffix if needed)
    if len(name) > max_len:
        name = name[:max_len - 3] + '...'

    return name if name else "Sheet"


def write_bom_compare_excel(
    result: BomCompareResult,
    filename: str,
    bom_a_name: str = "BOM A",
    bom_b_name: str = "BOM B",
    fmea_result_1: Optional['FmeaCoverageResult'] = None,
    fmea_result_2: Optional['FmeaCoverageResult'] = None,
) -> None:
    """Write BOM comparison results to Excel file.

    Args:
        result: BOM comparison result
        filename: Output file path
        bom_a_name: Display name for File 1
        bom_b_name: Display name for File 2
        fmea_result_1: Optional FMEA coverage result for File 1
        fmea_result_2: Optional FMEA coverage result for File 2
    """
    from openpyxl import Workbook
    from common.excel_styles import sanitize_for_excel

    wb = Workbook()

    # Sanitize display names to prevent XML corruption
    bom_a_name = sanitize_for_excel(bom_a_name)
    bom_b_name = sanitize_for_excel(bom_b_name)

    # Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary_data = [
        ["BOM Comparison Summary"],
        [],
        ["Metric", "Value"],
        [f"{bom_a_name} Total Rows", result.summary.get('bom_a_rows', 0)],
        [f"{bom_b_name} Total Rows", result.summary.get('bom_b_rows', 0)],
        [f"{bom_a_name} Unique RefDes", result.summary.get('bom_a_unique_refdes', 0)],
        [f"{bom_b_name} Unique RefDes", result.summary.get('bom_b_unique_refdes', 0)],
        [],
        [f"Only in {bom_a_name}", result.summary.get('only_in_a', 0)],
        [f"Only in {bom_b_name}", result.summary.get('only_in_b', 0)],
        ["In Both", result.summary.get('in_both', 0)],
        ["Column Differences", result.summary.get('differences', 0)],
        [],
        [f"RefDes with Duplicates in {bom_a_name}", result.summary.get('duplicates_a', 0)],
        [f"RefDes with Duplicates in {bom_b_name}", result.summary.get('duplicates_b', 0)],
        [f"Total Duplicate Rows in {bom_a_name}", result.summary.get('duplicate_rows_a', 0)],
        [f"Total Duplicate Rows in {bom_b_name}", result.summary.get('duplicate_rows_b', 0)],
        [],
        ["Part Usage Warnings", result.summary.get('part_usage_warnings', 0)],
        [f"Scope Warnings in {bom_a_name}", result.summary.get('scope_warnings_a', 0)],
        [f"Scope Warnings in {bom_b_name}", result.summary.get('scope_warnings_b', 0)],
        ["Scope Warnings (Total)", result.summary.get('scope_warnings_total', 0)],
    ]
    for row in summary_data:
        ws_summary.append(row)

    # Style the Summary sheet
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from common.excel_styles import PRESETS

    # Title row (row 1) - large, bold, primary color
    title_cell = ws_summary.cell(row=1, column=1)
    title_cell.font = Font(name=PRESETS.FONT_NAME, size=14, bold=True, color=PRESETS.colors.HEADER_BG)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    # Header row (row 3) - "Metric" and "Value"
    for col_idx in range(1, 3):
        cell = ws_summary.cell(row=3, column=col_idx)
        cell.font = PRESETS.header_font
        cell.fill = PRESETS.header_fill
        cell.alignment = PRESETS.header_alignment
        cell.border = PRESETS.thin_border

    # Data rows (rows 4 onwards) - apply data styling where content exists
    for row_idx in range(4, len(summary_data) + 1):
        for col_idx in range(1, 3):
            cell = ws_summary.cell(row=row_idx, column=col_idx)
            if cell.value not in (None, "", " "):
                cell.font = PRESETS.data_font
                cell.alignment = PRESETS.data_alignment
                cell.border = PRESETS.thin_border

    # Set column widths
    ws_summary.column_dimensions['A'].width = 40  # Metric column
    ws_summary.column_dimensions['B'].width = 15  # Value column

    # Only in A sheet (use sanitized sheet name to handle invalid chars and length)
    sheet_name_a = None  # Initialize for collision check below
    if result.only_in_a:
        sheet_name_a = sanitize_sheet_name(f"Only In {bom_a_name}")
        ws_a = wb.create_sheet(sheet_name_a)
        df_a = pd.DataFrame(result.only_in_a)
        write_df_to_sheet(ws_a, df_a)
        style_worksheet(ws_a, df_a, max_width=40, alternate_rows=True)

    # Only in B sheet (use sanitized sheet name to handle invalid chars and length)
    if result.only_in_b:
        sheet_name_b = sanitize_sheet_name(f"Only In {bom_b_name}")
        # Handle collision if both names sanitize to the same value
        if result.only_in_a and sheet_name_b == sheet_name_a:
            sheet_name_b = sanitize_sheet_name(f"Only In {bom_b_name} 2")
        ws_b = wb.create_sheet(sheet_name_b)
        df_b = pd.DataFrame(result.only_in_b)
        write_df_to_sheet(ws_b, df_b)
        style_worksheet(ws_b, df_b, max_width=40, alternate_rows=True)

    # Differences sheet
    if result.differences:
        ws_diff = wb.create_sheet("Differences")
        df_diff = pd.DataFrame(result.differences)
        write_df_to_sheet(ws_diff, df_diff)
        style_worksheet(ws_diff, df_diff, max_width=40, alternate_rows=True)

    # Duplicates sheet (if any duplicates found)
    all_duplicates = result.duplicates_a + result.duplicates_b
    if all_duplicates:
        ws_dups = wb.create_sheet("Duplicates")
        # Map internal Source values to display names
        source_map = {'BOM A': bom_a_name, 'BOM B': bom_b_name}
        # Flatten RowData for display - merge main fields with row data
        flattened_dups = []
        for dup in all_duplicates:
            internal_source = dup.get('Source', '')
            flat_row = {
                'RefDes': dup.get('RefDes'),
                'Source': source_map.get(internal_source, internal_source),
                'OccurrenceNum': dup.get('OccurrenceNum'),
                'Category': dup.get('Category'),
                'UsedForComparison': dup.get('UsedForComparison'),
            }
            # Add all row data fields (prefix with 'Row_' to distinguish)
            row_data = dup.get('RowData', {})
            for key, value in row_data.items():
                flat_row[f'Row_{key}'] = value
            flattened_dups.append(flat_row)
        df_dups = pd.DataFrame(flattened_dups)
        write_df_to_sheet(ws_dups, df_dups)
        style_worksheet(ws_dups, df_dups, max_width=40, alternate_rows=True)

    # Part Usage Warnings sheet (if any warnings found)
    if result.part_usage_warnings:
        ws_usage = wb.create_sheet("Part Usage")
        # Map File 1/2 source to display names
        source_map = {'File 1': bom_a_name, 'File 2': bom_b_name}
        usage_rows = []
        for w in result.part_usage_warnings:
            src = w.get('Source', '')
            usage_rows.append({
                'Source': source_map.get(src, src),
                'RefDes': w.get('RefDes'),
                'Base': w.get('Base'),
                'Usage': w.get('Usage'),
                'Expected': w.get('Expected'),
                'Count': w.get('Count'),
                'Reason Code': to_reason_code_label(w.get('ReasonCode', '')),
                'Reason': to_user_facing_text(w.get('Reason')),
            })
        df_usage = pd.DataFrame(usage_rows)
        write_df_to_sheet(ws_usage, df_usage)
        style_worksheet(ws_usage, df_usage, max_width=40, alternate_rows=True)

    # Failure Mode Ratio warnings sheet (custom-path check_fmr)
    if getattr(result, "fmr_warnings", None):
        # Same name as the group path's FMR sheet — one tool, one vocabulary.
        ws_fmr = wb.create_sheet("Failure Mode Ratio Errors")
        source_map = {'File 1': bom_a_name, 'File 2': bom_b_name}
        fmr_rows = []
        for w in result.fmr_warnings:
            src = w.get('Source', '')
            fmr_rows.append({
                'Source': source_map.get(src, src),
                'RefDes': w.get('RefDes'),
                'Sum': w.get('Sum'),
                # Same translation the group path applies — the raw
                # "FMR != 1.0" token is an unexpanded abbreviation.
                'Status': to_user_facing_text(w.get('Status')),
            })
        df_fmr = pd.DataFrame(fmr_rows)
        write_df_to_sheet(ws_fmr, df_fmr)
        style_worksheet(ws_fmr, df_fmr, max_width=40, alternate_rows=True)

    # Scope warnings sheet (for FMEA CB-vs-PP consistency warnings)
    if result.scope_warnings:
        ws_scope = wb.create_sheet("Scope Warnings")
        source_map = {'BOM A': bom_a_name, 'BOM B': bom_b_name}
        scope_rows = []
        for w in result.scope_warnings:
            src = w.get('Source', '')
            scope_rows.append({
                'Source': source_map.get(src, src),
                'RefDes': w.get('RefDes'),
                'Reason Code': to_reason_code_label(w.get('ReasonCode', '')),
                'Scope Status': to_user_facing_text(w.get('Scope Status')),
                'Cross-File': w.get('Cross-File'),
                'InCircuitBlock': w.get('InCircuitBlock'),
                'InPiecePart': w.get('InPiecePart'),
                'ScopeColumn': w.get('ScopeColumn'),
                'Details': to_user_facing_text(w.get('Details')),
            })
        df_scope = pd.DataFrame(scope_rows)
        write_df_to_sheet(ws_scope, df_scope)
        style_worksheet(ws_scope, df_scope, max_width=50, alternate_rows=True)

    # --- FMEA Coverage Sheets (if validation was run) ---
    def write_fmea_coverage_sheets(fmea_result, file_label: str):
        """Write FMEA coverage validation sheets for one file."""
        if fmea_result is None:
            return

        summary = fmea_result.summary
        # Use the file labels directly to avoid prefix collisions
        # Limit prefix to 22 chars so the longest suffix (" PP Index" = 9 chars)
        # stays under Excel's 31-char sheet-name cap
        raw_prefix = f"FMEA {file_label}"
        prefix = sanitize_sheet_name(raw_prefix[:22] if len(raw_prefix) > 22 else raw_prefix)

        # Summary sheet
        ws_summary = wb.create_sheet(f"{prefix} Summary")
        summary_rows = [
            ["FMEA Coverage Summary"],
            [],
            ["Metric", "Value"],
            ["Source File", fmea_result.source_file],
            ["FMEA Detected by Filename", str(fmea_result.detected_by_filename)],
            ["FMEA Level Column", fmea_result.fmea_level_column or "row_scan"],
            ["Level Column Validated", str(fmea_result.fmea_level_validated_by_values)],
            ["RefDes Column", fmea_result.refdes_column_detected or "per_row_best_cell"],
            ["RefDes Detection Method", fmea_result.refdes_detection_method],
            [],
            ["Total Block Rows", len(fmea_result.circuit_blocks)],
            ["Total Tokens", summary.total_tokens],
            ["Covered Tokens", summary.covered_tokens],
            ["Missing Tokens", summary.missing_tokens],
            ["Coverage %", f"{summary.coverage_pct:.1f}%"],
            [],
            ["Piece-Part Rows", len(fmea_result.piece_part_rows)],
            ["Classified by Level Column", fmea_result.classified_by_level_col_count],
            ["Classified by Row Scan", fmea_result.classified_by_row_scan_count],
        ]
        for row_data in summary_rows:
            ws_summary.append([sanitize_for_excel(str(v)) if v else "" for v in row_data])
        style_worksheet(ws_summary, pd.DataFrame(summary_rows[2:]), max_width=50)

        # Missing tokens sheet
        missing_rows = []
        for cb in fmea_result.circuit_blocks:
            for token in cb.tokens:
                if not token.has_piece_part:
                    missing_rows.append({
                        "BlockRowIndex": cb.row_index,
                        "BlockType": cb.fmea_level_value or "circuit_block",
                        "Token": token.token,
                        "OriginalCell": token.original_cell[:100] if token.original_cell else "",
                        "SourceColumn": token.source_column,
                    })

        if missing_rows:
            ws_missing = wb.create_sheet(f"{prefix} Missing")
            df_missing = pd.DataFrame(missing_rows)
            write_df_to_sheet(ws_missing, df_missing)
            style_worksheet(ws_missing, df_missing, max_width=40, alternate_rows=True)

        # All tokens sheet (detailed with CB_ and PP_FIRST_ columns)
        all_token_rows = []
        for cb in fmea_result.circuit_blocks:
            for token in cb.tokens:
                row = {
                    "BlockRowIndex": cb.row_index,
                    "BlockType": cb.fmea_level_value or "circuit_block",
                    "ClassificationSource": cb.classification_source,
                    "Token": token.token,
                    "HasPiecePart": token.has_piece_part,
                    "PiecePartRowCount": token.piece_part_row_count,
                    "PiecePartRowIndexes": ",".join(str(i) for i in token.piece_part_row_indexes[:5]),
                    "SourceColumn": token.source_column,
                }
                # Add CB_<col> columns from original circuit block row
                for col, val in cb.original_row_data.items():
                    row[f"CB_{col}"] = sanitize_for_excel(str(val)) if pd.notna(val) else ""
                # Add PP_FIRST_<col> columns from first matching piece-part row
                if token.first_piece_part_row_data:
                    for col, val in token.first_piece_part_row_data.items():
                        row[f"PP_FIRST_{col}"] = sanitize_for_excel(str(val)) if pd.notna(val) else ""
                all_token_rows.append(row)

        if all_token_rows:
            ws_tokens = wb.create_sheet(f"{prefix} Tokens")
            df_tokens = pd.DataFrame(all_token_rows)
            write_df_to_sheet(ws_tokens, df_tokens)
            style_worksheet(ws_tokens, df_tokens, max_width=40, alternate_rows=True)

        # Sheet 4: Piece-part index (for debugging)
        pp_index_rows = []
        for pp_class in fmea_result.piece_part_rows:
            row_data = fmea_result.piece_part_row_data.get(pp_class.row_index, {})
            row = {
                "PiecePartRowIndex": pp_class.row_index,
                "FmeaLevelValue": pp_class.fmea_level_value or "",
                "ClassificationSource": pp_class.classification_source,
            }
            # Add PP_<col> columns from original piece-part row
            for col, val in row_data.items():
                row[f"PP_{col}"] = sanitize_for_excel(str(val)) if pd.notna(val) else ""
            pp_index_rows.append(row)

        if pp_index_rows:
            ws_pp = wb.create_sheet(f"{prefix} PP Index")
            df_pp = pd.DataFrame(pp_index_rows)
            write_df_to_sheet(ws_pp, df_pp)
            style_worksheet(ws_pp, df_pp, max_width=40, alternate_rows=True)

    # Write FMEA coverage sheets for each file
    if fmea_result_1:
        write_fmea_coverage_sheets(fmea_result_1, bom_a_name)
    if fmea_result_2:
        write_fmea_coverage_sheets(fmea_result_2, bom_b_name)

    wb.save(filename)
    _logger.info(f"BOM comparison results saved to: {filename}")



