"""Tests for the Extraction Compare differ (bom_compare/extraction_compare.py).

Diffs two RefDes-extraction sheets into appeared / disappeared / moved-group
reports. The critical normalization traps: the (Verified)/(Unverified) split
must not read as two groups, and GROUP NOT DETECTED gap rows must never count
as extraction evidence — otherwise every rev-to-rev diff reports phantom
churn.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook, load_workbook

from bom_compare.extraction_compare import (
    ExtractionCompareResult,
    compare_extractions,
    load_component_groups,
    write_extraction_compare_excel,
)


def _sheet(rows: list[tuple[str, str]]) -> pd.DataFrame:
    return pd.DataFrame(
        [{"Group": group, "Failure Mode Causes": causes, "Component Count": 0, "Pages": ""}
         for group, causes in rows]
    )


def test_verified_unverified_split_reads_as_one_group() -> None:
    df = _sheet(
        [
            ("DIG-076 (Verified)", "U7, U9"),
            ("DIG-076 (Unverified)", "U12"),
        ]
    )
    groups = load_component_groups(df)
    assert groups == {"U7": "DIG-076", "U9": "DIG-076", "U12": "DIG-076"}


def test_gap_rows_are_not_extraction_evidence() -> None:
    df = _sheet(
        [
            ("DIG-076 (Verified)", "U7"),
            ("CPU-003 (GROUP NOT DETECTED)", "U99"),
        ]
    )
    groups = load_component_groups(df)
    assert "U99" not in groups
    assert groups == {"U7": "DIG-076"}


def test_ungrouped_bucket_is_a_named_group() -> None:
    df = _sheet([("UNGROUPED (IN BOM)", "U55")])
    assert load_component_groups(df) == {"U55": "UNGROUPED (IN BOM)"}


def test_compare_reports_appeared_disappeared_and_moved() -> None:
    rev_a = _sheet(
        [
            ("DIG-076 (Verified)", "U7, U9"),
            ("DIG-081 (Verified)", "C1"),
        ]
    )
    rev_b = _sheet(
        [
            ("DIG-076 (Verified)", "U7"),
            ("DIG-081 (Verified)", "U9, R5"),
        ]
    )
    result = compare_extractions(rev_a, rev_b)

    assert result.appeared == [{"component": "R5", "group": "DIG-081"}]
    assert result.disappeared == [{"component": "C1", "group": "DIG-081"}]
    assert result.moved == [
        {"component": "U9", "from_group": "DIG-076", "to_group": "DIG-081"}
    ]
    assert result.in_both_count == 2  # U7 unchanged, U9 moved
    assert result.total_a == 3 and result.total_b == 3


def test_identical_revisions_report_zero_churn() -> None:
    df = _sheet(
        [
            ("DIG-076 (Verified)", "U7"),
            ("DIG-076 (Unverified)", "U12"),
            ("CPU-003 (GROUP NOT DETECTED)", ""),
        ]
    )
    result = compare_extractions(df, df.copy())
    assert result.appeared == []
    assert result.disappeared == []
    assert result.moved == []


def test_missing_columns_raise_a_readable_error() -> None:
    df = pd.DataFrame([{"Something": "else"}])
    try:
        load_component_groups(df)
        raise AssertionError("expected ValueError")
    except ValueError as exc:
        assert "RefDes extraction sheet" in str(exc)


def test_report_workbook_has_four_styled_sheets(tmp_path) -> None:
    result = ExtractionCompareResult(
        appeared=[{"component": "R5", "group": "DIG-081"}],
        disappeared=[],
        moved=[{"component": "U9", "from_group": "DIG-076", "to_group": "DIG-081"}],
        in_both_count=2,
        total_a=3,
        total_b=3,
    )
    wb = Workbook()
    write_extraction_compare_excel(result, wb, name_a="revA", name_b="revB")
    path = tmp_path / "compare.xlsx"
    wb.save(path)

    loaded = load_workbook(path)
    assert loaded.sheetnames == ["Summary", "Appeared", "Disappeared", "Moved Groups"]
    appeared = loaded["Appeared"]
    assert appeared.cell(row=1, column=2).value == "Group in revB"
    assert appeared.cell(row=2, column=1).value == "R5"
    # Empty sheet still gets a styled, frozen header.
    disappeared = loaded["Disappeared"]
    assert disappeared.cell(row=1, column=1).value == "Component"
    assert disappeared.freeze_panes == "A2"
    moved = loaded["Moved Groups"]
    assert moved.cell(row=2, column=2).value == "DIG-076"
    assert moved.cell(row=2, column=3).value == "DIG-081"
