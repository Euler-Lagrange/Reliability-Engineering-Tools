"""Unit tests for bom_compare/unified_bom.py — the Unified BOM merge.

Spec: docs/superpowers/specs/2026-08-13-unified-bom-design.md
"""
from __future__ import annotations

import subprocess
import sys
import threading
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from bom_compare.unified_bom import (
    COL_NOTES,
    COL_SOURCE,
    COL_STATUS,
    SHEET_UNIFIED,
    STATUS_ADDED,
    STATUS_CARRIED,
    STATUS_CHANGED,
    STATUS_DELETE,
    STATUS_SUPERSEDED,
    UnifiedBomResult,
    _cell_text,
    _is_blank,
    _is_suffix_of,
    _plan_columns,
    _row_tokens,
    _values_equal,
    build_unified_bom,
    write_unified_bom_sheet,
)
from common import CancellationError


def test_plan_columns_union_order_and_tool_columns():
    old = pd.DataFrame(columns=["RefDes", "Sheet Number", "Commodity L1"])
    new = pd.DataFrame(columns=["RefDes", "Part Number"])
    cols, old_map, tool, notes = _plan_columns(old, new, None)
    # New file's columns first (original order), then old-only columns
    # (old-file order), then the three tool-authored columns.
    assert cols == [
        "RefDes", "Part Number", "Sheet Number", "Commodity L1",
        "Status", "Source", "Change Notes",
    ]
    assert old_map == {
        "RefDes": "RefDes",
        "Sheet Number": "Sheet Number",
        "Commodity L1": "Commodity L1",
    }
    assert tool == {"status": "Status", "source": "Source", "notes": "Change Notes"}
    assert notes == []


def test_plan_columns_case_insensitive_identity():
    old = pd.DataFrame(columns=["REFDES", "part number"])
    new = pd.DataFrame(columns=["RefDes", "Part Number"])
    cols, old_map, _tool, notes = _plan_columns(old, new, None)
    assert old_map == {"REFDES": "RefDes", "part number": "Part Number"}
    # Identity match must not duplicate the columns.
    assert cols[:2] == ["RefDes", "Part Number"]
    assert "REFDES" not in cols
    assert notes == []


def test_plan_columns_explicit_pairs_win():
    old = pd.DataFrame(columns=["Ref", "Desc"])
    new = pd.DataFrame(columns=["RefDes", "Part Description"])
    cols, old_map, _tool, notes = _plan_columns(
        old, new, [("Ref", "RefDes"), ("Desc", "Part Description")]
    )
    assert old_map == {"Ref": "RefDes", "Desc": "Part Description"}
    assert cols[:2] == ["RefDes", "Part Description"]
    assert "Ref" not in cols and "Desc" not in cols
    assert notes == []


def test_plan_columns_tool_name_collision_guard():
    old = pd.DataFrame(columns=["RefDes"])
    new = pd.DataFrame(columns=["RefDes", "Status"])
    cols, _old_map, tool, notes = _plan_columns(old, new, None)
    assert tool["status"] == "Merge Status"
    assert cols.count("Status") == 1
    assert cols.count("Merge Status") == 1
    assert notes == []


def test_plan_columns_collision_keeps_both_columns():
    old = pd.DataFrame(
        columns=["RefDes", "Part Number", "Part number", "Sheet Number"]
    )
    new = pd.DataFrame(columns=["RefDes", "Part Number"])
    cols, old_map, _tool, notes = _plan_columns(old, new, None)
    assert "Part Number" in cols
    assert "Part number (old)" in cols
    assert old_map["Part number"] == "Part number (old)"
    assert len(notes) == 1
    assert "Part number" in notes[0] and "Part Number" in notes[0]


def test_plan_columns_unresolved_pair_notes():
    old = pd.DataFrame(columns=["Ref"])
    new = pd.DataFrame(columns=["RefDes"])
    cols, old_map, _tool, notes = _plan_columns(old, new, [("Reff", "RefDes")])
    # Falls back to name matching: "Ref" doesn't fold-match "RefDes", so it
    # becomes its own old-only column.
    assert old_map == {"Ref": "Ref"}
    assert "Ref" in cols
    assert len(notes) == 1
    assert "was ignored" in notes[0]


def test_plan_columns_pairs_match_case_insensitively():
    old = pd.DataFrame(columns=["ref"])
    new = pd.DataFrame(columns=["RefDes"])
    cols, old_map, _tool, notes = _plan_columns(old, new, [("REF", "RefDes")])
    assert old_map == {"ref": "RefDes"}
    assert notes == []


def test_plan_columns_old_only_headers_are_trimmed():
    old = pd.DataFrame(columns=["  Sheet Number  "])
    new = pd.DataFrame(columns=["RefDes"])
    cols, old_map, _tool, _notes = _plan_columns(old, new, None)
    assert "Sheet Number" in cols
    assert old_map["  Sheet Number  "] == "Sheet Number"


def test_plan_columns_duplicate_pair_note():
    old = pd.DataFrame(columns=["Ref"])
    new = pd.DataFrame(columns=["RefDes", "Part Number"])
    cols, old_map, _tool, notes = _plan_columns(
        old, new, [("Ref", "RefDes"), ("Ref", "Part Number")]
    )
    # First pair wins; the second pair references an old column that's
    # already mapped, which is a truthfully different situation from an
    # unresolved pair.
    assert old_map == {"Ref": "RefDes"}
    assert len(notes) == 1
    assert "already mapped by an earlier pair" in notes[0]
    assert '"Ref"' in notes[0] and '"Part Number"' in notes[0]


def test_plan_columns_malformed_pair_note():
    old = pd.DataFrame(columns=["Ref"])
    new = pd.DataFrame(columns=["RefDes"])
    cols, old_map, _tool, notes = _plan_columns(old, new, [("Ref",)])
    assert len(notes) == 1
    assert notes[0] == "A malformed column pair entry was ignored"
    # The malformed pair doesn't block normal fold-identity fallback... but
    # "Ref" doesn't fold-match "RefDes" so it becomes its own old-only column.
    assert old_map == {"Ref": "Ref"}


def test_plan_columns_tool_collision_case_insensitive():
    old = pd.DataFrame(columns=["status"])
    new = pd.DataFrame(columns=["RefDes"])
    cols, _old_map, tool, _notes = _plan_columns(old, new, None)
    assert tool["status"] == "Merge Status"
    # The user's own "status" column must survive untouched in the union.
    assert "status" in cols
    assert "Status" not in cols


def test_plan_columns_readable_uniquification():
    old = pd.DataFrame(columns=["Qty", "QTY", "qty"])
    new = pd.DataFrame(columns=["Qty"])
    cols, old_map, _tool, _notes = _plan_columns(old, new, None)
    assert "QTY (old)" in cols
    assert "qty (old 2)" in cols
    # No triple-stacked "(old) (old) (old)" suffixes.
    assert not any("(old) (old)" in c for c in cols)
    assert old_map["QTY"] == "QTY (old)"
    assert old_map["qty"] == "qty (old 2)"


@pytest.mark.parametrize(
    ("a", "b", "equal"),
    [
        ("10", "10.0", True),
        (10, "10", True),
        ("RES 10K", "res 10k", True),
        ("  x ", "x", True),
        ("", None, True),
        (float("nan"), "", True),
        ("10", "10.5", False),
        ("A", "B", False),
        ("0123", "123", False),
        ("1e3", "1000", False),
        ("12345678901234567890", "12345678901234567891", False),
        ("1_000", "1000", False),
        ("+123", "123", True),
        ("10.50", "10.5", True),
        ("5", "5.00", True),
    ],
)
def test_values_equal(a, b, equal):
    assert _values_equal(a, b) is equal


def test_cell_text_handles_na_variants():
    assert _cell_text(pd.NA) == ""
    assert _cell_text(float("nan")) == ""
    assert _cell_text(None) == ""
    # A non-scalar cell must not crash pd.isna() and should str()-ify.
    assert _cell_text([1, 2, 3]) == "[1, 2, 3]"


def test_is_blank_basics():
    assert _is_blank("") is True
    assert _is_blank("  ") is True
    assert _is_blank(None) is True
    assert _is_blank(float("nan")) is True
    assert _is_blank("x") is False
    # 0 is a real value, not blank.
    assert _is_blank(0) is False


def test_row_tokens_variants():
    assert _row_tokens(pd.NA) == []
    assert _row_tokens(None) == []
    assert _row_tokens(float("nan")) == []
    assert _row_tokens("R1, R2") == ["R1", "R2"]
    assert _row_tokens("  r5 ") == ["R5"]


@pytest.mark.parametrize(
    ("token", "base", "expected"),
    [
        ("U2000-1", "U2000", True),
        ("U2000-D4", "U2000", True),
        ("J4-P1", "J4", True),
        ("J4-20", "J4", True),
        ("R1-R3", "R1", False),
        ("U1A", "U1", False),
        ("u2000-1", "U2000", True),
        ("u2000", "U2000", False),
        ("U2000", "U2000", False),
        ("", "U2000", False),
        ("J1-J5", "J1", False),
        ("P2-P4", "P2", False),
        ("CN1-CN3", "CN1", False),
        ("X1-X9", "X1", False),
    ],
)
def test_is_suffix_of_matrix(token, base, expected):
    assert _is_suffix_of(token, base) is expected


def _merge(old_rows, new_rows, **kwargs):
    """Shorthand: both files use a 'RefDes' key column unless overridden."""
    old_df = pd.DataFrame(old_rows)
    new_df = pd.DataFrame(new_rows)
    kwargs.setdefault("old_refdes_col", "RefDes")
    kwargs.setdefault("new_refdes_col", "RefDes")
    return build_unified_bom(old_df, new_df, **kwargs)


# ---------------------------------------------------------------------------
# Statuses, sources, styles
# ---------------------------------------------------------------------------

def test_added_row_is_green_and_noted():
    result = _merge(
        [{"RefDes": "R1", "Part Number": "PN-1"}],
        [{"RefDes": "R1", "Part Number": "PN-1"},
         {"RefDes": "R9", "Part Number": "PN-9"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["R1", "R9"]
    added = frame.iloc[1]
    assert added[COL_STATUS] == STATUS_ADDED
    assert added[COL_SOURCE] == "New BOM only"
    assert "add to the FMEAs" in added[COL_NOTES]
    assert result.row_styles[1] == "success"
    assert result.counts["added"] == 1


def test_unchanged_row_has_no_status_and_source_both():
    result = _merge([{"RefDes": "R1"}], [{"RefDes": "R1"}])
    row = result.frame.iloc[0]
    assert row[COL_STATUS] == ""
    assert row[COL_SOURCE] == "Both"
    assert row[COL_NOTES] == ""
    assert result.row_styles[0] == ""
    assert result.counts["unchanged"] == 1


def test_deleted_row_inserted_inline_after_anchor():
    result = _merge(
        [{"RefDes": "R1"}, {"RefDes": "R99"}, {"RefDes": "R2"}],
        [{"RefDes": "R1"}, {"RefDes": "R2"}],
    )
    assert list(result.frame["RefDes"]) == ["R1", "R99", "R2"]
    row = result.frame.iloc[1]
    assert row[COL_STATUS] == STATUS_DELETE
    assert row[COL_SOURCE] == "Old BOM only"
    assert "remove from the FMEAs" in row[COL_NOTES]
    assert result.row_styles[1] == "error"
    assert result.counts["deleted"] == 1


def test_deleted_first_row_goes_to_top():
    result = _merge(
        [{"RefDes": "R99"}, {"RefDes": "R1"}],
        [{"RefDes": "R1"}],
    )
    assert list(result.frame["RefDes"]) == ["R99", "R1"]


def test_consecutive_deletions_keep_old_file_order():
    result = _merge(
        [{"RefDes": "R1"}, {"RefDes": "R97"}, {"RefDes": "R98"}, {"RefDes": "R99"}],
        [{"RefDes": "R1"}],
    )
    assert list(result.frame["RefDes"]) == ["R1", "R97", "R98", "R99"]


# ---------------------------------------------------------------------------
# Conflict rule
# ---------------------------------------------------------------------------

def test_changed_value_new_wins_and_noted():
    result = _merge(
        [{"RefDes": "R1", "Part Description": "RES 10K 1%"}],
        [{"RefDes": "R1", "Part Description": "RES 10.5K 1%"}],
    )
    row = result.frame.iloc[0]
    assert row["Part Description"] == "RES 10.5K 1%"
    assert row[COL_STATUS] == STATUS_CHANGED
    assert row[COL_SOURCE] == "Both"
    assert (
        'Part Description changed from "RES 10K 1%" to "RES 10.5K 1%"'
        in row[COL_NOTES]
    )
    assert result.row_styles[0] == "warning"
    assert result.counts["changed"] == 1
    # Value changes are the feature's point, not data-integrity warnings.
    assert result.warning_count == 0


def test_numeric_equivalent_values_are_not_changes():
    result = _merge(
        [{"RefDes": "R1", "Qty": 10}],
        [{"RefDes": "R1", "Qty": "10.0"}],
    )
    row = result.frame.iloc[0]
    assert row[COL_STATUS] == ""
    assert row[COL_NOTES] == ""


def test_blank_new_value_keeps_old_and_warns():
    result = _merge(
        [{"RefDes": "R1", "Part Description": "RES 10K"}],
        [{"RefDes": "R1", "Part Description": ""}],
    )
    row = result.frame.iloc[0]
    assert row["Part Description"] == "RES 10K"
    assert 'Part Description: kept old value "RES 10K" (new was blank)' in row[COL_NOTES]
    assert row[COL_STATUS] == STATUS_CHANGED
    assert result.warning_count == 1


def test_old_blank_filled_by_new_value_is_silent():
    """J3 (accepted residual, docs/reviews/KNOWN_RESIDUALS.md): the mirror
    case of the test above. An old cell that was blank getting filled by a
    real new value is new data landing on the row, not old data being
    lost — no Changed note, no status, no warning."""
    result = _merge(
        [{"RefDes": "R1", "Part Description": ""}],
        [{"RefDes": "R1", "Part Description": "NEW DESC"}],
    )
    row = result.frame.iloc[0]
    assert row["Part Description"] == "NEW DESC"
    assert row[COL_STATUS] == ""
    assert row[COL_NOTES] == ""
    assert result.warning_count == 0


def test_old_only_column_carries_manual_data_silently():
    result = _merge(
        [{"RefDes": "R1", "Sheet Number": "12"}],
        [{"RefDes": "R1"}],
    )
    row = result.frame.iloc[0]
    assert row["Sheet Number"] == "12"
    # Old-only column carry is expected behavior — no note, no status.
    assert row[COL_STATUS] == ""


# ---------------------------------------------------------------------------
# Duplicates, DNP, multi-token rows, empties, NaN, plan notes, cancellation
# ---------------------------------------------------------------------------

def test_duplicate_new_refdes_first_wins_and_flagged():
    result = _merge(
        [{"RefDes": "R1", "Part Number": "PN-1"}],
        [{"RefDes": "R1", "Part Number": "PN-1"},
         {"RefDes": "R1", "Part Number": "PN-DUP"}],
    )
    dup = result.frame.iloc[1]
    assert "Duplicate RefDes R1" in dup[COL_NOTES]
    assert dup[COL_STATUS] == STATUS_CHANGED
    assert result.warning_count == 1


def test_dnp_old_only_row_not_emitted_as_delete():
    result = _merge(
        [{"RefDes": "R1", "Description": "Res"},
         {"RefDes": "R99", "Description": "DNP spare"}],
        [{"RefDes": "R1", "Description": "Res"}],
        old_desc_col="Description",
        ignore_dnp=True,
    )
    assert list(result.frame["RefDes"]) == ["R1"]
    assert result.counts["deleted"] == 0
    assert result.counts["dnp_skipped"] == 1


def test_dnp_old_only_row_emitted_when_ignore_dnp_off():
    result = _merge(
        [{"RefDes": "R1", "Description": "Res"},
         {"RefDes": "R99", "Description": "DNP spare"}],
        [{"RefDes": "R1", "Description": "Res"}],
        old_desc_col="Description",
        ignore_dnp=False,
    )
    assert list(result.frame["RefDes"]) == ["R1", "R99"]


def test_multi_token_old_row_partial_delete_names_origin():
    result = _merge(
        [{"RefDes": "R1, R2, R3", "Group": "CPU"}],
        [{"RefDes": "R1"}, {"RefDes": "R2"}],
    )
    # The old row survives via R1's match (entry 0), so its deleted token
    # anchors right after that backbone row.
    assert list(result.frame["RefDes"]) == ["R1", "R3", "R2"]
    deleted = result.frame.iloc[1]
    assert deleted[COL_STATUS] == STATUS_DELETE
    assert 'From an old row listing "R1, R2, R3"' in deleted[COL_NOTES]
    assert deleted["Group"] == "CPU"


def test_multi_token_new_row_mixed_status_is_changed():
    result = _merge(
        [{"RefDes": "R1", "Group": "CPU"}],
        [{"RefDes": "R1, R9", "Group": "CPU"}],
    )
    row = result.frame.iloc[0]
    assert row[COL_STATUS] == STATUS_CHANGED  # not Added: only R9 is new
    assert "R9 is not in Old BOM" in row[COL_NOTES]


def test_multi_token_new_row_all_new_is_added():
    result = _merge(
        [{"RefDes": "R1"}],
        [{"RefDes": "R8, R9"}],
    )
    # R1 is an unrelated, unmatched old-only row; per the "deletions with
    # no surviving anchor go at the top" rule it lands at iloc[0], so the
    # row under test is located by RefDes rather than assumed position.
    row = result.frame[result.frame["RefDes"] == "R8, R9"].iloc[0]
    assert row[COL_STATUS] == STATUS_ADDED


def test_empty_new_file_emits_all_deletes():
    result = _merge([{"RefDes": "R1"}], [])
    assert list(result.frame["RefDes"]) == ["R1"]
    assert result.frame.iloc[0][COL_STATUS] == STATUS_DELETE


def test_empty_old_file_all_added():
    result = _merge([], [{"RefDes": "R1"}])
    assert result.frame.iloc[0][COL_STATUS] == STATUS_ADDED


def test_nan_refdes_cells_do_not_crash():
    result = _merge(
        [{"RefDes": None, "X": "1"}, {"RefDes": "R1", "X": "2"}],
        [{"RefDes": "R1", "X": "2"}, {"RefDes": float("nan"), "X": "3"}],
    )
    assert len(result.frame) == 2  # NaN old row has no tokens -> no delete


def test_numeric_headers_do_not_crash():
    old = pd.DataFrame([["R1", "5"]], columns=[0, 1])
    new = pd.DataFrame([["R1", "6"]], columns=[0, 1])
    result = build_unified_bom(old, new, old_refdes_col="0", new_refdes_col="0")
    assert result.frame.iloc[0][COL_STATUS] == STATUS_CHANGED  # "5" -> "6" noted


def test_unlisted_prefix_suffix_becomes_flagged_delete():
    # R is not in the instance-notation prefix whitelist, so R1-1 never
    # carries — it must surface as a visible Delete with the review note.
    result = _merge(
        [{"RefDes": "R1-1", "Sheet Number": "4"}],
        [{"RefDes": "R1"}],
    )
    frame = result.frame
    deleted = frame[frame["RefDes"] == "R1-1"].iloc[0]
    assert deleted[COL_STATUS] == STATUS_DELETE
    assert "manual expansion was not auto-carried" in deleted[COL_NOTES]
    assert deleted["Sheet Number"] == "4"


def test_bare_old_row_flags_new_expansion_rows_when_deleted():
    """Coordinator ruling on QA judgment call 3 (reverse case): the old
    file's bare row has no match, but the NEW file lists its expansion/
    suffix rows instead. This is likely a legitimate suffix-row
    replacement, not a genuine deletion, so the Delete row's note must say
    so instead of the plain "remove from the FMEAs" wording."""
    result = _merge(
        [{"RefDes": "U2000"}],
        [{"RefDes": "U2000-1"}, {"RefDes": "U2000-2"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["U2000", "U2000-1", "U2000-2"]
    assert list(frame[COL_STATUS]) == [STATUS_DELETE, STATUS_ADDED, STATUS_ADDED]
    delete_row = frame.iloc[0]
    assert (
        "U2000 is not in New BOM as a bare row, but its expansion rows are"
        in delete_row[COL_NOTES]
    )
    assert "likely replaced by suffix rows" in delete_row[COL_NOTES]


def test_genuine_deletion_keeps_plain_delete_note():
    """Negative control for the reverse case above: when the old file's
    bare row truly has no trace in the new file (no bare match, no base
    match, no expansion/suffix rows either), the Delete row must keep the
    plain "remove from the FMEAs" wording — not the "expansion rows"/
    "auto-carried" review language reserved for likely suffix-row
    replacements."""
    result = _merge(
        [{"RefDes": "R7"}],
        [{"RefDes": "R1"}],
    )
    frame = result.frame
    delete_row = frame[frame["RefDes"] == "R7"].iloc[0]
    assert delete_row[COL_STATUS] == STATUS_DELETE
    assert delete_row[COL_NOTES] == "Not in New BOM — remove from the FMEAs"
    assert "expansion rows" not in delete_row[COL_NOTES]
    assert "auto-carried" not in delete_row[COL_NOTES]


def test_plan_notes_surface_in_result():
    result = _merge(
        [{"RefDes": "R1", "Part Number": "PN", "Part number": "pn2"}],
        [{"RefDes": "R1", "Part Number": "PN"}],
    )
    assert any("kept separately" in n for n in result.notes)
    assert result.warning_count >= 1


def test_summary_note_wording_and_counts():
    result = _merge(
        [{"RefDes": "R1"}, {"RefDes": "R99"}],
        [{"RefDes": "R1"}, {"RefDes": "R9"}],
    )
    assert result.counts["added"] == 1
    assert result.counts["deleted"] == 1
    note = result.notes[0]
    assert note.startswith("Unified BOM: ")
    assert "1 added" in note and "1 to delete" in note


# ---------------------------------------------------------------------------
# Adversarial-QA hardening batch (B1-B6)
# ---------------------------------------------------------------------------

def test_duplicate_old_refdes_surfaces_note_and_is_dropped():
    # B1: a duplicate old RefDes must never vanish silently — the surviving
    # (first-occurrence) row's notes must say so, and the second row's data
    # must genuinely be dropped (not silently overwrite the first).
    result = _merge(
        [{"RefDes": "R1", "Sheet Number": "10", "Note": "checked"},
         {"RefDes": "R1", "Sheet Number": "77", "Note": "SECOND"}],
        [{"RefDes": "R1"}],
    )
    row = result.frame.iloc[0]
    assert (
        "Duplicate RefDes R1 in Old BOM — old row 3 was not merged"
        in row[COL_NOTES]
    )
    assert row[COL_STATUS] == STATUS_CHANGED
    assert result.warning_count == 1
    assert row["Sheet Number"] == "10"
    assert row["Note"] == "checked"


def test_duplicate_old_refdes_note_reaches_delete_row():
    # B1: when the duplicated token is old-only (deleted, not matched by
    # the backbone), the duplicate note must land on the Delete row.
    result = _merge(
        [{"RefDes": "R5", "X": "a"}, {"RefDes": "R5", "X": "b"}],
        [{"RefDes": "R1"}],
    )
    delete_row = result.frame[result.frame["RefDes"] == "R5"].iloc[0]
    assert delete_row[COL_STATUS] == STATUS_DELETE
    assert (
        "Duplicate RefDes R5 in Old BOM — old row 3 was not merged"
        in delete_row[COL_NOTES]
    )
    assert result.warning_count == 1


def test_duplicate_old_refdes_leftover_when_first_occurrence_dnp_suppressed():
    # B1: if the first occurrence itself is DNP-suppressed (never emitted
    # as a Delete row and never matched by the backbone), the duplicate
    # note has nowhere to attach — it must still surface in result.notes.
    result = _merge(
        [{"RefDes": "R1", "Description": "DNP spare"},
         {"RefDes": "R1", "Description": "DNP spare 2"}],
        [],
        old_desc_col="Description",
        ignore_dnp=True,
    )
    assert any(
        "Duplicate RefDes R1 in Old BOM — old row 3 was not merged" in n
        for n in result.notes
    )
    assert result.warning_count == 1


def test_dnp_detection_strips_control_characters():
    # B2: a stray control character embedded in the RefDes/description text
    # must not defeat DNP detection — every sibling site strips these.
    result = _merge(
        [{"RefDes": "R1", "Description": "Res"},
         {"RefDes": "R99", "Description": "SPARE D\x00NP DO NOT LOAD"}],
        [{"RefDes": "R1", "Description": "Res"}],
        old_desc_col="Description",
        ignore_dnp=True,
    )
    assert result.counts["dnp_skipped"] == 1
    assert "R99" not in list(result.frame["RefDes"])


def test_duplicate_new_row_source_reflects_known_token():
    # B3: a duplicate new-file row whose token IS in the old file must not
    # claim "New BOM only" just because match-by-luck excluded it.
    result = _merge(
        [{"RefDes": "R1", "Part Number": "PN-1"}],
        [{"RefDes": "R1", "Part Number": "PN-1"},
         {"RefDes": "R1", "Part Number": "PN-DUP"}],
    )
    dup = result.frame.iloc[1]
    assert dup[COL_SOURCE] == "Both"


def test_invalid_dnp_regex_falls_back_with_note():
    # B4: an invalid custom DNP pattern falls back to the default silently
    # in effect, but must not be silent in the notes.
    result = _merge(
        [{"RefDes": "R1"}],
        [{"RefDes": "R1"}],
        dnp_regex="[unclosed",
    )
    assert (
        "Unified BOM: the custom DNP pattern was invalid — the default "
        "pattern was used." in result.notes
    )


def test_differing_old_values_old_only_column_is_a_warning():
    # B5: two old rows feeding one multi-token new row disagree on an
    # old-only column's value — that's a genuine data conflict, not free.
    result = _merge(
        [{"RefDes": "R1", "Sheet Number": "10"},
         {"RefDes": "R2", "Sheet Number": "20"}],
        [{"RefDes": "R1, R2"}],
    )
    row = result.frame.iloc[0]
    assert "differing old values" in row[COL_NOTES]
    assert result.warning_count == 1


def test_differing_old_values_shared_column_blank_new_is_a_warning():
    # B5: same conflict, but on a shared column where the new cell is
    # blank — the first old value is kept (flagged) and the second old
    # row's disagreement must ALSO be flagged, not silently dropped.
    result = _merge(
        [{"RefDes": "R1", "Part Description": "A"},
         {"RefDes": "R2", "Part Description": "B"}],
        [{"RefDes": "R1, R2", "Part Description": ""}],
    )
    row = result.frame.iloc[0]
    assert "differing old values" in row[COL_NOTES]
    # One blank-keep flag (R1's carry) + one differing-values flag (R2's
    # conflict with the value R1 already put in the cell).
    assert result.warning_count == 2


def test_differing_key_column_names_suppress_old_key_column():
    # B6: when the two files' key columns have different headers, the old
    # key column renders blank on every surviving row and filled only on
    # Delete rows — a broken-looking column. It must be dropped from the
    # union; Delete rows still carry their token in the new key column.
    old_df = pd.DataFrame([
        {"Reference Designator": "R1"},
        {"Reference Designator": "R99"},
    ])
    new_df = pd.DataFrame([{"RefDes": "R1"}])
    result = build_unified_bom(
        old_df, new_df,
        old_refdes_col="Reference Designator",
        new_refdes_col="RefDes",
    )
    assert "Reference Designator" not in result.frame.columns
    deleted = result.frame[result.frame[COL_STATUS] == STATUS_DELETE].iloc[0]
    assert deleted["RefDes"] == "R99"


def test_differing_key_column_names_keep_old_key_when_new_file_empty():
    # B6: with no new file at all, there is no new key column to speak
    # of — the old key column must remain as the output key, unsuppressed.
    old_df = pd.DataFrame([{"Reference Designator": "R1"}])
    new_df = pd.DataFrame([])
    result = build_unified_bom(
        old_df, new_df,
        old_refdes_col="Reference Designator",
        new_refdes_col="RefDes",
    )
    assert "Reference Designator" in result.frame.columns


def test_differing_key_column_names_do_not_suppress_real_new_file_column():
    # B6 regression: the OLD key's header text can coincide with a REAL
    # data column the NEW file owns, distinct from the new file's own key
    # column ("Reference Designator" here). Suppression must only remove
    # a column the old key actually OWNS — never a column the new file
    # populates itself, or that column's real values silently vanish via
    # the DataFrame(..., columns=unified_cols) projection.
    old_df = pd.DataFrame([{"RefDes": "R1", "Sheet": "3"}])
    new_df = pd.DataFrame([
        {"Reference Designator": "R1", "RefDes": "ALT-1", "PN": "PN-1"},
    ])
    result = build_unified_bom(
        old_df, new_df,
        old_refdes_col="RefDes",
        new_refdes_col="Reference Designator",
    )
    assert "RefDes" in result.frame.columns
    assert result.frame.iloc[0]["RefDes"] == "ALT-1"


def test_duplicate_old_note_pops_once_when_row_is_both_matched_and_delete_source():
    # Pin: the SAME old row can supply both a backbone match (via one
    # token on a multi-token cell) and a Delete row (via another token on
    # that same cell) while also carrying a duplicate-old-RefDes note.
    # The note must land exactly once — on the backbone match, which
    # processes (and pops) it first — never duplicated onto the Delete row.
    result = _merge(
        [{"RefDes": "R1, D9"}, {"RefDes": "R1"}],
        [{"RefDes": "R1"}],
    )
    matched = result.frame[result.frame["RefDes"] == "R1"].iloc[0]
    deleted = result.frame[result.frame["RefDes"] == "D9"].iloc[0]
    dup_note = "Duplicate RefDes R1 in Old BOM — old row 3 was not merged"
    assert dup_note in matched[COL_NOTES]
    assert dup_note not in deleted[COL_NOTES]
    assert result.warning_count == 1


def test_duplicate_old_note_attaches_to_exactly_one_of_two_delete_rows():
    # Pin: when a dup-noted old row produces MULTIPLE deleted tokens, the
    # single pop() call only fires for the first token processed — the
    # second Delete row sourced from the same old row must not also carry
    # the note (no double-count), but the note must not vanish either.
    result = _merge(
        [{"RefDes": "D8, D9"}, {"RefDes": "D8"}],
        [{"RefDes": "R1"}],
    )
    d8 = result.frame[result.frame["RefDes"] == "D8"].iloc[0]
    d9 = result.frame[result.frame["RefDes"] == "D9"].iloc[0]
    dup_note = "Duplicate RefDes D8 in Old BOM — old row 3 was not merged"
    has_dup = [dup_note in d8[COL_NOTES], dup_note in d9[COL_NOTES]]
    assert has_dup.count(True) == 1
    assert result.warning_count == 1


def test_cancellation_via_stop_event():
    stop = threading.Event()
    stop.set()
    rows = [{"RefDes": f"R{i}"} for i in range(10)]
    with pytest.raises(CancellationError):
        build_unified_bom(
            pd.DataFrame(rows),
            pd.DataFrame(rows),
            old_refdes_col="RefDes",
            new_refdes_col="RefDes",
            stop_event=stop,
        )


# ---------------------------------------------------------------------------
# Task 3: suffix/pin carry-forward with supersession
# ---------------------------------------------------------------------------

def test_suffix_carry_forward_supersedes_bare_row():
    result = _merge(
        [{"RefDes": "U2000-1", "Part Number": "PN-123", "Sheet Number": "5"},
         {"RefDes": "U2000-2", "Part Number": "PN-123", "Sheet Number": "6"}],
        [{"RefDes": "U2000", "Part Number": "PN-123"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["U2000", "U2000-1", "U2000-2"]
    bare = frame.iloc[0]
    assert bare[COL_STATUS] == STATUS_SUPERSEDED
    assert bare[COL_SOURCE] == "New BOM (superseded)"
    assert (
        "Superseded by 2 carried rows (U2000-1 … U2000-2) — delete this row"
        in bare[COL_NOTES]
    )
    assert result.row_styles[0] == "highlight"
    for i in (1, 2):
        assert frame.iloc[i][COL_STATUS] == STATUS_CARRIED
        assert frame.iloc[i][COL_SOURCE] == "Old BOM (carried)"
        assert result.row_styles[i] == "warning"
    assert frame.iloc[1]["Sheet Number"] == "5"
    assert frame.iloc[2]["Sheet Number"] == "6"
    assert result.counts["superseded"] == 1
    assert result.counts["carried"] == 2
    assert result.counts["deleted"] == 0


def test_pin_style_rows_carry_exactly_like_suffix_rows():
    result = _merge(
        [{"RefDes": "J4-P1", "Signal": "GND"},
         {"RefDes": "J4-20", "Signal": "VCC"}],
        [{"RefDes": "J4"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["J4", "J4-P1", "J4-20"]
    assert frame.iloc[0][COL_STATUS] == STATUS_SUPERSEDED
    assert frame.iloc[1][COL_STATUS] == STATUS_CARRIED
    assert frame.iloc[1]["Signal"] == "GND"
    assert frame.iloc[2]["Signal"] == "VCC"


def test_uniform_column_rule_updates_pn_but_not_per_row_data():
    result = _merge(
        [{"RefDes": "U2000-1", "Part Number": "123-456", "Sheet Number": "5"},
         {"RefDes": "U2000-2", "Part Number": "123-456", "Sheet Number": "6"}],
        [{"RefDes": "U2000", "Part Number": "123-789", "Sheet Number": "9"}],
    )
    frame = result.frame
    carried = frame.iloc[1:3]
    assert list(carried["Part Number"]) == ["123-789", "123-789"]
    assert (
        'Part Number updated from "123-456" to "123-789" by New BOM'
        in frame.iloc[1][COL_NOTES]
    )
    assert list(carried["Sheet Number"]) == ["5", "6"]
    assert result.warning_count == 1


def test_blank_new_value_never_overwrites_carried_rows():
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "PN-A"},
         {"RefDes": "U1-2", "Part Number": "PN-A"}],
        [{"RefDes": "U1", "Part Number": ""}],
    )
    carried = result.frame.iloc[1:3]
    assert list(carried["Part Number"]) == ["PN-A", "PN-A"]


def test_new_only_columns_fill_carried_rows_from_bare_row():
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "PN-A"},
         {"RefDes": "U1-2", "Part Number": "PN-A"}],
        [{"RefDes": "U1", "Part Number": "PN-A", "Supplier": "ACME"}],
    )
    carried = result.frame.iloc[1:3]
    assert list(carried["Supplier"]) == ["ACME", "ACME"]


def test_old_bare_and_suffix_rows_supersession_still_wins():
    result = _merge(
        [{"RefDes": "U1", "Part Number": "PN-A"},
         {"RefDes": "U1-1", "Part Number": "PN-A"}],
        [{"RefDes": "U1", "Part Number": "PN-A"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["U1", "U1-1"]
    assert frame.iloc[0][COL_STATUS] == STATUS_SUPERSEDED
    assert frame.iloc[1][COL_STATUS] == STATUS_CARRIED


def test_suffixes_present_in_both_files_match_exactly():
    result = _merge(
        [{"RefDes": "U1-1"}, {"RefDes": "U1-2"}],
        [{"RefDes": "U1-1"}, {"RefDes": "U1-2"}],
    )
    assert list(result.frame[COL_STATUS]) == ["", ""]
    assert result.counts["carried"] == 0


def test_carry_skips_suffixes_the_new_file_already_lists():
    result = _merge(
        [{"RefDes": "U1-1", "Sheet Number": "5"},
         {"RefDes": "U1-2", "Sheet Number": "6"}],
        [{"RefDes": "U1"},
         {"RefDes": "U1-1", "Sheet Number": "50"}],
    )
    frame = result.frame
    # U1-2 carries under the bare row; U1-1 exact-matches its own new row.
    assert list(frame["RefDes"]) == ["U1", "U1-2", "U1-1"]
    assert frame.iloc[0][COL_STATUS] == STATUS_SUPERSEDED
    assert frame.iloc[1][COL_STATUS] == STATUS_CARRIED
    assert frame.iloc[2][COL_STATUS] == STATUS_CHANGED  # 5 -> 50 noted


# ---------------------------------------------------------------------------
# Adversarial-QA hardening batch on Task 3 (Fixes 1-4)
# ---------------------------------------------------------------------------

def test_duplicate_bare_row_carries_suffix_block_only_once():
    # Fix 1: a second new-file row bearing the same bare base must not
    # re-supersede/re-carry the suffix group — the block carries exactly
    # once, and the duplicate bare row falls through to the ordinary
    # duplicate-token path (Changed/Both/dup note).
    result = _merge(
        [{"RefDes": "U2000-1", "Sheet Number": "1"},
         {"RefDes": "U2000-2", "Sheet Number": "2"}],
        [{"RefDes": "U2000"}, {"RefDes": "U2000"}],
    )
    frame = result.frame
    assert list(frame["RefDes"]) == ["U2000", "U2000-1", "U2000-2", "U2000"]
    assert frame.iloc[0][COL_STATUS] == STATUS_SUPERSEDED
    assert frame.iloc[1][COL_STATUS] == STATUS_CARRIED
    assert frame.iloc[2][COL_STATUS] == STATUS_CARRIED
    dup_row = frame.iloc[3]
    assert dup_row[COL_STATUS] == STATUS_CHANGED
    assert dup_row[COL_SOURCE] == "Both"
    assert "Duplicate RefDes U2000" in dup_row[COL_NOTES]
    assert result.counts["carried"] == 2
    assert result.counts["superseded"] == 1


def test_uniform_rule_overwrites_when_group_has_one_non_blank_value():
    # Fix 2 (a): non-blank consensus among the group's populated cells
    # (blanks excluded), bare row differs -> every carried row's cell
    # (blank or not) is overwritten, with a note on each, flagged once.
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "OLD-PN"},
         {"RefDes": "U1-2", "Part Number": ""},
         {"RefDes": "U1-3", "Part Number": ""}],
        [{"RefDes": "U1", "Part Number": "NEW-PN"}],
    )
    carried = result.frame.iloc[1:4]
    assert list(carried["Part Number"]) == ["NEW-PN", "NEW-PN", "NEW-PN"]
    for _, row in carried.iterrows():
        assert (
            'Part Number updated from "OLD-PN" to "NEW-PN" by New BOM'
            in row[COL_NOTES]
        )
    assert result.warning_count == 1


def test_uniform_rule_blank_fills_with_consensus_when_bare_matches():
    # Fix 2 (b): same group, but the bare row's value equals the group's
    # consensus -> no overwrite/no note; blank cells fill with the
    # consensus text instead (silent), non-blank cells are untouched.
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "OLD-PN"},
         {"RefDes": "U1-2", "Part Number": ""},
         {"RefDes": "U1-3", "Part Number": ""}],
        [{"RefDes": "U1", "Part Number": "OLD-PN"}],
    )
    carried = result.frame.iloc[1:4]
    assert list(carried["Part Number"]) == ["OLD-PN", "OLD-PN", "OLD-PN"]
    for _, row in carried.iterrows():
        assert "updated from" not in row[COL_NOTES]
    assert result.warning_count == 0


def test_genuine_disagreement_leaves_every_cell_verbatim():
    # Fix 2 (c): the old suffix rows genuinely disagree (not just blanks)
    # -> every cell (including blanks) stays exactly as the old data has
    # it; no bare-row backfill, no note, no flag.
    result = _merge(
        [{"RefDes": "U1-1", "Sheet Number": "1"},
         {"RefDes": "U1-2", "Sheet Number": "2"},
         {"RefDes": "U1-3", "Sheet Number": ""}],
        [{"RefDes": "U1", "Sheet Number": "9"}],
    )
    carried = result.frame.iloc[1:4]
    assert list(carried["Sheet Number"]) == ["1", "2", ""]
    for _, row in carried.iterrows():
        assert "Sheet Number" not in row[COL_NOTES]
    assert result.warning_count == 0


def test_uniform_rule_single_row_group_still_overwrites_and_flags():
    # Fix 2 (d): pin the trivial single-row-group case — consensus of one
    # is still a consensus, so a genuine bare-row difference still
    # overwrites and flags exactly like a multi-row group would.
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "OLD"}],
        [{"RefDes": "U1", "Part Number": "NEW"}],
    )
    carried = result.frame.iloc[1]
    assert carried["Part Number"] == "NEW"
    assert 'Part Number updated from "OLD" to "NEW" by New BOM' in carried[COL_NOTES]
    assert result.warning_count == 1


def test_uniform_rule_counts_once_per_column_not_per_row():
    # Fix 2 (e): two independently-uniform-and-changed columns in the same
    # carried group must flag twice (once per column), proving the
    # integrity count is column-scoped, not row-scoped.
    result = _merge(
        [{"RefDes": "U1-1", "Part Number": "OLD-PN", "Value": "10K"},
         {"RefDes": "U1-2", "Part Number": "OLD-PN", "Value": "10K"}],
        [{"RefDes": "U1", "Part Number": "NEW-PN", "Value": "22K"}],
    )
    carried = result.frame.iloc[1:3]
    assert list(carried["Part Number"]) == ["NEW-PN", "NEW-PN"]
    assert list(carried["Value"]) == ["22K", "22K"]
    assert result.warning_count == 2


def test_duplicate_old_suffix_row_note_lands_on_carried_row():
    # Fix 3: a duplicated old RefDes that's ALSO a manual suffix row must
    # have its duplicate note anchor to the visible carried row it
    # produces, not fall through silently into result.notes.
    result = _merge(
        [{"RefDes": "U2000-1", "Sheet Number": "FIRST"},
         {"RefDes": "U2000-1", "Sheet Number": "SECOND"},
         {"RefDes": "U2000-2"}],
        [{"RefDes": "U2000"}],
    )
    frame = result.frame
    carried = frame[frame["RefDes"] == "U2000-1"].iloc[0]
    assert "Duplicate RefDes U2000-1" in carried[COL_NOTES]
    assert carried["Sheet Number"] == "FIRST"
    assert not any("Duplicate RefDes U2000-1" in n for n in result.notes)
    assert result.warning_count == 1


# ---------------------------------------------------------------------------
# Tasks 4+5: styled sheet writer, facade re-export, writer integration
# ---------------------------------------------------------------------------

def test_write_unified_bom_sheet_styles_and_layout(tmp_path):
    result = _merge(
        [{"RefDes": "R1", "Part Description": "RES 10K"},
         {"RefDes": "R99", "Part Description": "OLD ONLY"},
         {"RefDes": "U1-1", "Part Description": "SECTION"},
         {"RefDes": "R2", "Part Description": "SAME"}],
        [{"RefDes": "R1", "Part Description": "RES 12K"},
         {"RefDes": "R5", "Part Description": "NEW PART"},
         {"RefDes": "U1", "Part Description": "SECTION"},
         {"RefDes": "R2", "Part Description": "SAME"}],
    )
    wb = Workbook()
    write_unified_bom_sheet(wb, result)
    out = tmp_path / "unified.xlsx"
    wb.save(out)

    read = load_workbook(out)
    assert SHEET_UNIFIED in read.sheetnames
    ws = read[SHEET_UNIFIED]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None

    headers = [c.value for c in ws[1]]
    status_idx = headers.index(COL_STATUS) + 1
    fills_by_status = {}
    unchanged_row_idx = None
    for row_idx in range(2, ws.max_row + 1):
        status = ws.cell(row=row_idx, column=status_idx).value
        if status:
            fills_by_status[status] = str(
                ws.cell(row=row_idx, column=1).fill.start_color.rgb
            )
        elif unchanged_row_idx is None:
            unchanged_row_idx = row_idx
    assert fills_by_status[STATUS_CHANGED].endswith("FFEB9C")     # yellow
    assert fills_by_status[STATUS_ADDED].endswith("C6EFCE")       # green
    assert fills_by_status[STATUS_DELETE].endswith("FFC7CE")      # red
    assert fills_by_status[STATUS_SUPERSEDED].endswith("BDD7EE")  # blue
    assert fills_by_status[STATUS_CARRIED].endswith("FFEB9C")     # yellow

    # An unchanged row (blank Status) keeps default styling — no solid fill.
    assert unchanged_row_idx is not None
    assert ws.cell(row=unchanged_row_idx, column=1).fill.patternType is None


def test_write_unified_bom_sheet_empty_frame_headers_only():
    result = build_unified_bom(
        pd.DataFrame(), pd.DataFrame(),
        old_refdes_col="RefDes", new_refdes_col="RefDes",
    )
    wb = Workbook()
    write_unified_bom_sheet(wb, result)
    ws = wb[SHEET_UNIFIED]
    assert ws.freeze_panes == "A2"
    headers = [c.value for c in ws[1]]
    assert COL_STATUS in headers
    assert COL_SOURCE in headers
    assert COL_NOTES in headers


def test_facade_reexports_unified_names():
    from bom_compare import bom_compare_logic
    assert bom_compare_logic.build_unified_bom is build_unified_bom
    assert bom_compare_logic.SHEET_UNIFIED == SHEET_UNIFIED
    assert bom_compare_logic.UnifiedBomResult is UnifiedBomResult
    assert bom_compare_logic.write_unified_bom_sheet is write_unified_bom_sheet


def test_submodule_first_import_is_order_safe():
    """The facade re-export must not make `import bom_compare.unified_bom`
    order-dependent (adversarial-QA finding from Task 1: the cycle is broken
    only by the lazy DEFAULT_DNP_REGEX import — pin that forever)."""
    proc = subprocess.run(
        [sys.executable, "-c",
         "import bom_compare.unified_bom; import bom_compare.bom_compare_logic"],
        cwd=str(Path(__file__).resolve().parents[1] / "python"),
        capture_output=True, text=True, timeout=120,
    )
    assert proc.returncode == 0, proc.stderr


def test_group_report_includes_unified_tab(tmp_path):
    from bom_compare.bom_compare_logic import AnalyzeResults, write_excel_report

    unified = _merge([{"RefDes": "R1"}], [{"RefDes": "R1"}])
    results = AnalyzeResults(
        summary=pd.DataFrame([("Files", "2")], columns=["Item", "Value"])
    )
    out = tmp_path / "group.xlsx"
    write_excel_report(results, str(out), unified=unified)
    assert SHEET_UNIFIED in load_workbook(out).sheetnames


def test_group_report_without_unified_is_unchanged(tmp_path):
    from bom_compare.bom_compare_logic import AnalyzeResults, write_excel_report

    results = AnalyzeResults(
        summary=pd.DataFrame([("Files", "2")], columns=["Item", "Value"])
    )
    out = tmp_path / "group_plain.xlsx"
    write_excel_report(results, str(out))
    assert SHEET_UNIFIED not in load_workbook(out).sheetnames


def test_custom_report_includes_unified_tab(tmp_path):
    from bom_compare.bom_compare_logic import BomCompareResult, write_bom_compare_excel

    unified = _merge([{"RefDes": "R1"}], [{"RefDes": "R1"}])
    out = tmp_path / "custom.xlsx"
    write_bom_compare_excel(BomCompareResult(), str(out), unified=unified)
    assert SHEET_UNIFIED in load_workbook(out).sheetnames


def test_custom_report_without_unified_is_unchanged(tmp_path):
    from bom_compare.bom_compare_logic import BomCompareResult, write_bom_compare_excel

    out = tmp_path / "custom_plain.xlsx"
    write_bom_compare_excel(BomCompareResult(), str(out))
    assert SHEET_UNIFIED not in load_workbook(out).sheetnames
