"""Read-layer parity tests (Tier-1 fixes #10 and #11).

#10: literal 'NA' / 'N/A' text must NOT be coerced to NaN when read from .xlsx
     (it already isn't for .csv), while genuinely-empty cells still read as NaN.
#11: duplicate header names must be deduped with pandas' '.1' suffix scheme so
     the inspect dropdown matches the column names the runtime binds against.
"""
from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from common.utils import try_read_table
from sidecar_main import _finalize_headers


@pytest.fixture
def na_workbook(tmp_path):
    p = tmp_path / "na.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["RefDes", "Note"])
    ws.append(["U1", None])      # truly empty
    ws.append(["U2", "NA"])
    ws.append(["U3", "N/A"])
    ws.append(["U4", "normal"])
    wb.save(p)
    return str(p)


def test_excel_preserves_na_literal_text_but_keeps_empty_nan(na_workbook):
    """'NA'/'N/A' survive as text; the genuinely-empty cell still reads NaN."""
    df = try_read_table(na_workbook)
    notes = list(df["Note"])
    assert pd.isna(notes[0])      # empty cell stays NaN (behavior preserved)
    assert notes[1] == "NA"       # literal text preserved (was silently NaN before)
    assert notes[2] == "N/A"
    assert notes[3] == "normal"


def test_csv_and_excel_agree_on_na_literal(tmp_path, na_workbook):
    """The same NA-literal data reads identically from .xlsx and .csv."""
    csv = tmp_path / "na.csv"
    csv.write_text("RefDes,Note\nU1,\nU2,NA\nU3,N/A\nU4,normal\n", encoding="utf-8")
    df_xlsx = try_read_table(na_workbook)
    df_csv = try_read_table(str(csv))
    assert list(df_xlsx["Note"])[1:] == ["NA", "N/A", "normal"]
    assert list(df_csv["Note"])[1:] == ["NA", "N/A", "normal"]


def test_finalize_headers_dedupes_like_pandas():
    """Duplicate headers get pandas' '.1'/'.2' suffixes; uniques are untouched."""
    assert _finalize_headers(["Part Number", "Part Number", "Desc"]) == [
        "Part Number", "Part Number.1", "Desc",
    ]
    assert _finalize_headers(["A", "A", "A"]) == ["A", "A.1", "A.2"]
    assert _finalize_headers(["RefDes", "Qty"]) == ["RefDes", "Qty"]


def test_finalize_headers_matches_pandas_runtime_frame(tmp_path):
    """The inspect dropdown names equal the column names the runtime binds."""
    p = tmp_path / "dup.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["RefDes", "Part Number", "Desc", "Part Number"])
    ws.append(["U1", "PN-A", "x", "PN-B"])
    wb.save(p)
    runtime_cols = list(try_read_table(str(p)).columns)
    inspect_cols = _finalize_headers(["RefDes", "Part Number", "Desc", "Part Number"])
    assert inspect_cols == runtime_cols


def test_finalize_headers_collision_with_preexisting_dotted_header():
    """A pre-existing 'Name.1' alongside duplicate 'Name's must still match
    pandas' mangle and produce NO duplicate names (QA collision case)."""
    assert _finalize_headers(["A", "A", "A", "B", "A.1"]) == ["A", "A.2", "A.3", "B", "A.1"]
    assert _finalize_headers(["Part", "Part", "Part.1"]) == ["Part", "Part.2", "Part.1"]
    out = _finalize_headers(["A", "A", "A", "B", "A.1"])
    assert len(out) == len(set(out))  # no duplicate names


def test_finalize_headers_matches_pandas_on_collision(tmp_path):
    """inspect names == runtime frame names even with a pre-existing dotted header."""
    p = tmp_path / "collide.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["A", "A", "A", "B", "A.1"]
    ws.append(headers)
    ws.append(["x"] * len(headers))
    wb.save(p)
    runtime_cols = list(try_read_table(str(p)).columns)
    assert _finalize_headers(headers) == runtime_cols


def test_canonicalize_refdes_family_guards_nan_and_none() -> None:
    """Final-day audit F3: a NaN/None cell fed to the canonicalize family
    must yield an empty string, not a spurious "NAN"/"NONE" RefDes key
    (split_refdes_list already guards this way)."""
    import pandas as pd

    from common.refdes_utils import (
        canonicalize_refdes,
        get_base_refdes,
        get_usage_base_refdes,
    )

    assert canonicalize_refdes(float("nan")) == ""
    assert canonicalize_refdes(None) == ""
    assert canonicalize_refdes(pd.NA) == ""
    assert get_base_refdes(float("nan")) == ""
    assert get_usage_base_refdes(None) == ""
    # Real tokens are unaffected.
    assert canonicalize_refdes("r100") == "R100"


def test_get_prefix_strips_invisible_characters() -> None:
    """Final-day audit F4: get_prefix must strip the same zero-width /
    invisible characters canonicalize_refdes strips — a PDF-pasted
    "​R1" is still an R-prefixed component."""
    from common.refdes_utils import get_prefix

    assert get_prefix("​R1") == "R"
    assert get_prefix("﻿C22") == "C"
    assert get_prefix("R1") == "R"


def test_style_array_stamp_is_not_shared(tmp_path) -> None:
    """Perf-cache guard: style_worksheet stamps cached StyleArray COPIES.
    openpyxl mutates a cell''s _style in place, so a shared array would let
    a later number_format assignment bleed across every same-styled cell
    (the FMEA fraction/scientific formats are applied exactly that way)."""
    import pandas as pd
    from openpyxl import Workbook, load_workbook

    from common.excel_styles import style_worksheet, write_df_to_sheet

    df = pd.DataFrame(
        [
            {"RefDes": "R1", "Part Usage": 0.5},
            {"RefDes": "R2", "Part Usage": 0.25},
        ]
    )
    wb = Workbook()
    ws = wb.active
    write_df_to_sheet(ws, df)
    style_worksheet(ws, df, alternate_rows=False)

    # Both data cells in the Part Usage column share the cached style combo.
    ws.cell(row=2, column=2).number_format = "# ???/???"

    assert ws.cell(row=2, column=2).number_format == "# ???/???"
    assert ws.cell(row=3, column=2).number_format != "# ???/???"

    # The styling itself survives a real save/load round-trip.
    out = tmp_path / "styled.xlsx"
    wb.save(out)
    wb.close()
    rt = load_workbook(out)
    try:
        assert rt.active.cell(row=2, column=1).font.name == ws.cell(row=2, column=1).font.name
        assert rt.active.cell(row=2, column=2).number_format == "# ???/???"
        assert rt.active.cell(row=3, column=2).number_format != "# ???/???"
    finally:
        rt.close()
