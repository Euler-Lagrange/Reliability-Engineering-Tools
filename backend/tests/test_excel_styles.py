"""excel_styles writer defensiveness (deferred LOW fix C).

`write_df_to_sheet` / `write_styled_excel` call `pd.isna(value)` per cell.
On a list-/ndarray-valued cell that raises
``ValueError: The truth value of an array ... is ambiguous`` and aborts the
whole sheet write. The writers must tolerate non-scalar cells by stringifying
them instead of crashing.
"""
from __future__ import annotations

import openpyxl
import pandas as pd

from common.excel_styles import write_df_to_sheet, write_styled_excel


def test_write_df_to_sheet_handles_list_valued_cell(tmp_path):
    """A list-valued cell round-trips to a written workbook without raising."""
    df = pd.DataFrame(
        {
            "RefDes": ["U1", "U2"],
            "Pins": [["1", "2", "3"], []],
        }
    )
    wb = openpyxl.Workbook()
    ws = wb.active

    # Must not raise ValueError("truth value of an array is ambiguous").
    write_df_to_sheet(ws, df)

    out = tmp_path / "list_cells.xlsx"
    wb.save(out)

    reloaded = openpyxl.load_workbook(out)
    rs = reloaded.active
    assert rs.cell(row=2, column=1).value == "U1"
    # The list cell is stringified rather than crashing the write.
    assert rs.cell(row=2, column=2).value == str(["1", "2", "3"])
    assert rs.cell(row=3, column=2).value == str([])


def test_write_styled_excel_handles_list_valued_cell(tmp_path):
    """The styled writer's data pass tolerates non-scalar cells too."""
    df = pd.DataFrame(
        {
            "RefDes": ["U1"],
            "Notes": [["a", "b"]],
        }
    )
    out = tmp_path / "styled_list_cells.xlsx"

    # Must not raise during the data-write pass.
    write_styled_excel(df, str(out), sheet_name="Data")

    reloaded = openpyxl.load_workbook(out)
    rs = reloaded["Data"]
    assert rs.cell(row=2, column=1).value == "U1"
    assert rs.cell(row=2, column=2).value == str(["a", "b"])


def test_write_df_to_sheet_preserves_scalar_nan_and_values(tmp_path):
    """Scalar cells keep their existing behavior: NaN -> blank, values kept."""
    df = pd.DataFrame(
        {
            "RefDes": ["U1", "U2"],
            "Value": [1.5, float("nan")],
        }
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    write_df_to_sheet(ws, df)
    out = tmp_path / "scalars.xlsx"
    wb.save(out)

    reloaded = openpyxl.load_workbook(out)
    rs = reloaded.active
    assert rs.cell(row=2, column=2).value == 1.5
    # NaN becomes an empty string cell (openpyxl reads back as None).
    assert rs.cell(row=3, column=2).value in (None, "")
