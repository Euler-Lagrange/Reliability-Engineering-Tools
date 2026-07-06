"""Phase D tests for the FMEA Generator backend.

Covers:
- BOM inheritance: pin/variant RefDes (e.g., U200-X) inherits from base U200
- BOM_Additions tracking: each inherited row recorded with usage_fraction = 1/N
- functional_to_piecepart workflow: detects circuit blocks in functional FMEA
  and inserts piece-part rows beneath each
- FMD-91 vs FMD-2016 standard selection: drives output column headers
- Validation: missing failureModesStandard rejected
- Cascade prevention: variant whose base is also missing falls through cleanly

These tests exercise the FMEAProcessor and runtime.validate_run_request directly
rather than spawning the sidecar subprocess, because they're focused unit tests
for the Phase D logic.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from common.exceptions import ValidationError
from fmea.fmea_generator_logic import FMEAProcessor, _index_to_suffix, write_excel_report
from fmea.runtime import execute_run_request, validate_run_request


def test_index_to_suffix_is_bijective_and_unique():
    """FMEA-ID overflow suffixes must stay unique past 26 failure modes.

    The previous inline scheme (``letters[idx] if idx < 26 else f"Z{idx}"``) was
    not bijective; ``_index_to_suffix`` uses base-26 bijective numeration
    (A..Z, AA..AZ, BA..) so every index maps to a distinct suffix.
    """
    assert _index_to_suffix(0) == "A"
    assert _index_to_suffix(25) == "Z"
    assert _index_to_suffix(26) == "AA"
    assert _index_to_suffix(27) == "AB"
    assert _index_to_suffix(51) == "AZ"
    assert _index_to_suffix(52) == "BA"
    # No collisions across a wide range (the old ``Z{idx}`` scheme reused "Z").
    suffixes = [_index_to_suffix(i) for i in range(1000)]
    assert len(set(suffixes)) == 1000


# ----- helpers ----------------------------------------------------------------


def _write_fixture(
    tmp_path: Path,
    *,
    bom_rows: list[dict],
    grouping_rows: list[dict] | None = None,
    fm_rows: list[dict] | None = None,
    functional_rows: list[dict] | None = None,
) -> dict:
    """Build standard test fixtures and return a dict of paths."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    bom_path = tmp_path / "bom.xlsx"
    pd.DataFrame(bom_rows).to_excel(bom_path, index=False)
    paths["bom"] = bom_path

    if grouping_rows is not None:
        grouping_path = tmp_path / "grouping.xlsx"
        pd.DataFrame(grouping_rows).to_excel(grouping_path, index=False)
        paths["grouping"] = grouping_path

    fm_path = tmp_path / "failure_modes.xlsx"
    pd.DataFrame(
        fm_rows
        or [
            {
                "FMD-2016 Commodity Type 1": "Microcircuit",
                "FMD-2016 Commodity Type 2": "Digital",
                "Failure Mode": "Stuck high",
                "Failure Mode Ratio": 0.5,
            },
            {
                "FMD-2016 Commodity Type 1": "Microcircuit",
                "FMD-2016 Commodity Type 2": "Digital",
                "Failure Mode": "Stuck low",
                "Failure Mode Ratio": 0.5,
            },
        ]
    ).to_excel(fm_path, index=False)
    paths["fm"] = fm_path

    if functional_rows is not None:
        func_path = tmp_path / "functional.xlsx"
        pd.DataFrame(functional_rows).to_excel(func_path, index=False)
        paths["func"] = func_path

    return paths


# ----- D2/D5/D10: Inheritance + BOM Additions sheet ---------------------------


def test_inheritance_creates_bom_addition(tmp_path: Path) -> None:
    """U200-A/B/C reference a base U200 that exists in BOM but the variants
    do not. Each variant should generate a piece-part row that inherits
    Part Number/Description/HDA/FMD from U200, and each should appear in
    proc.bom_additions with usage_fraction == "1/3".
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/3",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200-A, U200-B, U200-C",
                "Function Description": "Processor variants",
                "Schematic Page": "12",
            }
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "template_preserve": False,
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )

    assert not df.empty
    pp_rows = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    refdes_seen = set(pp_rows["Failure Mode Causes"].dropna().tolist())
    assert refdes_seen == {"U200-A", "U200-B", "U200-C"}, refdes_seen

    # All three variants should inherit Part Number IC-1234 from U200
    inherited_pns = set(pp_rows["Component Part Number"].dropna().tolist())
    assert inherited_pns == {"IC-1234"}, inherited_pns

    # bom_additions captured all three variants with usage_fraction = 1/3
    assert len(proc.bom_additions) == 3
    assert {e["ref_des"] for e in proc.bom_additions} == {"U200-A", "U200-B", "U200-C"}
    assert {e["base_refdes"] for e in proc.bom_additions} == {"U200"}
    assert {e["usage_fraction"] for e in proc.bom_additions} == {"1/3"}
    assert {e["part_number"] for e in proc.bom_additions} == {"IC-1234"}
    assert all(e["source_workflow"] == "piece_part_generate" for e in proc.bom_additions)


def test_inheritance_writes_bom_additions_sheet(tmp_path: Path) -> None:
    """write_excel_report should emit a "FMEA Gen New RefDes" sheet whenever
    proc.bom_additions is non-empty (Phase 4 / A9 rename of BOM_Additions)."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U500",
                "Part Number": "PN-9999",
                "Description": "FPGA",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/2",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "FPGA-001",
                "Reference Designator": "U500-X, U500-Y",
                "Function Description": "FPGA pins",
                "Schematic Page": "5",
            }
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    out_path = tmp_path / "report.xlsx"
    write_excel_report(df, out_path, proc)
    assert out_path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    try:
        # Phase 4 / A9: the user-visible sheet title was renamed from
        # "BOM_Additions" to "FMEA Gen New RefDes".
        assert "FMEA Gen New RefDes" in wb.sheetnames, wb.sheetnames
        sheet = wb["FMEA Gen New RefDes"]
        # Phase D banner row at row 1 + header row at row 2 + 2 data rows = 4
        assert sheet.max_row >= 4
        # Row 1 is the explanatory banner (merged across the column range)
        banner = sheet.cell(row=1, column=1).value or ""
        assert "RefDes variants" in banner, f"Unexpected banner: {banner!r}"
        assert "inherited" in banner.lower()
        # Row 2 carries the actual column headers
        headers = [sheet.cell(row=2, column=col).value for col in range(1, sheet.max_column + 1)]
        assert "RefDes" in headers
        assert "Base RefDes" in headers
        assert "Usage" in headers
    finally:
        wb.close()


def test_variant_in_bom_does_not_trigger_inheritance(tmp_path: Path) -> None:
    """Phase D regression: when BOTH the variant RefDes AND its base are
    in the BOM, the exact match must win. The variant should get its own
    BOM row (not inherit from the base), and bom_additions must be empty.
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200",
                "Part Number": "IC-BASE",
                "Description": "Base component",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1",
            },
            {
                "Reference Designator": "U200-X",
                "Part Number": "IC-VARIANT",
                "Description": "Variant component",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "G-001",
                "Reference Designator": "U200-X",
                "Function Description": "Variant use",
                "Schematic Page": "1",
            }
        ],
    )
    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    # bom_additions must be empty — U200-X is a direct BOM row, no inheritance
    assert proc.bom_additions == [], proc.bom_additions
    # The generated piece-part row must carry IC-VARIANT (from the exact
    # match), NOT IC-BASE (which would indicate incorrect inheritance).
    pp = df[
        (df["Failure Mode Causes"] == "U200-X")
        & (df["_row_type"].isin(("piece_part", "validation_warning", "piece_part_no_match")))
    ]
    assert len(pp) >= 1, df
    part_numbers = set(pp["Component Part Number"].dropna().tolist())
    assert part_numbers == {"IC-VARIANT"}, (
        f"Expected IC-VARIANT, got {part_numbers}"
    )


def test_variant_base_also_missing_falls_through(tmp_path: Path) -> None:
    """When neither the variant nor its base is in the BOM, no inheritance
    occurs and no bom_additions entry is created."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C100",
                "Part Number": "CAP-1",
                "Description": "Capacitor",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "GAP-001",
                "Reference Designator": "U999-A",
                "Function Description": "Missing variant",
                "Schematic Page": "1",
            }
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )

    assert proc.bom_additions == [], proc.bom_additions
    # The missing variant must end up in group_missing_in_bom
    assert any(ref == "U999-A" for _, ref in proc.group_missing_in_bom)


# ----- D4: FMD-91 vs FMD-2016 standard selection ------------------------------


def test_failure_modes_standard_drives_headers(tmp_path: Path) -> None:
    """Output column headers should match the chosen FMD standard."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "FMD-91 Commodity Type 1": "Resistor",
                "FMD-91 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-91",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    assert "FMD-91 Commodity Type 1" in df.columns
    assert "FMD-91 Commodity Type 2" in df.columns
    assert "FMD-2016 Commodity Type 1" not in df.columns


def test_validate_run_accepts_fill_gaps_with_preserve_formatting(tmp_path: Path) -> None:
    """Phase E: validate_run_request should now accept fill_gaps with the
    preserve-formatting strategy. The legacy block that rejected it
    ("Existing workbook and preserve-formatting paths stay in analysis-only
    mode") was removed in Phase D."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
    )
    # Existing FMEA file (any valid xlsx works for validation level)
    existing_fmea = tmp_path / "existing.xlsx"
    pd.DataFrame([{"FMEA-ID": "X-001", "FMEA Level": "Piece Part"}]).to_excel(
        existing_fmea, index=False,
    )
    target = tmp_path / "target.xlsx"
    pd.DataFrame([{"FMEA-ID": "T-001", "FMEA Level": "Circuit Block"}]).to_excel(
        target, index=False,
    )

    def state(role: str, path: Path, sheet: str = "Sheet1") -> dict:
        return {
            "role": role,
            "label": role,
            "path": str(path),
            "selectedSheet": sheet,
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "s1", "label": sheet}],
        }

    body = {
        "workflowId": "fill_gaps",
        "outputStrategyId": "existing_workbook_preserve_formatting",
        "options": {
            "failureModesStandard": "FMD-2016",
        },
        "inputs": [
            state("existingFmea", existing_fmea),
            state("bom", paths["bom"]),
            state("failureModes", paths["fm"]),
            state("targetWorkbook", target),
        ],
        # Fix C2: merge modes require Failure Mode Causes mapping.
        "mappings": [
            {
                "canonical": "Failure Mode Causes",
                "mappedTo": "Failure Mode Causes",
                "status": "mapped",
            },
        ],
    }
    result = validate_run_request(body)
    assert result["ok"] is True, result


def test_failure_modes_standard_required_in_validation() -> None:
    """validate_run_request should reject FMEA bodies without failureModesStandard."""
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "inputs": [
            {
                "role": "grouping",
                "label": "Grouping",
                "path": "/tmp/g.xlsx",
                "selectedSheet": "Sheet1",
                "sheets": [{"id": "s1", "label": "Sheet1"}],
            },
            {
                "role": "bom",
                "label": "BOM",
                "path": "/tmp/b.xlsx",
                "selectedSheet": "Sheet1",
                "sheets": [{"id": "s1", "label": "Sheet1"}],
            },
            {
                "role": "failureModes",
                "label": "FM",
                "path": "/tmp/f.xlsx",
                "selectedSheet": "Sheet1",
                "sheets": [{"id": "s1", "label": "Sheet1"}],
            },
        ],
        "mappings": [],
        # Note: no "options" key, no failureModesStandard.
    }
    result = validate_run_request(body)
    assert result["ok"] is False
    assert result["reason_code"] == "missing_failure_modes_standard"


# ----- D7: functional_to_piecepart workflow -----------------------------------


def test_usage_fraction_uses_source_variant_count_not_addition_count(tmp_path: Path) -> None:
    """Phase D C5 regression: usage_fraction must reflect the source-derived
    variant count, not the length of bom_additions. When U200 IS in BOM and
    U200-A is the only variant referenced in the grouping file, the variant
    count for U200 is 1 (just U200-A), so usage_fraction == '1/1'. If the
    code were using len(bom_additions), this would also be '1/1', so to
    discriminate we need a second variant that ISN'T in bom_additions.
    We put U200-A in grouping (missing from BOM → inheritance kicks in) AND
    U200 in the grouping file (present in BOM → no inheritance). That gives
    variant_counts_by_base[U200] == 2 but bom_additions has only 1 entry
    (U200-A). Expected usage_fraction == '1/2', NOT '1/1'."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/2",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200, U200-A",
                "Function Description": "CPU + variant",
                "Schematic Page": "12",
            }
        ],
    )
    proc = FMEAProcessor()
    proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    # Only U200-A inherits (U200 is directly in BOM).
    assert len(proc.bom_additions) == 1, proc.bom_additions
    assert proc.bom_additions[0]["ref_des"] == "U200-A"
    # Variant count for U200 base == 2 (U200 + U200-A in grouping).
    # Therefore usage_fraction for the inherited row must be "1/2".
    assert proc.bom_additions[0]["usage_fraction"] == "1/2", proc.bom_additions[0]


def test_inherited_variant_does_not_get_false_usage_warning(tmp_path: Path) -> None:
    """Phase D C5 regression: inherited variants must NOT be flagged as
    "Usage mismatch" validation warnings just because the BOM base has
    count=1. With U200 in BOM (Part Usage=1/3) and 3 variants in grouping,
    the inherited variants should see expected = 1/3 (from variant count),
    NOT 1/1 (from BOM usage_base_count)."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/3",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200-A, U200-B, U200-C",
                "Function Description": "Processor variants",
                "Schematic Page": "12",
            }
        ],
    )
    proc = FMEAProcessor()
    proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    # 3 inherited variants expected
    assert len(proc.bom_additions) == 3
    # ZERO usage warnings — inherited rows use variant_count, not BOM count
    inherited_warnings = [
        w for w in proc.usage_warnings if w.get("ReasonCode") == "PU_INHERITED_MISMATCH"
    ]
    assert inherited_warnings == [], inherited_warnings
    legacy_warnings = [
        w for w in proc.usage_warnings if w.get("ReasonCode") == "PU_EXPECTED_MISMATCH_BASIC"
    ]
    assert legacy_warnings == [], legacy_warnings


def test_legacy_enrichment_payload_is_rejected(tmp_path: Path) -> None:
    """Phase D C6: a stale client sending enrichments.functional=true must
    hit a hard ValidationError, not get silently migrated into the new
    primary-workflow path (which would produce output without the merge
    the user requested)."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
    )

    def state(role: str, path: Path) -> dict:
        return {
            "role": role,
            "label": role,
            "path": str(path),
            "selectedSheet": "Sheet1",
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        # Legacy enrichment toggle — should be rejected.
        "enrichments": {"functional": True, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            state("grouping", paths["grouping"]),
            state("bom", paths["bom"]),
            state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    with pytest.raises(ValidationError, match="enrichment toggles have been removed"):
        execute_run_request(body)


def test_failure_modes_standard_filters_mixed_standard_file(tmp_path: Path) -> None:
    """Phase D L6: FM file with a 'Standard' column and mixed FMD-91 +
    FMD-2016 rows should filter to only the selected standard when the
    user picks FMD-91."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "Standard": "FMD-91",
                "Commodity Type 1": "Resistor",
                "Commodity Type 2": "Chip",
                "Failure Mode": "Open (FMD-91)",
                "Failure Mode Ratio": 1.0,
            },
            {
                "Standard": "FMD-2016",
                "Commodity Type 1": "Resistor",
                "Commodity Type 2": "Chip",
                "Failure Mode": "Open (FMD-2016)",
                "Failure Mode Ratio": 1.0,
            },
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-91",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )
    # Only the FMD-91 failure mode should appear in the piece-part rows
    pp_rows = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    failure_modes = set(pp_rows["Failure Mode"].dropna().tolist())
    assert "Open (FMD-91)" in failure_modes
    assert "Open (FMD-2016)" not in failure_modes


def test_functional_to_piecepart_preserves_functional_rows(tmp_path: Path) -> None:
    """A 3-row functional FMEA with one circuit-block row containing
    "R100, C200" should produce: 3 functional rows preserved + 2 piece-part
    rows generated under the block."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            },
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Capacitor",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            },
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            },
        ],
        functional_rows=[
            # Row 1: a function header / not a circuit block
            {
                "FMEA Level": "Function",
                "FMEA-ID": "FN-001",
                "Failure Mode Causes": "",
                "Function Description": "Power conditioning",
                "Schematic Page": "1",
            },
            # Row 2: a circuit block with two RefDes
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CB-001",
                "Failure Mode Causes": "R100, C200",
                "Function Description": "RC filter",
                "Schematic Page": "2",
            },
            # Row 3: another function header
            {
                "FMEA Level": "Function",
                "FMEA-ID": "FN-002",
                "Failure Mode Causes": "",
                "Function Description": "Output stage",
                "Schematic Page": "3",
            },
        ],
    )

    proc = FMEAProcessor()
    df = proc.process_functional_to_piecepart(
        {
            "func": str(paths["func"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "func_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
        }
    )

    assert not df.empty
    # All 3 functional passthrough rows should be preserved AS-IS.
    # They're tagged with their classification rtype (circuit_block for
    # row 2; whatever classify_fmea_rows returns for the others — most
    # likely 'other' since they don't have level/refdes indicators).
    func_rows = df[~df["_row_type"].isin(("piece_part", "piece_part_no_match", "validation_warning"))]
    assert len(func_rows) == 3, (
        f"Expected 3 functional passthrough rows, got {len(func_rows)}: "
        f"{func_rows['_row_type'].tolist()}"
    )
    # Original column names should be preserved on passthrough rows
    assert "FMEA Level" in df.columns
    assert "FMEA-ID" in df.columns
    assert "Function Description" in df.columns
    # Exactly 2 piece-part rows should be generated (one per RefDes),
    # since each FM file row has ratio 1.0 (no multi-mode multiplication).
    pp_rows = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    assert len(pp_rows) == 2, f"Expected 2 piece-part rows, got {len(pp_rows)}"
    refdes_seen = set(pp_rows["Failure Mode Causes"].dropna().tolist())
    assert refdes_seen == {"R100", "C200"}, refdes_seen


# ----- Phase 4 / A6: CCA prefix plumbing --------------------------------------


def _state(role: str, path: Path, sheet: str = "Sheet1") -> dict:
    return {
        "role": role,
        "label": role,
        "path": str(path),
        "selectedSheet": sheet,
        "source": "desktop-bridge",
        "isResolvingSheets": False,
        "isAnalyzing": False,
        "resolutionError": None,
        "sheets": [{"id": "s1", "label": sheet}],
    }


def test_bom_only_requires_cca_prefix(tmp_path: Path) -> None:
    """Phase 4 / A6: BOM-Only validation must reject a run with no
    ccaPrefix in options."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
    )
    body = {
        "workflowId": "bom_only",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_cca_prefix", result


def test_bom_only_rejects_invalid_cca_prefix(tmp_path: Path) -> None:
    """Phase 4 / A6: a lowercase / space-containing ccaPrefix must be
    rejected with invalid_cca_prefix."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
    )
    body = {
        "workflowId": "bom_only",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "ccaPrefix": "p S u",  # lowercase + space
        },
        "inputs": [
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "invalid_cca_prefix", result


def test_bom_only_uses_cca_prefix_in_fmea_ids(tmp_path: Path) -> None:
    """Phase 4 / A6: a successful BOM-Only run with ccaPrefix='PSU'
    should produce output FMEA-IDs that start with 'PSU-' (not 'BOM-')."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    body = {
        "workflowId": "bom_only",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "ccaPrefix": "PSU",
        },
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = execute_run_request(body)
    assert result["status"] == "success", result
    # Output workbook should exist and contain FMEA IDs prefixed by PSU.
    output_file = Path(result["output_file"])
    assert output_file.exists(), output_file
    from openpyxl import load_workbook

    wb = load_workbook(output_file)
    try:
        ws = wb["FMEA"]
        header_row = [c.value for c in ws[1]]
        id_col_idx = header_row.index("FMEA-ID") + 1
        seen_ids = []
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=id_col_idx).value
            if val:
                seen_ids.append(str(val))
        assert seen_ids, "No FMEA-IDs found in output"
        assert any(i.startswith("PSU-") for i in seen_ids), seen_ids
        # And critically, NONE of them should use the legacy "BOM-" prefix
        assert not any(i.startswith("BOM-") for i in seen_ids), seen_ids
        # Fix A3: piece-part FMEA-IDs in BOM-Only mode must use the
        # hyphenated "<PREFIX>-<REFDES>-<SUFFIX>" shape, not the
        # buggy "<PREFIX>-<REFDES><SUFFIX>" that dropped the hyphen.
        import re as _re
        piece_part_ids = [i for i in seen_ids if i.startswith("PSU-")]
        # Every piece-part ID should match e.g. "PSU-C200-A".
        pattern = _re.compile(r"^PSU-[A-Z0-9]+-[A-Z]+$")
        unmatched = [i for i in piece_part_ids if not pattern.match(i)]
        assert not unmatched, (
            f"Expected every PSU-... ID to match PSU-<REFDES>-<SUFFIX> "
            f"(fix A3); unmatched: {unmatched}"
        )
    finally:
        wb.close()


# ----- Phase 4 / A7: explicit outputDirectory honored -------------------------


def test_explicit_output_directory_honored(tmp_path: Path) -> None:
    """Phase 4 / A7: when body['outputDirectory'] is set, the generated
    workbook should land under that directory rather than the heuristic
    fallback (parent of the first input file)."""
    inputs_dir = tmp_path / "inputs"
    outputs_dir = tmp_path / "custom-outputs"
    outputs_dir.mkdir(parents=True, exist_ok=True)
    paths = _write_fixture(
        inputs_dir,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(outputs_dir),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = execute_run_request(body)
    assert result["status"] == "success", result
    out_path = Path(result["output_file"])
    assert out_path.exists(), out_path
    # Crucial: the file should be under the explicit outputs_dir, NOT
    # under inputs_dir (the heuristic would pick the BOM's parent).
    assert out_path.parent.resolve() == outputs_dir.resolve(), (
        f"Expected output under {outputs_dir}, got {out_path.parent}"
    )


# ----- Phase 4 / A8: Part Usage discrepancy ------------------------------------


def test_part_usage_mismatch_flags_row_and_logs_diagnostic(tmp_path: Path) -> None:
    """Phase 4 / A8: when BOM's mapped Part Usage disagrees with the
    computed (1/N) value, the row should be flagged and a diagnostic
    entry should be recorded."""
    # Three BOM rows for R100 would give expected usage = 1/3, but we
    # mark it as 1/2 — the generator should detect the mismatch.
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1/2",  # Mapped value
            },
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1/2",
            },
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1/2",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    proc = FMEAProcessor()
    df = proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            # Fix A5: Part Usage discrepancy capture is now gated on an
            # explicit mapping. Simulate the user picking the BOM's
            # Part Usage column by setting the canonical mapping key.
            "column_overrides": {"Part Usage": "Part Usage"},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )

    # At least one discrepancy entry should have been captured.
    assert proc.part_usage_discrepancies, proc.part_usage_discrepancies
    first = proc.part_usage_discrepancies[0]
    assert first["refdes"] == "R100", first
    # Fix A5: new schema uses consistent count semantics.
    assert {"refdes", "mapped_count", "computed_count", "diff"} <= set(first.keys()), first
    # BOM says Part Usage=1/2 → mapped_count=2. Three R100 rows exist
    # in the BOM → computed_count=3. diff = 3 - 2 = 1.
    assert first["mapped_count"] == 2, first
    assert first["computed_count"] == 3, first
    assert first["diff"] == 1, first
    # Fix R2-H1: the previous "mapped + diff == computed" assertion was
    # tautological and never caught anything. The new sanity checks are
    # non-negativity of both counts — test that each entry honors them.
    for entry in proc.part_usage_discrepancies:
        assert entry["mapped_count"] >= 0, entry
        assert entry["computed_count"] >= 0, entry

    # At least one row for R100 should be tagged as a validation warning
    # (yellow fill via fmea_row_style).
    pp = df[
        (df["Failure Mode Causes"] == "R100")
        & (df["_row_type"] == "validation_warning")
    ]
    assert not pp.empty, df

    # Write the workbook and verify the diagnostics sheet shows up.
    out_path = tmp_path / "diag_out.xlsx"
    write_excel_report(df, out_path, proc)
    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    try:
        assert "Part Usage Diagnostics" in wb.sheetnames, wb.sheetnames
        ws = wb["Part Usage Diagnostics"]
        # Batch 2: row 1 is now the explanatory banner; headers moved to row 2.
        banner = ws.cell(row=1, column=1).value
        assert banner and "Mapped Count" in str(banner), banner
        headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        assert "RefDes" in headers
        assert "Mapped Count" in headers
        assert "Computed Count" in headers
        assert "Diff" in headers
    finally:
        wb.close()


def test_part_usage_suspicious_mapped_value_increments_counter_but_still_records_entry(
    tmp_path: Path,
) -> None:
    """Fix R2-H1 / R3-M2: a suspiciously large mapped_count (Part Usage
    < 1e-6, implying > 1,000,000 instances) is almost certainly a BOM
    data-entry typo. R3-M2 changes the behavior:

    1. The discrepancy entry is STILL recorded in
       ``part_usage_discrepancies`` so the user sees it in the
       "Part Usage Diagnostics" sheet. Previously we dropped the entry
       entirely, which made the warning useless for triage.
    2. A per-processor counter (``part_usage_suspicious_count``) tracks
       how many rows tripped the threshold.
    3. ONE aggregated WARNING is emitted after the generator loop
       completes, not one per row (prevents log spam for BOMs with many
       suspicious values).
    """
    log_lines: list[str] = []

    def _log(message: str) -> None:
        log_lines.append(message)

    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R200",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                # Part Usage = 5e-8 → round(1/5e-8) = 20,000,000. Way
                # above the new 1,000,000 suspicious threshold. The
                # real base count is 2 (two BOM rows for R200) so the
                # usage is obviously wrong, triggering the mismatch
                # branch; the suspicious counter should increment and
                # the entry should still be recorded.
                "Part Usage": "0.00000005",
            },
            {
                "Reference Designator": "R200",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "0.00000005",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "RES-002",
                "Reference Designator": "R200",
                "Function Description": "Pull-down",
                "Schematic Page": "4",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    proc = FMEAProcessor(log_callback=_log)
    proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {"Part Usage": "Part Usage"},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )

    # Fix R3-M2: the discrepancy entry IS present (not skipped).
    refdes_in_discrepancies = [
        entry["refdes"] for entry in proc.part_usage_discrepancies
    ]
    assert "R200" in refdes_in_discrepancies, proc.part_usage_discrepancies
    # Two BOM rows referencing R200 both trip the threshold.
    assert proc.part_usage_suspicious_count >= 1, (
        proc.part_usage_suspicious_count
    )

    # Fix R3-M2: ONE aggregated WARNING was logged after the generator
    # loop, mentioning the > 1,000,000 threshold. No per-row
    # "suspiciously large" spam from the old implementation should
    # remain.
    warnings = [line for line in log_lines if "WARNING" in line]
    assert any("1,000,000" in w for w in warnings), warnings
    assert any(
        "Part Usage Diagnostics" in w for w in warnings
    ), warnings
    # The old per-row message is gone.
    assert not any(
        "suspiciously large" in w for w in warnings
    ), warnings


def test_part_usage_large_legitimate_mapped_count_is_recorded_normally(
    tmp_path: Path,
) -> None:
    """Fix R3-M2: a legitimately large mapped count (e.g. 15,000
    instances of a single decoupling-cap variant) must be recorded as
    a normal discrepancy entry without tripping the suspicious-count
    aggregator. Production dense SMD PCBs routinely have thousands of
    instances of the same cap line; the old 10,000 threshold dropped
    these silently, which was wrong. The new 1,000,000 threshold
    leaves them untouched."""
    log_lines: list[str] = []

    def _log(message: str) -> None:
        log_lines.append(message)

    # Part Usage = 1/15000 ≈ 6.667e-5 → round(1/6.667e-5) = 15,000.
    # Well below the new 1,000,000 threshold.
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C400",
                "Part Number": "CAP-1",
                "Description": "Decoupling Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": str(1.0 / 15000.0),
            },
            {
                "Reference Designator": "C400",
                "Part Number": "CAP-1",
                "Description": "Decoupling Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": str(1.0 / 15000.0),
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CAP-DEC",
                "Reference Designator": "C400",
                "Function Description": "Decoupling",
                "Schematic Page": "7",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    proc = FMEAProcessor(log_callback=_log)
    proc.process(
        {
            "group": str(paths["grouping"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "func": None,
            "piecepart_fmea": None,
            "out_folder": str(tmp_path),
            "out_name": "out",
            "bom_only_mode": False,
            "use_func": False,
            "use_piecepart_merge": False,
            "verbose": False,
            "column_overrides": {"Part Usage": "Part Usage"},
            "failure_modes_standard": "FMD-2016",
            "group_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "func_sheet": None,
            "piecepart_fmea_sheet": None,
        }
    )

    # The discrepancy entry IS present — usage mismatch (expected 1/2
    # vs listed 1/15000) triggers the mismatch branch for both rows.
    refdes_in_discrepancies = [
        entry["refdes"] for entry in proc.part_usage_discrepancies
    ]
    assert "C400" in refdes_in_discrepancies, proc.part_usage_discrepancies

    # R3-M2: counter is ZERO because 15,000 < 1,000,000 threshold.
    assert proc.part_usage_suspicious_count == 0, (
        proc.part_usage_suspicious_count
    )

    # No aggregated suspicious-count warning should appear.
    warnings = [line for line in log_lines if "WARNING" in line]
    assert not any(
        "1,000,000" in w for w in warnings
    ), warnings


# ----- Phase 4 / A5: Group-level union merge ----------------------------------


def _merge_fixture(
    tmp_path: Path,
    *,
    old_fmea_rows: list[dict],
    grouping_rows: list[dict] | None,
    bom_rows: list[dict],
    fm_rows: list[dict] | None = None,
) -> dict:
    """Build fixture files for a union-merge test run.

    Produces an existing FMEA workbook + (optional) grouping + BOM + FM
    and returns their paths.
    """
    tmp_path.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    bom_path = tmp_path / "bom.xlsx"
    pd.DataFrame(bom_rows).to_excel(bom_path, index=False)
    paths["bom"] = bom_path

    old_path = tmp_path / "old_fmea.xlsx"
    pd.DataFrame(old_fmea_rows).to_excel(old_path, index=False)
    paths["fmea"] = old_path

    if grouping_rows is not None:
        grouping_path = tmp_path / "grouping.xlsx"
        pd.DataFrame(grouping_rows).to_excel(grouping_path, index=False)
        paths["group"] = grouping_path

    fm_path = tmp_path / "failure_modes.xlsx"
    pd.DataFrame(
        fm_rows
        or [
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Inductor",
                "FMD-2016 Commodity Type 2": "SMD",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
        ]
    ).to_excel(fm_path, index=False)
    paths["fm"] = fm_path
    return paths


def _merge_bom_rows() -> list[dict]:
    return [
        {
            "Reference Designator": "L100",
            "Part Number": "IND-1",
            "Description": "Inductor",
            "BAE HDA Commodity I": "Inductor",
            "BAE HDA Commodity II": "SMD",
            "Part Usage": "1",
        },
        {
            "Reference Designator": "C200",
            "Part Number": "CAP-1",
            "Description": "Capacitor",
            "BAE HDA Commodity I": "Capacitor",
            "BAE HDA Commodity II": "Ceramic",
            "Part Usage": "1",
        },
        {
            "Reference Designator": "R34",
            "Part Number": "RES-1",
            "Description": "Resistor",
            "BAE HDA Commodity I": "Resistor",
            "BAE HDA Commodity II": "Chip",
            "Part Usage": "1",
        },
    ]


def _run_gaps_merge(
    tmp_path: Path,
    paths: dict,
    *,
    include_grouping: bool = True,
) -> tuple[FMEAProcessor, pd.DataFrame]:
    proc = FMEAProcessor()
    df = proc.process_gaps(
        {
            "fmea": str(paths["fmea"]),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "group": str(paths["group"]) if include_grouping and "group" in paths else None,
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "fmea_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "fm_sheet": "Sheet1",
            "hda_sheet": None,
            "group_sheet": "Sheet1" if include_grouping and "group" in paths else None,
        }
    )
    return proc, df


def test_union_merge_both_sources_no_diagnostic(tmp_path: Path) -> None:
    """Phase 4 / A5: when a component exists in BOTH old FMEA and
    grouping file, the generated row must not carry a merge diagnostic."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100, C200",
                "Function Description": "Power rail filter",
                "Schematic Page": "1",
                "Local Effect": "",
                "Next Higher Effect": "",
                "End Effect": "",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-200",
                "Reference Designator": "L100, C200",
                "Function Description": "Power rail filter",
                "Schematic Page": "1",
            }
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    # The circuit-block row + 2 piece-part rows
    pp = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    refdes = set(pp["Failure Mode Causes"].dropna().tolist())
    assert refdes == {"L100", "C200"}, refdes
    # No merge diagnostics recorded at all
    assert proc.group_merge_diagnostics == [], proc.group_merge_diagnostics


def test_fill_gaps_inheritance_without_count_source_flags_guess(tmp_path: Path) -> None:
    """Tier-1 #7: fill_gaps inheritance with NO grouping/count source must emit
    usage_fraction '1/1' as a FLAGGED best guess, not a silent assertion.

    The old FMEA lists variant R34-A (its base R34 is in the BOM, the variant
    is not), and no grouping file is supplied, so variant_counts_by_base is
    empty. The inherited BOM_Additions row must carry usage '1/1' AND a review
    note (previously '1/1' was emitted with an empty note, silently implying the
    base has exactly one instance).
    """
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "BLK-001",
                "Failure Mode Causes": "R34-A",
                "Function Description": "Bias network",
                "Schematic Page": "3",
                "Local Effect": "",
                "Next Higher Effect": "",
                "End Effect": "",
            },
        ],
        grouping_rows=None,
        bom_rows=[
            {
                "Reference Designator": "R34",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            },
        ],
    )
    proc, df = _run_gaps_merge(tmp_path, paths, include_grouping=False)

    inherited = [e for e in proc.bom_additions if e["base_refdes"] == "R34"]
    assert inherited, proc.bom_additions
    assert inherited[0]["usage_fraction"] == "1/1"
    note = inherited[0]["notes"].lower()
    assert "best guess" in note and "verify" in note, inherited[0]["notes"]


def test_union_merge_old_only_flags_missing_from_grouping(tmp_path: Path) -> None:
    """Phase 4 / A5: when a component is only in the old FMEA, its row
    must carry the 'missing from Grouping File' diagnostic."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100, C200, R34",
                "Function Description": "Power rail",
                "Schematic Page": "1",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-200",
                "Reference Designator": "L100, C200",  # R34 absent
                "Function Description": "Power rail",
                "Schematic Page": "1",
            }
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    r34_rows = df[
        (df["Failure Mode Causes"] == "R34")
        & (df["_row_type"].isin(("piece_part", "validation_warning", "piece_part_no_match")))
    ]
    assert not r34_rows.empty, df
    diagnostics = r34_rows["Diagnostic"].dropna().tolist()
    assert any(
        FMEAProcessor.MERGE_DIAG_OLD_ONLY in d for d in diagnostics
    ), diagnostics
    # And a structured entry was recorded
    matches = [
        d for d in proc.group_merge_diagnostics
        if d.get("refdes") == "R34"
        and d.get("diagnostic") == FMEAProcessor.MERGE_DIAG_OLD_ONLY
    ]
    assert matches, proc.group_merge_diagnostics


def test_union_merge_grouping_only_flags_missing_from_merged(tmp_path: Path) -> None:
    """Phase 4 / A5: when a component is only in the grouping file, its
    row must carry the 'missing from Merged FMEA' diagnostic."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100, C200",  # R34 absent
                "Function Description": "Power rail",
                "Schematic Page": "1",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-200",
                "Reference Designator": "L100, C200, R34",
                "Function Description": "Power rail",
                "Schematic Page": "1",
            }
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    r34_rows = df[
        (df["Failure Mode Causes"] == "R34")
        & (df["_row_type"].isin(("piece_part", "validation_warning", "piece_part_no_match")))
    ]
    assert not r34_rows.empty, df
    diagnostics = r34_rows["Diagnostic"].dropna().tolist()
    assert any(
        FMEAProcessor.MERGE_DIAG_GROUPING_ONLY in d for d in diagnostics
    ), diagnostics
    matches = [
        d for d in proc.group_merge_diagnostics
        if d.get("refdes") == "R34"
        and d.get("diagnostic") == FMEAProcessor.MERGE_DIAG_GROUPING_ONLY
    ]
    assert matches, proc.group_merge_diagnostics


def test_union_merge_group_only_in_old_fmea(tmp_path: Path) -> None:
    """Phase 4 / A5: a function group only in the old FMEA should be
    preserved with a group-level diagnostic."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100",
                "Function Description": "Legacy group",
                "Schematic Page": "1",
            },
        ],
        grouping_rows=[
            # Different group only — CPU-200 absent from grouping.
            {
                "Component Group": "OTHER",
                "Reference Designator": "C200",
                "Function Description": "Other",
                "Schematic Page": "2",
            }
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    # The old-only group CPU-200 should have a group-level diagnostic
    group_level = [
        d for d in proc.group_merge_diagnostics
        if d.get("source") == "group"
        and d.get("group") == "CPU-200"
    ]
    assert group_level, proc.group_merge_diagnostics
    assert any(
        d.get("diagnostic") == FMEAProcessor.MERGE_DIAG_GROUP_OLD_ONLY
        for d in group_level
    ), group_level
    # The L100 row belongs to CPU-200 and should inherit the diagnostic
    l100_rows = df[df["Failure Mode Causes"] == "L100"]
    assert not l100_rows.empty, df


def test_union_merge_group_only_in_grouping(tmp_path: Path) -> None:
    """Phase 4 / A5: a group only in the grouping file should be
    generated with the 'New function group' group diagnostic."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "OLD-GRP",
                "Failure Mode Causes": "L100",
                "Function Description": "Old only",
                "Schematic Page": "1",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "OLD-GRP",
                "Reference Designator": "L100",
                "Function Description": "Old only",
                "Schematic Page": "1",
            },
            {
                "Component Group": "NEW-GRP",
                "Reference Designator": "C200",
                "Function Description": "Brand new",
                "Schematic Page": "2",
            },
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    new_group_diags = [
        d for d in proc.group_merge_diagnostics
        if d.get("source") == "group" and d.get("group") == "NEW-GRP"
    ]
    assert new_group_diags, proc.group_merge_diagnostics
    assert any(
        d.get("diagnostic") == FMEAProcessor.MERGE_DIAG_GROUP_NEW
        for d in new_group_diags
    ), new_group_diags


def test_union_merge_inherits_local_effect_from_old_fmea(tmp_path: Path) -> None:
    """Phase 4 / A5: Local Effect on the old FMEA's circuit-block row
    should flow into the generated piece-part rows beneath it."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100, C200",
                "Function Description": "Power rail filter",
                "Schematic Page": "1",
                "Local Effect": "No voltage",
                "Next Higher Effect": "Power supply fails",
                "End Effect": "System down",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-200",
                "Reference Designator": "L100, C200",
                "Function Description": "Power rail filter",
                "Schematic Page": "1",
            }
        ],
        bom_rows=_merge_bom_rows(),
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    # Every generated piece-part row for CPU-200 should carry the
    # inherited effect values.
    pp = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    assert not pp.empty
    local_effects = set(pp["Local Effect"].dropna().tolist())
    next_effects = set(pp["Next Higher Effect"].dropna().tolist())
    end_effects = set(pp["End Effect"].dropna().tolist())
    assert "No voltage" in local_effects, local_effects
    assert "Power supply fails" in next_effects, next_effects
    assert "System down" in end_effects, end_effects


# ----- Fix E1 / E2: grouping file hygiene warnings ---------------------------


def test_duplicate_grouping_file_group_emits_warning(tmp_path: Path) -> None:
    """Fix E1: two rows in the grouping file that share the same
    component_group but disagree on description or schematic page
    should emit a WARNING log so the user notices the data-entry
    mismatch. Components are still unioned (first-row-wins for
    metadata, components unioned across both)."""
    log_lines: list[str] = []

    # FMEAProcessor.log() formats the message as
    #   "[HH:MM:SS] LEVEL: <body>"
    # then calls log_callback with that single string.
    def _log(message: str) -> None:
        log_lines.append(message)

    proc = FMEAProcessor(log_callback=_log)
    group_df = pd.DataFrame(
        [
            {
                "component_group": "PSU-1",
                "description": "Primary power",
                "schematic_page": "3",
                "ref_des": "R1, R2",
            },
            {
                # Same label, conflicting description AND schematic page.
                "component_group": "PSU-1",
                "description": "Backup power",
                "schematic_page": "9",
                "ref_des": "R3",
            },
        ]
    )
    groups = proc._parse_grouping_file_groups(group_df)

    # Components are unioned across rows.
    assert groups["PSU-1"]["components"] == {"R1", "R2", "R3"}
    # First-row-wins for metadata.
    assert groups["PSU-1"]["description"] == "Primary power"
    assert groups["PSU-1"]["schematic_page"] == "3"
    # Two WARNING log entries — one for description drift, one for page.
    warnings = [line for line in log_lines if "WARNING" in line]
    assert any("description differs" in w for w in warnings), warnings
    assert any("schematic page differs" in w for w in warnings), warnings


def test_local_effect_user_mapping_reaches_parser(tmp_path: Path) -> None:
    """Fix R2-H2: when the user explicitly maps "Local Effect" to a
    non-standard column name (e.g. "Component Impact"), the parser must
    honor that pick instead of falling through to the heuristic resolver
    (which will not find "Component Impact" in the synonym list)."""
    proc = FMEAProcessor()
    # Seed the flat column_overrides that the frontend would normally
    # provide via _build_column_overrides.
    proc.column_overrides = {
        "Local Effect": "Component Impact",
        "Next Higher Effect": "Subsystem Impact",
        "End Effect": "System Impact",
    }

    fmea_df = pd.DataFrame(
        [
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "GRP-1",
                "Failure Mode Causes": "R1, R2",
                "Function Description": "Filter",
                "Schematic Page": "1",
                # Non-standard column names that the synonym resolver
                # would otherwise miss.
                "Component Impact": "Local impact text",
                "Subsystem Impact": "Next higher impact text",
                "System Impact": "End impact text",
            },
            {
                "FMEA Level": "Piece Part",
                "FMEA-ID": "GRP-1",
                "Failure Mode Causes": "R1",
                "Function Description": "",
                "Schematic Page": "1",
                "Component Impact": "",
                "Subsystem Impact": "",
                "System Impact": "",
            },
        ]
    )
    groups = proc._parse_old_fmea_groups(fmea_df)

    # The parser should have read the Local/Next/End Effect values from
    # the user-mapped columns, not the synonyms-based heuristic.
    assert "GRP-1" in groups, groups
    entry = groups["GRP-1"]
    assert entry["local_effect"] == "Local impact text", entry
    assert entry["next_higher_effect"] == "Next higher impact text", entry
    assert entry["end_effect"] == "End impact text", entry


def test_failure_mode_causes_mapping_reaches_grouping_parser(
    tmp_path: Path,
) -> None:
    """Fix R2-L4: when the user maps "Failure Mode Causes" to a
    non-standard column name (e.g. "MyRefs") on the grouping file /
    existing FMEA, ``_parse_old_fmea_groups`` must honor that mapping
    and read components from "MyRefs" instead of the heuristic detector
    that would otherwise try common synonyms and fail."""
    proc = FMEAProcessor()
    proc.column_overrides = {
        "Failure Mode Causes": "MyRefs",
    }

    fmea_df = pd.DataFrame(
        [
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "GRP-1",
                # Put a deliberately-wrong-looking value in the typical
                # refdes columns to make sure the parser ignores them.
                "Failure Mode Causes": "SHOULD_NOT_BE_READ",
                "MyRefs": "R1, R2, R3",
                "Function Description": "Bias network",
                "Schematic Page": "1",
            },
        ]
    )
    groups = proc._parse_old_fmea_groups(fmea_df)
    assert "GRP-1" in groups
    components = groups["GRP-1"]["components"]
    # The parser should have read R1/R2/R3 from "MyRefs".
    assert "R1" in components
    assert "R2" in components
    assert "R3" in components
    # And not seen the typical-column sentinel.
    assert "SHOULD_NOT_BE_READ" not in components


def test_blank_fmea_id_circuit_block_row_is_skipped(tmp_path: Path) -> None:
    """Fix E2: a circuit-block row with a blank FMEA-ID must be SKIPPED
    with a WARNING log. Previously the parser synthesized ``GROUP-{pos+1}``
    which was non-deterministic across runs and could collide with a
    real group literally named ``GROUP-2`` etc."""
    log_lines: list[str] = []

    def _log(message: str) -> None:
        log_lines.append(message)

    proc = FMEAProcessor(log_callback=_log)
    fmea_df = pd.DataFrame(
        [
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "",  # blank — should trigger skip
                "Failure Mode Causes": "C1, C2",
                "Function Description": "Filter",
                "Schematic Page": "1",
            },
            {
                "FMEA Level": "Piece Part",
                "FMEA-ID": "",  # piece-parts under the skipped group
                "Failure Mode Causes": "C1",
                "Function Description": "",
                "Schematic Page": "1",
            },
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "GRP-2",  # real label — should survive
                "Failure Mode Causes": "R1",
                "Function Description": "Bias",
                "Schematic Page": "2",
            },
        ]
    )
    groups = proc._parse_old_fmea_groups(fmea_df)

    # The blank-labeled circuit-block row was skipped; no GROUP-1 / GROUP-2
    # synthetic key contaminates the output.
    assert "" not in groups
    assert not any(label.startswith("GROUP-") for label in groups), list(
        groups.keys()
    )
    # The real GRP-2 label survives.
    assert "GRP-2" in groups
    # The WARNING log was emitted mentioning the blank row.
    warnings = [line for line in log_lines if "WARNING" in line]
    assert any("blank FMEA-ID" in w for w in warnings), warnings
    # Fix R2-M2: a second WARNING should report the total count of
    # piece-part rows orphaned under the skipped circuit-block so the
    # user knows the full scope of the data loss.
    orphan_warnings = [w for w in warnings if "orphaned" in w]
    assert orphan_warnings, warnings
    # The fixture has 1 piece-part row under the skipped circuit-block.
    assert "1 piece-part row" in orphan_warnings[0], orphan_warnings


def test_piece_parts_before_first_circuit_block_warn_separately(
    tmp_path: Path,
) -> None:
    """Fix R3-L1: a piece-part row that appears BEFORE any circuit-block
    row in the source workbook is a DIFFERENT root cause than a
    piece-part orphaned under a skipped blank-id block. The user can
    fix the latter by filling in the blank FMEA-ID cell, but the
    former is a structural problem in the source file that requires
    re-ordering rows. The two cases must produce distinct warnings so
    the user knows what to actually fix."""
    log_lines: list[str] = []

    def _log(message: str) -> None:
        log_lines.append(message)

    proc = FMEAProcessor(log_callback=_log)
    fmea_df = pd.DataFrame(
        [
            # Two piece-part rows BEFORE any circuit-block — structural
            # issue, not a blank-id skip.
            {
                "FMEA Level": "Piece Part",
                "FMEA-ID": "",
                "Failure Mode Causes": "C1",
                "Function Description": "",
                "Schematic Page": "1",
            },
            {
                "FMEA Level": "Piece Part",
                "FMEA-ID": "",
                "Failure Mode Causes": "C2",
                "Function Description": "",
                "Schematic Page": "1",
            },
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "GRP-1",
                "Failure Mode Causes": "R1",
                "Function Description": "Bias",
                "Schematic Page": "2",
            },
        ]
    )
    groups = proc._parse_old_fmea_groups(fmea_df)

    # The real group survives.
    assert "GRP-1" in groups

    warnings = [line for line in log_lines if "WARNING" in line]
    # The "before any circuit-block" warning should fire with count=2.
    structural_warnings = [
        w for w in warnings if "before any circuit-block row" in w
    ]
    assert structural_warnings, warnings
    assert "2 piece-part row" in structural_warnings[0], structural_warnings
    # The "orphaned under skipped" warning should NOT fire — no blank
    # circuit-block row was skipped in this fixture.
    assert not any(
        "orphaned under skipped" in w for w in warnings
    ), warnings


# ----- Fix A2: mappings -> column_overrides plumbing -------------------------


def test_validate_run_mappings_become_column_overrides(tmp_path: Path) -> None:
    """Fix A2: body['mappings'] must be converted into the processor's
    column_overrides dict so the user's explicit column picks are honored
    by the backend instead of silently discarded."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "GRP-1",
                "Reference Designator": "C200",
                "Function Description": "Filter",
                "Schematic Page": "1",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    captured: dict = {}

    def _capture_processor(proc: FMEAProcessor) -> None:
        # Attach a post-run hook. The runtime calls processor_ready_callback
        # BEFORE process() runs, but we want the final column_overrides
        # that were actually stashed on the processor by the run. We can
        # read it immediately after execute_run_request returns by
        # referencing the same processor instance.
        captured["proc"] = proc

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [
            {
                "canonical": "Failure Mode",
                "mappedTo": "fm_col",
                "status": "mapped",
            },
            {
                "canonical": "Part Usage",
                "mappedTo": "Part Usage",
                "status": "mapped",
            },
            # Empty mappedTo should be skipped
            {"canonical": "Local Effect", "mappedTo": "", "status": "unmapped"},
            # DO_NOT_MAP sentinel should be skipped
            {
                "canonical": "End Effect",
                "mappedTo": "__do_not_map__",
                "status": "manual",
            },
        ],
    }
    execute_run_request(body, processor_ready_callback=_capture_processor)
    proc = captured["proc"]
    assert proc is not None
    # The processor's column_overrides should contain the two mapped rows
    # as flat entries and neither the empty nor the DO_NOT_MAP sentinel row.
    # Fix D: the dict now ALSO carries nested per-file shapes derived
    # from FRONTEND_TO_BACKEND_MAPPING, so map_columns can use them.
    assert proc.column_overrides["Failure Mode"] == "fm_col", proc.column_overrides
    assert proc.column_overrides["Part Usage"] == "Part Usage", proc.column_overrides
    assert "Local Effect" not in proc.column_overrides, proc.column_overrides
    assert "End Effect" not in proc.column_overrides, proc.column_overrides
    # Nested shapes: Failure Mode belongs to FAILURE_MODES file; Part
    # Usage belongs to BOM file.
    assert proc.column_overrides.get("FAILURE_MODES") == {
        "failure_mode": "fm_col",
    }, proc.column_overrides
    assert proc.column_overrides.get("BOM") == {
        "part_usage": "Part Usage",
    }, proc.column_overrides


# ----- Fix D: flat-to-nested column_overrides translation --------------------


def test_flat_mappings_translated_to_nested_col_overrides(tmp_path: Path) -> None:
    """Fix D: ``_build_column_overrides`` should produce BOTH flat and
    nested shapes so ``map_columns`` can override its heuristic resolver
    with the user's explicit picks."""
    from fmea.runtime import _build_column_overrides

    body_mappings = [
        {"canonical": "Failure Mode", "mappedTo": "FailureMode", "status": "mapped"},
        {
            "canonical": "Component Part Description",
            "mappedTo": "Desc",
            "status": "mapped",
        },
        {"canonical": "Part Usage", "mappedTo": "Qty", "status": "mapped"},
        {"canonical": "FMEA-ID", "mappedTo": "GroupID", "status": "mapped"},
        {
            "canonical": "BAE HDA Commodity Level 1",
            "mappedTo": "HDA1",
            "status": "mapped",
        },
        # FMD standard-specific label — the frontend sends the active
        # label verbatim; both FMD-91 and FMD-2016 point at fmd_type1.
        {
            "canonical": "FMD-91 Commodity Type 1",
            "mappedTo": "FmdT1",
            "status": "mapped",
        },
        {"canonical": "Failure Mode Ratio", "mappedTo": "Ratio", "status": "mapped"},
        {"canonical": "Failure Mode Causes", "mappedTo": "Refs", "status": "mapped"},
        # Unmapped / sentinel rows should be ignored entirely.
        {"canonical": "Local Effect", "mappedTo": "", "status": "unmapped"},
        {"canonical": "End Effect", "mappedTo": "__do_not_map__", "status": "manual"},
    ]
    overrides = _build_column_overrides(body_mappings)

    # Flat entries — all mapped rows present under their frontend canonical.
    assert overrides["Failure Mode"] == "FailureMode"
    assert overrides["Component Part Description"] == "Desc"
    assert overrides["Part Usage"] == "Qty"
    assert overrides["FMEA-ID"] == "GroupID"
    assert overrides["BAE HDA Commodity Level 1"] == "HDA1"
    assert overrides["FMD-91 Commodity Type 1"] == "FmdT1"
    assert overrides["Failure Mode Ratio"] == "Ratio"
    assert overrides["Failure Mode Causes"] == "Refs"
    # Sentinel / empty rows are skipped entirely.
    assert "Local Effect" not in overrides
    assert "End Effect" not in overrides

    # Nested per-file entries — derived from FRONTEND_TO_BACKEND_MAPPING.
    assert overrides["FAILURE_MODES"] == {
        "failure_mode": "FailureMode",
        "ratio": "Ratio",
    }
    assert overrides["BOM"] == {
        "description": "Desc",
        "part_usage": "Qty",
    }
    assert overrides["HDA"] == {
        "commodity_level1": "HDA1",
        "fmd_type1": "FmdT1",
    }
    assert overrides["COMPONENT_GROUPING"] == {
        "component_group": "GroupID",
        "ref_des": "Refs",
    }


def test_reserved_file_type_key_collision_raises(tmp_path: Path) -> None:
    """Fix R2-M3: ``_build_column_overrides`` carries both flat canonical
    keys and nested file-type buckets (``BOM``, ``HDA``, ``FAILURE_MODES``,
    ``COMPONENT_GROUPING``) at the same top level. If a future frontend
    canonical literally matches one of those reserved names, the flat
    write would silently clobber the nested sub-dict. Guard against that
    by raising ValidationError at build time."""
    from fmea.runtime import _build_column_overrides
    from common.exceptions import ValidationError
    import pytest as _pytest

    body_mappings = [
        # A hypothetical future canonical that collides with the BOM
        # file-type bucket name.
        {"canonical": "BOM", "mappedTo": "SomeCol", "status": "mapped"},
    ]

    with _pytest.raises(ValidationError):
        _build_column_overrides(body_mappings)


def test_unknown_canonical_keeps_flat_only(tmp_path: Path) -> None:
    """Fix D: canonicals that are NOT in FRONTEND_TO_BACKEND_MAPPING
    must still populate the flat dict (forward-compat with future
    frontend additions) but must NOT contaminate any nested sub-dict."""
    from fmea.runtime import _build_column_overrides

    body_mappings = [
        {"canonical": "Future Column", "mappedTo": "xcol", "status": "mapped"},
        {"canonical": "Failure Mode", "mappedTo": "fm", "status": "mapped"},
    ]
    overrides = _build_column_overrides(body_mappings)

    # Flat dict carries both.
    assert overrides["Future Column"] == "xcol"
    assert overrides["Failure Mode"] == "fm"
    # The known canonical got nested; the unknown one did not.
    assert overrides["FAILURE_MODES"] == {"failure_mode": "fm"}
    for file_type in ("BOM", "HDA", "COMPONENT_GROUPING"):
        if file_type in overrides:
            # "Future Column" must not appear under any nested bucket.
            bucket = overrides[file_type]
            assert isinstance(bucket, dict)
            assert "xcol" not in bucket.values(), bucket


def test_nested_col_overrides_reach_map_columns(tmp_path: Path) -> None:
    """Fix D: when the user explicitly maps Part Usage to a non-standard
    column header (``Qty``), the nested ``BOM`` sub-dict should reach
    ``FMEAProcessor.map_columns()`` and rename that column to
    ``part_usage`` so the generator honors the user's pick instead of
    guessing via synonyms. Integration-test: run the processor end to
    end and verify usage ends up on the output row."""
    # Fabricate a BOM with a custom Part Usage header name that would
    # NOT match the synonym list — the only way the processor can see
    # it is via the explicit override.
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                # Intentionally obscure header: 'Assembly Count' is NOT in
                # the ``part_usage`` synonym list, so the heuristic resolver
                # will return None and the processor will default to 1.
                # When the override is applied, the processor will rename
                # this column to ``part_usage`` and the custom value wins.
                "Assembly Count": "1/2",
            },
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Assembly Count": "1/2",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "GRP-1",
                "Reference Designator": "C200",
                "Function Description": "Filter",
                "Schematic Page": "1",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )

    captured: dict = {}

    def _capture_processor(proc: FMEAProcessor) -> None:
        captured["proc"] = proc

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [
            # Override: the BOM's Part Usage column is named "Assembly Count".
            {
                "canonical": "Part Usage",
                "mappedTo": "Assembly Count",
                "status": "manual",
            },
        ],
    }
    execute_run_request(body, processor_ready_callback=_capture_processor)
    proc = captured["proc"]
    assert proc is not None
    # The processor should have received a nested BOM override.
    assert proc.column_overrides.get("BOM") == {
        "part_usage": "Assembly Count",
    }, proc.column_overrides
    # And the BOM's per-RefDes usage_base_counts should reflect the two
    # C200 rows — which only happens if the rename succeeded and the
    # index build was able to read the per-row usage value.
    assert "C200" in proc.usage_base_counts
    assert proc.usage_base_counts["C200"] == 2, proc.usage_base_counts


# ----- Fix A4: missing-from-BOM union component emits placeholder row --------


def test_fill_gaps_requires_failure_mode_causes_mapping(tmp_path: Path) -> None:
    """Fix C2: fill_gaps without a Failure Mode Causes mapping should
    be rejected with reason_code='missing_failure_mode_causes_mapping'.
    The backend parses that column as a comma-separated RefDes list; an
    unmapped value produces empty component sets and a blank FMEA."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
    )
    existing_fmea = tmp_path / "existing.xlsx"
    pd.DataFrame([{"FMEA-ID": "X-001"}]).to_excel(existing_fmea, index=False)

    body = {
        "workflowId": "fill_gaps",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("existingFmea", existing_fmea),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_failure_mode_causes_mapping", result


def test_functional_to_piecepart_requires_failure_mode_causes_mapping(
    tmp_path: Path,
) -> None:
    """Fix C2: functional_to_piecepart has the same requirement as
    fill_gaps — Failure Mode Causes must be explicitly mapped."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
    )
    functional = tmp_path / "functional.xlsx"
    pd.DataFrame([{"FMEA-ID": "F-001"}]).to_excel(functional, index=False)

    body = {
        "workflowId": "functional_to_piecepart",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("functionalFmea", functional),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [
            # Explicitly set to DO_NOT_MAP → should be treated as missing.
            {
                "canonical": "Failure Mode Causes",
                "mappedTo": "__do_not_map__",
                "status": "manual",
            },
        ],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_failure_mode_causes_mapping", result


def test_functional_to_piecepart_honors_failure_mode_causes_mapping(
    tmp_path: Path,
) -> None:
    """Fix R3-M1: when the functional FMEA uses a non-standard RefDes
    column header (e.g. ``FuncRefs``) and the user explicitly maps
    ``Failure Mode Causes`` → ``FuncRefs`` in the frontend, the
    processor must honor that mapping instead of falling through to
    ``detect_refdes_column_for_fmea`` (which doesn't know about the
    custom header and raises ``ColumnMappingError``). Mirrors the
    R2-H2 override pattern that was already applied to the sibling
    ``_parse_old_fmea_groups`` path."""
    # Functional FMEA with a NON-STANDARD refdes column — no synonym
    # of "Reference Designator" / "RefDes" / "Failure Mode Causes" is
    # present, so the heuristic detector would return None. The
    # "FMEA Level" column marks the row as a circuit block so
    # classify_fmea_rows picks it up for piece-part generation.
    func_rows = [
        {
            "FMEA Level": "Circuit Block",
            "MyGroupID": "CPU-100",
            "MyDescription": "Clock generator",
            "MyPage": "3",
            "FuncRefs": "R100, C100",
        },
    ]
    func_path = tmp_path / "functional.xlsx"
    pd.DataFrame(func_rows).to_excel(func_path, index=False)

    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            },
            {
                "Reference Designator": "C100",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            },
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            },
        ],
    )

    # Call process_functional_to_piecepart directly with the flat
    # column_overrides shape the runtime constructs from frontend
    # mappings. The key observation: without R3-M1, this run raises
    # ColumnMappingError on the refdes column because
    # detect_refdes_column_for_fmea can't find "FuncRefs" via synonyms.
    proc = FMEAProcessor()
    df = proc.process_functional_to_piecepart(
        {
            "func": str(func_path),
            "bom": str(paths["bom"]),
            "fm": str(paths["fm"]),
            "hda": None,
            "group": None,
            "verbose": False,
            "column_overrides": {
                "Failure Mode Causes": "FuncRefs",
                "FMEA-ID": "MyGroupID",
                "Function Description": "MyDescription",
                "Schematic Page": "MyPage",
            },
            "failure_modes_standard": "FMD-2016",
            "func_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "group_sheet": None,
        }
    )

    # The run succeeded and emitted piece-part rows for both R100 and C100.
    assert not df.empty
    # Piece-part rows should reference R100 and C100 via the
    # 'Failure Mode Causes' output column (this is the canonical output
    # header regardless of the input header name).
    pp_rows = df[df["_row_type"].isin(("piece_part", "validation_warning"))]
    refdes_seen = set(pp_rows["Failure Mode Causes"].dropna().tolist())
    assert {"R100", "C100"} <= refdes_seen, refdes_seen


def test_validate_run_fmc_mapping_is_whitespace_insensitive(
    tmp_path: Path,
) -> None:
    """Fix R2-H3: the validator must normalize the canonical key the
    same way ``_build_column_overrides`` does (``.strip()``), otherwise
    a padded canonical (``' Failure Mode Causes '``) bypasses the
    validator but still lands in the overrides dict under a different
    key, producing a silent mismatch downstream."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
    )
    existing_fmea = tmp_path / "existing.xlsx"
    pd.DataFrame([{"FMEA-ID": "X-001"}]).to_excel(existing_fmea, index=False)

    body = {
        "workflowId": "fill_gaps",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("existingFmea", existing_fmea),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [
            # Padded canonical — should still be treated as a valid
            # "Failure Mode Causes" mapping after the stripping fix.
            {
                "canonical": " Failure Mode Causes ",
                "mappedTo": "RefDes",
                "status": "mapped",
            },
        ],
    }
    result = validate_run_request(body)
    # The FMC mapping gate should PASS (reason code must not match the
    # missing-FMC one). Any remaining validation failures are for other
    # reasons (missing files, etc.) which this test doesn't care about.
    assert result["reason_code"] != "missing_failure_mode_causes_mapping", result


def test_piece_part_generate_does_not_require_failure_mode_causes_mapping(
    tmp_path: Path,
) -> None:
    """Fix C2: non-merge modes (piece_part_generate, bom_only) must NOT
    require the Failure Mode Causes mapping. They derive components from
    the grouping file or BOM directly, so the column is informational."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "G1",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "1",
            }
        ],
    )
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],  # deliberately empty
    }
    result = validate_run_request(body)
    # Should NOT reject for missing Failure Mode Causes.
    assert result["reason_code"] != "missing_failure_mode_causes_mapping", result
    assert result["ok"] is True, result


def test_functional_to_piecepart_with_grouping_emits_union_merge_diagnostics(
    tmp_path: Path,
) -> None:
    """Fix C1: functional_to_piecepart now accepts an optional grouping
    file. When supplied, components present in the grouping file but NOT
    referenced by the functional FMEA should be emitted via
    process_union_merge with the MERGE_DIAG_GROUPING_ONLY diagnostic
    attached. Components that appear in both are preserved from the
    functional expansion (no duplicate rows)."""
    # Functional FMEA references L100 and C200 under group CPU-200.
    # Grouping file lists CPU-200 = [L100, C200, R34] — R34 is the
    # extra row that should show up with the diagnostic.
    inputs_dir = tmp_path
    inputs_dir.mkdir(parents=True, exist_ok=True)

    func_rows = [
        {
            "FMEA Level": "Circuit Block",
            "FMEA-ID": "CPU-200",
            "Failure Mode Causes": "L100, C200",
            "Function Description": "Power rail",
            "Schematic Page": "1",
        },
    ]
    grouping_rows = [
        {
            "Component Group": "CPU-200",
            "Reference Designator": "L100, C200, R34",
            "Function Description": "Power rail",
            "Schematic Page": "1",
        }
    ]
    bom_rows = _merge_bom_rows()

    fm_path = inputs_dir / "failure_modes.xlsx"
    pd.DataFrame(
        [
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            },
            {
                "FMD-2016 Commodity Type 1": "Inductor",
                "FMD-2016 Commodity Type 2": "SMD",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
        ]
    ).to_excel(fm_path, index=False)

    func_path = inputs_dir / "functional.xlsx"
    pd.DataFrame(func_rows).to_excel(func_path, index=False)

    bom_path = inputs_dir / "bom.xlsx"
    pd.DataFrame(bom_rows).to_excel(bom_path, index=False)

    group_path = inputs_dir / "grouping.xlsx"
    pd.DataFrame(grouping_rows).to_excel(group_path, index=False)

    proc = FMEAProcessor()
    df = proc.process_functional_to_piecepart(
        {
            "func": str(func_path),
            "bom": str(bom_path),
            "fm": str(fm_path),
            "hda": None,
            "group": str(group_path),
            "verbose": False,
            "column_overrides": {},
            "failure_modes_standard": "FMD-2016",
            "func_sheet": "Sheet1",
            "bom_sheet": "Sheet1",
            "hda_sheet": None,
            "fm_sheet": "Sheet1",
            "group_sheet": "Sheet1",
        }
    )

    # R34 should appear in the output with the grouping-only diagnostic.
    r34_rows = df[df["Failure Mode Causes"] == "R34"]
    assert not r34_rows.empty, (
        f"Expected R34 from the grouping file to show up in the "
        f"functional output. Columns: {list(df.columns)}"
    )
    diagnostics = " ".join(
        str(d) for d in r34_rows.get("Diagnostic", pd.Series(dtype=str)).dropna().tolist()
    )
    assert FMEAProcessor.MERGE_DIAG_GROUPING_ONLY in diagnostics, diagnostics


def test_invalid_output_directory_logs_warning_and_falls_back(tmp_path: Path) -> None:
    """Fix B3: when outputDirectory points to a nonexistent path, the
    runtime must log a WARNING about the fallback and then drop the
    output into the input-file heuristic location. Previously the
    fallback was silent, leaving users confused about where their
    workbook actually landed."""
    inputs_dir = tmp_path / "inputs"
    inputs_dir.mkdir(parents=True, exist_ok=True)
    paths = _write_fixture(
        inputs_dir,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )
    bogus_dir = tmp_path / "does_not_exist" / "nested"
    # Intentionally do NOT create bogus_dir.
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(bogus_dir),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }

    captured_logs: list[str] = []

    def _log_cb(msg: str) -> None:
        captured_logs.append(msg)

    result = execute_run_request(body, log_callback=_log_cb)
    assert result["status"] == "success", result
    out_path = Path(result["output_file"])
    # Output should have landed in the fallback directory (input-file
    # heuristic → BOM's parent == inputs_dir), NOT under bogus_dir.
    assert out_path.parent.resolve() == inputs_dir.resolve(), (
        f"Expected fallback to {inputs_dir}, got {out_path.parent}"
    )
    # And we must have surfaced a WARNING about the fallback.
    warnings = [m for m in captured_logs if "WARNING" in m and "outputDirectory" in m]
    assert warnings, captured_logs


def test_unwritable_output_directory_logs_warning_and_falls_back(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Explicit output directories must be writable, not merely present.

    Regression for the unwritable-directory path: if the chosen output
    folder exists but cannot accept new files, the runtime should warn
    and fall back to the input-file heuristic instead of failing later
    during workbook write.
    """
    inputs_dir = tmp_path / "inputs"
    inputs_dir.mkdir(parents=True, exist_ok=True)
    paths = _write_fixture(
        inputs_dir,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "RES-001",
                "Reference Designator": "R100",
                "Function Description": "Pull-up",
                "Schematic Page": "3",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )
    locked_dir = tmp_path / "locked_output"
    locked_dir.mkdir(parents=True, exist_ok=True)

    monkeypatch.setattr("common.utils.is_writable_directory", lambda path: False)

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(locked_dir),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }

    captured_logs: list[str] = []

    def _log_cb(msg: str) -> None:
        captured_logs.append(msg)

    result = execute_run_request(body, log_callback=_log_cb)
    assert result["status"] == "success", result
    out_path = Path(result["output_file"])
    assert out_path.parent.resolve() == inputs_dir.resolve(), (
        f"Expected fallback to {inputs_dir}, got {out_path.parent}"
    )
    warnings = [m for m in captured_logs if "WARNING" in m and "not writable" in m]
    assert warnings, captured_logs


def test_union_merge_missing_from_bom_emits_placeholder_with_diagnostic(
    tmp_path: Path,
) -> None:
    """Fix A4: when a union-set component is missing from the BOM (and its
    base variant is also missing), process_union_merge must emit a
    synthetic placeholder row so the merge diagnostic still has something
    to attach to. Without this, the component vanishes from the output
    entirely (no piece-part row, no diagnostic visible to the user)."""
    paths = _merge_fixture(
        tmp_path,
        old_fmea_rows=[
            {
                "FMEA Level": "Circuit Block",
                "FMEA-ID": "CPU-200",
                "Failure Mode Causes": "L100, C200",
                "Function Description": "Power rail",
                "Schematic Page": "1",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-200",
                # Z999 is in the grouping file but NOT in the BOM.
                "Reference Designator": "L100, C200, Z999",
                "Function Description": "Power rail",
                "Schematic Page": "1",
            }
        ],
        bom_rows=_merge_bom_rows(),  # no Z999
    )
    proc, df = _run_gaps_merge(tmp_path, paths)
    # The placeholder row for Z999 must exist.
    z999_rows = df[df["Failure Mode Causes"] == "Z999"]
    assert not z999_rows.empty, (
        f"Expected a placeholder row for Z999 (missing from BOM). "
        f"Columns available: {list(df.columns)}"
    )
    # And it must carry the grouping-only merge diagnostic so the user
    # sees BOTH "missing from BOM" (the placeholder message) and the
    # "present in Grouping File but missing from Merged FMEA" diagnostic.
    diagnostics = " ".join(
        str(d) for d in z999_rows["Diagnostic"].dropna().tolist()
    )
    assert "missing from BOM" in diagnostics, diagnostics
    assert FMEAProcessor.MERGE_DIAG_GROUPING_ONLY in diagnostics, diagnostics


# ---------------------------------------------------------------------------
# Bug 3: an explicit column override naming a column NOT present in the
# dataframe is silently discarded (heuristic detection runs instead). Add a
# WARNING log so the silent discard becomes diagnosable. No behavior change.
# ---------------------------------------------------------------------------


def test_map_columns_warns_when_override_column_missing(caplog) -> None:
    """A user override pointing at a non-existent column logs a WARNING.

    ``map_columns`` falls back to heuristic detection when the override names
    a column that isn't in the dataframe. Before the fix this fallback was
    silent; now it emits a WARNING via the module logger so the run log shows
    why the explicit pick was ignored.
    """
    proc = FMEAProcessor(log_callback=None)
    df = pd.DataFrame({"Reference Designator": ["R1"], "Description": ["Resistor"]})
    config_section = {"description": ["Description", "Component Description"]}

    with caplog.at_level("WARNING", logger="fmea_generator"):
        mapped = proc.map_columns(
            df,
            config_section,
            required_list=["description"],
            source_name="BOM file",
            overrides={"description": "Nonexistent Column"},
        )

    # Heuristic fallback still resolves the column (no behavior change).
    assert mapped["description"] == "Description"
    # And a WARNING names the missing override column.
    warnings = " ".join(
        rec.getMessage() for rec in caplog.records if rec.levelname == "WARNING"
    )
    assert "Nonexistent Column" in warnings
    assert "description" in warnings


# ---------------------------------------------------------------------------
# Contract fix: backend honors the ``hdaSource`` flag.
#
# The frontend sends ``options.hdaSource`` ("inline" | "separate") to declare
# whether the run should read HDA data from a dedicated workbook or detect it
# inline from BOM columns. Previously the backend ignored the flag entirely and
# decided purely on whether an ``hda`` input path was present, so:
#
#   * Picking "Separate HDA file" but never attaching one silently fell back to
#     inline detection instead of failing — the user's explicit choice was lost.
#   * A stray ``hda`` path left over from a previous toggle (or sent by a client
#     that ships all roles) would be loaded even though the user selected inline.
#
# The flag is now authoritative in BOTH directions:
#   1. ``hdaSource == "separate"`` with no usable hda path → validation FAILS.
#   2. ``hdaSource == "inline"`` IGNORES any stray hda path (inline detection).
#   3. ``hdaSource`` absent (legacy client) → unchanged: path presence decides.
#
# The seam is in ``fmea/runtime.py``: validation gates case (1), and the execute
# path resolves the effective hda path/sheet from the flag before building
# ``run_inputs`` — so ``fmea_generator_logic.py`` stays flag-free and keeps its
# simple ``if hda_path:`` branch.
# ---------------------------------------------------------------------------


def _hda_separate_fixture(tmp_path: Path) -> dict:
    """Fixture where the BOM has NO inline commodity columns but a dedicated
    HDA workbook supplies them. This lets a test distinguish the
    "loaded separate HDA" branch from the "inline detection" branch:
    inline detection on this BOM finds no commodity columns and warns.
    """
    tmp_path.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # BOM intentionally omits BAE HDA Commodity columns.
    bom_path = tmp_path / "bom.xlsx"
    pd.DataFrame(
        [
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "Part Usage": "1",
            }
        ]
    ).to_excel(bom_path, index=False)
    paths["bom"] = bom_path

    grouping_path = tmp_path / "grouping.xlsx"
    pd.DataFrame(
        [
            {
                "Component Group": "GRP-1",
                "Reference Designator": "C200",
                "Function Description": "Filter",
                "Schematic Page": "1",
            }
        ]
    ).to_excel(grouping_path, index=False)
    paths["grouping"] = grouping_path

    # Dedicated HDA workbook keyed by Part Number.
    hda_path = tmp_path / "hda.xlsx"
    pd.DataFrame(
        [
            {
                "Part Number": "CAP-1",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
            }
        ]
    ).to_excel(hda_path, index=False)
    paths["hda"] = hda_path

    fm_path = tmp_path / "failure_modes.xlsx"
    pd.DataFrame(
        [
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ]
    ).to_excel(fm_path, index=False)
    paths["fm"] = fm_path

    return paths


def test_hda_source_separate_without_path_blocks_validation(tmp_path: Path) -> None:
    """(a) hdaSource=separate with no hda path → validation FAILS with an
    actionable reason code and message naming the HDA workbook."""
    paths = _hda_separate_fixture(tmp_path)
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "hdaSource": "separate",
        },
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
            # NOTE: no hda input attached.
        ],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_separate_hda", result
    detail = (result["toast_text"] or "").lower()
    assert "hda" in detail
    # Message must offer both remedies (attach a file OR switch to inline).
    assert "inline" in detail
    # The structured validations list should also surface the block.
    blocked = [
        v
        for v in result["validations"]
        if v.get("severity") == "error" and "hda" in (v.get("detail", "").lower())
    ]
    assert blocked, result["validations"]


def test_hda_source_separate_without_path_blocks_execute(tmp_path: Path) -> None:
    """(3) execute_run_request re-validates, so the separate-but-missing case
    is also blocked at execute time with a ValidationError."""
    paths = _hda_separate_fixture(tmp_path)
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "hdaSource": "separate",
        },
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    with pytest.raises(ValidationError):
        execute_run_request(body)


def test_hda_source_separate_with_path_loads_separate_file(tmp_path: Path) -> None:
    """(b) hdaSource=separate WITH an hda path → validation ok and execute
    loads the dedicated HDA file (proven by the "Loading dedicated HDA file"
    log line, which only the separate branch emits)."""
    paths = _hda_separate_fixture(tmp_path)
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "hdaSource": "separate",
        },
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
            _state("hda", paths["hda"]),
        ],
        "mappings": [],
    }
    # Validation passes when the hda file is attached.
    validation = validate_run_request(body)
    assert validation["ok"] is True, validation

    logs: list[str] = []
    result = execute_run_request(body, log_callback=logs.append)
    assert result["status"] == "success", result
    joined = " ".join(logs)
    assert "Loading dedicated HDA file" in joined, logs
    assert "Detecting inline HDA columns" not in joined, logs


def test_hda_source_inline_ignores_stray_hda_path(tmp_path: Path) -> None:
    """(c) + (4) hdaSource=inline WITH a stray hda path → the hda file is NOT
    loaded; inline detection runs instead. This guards against a hidden hda
    role being sent with a path while the user has toggled back to inline.
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                # BOM carries inline commodity columns here.
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "GRP-1",
                "Reference Designator": "C200",
                "Function Description": "Filter",
                "Schematic Page": "1",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )
    # A stray HDA workbook that, if loaded, would emit "Loading dedicated
    # HDA file". With hdaSource=inline it must be ignored.
    stray_hda = tmp_path / "stray_hda.xlsx"
    pd.DataFrame(
        [
            {
                "Part Number": "CAP-1",
                "BAE HDA Commodity I": "WRONG",
                "BAE HDA Commodity II": "WRONG",
            }
        ]
    ).to_excel(stray_hda, index=False)

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {
            "failureModesStandard": "FMD-2016",
            "hdaSource": "inline",
        },
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
            # Stray hda input present even though source is inline.
            _state("hda", stray_hda),
        ],
        "mappings": [],
    }
    # Validation passes (inline source never requires an hda file).
    validation = validate_run_request(body)
    assert validation["ok"] is True, validation

    logs: list[str] = []
    result = execute_run_request(body, log_callback=logs.append)
    assert result["status"] == "success", result
    joined = " ".join(logs)
    # The stray dedicated HDA file must NOT have been loaded.
    assert "Loading dedicated HDA file" not in joined, logs
    assert "Detecting inline HDA columns" in joined, logs


def test_hda_source_absent_legacy_client_path_presence_decides(tmp_path: Path) -> None:
    """(d) options WITHOUT hdaSource (legacy client) → behavior is unchanged:
    an attached hda path is loaded as a dedicated file (path presence decides).
    """
    paths = _hda_separate_fixture(tmp_path)
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        # No hdaSource key in options.
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
            _state("hda", paths["hda"]),
        ],
        "mappings": [],
    }
    validation = validate_run_request(body)
    assert validation["ok"] is True, validation

    logs: list[str] = []
    result = execute_run_request(body, log_callback=logs.append)
    assert result["status"] == "success", result
    joined = " ".join(logs)
    # Legacy behavior: a present hda path loads the dedicated file.
    assert "Loading dedicated HDA file" in joined, logs


def test_hda_source_absent_legacy_no_hda_path_uses_inline(tmp_path: Path) -> None:
    """(d) options WITHOUT hdaSource and NO hda path → inline detection, exactly
    as before the flag existed."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "C200",
                "Part Number": "CAP-1",
                "Description": "Cap",
                "BAE HDA Commodity I": "Capacitor",
                "BAE HDA Commodity II": "Ceramic",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "GRP-1",
                "Reference Designator": "C200",
                "Function Description": "Filter",
                "Schematic Page": "1",
            }
        ],
        fm_rows=[
            {
                "FMD-2016 Commodity Type 1": "Capacitor",
                "FMD-2016 Commodity Type 2": "Ceramic",
                "Failure Mode": "Short",
                "Failure Mode Ratio": 1.0,
            }
        ],
    )
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "outputDirectory": str(tmp_path),
        "inputs": [
            _state("grouping", paths["grouping"]),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    logs: list[str] = []
    result = execute_run_request(body, log_callback=logs.append)
    assert result["status"] == "success", result
    joined = " ".join(logs)
    assert "Detecting inline HDA columns" in joined, logs
    assert "Loading dedicated HDA file" not in joined, logs


# ----- Tier-1 #1: Part Usage compute-or-blank+flag ----------------------------
#
# Decision: Part Usage = 1 / (number of INSTANCES of the part), where instances
# are distinct RefDes sharing a usage-base. When the BOM carries NO explicit,
# parseable Part Usage value for a row, the generator MUST compute 1/N from the
# instance count in the source-of-truth (grouping/BOM/functional) rather than
# silently defaulting to 1 (which overstates multi-instance parts). When the
# count is unknown (count==0), the cell is left BLANK and flagged.


def _process_piecepart(proc: FMEAProcessor, paths: dict, tmp_path: Path, **overrides):
    """Run a standard piece_part_generate process() with sensible defaults."""
    request = {
        "group": str(paths["grouping"]) if "grouping" in paths else None,
        "bom": str(paths["bom"]),
        "fm": str(paths["fm"]),
        "hda": None,
        "func": None,
        "piecepart_fmea": None,
        "out_folder": str(tmp_path),
        "out_name": "out",
        "bom_only_mode": False,
        "use_func": False,
        "use_piecepart_merge": False,
        "verbose": False,
        "column_overrides": {},
        "failure_modes_standard": "FMD-2016",
        "group_sheet": "Sheet1",
        "bom_sheet": "Sheet1",
        "hda_sheet": None,
        "fm_sheet": "Sheet1",
        "func_sheet": None,
        "piecepart_fmea_sheet": None,
    }
    request.update(overrides)
    return proc.process(request)


def _usage_cells_for(df: pd.DataFrame, refdes: str) -> list:
    """Return the Part Usage cell value(s) for every PIECE-PART row whose cause
    is refdes (excludes circuit-block / group-header rows)."""
    pp_types = ("piece_part", "validation_warning", "piece_part_no_match")
    rows = df[
        (df["Failure Mode Causes"] == refdes)
        & (df["_row_type"].isin(pp_types))
    ]
    return list(rows["Part Usage"])


def test_multi_instance_base_no_explicit_usage_computes_one_over_n(tmp_path: Path) -> None:
    """U200-1 and U200-2 are distinct BOM instances of base U200 with NO
    explicit Part Usage. Every generated row for both instances must get
    Part Usage 1/2 (the Excel ratio formula "=1/2"), NOT the silent default 1.
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            # NOTE: no "Part Usage" column at all -> absent/blank path.
            {
                "Reference Designator": "U200-1",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
            },
            {
                "Reference Designator": "U200-2",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200-1, U200-2",
                "Function Description": "Processor instances",
                "Schematic Page": "12",
            }
        ],
    )
    proc = FMEAProcessor()
    df = _process_piecepart(proc, paths, tmp_path)

    u1_cells = _usage_cells_for(df, "U200-1")
    u2_cells = _usage_cells_for(df, "U200-2")
    assert u1_cells, df["Failure Mode Causes"].tolist()
    assert u2_cells
    # Both instances -> 1/2, rendered as the Excel ratio formula "=1/2".
    assert set(u1_cells) == {"=1/2"}, u1_cells
    assert set(u2_cells) == {"=1/2"}, u2_cells
    # This is a real instance count, NOT a guess -> no PU_GUESSED flag.
    guessed = [
        w for w in proc.usage_warnings
        if w.get("ReasonCode") == "PU_GUESSED_NO_COUNT_SOURCE"
    ]
    assert guessed == [], guessed


def test_part_usage_counts_distinct_instances_not_occurrences(tmp_path: Path) -> None:
    """A RefDes appearing in MULTIPLE source rows must count as ONE instance,
    not once per occurrence. U200-1 (listed in two groups) + U200-2 = 2 DISTINCT
    instances of base U200 -> Part Usage 1/2, NOT 1/3 (the occurrence count 3).
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {"Reference Designator": "U200-1", "Part Number": "IC-1234", "Description": "MCU",
             "BAE HDA Commodity I": "Microcircuit", "BAE HDA Commodity II": "Digital"},
            {"Reference Designator": "U200-2", "Part Number": "IC-1234", "Description": "MCU",
             "BAE HDA Commodity I": "Microcircuit", "BAE HDA Commodity II": "Digital"},
        ],
        grouping_rows=[
            {"Component Group": "CPU-001", "Reference Designator": "U200-1, U200-2",
             "Function Description": "Processor", "Schematic Page": "12"},
            # U200-1 ALSO appears in a second group -> 3 OCCURRENCES of base U200,
            # but still only 2 DISTINCT instances.
            {"Component Group": "CPU-002", "Reference Designator": "U200-1",
             "Function Description": "Processor support", "Schematic Page": "13"},
        ],
    )
    proc = FMEAProcessor()
    df = _process_piecepart(proc, paths, tmp_path)

    u1_cells = _usage_cells_for(df, "U200-1")
    u2_cells = _usage_cells_for(df, "U200-2")
    assert u1_cells and u2_cells
    assert set(u1_cells) == {"=1/2"}, u1_cells   # distinct instances = 2, NOT "=1/3"
    assert set(u2_cells) == {"=1/2"}, u2_cells


def test_single_instance_base_no_explicit_usage_computes_one(tmp_path: Path) -> None:
    """A lone instance of a base (count==1) with no explicit Part Usage gets 1."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R10",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "G-001",
                "Reference Designator": "R10",
                "Function Description": "Single resistor",
                "Schematic Page": "1",
            }
        ],
    )
    proc = FMEAProcessor()
    df = _process_piecepart(proc, paths, tmp_path)

    cells = _usage_cells_for(df, "R10")
    assert cells, df["Failure Mode Causes"].tolist()
    assert set(cells) == {1}, cells


def test_uncountable_base_no_explicit_usage_blanks_and_flags(tmp_path: Path) -> None:
    """fill_gaps-style: a RefDes whose base is NOT in the source-of-truth count
    (count==0) and that has no explicit Part Usage must leave Part Usage BLANK
    and emit a PU_GUESSED_NO_COUNT_SOURCE flag, rather than asserting usage 1.

    Trigger (matches "fill_gaps without a grouping file"): the variant U900-7 is
    absent from the BOM but its base U900 is present, so it hits the inheritance
    path. With NO grouping/count source, variant_counts_by_base has no entry for
    U900, so the usage-read site sees count==0 and must blank+flag instead of
    asserting 1. We drive _generate_component_rows directly with that exact
    state (empty variant_counts_by_base) to prove the branch deterministically.
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U900",
                "Part Number": "IC-9",
                "Description": "FPGA",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                # No Part Usage column at all.
            }
        ],
    )

    proc = FMEAProcessor()
    # Build the BOM/HDA/FM indexes the way process_gaps does, then run the
    # union-merge component generator with NO grouping count source.
    import pandas as _pd
    from common.utils import normalize_df_columns as _norm  # noqa: F401

    bom_raw = _pd.read_excel(paths["bom"])
    fm_df = _pd.read_excel(paths["fm"])
    # Map BOM/FM columns to canonical names exactly like process_gaps.
    from fmea.fmea_generator_logic import HEADER_CONFIG, REQUIRED_COLS

    # Resolve inline HDA from the BOM's commodity columns (no dedicated HDA file).
    hda_df = proc._resolve_hda_dataframe(None, bom_raw, None, None)
    b_map = proc.map_columns(
        bom_raw, HEADER_CONFIG["BOM"], REQUIRED_COLS["BOM"], source_name="BOM file",
    )
    bom_df = bom_raw.rename(columns={v: k for k, v in b_map.items()})
    f_map = proc.map_columns(
        fm_df, HEADER_CONFIG["FAILURE_MODES"], REQUIRED_COLS["FAILURE_MODES"],
        source_name="Failure Modes file",
    )
    fm_df = fm_df.rename(columns={v: k for k, v in f_map.items()})
    proc.failure_modes_standard = "FMD-2016"
    proc.column_overrides = {}
    proc._build_indexes(bom_df, hda_df, fm_df)
    # CRITICAL: no grouping file -> variant_counts_by_base has NO entry for U900.
    proc.variant_counts_by_base.clear()

    group_stub = {
        "component_group": "GAP-001",
        "description": "fill gap",
        "schematic_page": "1",
    }
    rows = proc._generate_component_rows(
        "U900-7", group_stub, fm_df, mode="standard", source_workflow="fill_gaps",
    )
    assert rows, "expected at least one generated row for U900-7"
    usage_cells = [r.get("Part Usage") for r in rows]
    # Uncountable -> blank cell on every generated row.
    assert all((c == "" or c is None) for c in usage_cells), usage_cells
    # And a PU_GUESSED flag was recorded.
    guessed = [
        w for w in proc.usage_warnings
        if w.get("ReasonCode") == "PU_GUESSED_NO_COUNT_SOURCE"
    ]
    assert guessed, proc.usage_warnings


def test_explicit_bom_usage_is_preserved_not_overridden_by_count(tmp_path: Path) -> None:
    """When the BOM carries an EXPLICIT Part Usage value, the data-file value
    WINS — the computed instance count must NOT override it. U200-1/U200-2 are
    two instances (count 2 -> would compute 1/2) but the BOM says 1/4, so every
    row must show the explicit 1/4 ("=1/4"), not 1/2.
    """
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200-1",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/4",
            },
            {
                "Reference Designator": "U200-2",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/4",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200-1, U200-2",
                "Function Description": "Processor instances",
                "Schematic Page": "12",
            }
        ],
    )
    proc = FMEAProcessor()
    df = _process_piecepart(proc, paths, tmp_path)

    u1_cells = _usage_cells_for(df, "U200-1")
    u2_cells = _usage_cells_for(df, "U200-2")
    assert set(u1_cells) == {"=1/4"}, u1_cells
    assert set(u2_cells) == {"=1/4"}, u2_cells
    # No PU_GUESSED flag — the value was explicit.
    guessed = [
        w for w in proc.usage_warnings
        if w.get("ReasonCode") == "PU_GUESSED_NO_COUNT_SOURCE"
    ]
    assert guessed == [], guessed


# ----- Batch 1 (2026-07 FMEA deep dive): diagnostic-sheet parity, internal
# ----- column hygiene, and required-mapping enforcement ----------------------


def test_build_summary_frames_covers_all_diagnostic_sheets() -> None:
    """H1: build_summary_frames() is the single source of truth for the
    diagnostic summary sheets. Populated processor collections must each
    produce their sheet frame; empty collections must be absent."""
    from fmea.fmea_generator_logic import (
        NEW_REFDES_SHEET_NAME,
        PART_USAGE_DIAGNOSTICS_SHEET_NAME,
        build_summary_frames,
    )

    proc = FMEAProcessor()
    proc.unmatched_hda.append(("R1", "PN-1"))
    proc.fmr_warnings.append(
        {
            "RefDes": "R1",
            "Sum": 1.1,
            "Ratios": "0.5, 0.6",
            "ReasonCode": "FMR_SUM_MISMATCH",
        }
    )
    proc.usage_warnings.append(
        {
            "RefDes": "R1",
            "Base": "R1",
            "Usage": 0.3333,
            "Expected": 1.0,
            "Count": 1,
            "ReasonCode": "PU_EXPECTED_MISMATCH_BASIC",
            "Reason": "Usage mismatch",
        }
    )
    proc.bom_additions.append(
        {
            "ref_des": "U200-A",
            "base_refdes": "U200",
            "usage_fraction": "1/2",
            "part_number": "IC-1234",
            "description": "Microcontroller",
            "hda1": "Microcircuit",
            "hda2": "Digital",
            "fmd1": "Microcircuit",
            "fmd2": "Digital",
            "source_workflow": "piece_part_generate",
            "notes": "",
        }
    )
    proc.part_usage_discrepancies.append(
        {"refdes": "R1", "mapped_count": 3, "computed_count": 1, "diff": -2}
    )

    frames = build_summary_frames(proc)

    assert set(frames.keys()) == {
        "Missing_HDA",
        "Validation_Warnings",
        NEW_REFDES_SHEET_NAME,
        PART_USAGE_DIAGNOSTICS_SHEET_NAME,
    }, sorted(frames.keys())
    for name, frame in frames.items():
        assert not frame.empty, f"{name} frame should not be empty"
    assert list(frames[PART_USAGE_DIAGNOSTICS_SHEET_NAME].columns) == [
        "RefDes",
        "Mapped Count",
        "Computed Count",
        "Diff",
    ]


def test_template_preserve_writes_diagnostic_summary_sheets(tmp_path: Path) -> None:
    """H1: the preserve-formatting writer must emit the same diagnostic
    sheets as the new-workbook writer. Before this fix, Validation_Warnings,
    FMEA Gen New RefDes, and Part Usage Diagnostics were silently dropped in
    preserve mode while the success note claimed they were captured."""
    from openpyxl import Workbook, load_workbook

    from fmea.fmea_generator_logic import NEW_REFDES_SHEET_NAME

    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U200",
                "Part Number": "IC-1234",
                "Description": "Microcontroller",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "1/2",
            },
            {
                "Reference Designator": "R300",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1/3",
            },
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "U200-A, U200-B",
                "Function Description": "Processor variants",
                "Schematic Page": "12",
            },
            {
                "Component Group": "CPU-002",
                "Reference Designator": "R300",
                "Function Description": "Support resistor",
                "Schematic Page": "13",
            },
        ],
        fm_rows=[
            # Ratios sum to 1.1 -> FMR_SUM_MISMATCH warnings.
            {
                "FMD-2016 Commodity Type 1": "Microcircuit",
                "FMD-2016 Commodity Type 2": "Digital",
                "Failure Mode": "Stuck high",
                "Failure Mode Ratio": 0.5,
            },
            {
                "FMD-2016 Commodity Type 1": "Microcircuit",
                "FMD-2016 Commodity Type 2": "Digital",
                "Failure Mode": "Stuck low",
                "Failure Mode Ratio": 0.6,
            },
            {
                "FMD-2016 Commodity Type 1": "Resistor",
                "FMD-2016 Commodity Type 2": "Chip",
                "Failure Mode": "Open",
                "Failure Mode Ratio": 1.0,
            },
        ],
    )

    target_path = tmp_path / "target_fmea.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "FMEA"
    ws.append(
        [
            "FMEA-ID",
            "FMEA Level",
            "Failure Mode Causes",
            "Function Description",
            "Failure Mode",
            "Failure Mode Ratio",
        ]
    )
    ws.append(["CPU-001", "Circuit Block", "U200-A, U200-B", "Processor variants", "", ""])
    ws.append(["CPU-002", "Circuit Block", "R300", "Support resistor", "", ""])
    wb.save(target_path)

    def state(role: str, path: Path, sheet: str = "Sheet1") -> dict:
        return {
            "role": role,
            "label": role,
            "path": str(path),
            "selectedSheet": sheet,
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "s1", "label": sheet}],
        }

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "existing_workbook_preserve_formatting",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            state("grouping", paths["grouping"]),
            state("bom", paths["bom"]),
            state("failureModes", paths["fm"]),
            state("targetWorkbook", target_path, sheet="FMEA"),
        ],
        # Part Usage explicitly mapped -> discrepancy tracking is active.
        "mappings": [
            {"canonical": "Part Usage", "mappedTo": "Part Usage", "status": "mapped"},
        ],
    }

    result = execute_run_request(body)
    assert result["status"] == "success", result
    assert result["warning_count"] > 0, result

    output = load_workbook(Path(result["output_file"]))
    try:
        sheetnames = output.sheetnames
        assert "Validation_Warnings" in sheetnames, sheetnames
        assert NEW_REFDES_SHEET_NAME in sheetnames, sheetnames
        assert "Part Usage Diagnostics" in sheetnames, sheetnames
        # The paste-back banner must survive the preserve path too.
        banner_cell = output[NEW_REFDES_SHEET_NAME].cell(row=1, column=1).value
        assert banner_cell and str(banner_cell).startswith(
            "These rows are RefDes variants"
        ), banner_cell
        # Batch 2: the Part Usage Diagnostics banner must too.
        pu_banner = output["Part Usage Diagnostics"].cell(row=1, column=1).value
        assert pu_banner and "Mapped Count" in str(pu_banner), pu_banner
    finally:
        output.close()


def test_write_excel_report_drops_internal_style_columns(tmp_path: Path) -> None:
    """H2: internal underscore-prefixed columns (_row_type, _style_hint)
    must never leak into the written FMEA sheet."""
    from openpyxl import load_workbook

    proc = FMEAProcessor()
    df = pd.DataFrame(
        [
            {
                "RefDes": "R1",
                "Failure Mode": "Open",
                "_row_type": "piece_part_no_match",
                "_style_hint": "error",
            }
        ]
    )
    out = tmp_path / "internal_cols.xlsx"
    write_excel_report(df, out, proc)

    wb = load_workbook(out)
    try:
        headers = [cell.value for cell in wb["FMEA"][1]]
    finally:
        wb.close()
    leaked = [h for h in headers if h is not None and str(h).startswith("_")]
    assert leaked == [], f"internal columns leaked into output: {leaked}"


def _validate_body_with_mapping(
    tmp_path: Path, workflow_id: str, mappings: list[dict], options: dict | None = None
) -> dict:
    """Shared fixture for the required-mapping validation tests."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-001",
                "Reference Designator": "R100",
                "Function Description": "Support",
                "Schematic Page": "1",
            }
        ],
    )

    def state(role: str, path: Path) -> dict:
        return {
            "role": role,
            "label": role,
            "path": str(path),
            "selectedSheet": "Sheet1",
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    inputs = [state("bom", paths["bom"]), state("failureModes", paths["fm"])]
    if workflow_id == "piece_part_generate":
        inputs.insert(0, state("grouping", paths["grouping"]))

    return {
        "workflowId": workflow_id,
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016", **(options or {})},
        "inputs": inputs,
        "mappings": mappings,
    }


def test_validate_run_rejects_do_not_map_on_required_mapping(tmp_path: Path) -> None:
    """H3: explicitly selecting 'Do Not Map' for a required column (here
    'Failure Mode') must block at validate time with invalid_do_not_map —
    not silently fall back to heuristic column detection."""
    from shared.pre_run_validation import DO_NOT_MAP_SENTINEL

    body = _validate_body_with_mapping(
        tmp_path,
        "piece_part_generate",
        mappings=[
            {
                "canonical": "Failure Mode",
                "mappedTo": DO_NOT_MAP_SENTINEL,
                "status": "not_mapped",
            }
        ],
    )
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "invalid_do_not_map", result
    assert "Failure Mode" in result["toast_text"], result


def test_validate_run_allows_do_not_map_on_optional_mapping(tmp_path: Path) -> None:
    """H3 companion: 'Do Not Map' on a non-required column must NOT block."""
    from shared.pre_run_validation import DO_NOT_MAP_SENTINEL

    body = _validate_body_with_mapping(
        tmp_path,
        "piece_part_generate",
        mappings=[
            {
                "canonical": "Component Part Description",
                "mappedTo": DO_NOT_MAP_SENTINEL,
                "status": "not_mapped",
            }
        ],
    )
    result = validate_run_request(body)
    assert result["ok"] is True, result


def test_validate_run_allows_do_not_map_on_fmea_id_in_bom_only(tmp_path: Path) -> None:
    """H3 companion: FMEA-ID is hidden (and derived) in BOM-Only mode, so a
    stale 'Do Not Map' for it must not block that workflow."""
    from shared.pre_run_validation import DO_NOT_MAP_SENTINEL

    body = _validate_body_with_mapping(
        tmp_path,
        "bom_only",
        mappings=[
            {
                "canonical": "FMEA-ID",
                "mappedTo": DO_NOT_MAP_SENTINEL,
                "status": "not_mapped",
            }
        ],
        options={"ccaPrefix": "PSU1"},
    )
    result = validate_run_request(body)
    assert result["ok"] is True, result


# ----- Batch 2 (2026-07 FMEA deep dive): diagnostics language ----------------


def test_reason_code_labels_cover_all_emitted_codes() -> None:
    """Lockstep guard: every UPPER_SNAKE ReasonCode literal emitted anywhere
    in the backend must have a plain-language REASON_CODE_LABELS entry, so a
    raw token-case string can never reach the Validation_Warnings sheet."""
    import re as _re

    from common.user_facing_labels import REASON_CODE_LABELS

    backend_root = Path(__file__).resolve().parents[1] / "python"
    pattern = _re.compile(
        r"['\"]ReasonCode['\"]\s*[:=]\s*['\"]([A-Z][A-Z_0-9]+)['\"]"
        r"|reason_code\s*=\s*['\"]([A-Z][A-Z_0-9]+)['\"]"
    )
    emitted: set[str] = set()
    for py_file in backend_root.rglob("*.py"):
        text = py_file.read_text(encoding="utf-8", errors="ignore")
        for match in pattern.finditer(text):
            emitted.add(match.group(1) or match.group(2))

    assert emitted, "scanner found no emitted ReasonCodes - pattern broken?"
    missing = sorted(code for code in emitted if code not in REASON_CODE_LABELS)
    assert missing == [], (
        f"ReasonCodes emitted without a REASON_CODE_LABELS entry: {missing}"
    )


def test_fmea_parse_failure_reason_code_says_replaced_with_count(
    tmp_path: Path,
) -> None:
    """The FMEA path replaces an unparseable Part Usage with the instance
    count 1/N - it never defaults to 1.0. Its reason code and label must say
    so; PU_PARSE_DEFAULTED remains accurate only for the BOM-Compare sites
    that genuinely default to 1.0."""
    from common.user_facing_labels import to_reason_code_label

    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "U500",
                "Part Number": "IC-9",
                "Description": "Micro",
                "BAE HDA Commodity I": "Microcircuit",
                "BAE HDA Commodity II": "Digital",
                "Part Usage": "garbage-value",
            }
        ],
        grouping_rows=[
            {
                "Component Group": "CPU-009",
                "Reference Designator": "U500",
                "Function Description": "Core",
                "Schematic Page": "2",
            }
        ],
    )
    proc = FMEAProcessor()
    _process_piecepart(proc, paths, tmp_path)

    parse_warnings = [
        w for w in proc.usage_warnings
        if "could not be parsed" in str(w.get("Reason", ""))
    ]
    assert parse_warnings, proc.usage_warnings
    code = parse_warnings[0]["ReasonCode"]
    assert code == "PU_PARSE_REPLACED_WITH_COUNT", parse_warnings[0]
    label = to_reason_code_label(code)
    assert "1.0" not in label, label
    assert "instance count" in label.lower(), label


def test_part_usage_diagnostics_sheet_has_explanatory_banner(
    tmp_path: Path,
) -> None:
    """The Part Usage Diagnostics sheet must open with a banner explaining
    Mapped Count / Computed Count / Diff, mirroring the New-RefDes sheet."""
    from openpyxl import load_workbook

    from fmea.fmea_generator_logic import PART_USAGE_DIAGNOSTICS_SHEET_NAME

    proc = FMEAProcessor()
    proc.part_usage_discrepancies.append(
        {"refdes": "R1", "mapped_count": 3, "computed_count": 1, "diff": -2}
    )
    df = pd.DataFrame([{"RefDes": "R1", "Failure Mode": "Open"}])
    out = tmp_path / "pu_banner.xlsx"
    write_excel_report(df, out, proc)

    wb = load_workbook(out)
    try:
        ws = wb[PART_USAGE_DIAGNOSTICS_SHEET_NAME]
        banner = ws.cell(row=1, column=1).value
        headers = [cell.value for cell in ws[2]]
    finally:
        wb.close()
    assert banner and "Mapped Count" in str(banner), banner
    assert "Computed Count" in str(banner), banner
    assert headers[0] == "RefDes", headers


def test_template_merge_summary_explains_diagnostic_flags() -> None:
    """Template_Merge_Summary must carry a legend for the two Diagnostic
    flag values so 'NOT IN BOM - Review' / 'NEW - Added by generator' are
    never bare tokens, and the flagged metric must name the flag it counts."""
    from openpyxl import Workbook

    from fmea.fmea_template_writer import TemplateWriteResult, _write_summary_sheets

    wb = Workbook()
    result = TemplateWriteResult(groups_matched=1, pp_rows_flagged=2)
    _write_summary_sheets(wb, FMEAProcessor(), result, lambda msg: None)

    ws = wb["Template_Merge_Summary"]
    rows = [
        (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
        for r in range(1, ws.max_row + 1)
    ]
    labels = [str(label) for label, _ in rows if label]

    flagged_metric = [m for m in labels if "flagged" in m.lower()]
    assert flagged_metric and "NOT IN BOM - Review" in flagged_metric[0], labels

    def legend_value(flag: str) -> str:
        for label, value in rows:
            if label and str(label).startswith("Diagnostic flag") and flag in str(label):
                return str(value or "")
        return ""

    assert len(legend_value("NOT IN BOM - Review")) > 20, rows
    assert len(legend_value("NEW - Added by generator")) > 20, rows


def test_unsupported_execution_path_toast_names_the_combo(tmp_path: Path) -> None:
    """The unsupported-path toast must say WHICH workflow/strategy combo was
    rejected and what is supported - not just 'not yet available'."""
    body = {
        "workflowId": "bogus_flow",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "unsupported_execution_path", result
    assert "bogus_flow" in result["toast_text"], result["toast_text"]
    assert "piece_part_generate" in result["toast_text"], result["toast_text"]


def test_merge_fmc_gate_defers_to_missing_files(tmp_path: Path) -> None:
    """File checks come first: a fill_gaps body with no files loaded must
    report missing_files, not the Failure Mode Causes mapping gate."""
    body = {
        "workflowId": "fill_gaps",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_files", result


def test_merge_fmc_gate_builds_validation_cards(tmp_path: Path) -> None:
    """The FMC gate must flow through the standard response builder so the
    UI gets a structured 'Run is blocked' card, not an empty validations
    list."""
    paths = _write_fixture(
        tmp_path,
        bom_rows=[
            {
                "Reference Designator": "R100",
                "Part Number": "RES-1",
                "Description": "Resistor",
                "BAE HDA Commodity I": "Resistor",
                "BAE HDA Commodity II": "Chip",
                "Part Usage": "1",
            }
        ],
    )
    existing_fmea = tmp_path / "existing.xlsx"
    pd.DataFrame([{"FMEA-ID": "X-001"}]).to_excel(existing_fmea, index=False)

    body = {
        "workflowId": "fill_gaps",
        "outputStrategyId": "new_workbook_standard",
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            _state("existingFmea", existing_fmea),
            _state("bom", paths["bom"]),
            _state("failureModes", paths["fm"]),
        ],
        "mappings": [],
    }
    result = validate_run_request(body)
    assert result["ok"] is False, result
    assert result["reason_code"] == "missing_failure_mode_causes_mapping", result
    blocked = [
        message
        for message in result["validations"]
        if message.get("severity") == "error"
        and "Failure Mode Causes" in str(message.get("detail", ""))
    ]
    assert blocked, result["validations"]
