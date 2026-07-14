from __future__ import annotations

import json
import os
import queue
import subprocess
import sys
import threading
import textwrap
import time
from collections.abc import Callable
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SIDECAR = ROOT / "backend" / "python" / "sidecar_main.py"
# Suppress heartbeat during tests and isolate log files to prevent contention
import tempfile as _tempfile
_TEST_LOG_DIR = _tempfile.mkdtemp(prefix="sidecar_test_logs_")
SIDECAR_ENV = {
    **os.environ,
    "SIDECAR_HEARTBEAT_INTERVAL": "9999",
    "RELIABILITY_TOOLS_LOG_DIR": _TEST_LOG_DIR,
    "PYTHONDONTWRITEBYTECODE": "1",
    "SIDECAR_LOG_LEVEL": "CRITICAL",  # Suppress file logging to avoid Windows PermissionError on rotation
}


class _SidecarReader:
    """Thread-safe line reader that uses a single persistent reader thread per process."""

    def __init__(self, process: subprocess.Popen[str]) -> None:
        self._process = process
        self._queue: queue.Queue[str] = queue.Queue()
        self._thread = threading.Thread(target=self._reader, args=(process,), daemon=True)
        self._thread.start()

    def _reader(self, process: subprocess.Popen[str]) -> None:
        try:
            while True:
                line = process.stdout.readline()
                if not line:
                    break
                self._queue.put(line)
        except Exception:
            pass

    def read_line(self, timeout: float = 15.0) -> str:
        try:
            return self._queue.get(timeout=timeout)
        except queue.Empty as exc:
            raise AssertionError(f"Timed out waiting for sidecar output after {timeout:.2f}s") from exc

    def read_until(
        self,
        *,
        request_id: str | None = None,
        run_id: str | None = None,
        kind: str | None = None,
        timeout: float = 30.0,
    ) -> dict:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            remaining = max(0.1, deadline - time.monotonic())
            try:
                line = self._queue.get(timeout=remaining).strip()
            except queue.Empty:
                continue
            if not line:
                continue
            message = json.loads(line)
            if request_id is not None and message.get("request_id") != request_id:
                continue
            if run_id is not None and message.get("run_id") != run_id:
                continue
            if kind is not None and message.get("kind") != kind:
                continue
            return message
        raise AssertionError(f"Timed out waiting for kind={kind!r}, request_id={request_id!r}, run_id={run_id!r}")


def _start_sidecar() -> tuple[subprocess.Popen[str], _SidecarReader]:
    """Start the sidecar and return (process, reader). Waits for the ready message."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    reader = _SidecarReader(process)
    ready_line = reader.read_line().strip()
    assert ready_line
    ready_msg = json.loads(ready_line)
    assert ready_msg["kind"] == "ready"
    return process, reader


def _start_scripted_sidecar(script: str) -> tuple[subprocess.Popen[str], _SidecarReader]:
    """Start the real protocol loop with test-controlled route functions."""
    process = subprocess.Popen(
        [sys.executable, "-c", textwrap.dedent(script)],
        cwd=SIDECAR.parent,
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    reader = _SidecarReader(process)
    ready_line = reader.read_line().strip()
    assert ready_line
    ready_msg = json.loads(ready_line)
    assert ready_msg["kind"] == "ready"
    return process, reader


def _read_ready_line(process: subprocess.Popen[str]) -> dict:
    """Read the ready message, initializing the shared reader for this process."""
    line = _get_reader(process).read_line().strip()
    assert line
    message = json.loads(line)
    assert message["kind"] == "ready"
    return message


def _write_command(
    process: subprocess.Popen[str],
    request_id: str,
    command: str,
    body: dict,
) -> None:
    process.stdin.write(
        json.dumps(
            {
                "protocol_version": "0.1.0",
                "id": f"cmd_{request_id}",
                "kind": "command",
                "request_id": request_id,
                "run_id": None,
                "timestamp": "2026-04-03T00:00:00Z",
                "payload": {
                    "command": command,
                    "body": body,
                },
            }
        )
        + "\n"
    )
    process.stdin.flush()


def _send_command(process: subprocess.Popen[str], request_id: str, command: str, body: dict) -> dict:
    _write_command(process, request_id, command, body)
    return _read_until(process, request_id=request_id)


def _read_messages_until(
    reader: _SidecarReader,
    predicate: Callable[[dict], bool],
    *,
    timeout: float = 30.0,
) -> tuple[list[dict], dict]:
    """Retain every ordered message while waiting for a matching envelope."""
    deadline = time.monotonic() + timeout
    messages: list[dict] = []
    while time.monotonic() < deadline:
        remaining = max(0.1, deadline - time.monotonic())
        line = reader.read_line(timeout=remaining).strip()
        if not line:
            continue
        message = json.loads(line)
        messages.append(message)
        if predicate(message):
            return messages, message
    raise AssertionError("Timed out waiting for a matching sidecar message")


_READERS: dict[int, _SidecarReader] = {}

import atexit
atexit.register(lambda: _READERS.clear())


def _get_reader(process: subprocess.Popen[str]) -> _SidecarReader:
    """Get or create a single reader per process (keyed by id to avoid PID reuse)."""
    key = id(process)
    existing = _READERS.get(key)
    if existing is not None and existing._process is not process:
        del _READERS[key]
        existing = None
    if existing is None:
        _READERS[key] = _SidecarReader(process)
    return _READERS[key]


def _read_until(
    process: subprocess.Popen[str],
    *,
    request_id: str | None = None,
    run_id: str | None = None,
    kind: str | None = None,
    timeout: float = 30.0,
) -> dict:
    """Read until a matching message. Reuses a single reader per process."""
    return _get_reader(process).read_until(request_id=request_id, run_id=run_id, kind=kind, timeout=timeout)


def _build_phase4_run_body(tmp_path: Path, group_rows: int = 1) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    grouping_path = tmp_path / "grouping.xlsx"
    bom_path = tmp_path / "bom.xlsx"
    failure_modes_path = tmp_path / "failure_modes.xlsx"

    grouping_rows = [
        {
            "Component Group": f"CPU-{index:03d}",
            "Reference Designator": f"R{200 + index}",
            "Function Description": "Processor support",
            "Schematic Page": "12",
        }
        for index in range(group_rows)
    ]
    bom_rows = [
        {
            "Reference Designator": f"R{200 + index}",
            "Part Number": f"PN-{index:04d}",
            "Description": "Resistor",
            "BAE HDA Commodity I": "Resistor",
            "BAE HDA Commodity II": "Chip",
            "Part Usage": "1",
        }
        for index in range(group_rows)
    ]

    pd.DataFrame(grouping_rows).to_excel(grouping_path, index=False)
    pd.DataFrame(bom_rows).to_excel(bom_path, index=False)
    pd.DataFrame(
        {
            "FMD-2016 Commodity Type 1": ["Resistor", "Resistor"],
            "FMD-2016 Commodity Type 2": ["Chip", "Chip"],
            "Failure Mode": ["Open", "Drift"],
            "Failure Mode Ratio": [0.5, 0.5],
        }
    ).to_excel(failure_modes_path, index=False)

    def input_state(role: str, label: str, path: Path) -> dict:
        return {
            "role": role,
            "label": label,
            "path": str(path),
            "selectedSheet": "Sheet1",
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "sheet1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            input_state("grouping", "Grouping workbook", grouping_path),
            input_state("bom", "BOM workbook", bom_path),
            input_state("failureModes", "Failure modes workbook", failure_modes_path),
        ],
        "mappings": [],
    }


def _build_fill_gaps_run_body(tmp_path: Path) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    fmea_path = tmp_path / "existing_fmea.xlsx"
    bom_path = tmp_path / "bom.xlsx"
    failure_modes_path = tmp_path / "failure_modes.xlsx"

    # Existing FMEA has R200 but NOT R201
    fmea_rows = [
        {
            "FMEA Level": "Circuit Block",
            "FMEA-ID": "CPU-000",
            "Failure Mode Causes": "R200",
            "Function Description": "Processor support",
        },
        {
            "FMEA Level": "Piece Part",
            "FMEA-ID": "CPU-000-R200-01",
            "Failure Mode Causes": "R200",
            "Function Description": "Resistor",
            "Failure Mode": "Open",
            "Failure Mode Ratio": 0.5,
        },
        {
            "FMEA Level": "Piece Part",
            "FMEA-ID": "CPU-000-R200-02",
            "Failure Mode Causes": "R200",
            "Function Description": "Resistor",
            "Failure Mode": "Drift",
            "Failure Mode Ratio": 0.5,
        },
    ]

    # BOM has both R200 and R201 — R201 is the gap
    bom_rows = [
        {
            "Reference Designator": "R200",
            "Part Number": "PN-0000",
            "Description": "Resistor",
            "BAE HDA Commodity I": "Resistor",
            "BAE HDA Commodity II": "Chip",
            "Part Usage": "1",
        },
        {
            "Reference Designator": "R201",
            "Part Number": "PN-0001",
            "Description": "Resistor",
            "BAE HDA Commodity I": "Resistor",
            "BAE HDA Commodity II": "Chip",
            "Part Usage": "1",
        },
    ]

    pd.DataFrame(fmea_rows).to_excel(fmea_path, index=False)
    pd.DataFrame(bom_rows).to_excel(bom_path, index=False)
    pd.DataFrame(
        {
            "FMD-2016 Commodity Type 1": ["Resistor", "Resistor"],
            "FMD-2016 Commodity Type 2": ["Chip", "Chip"],
            "Failure Mode": ["Open", "Drift"],
            "Failure Mode Ratio": [0.5, 0.5],
        }
    ).to_excel(failure_modes_path, index=False)

    def input_state(role: str, label: str, path: Path) -> dict:
        return {
            "role": role,
            "label": label,
            "path": str(path),
            "selectedSheet": "Sheet1",
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "sheet1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "fill_gaps",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            input_state("existingFmea", "Existing FMEA workbook", fmea_path),
            input_state("bom", "BOM workbook", bom_path),
            input_state("failureModes", "Failure modes workbook", failure_modes_path),
        ],
        # Fix C2: merge modes require an explicit Failure Mode Causes
        # mapping. Without this the validator rejects the run.
        "mappings": [
            {
                "canonical": "Failure Mode Causes",
                "mappedTo": "Failure Mode Causes",
                "status": "mapped",
            },
        ],
    }


def test_sidecar_health_check_round_trip() -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_1",
                    "kind": "command",
                    "request_id": "req_1",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "health_check",
                        "body": {"app": "reliability_tools_tauri"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, request_id="req_1", kind="result")
        assert result["payload"]["status"] == "ok"
    finally:
        process.kill()


def test_sidecar_emits_heartbeat_within_interval() -> None:
    """Verify the sidecar emits at least one heartbeat message within its interval."""
    heartbeat_env = {**os.environ, "SIDECAR_HEARTBEAT_INTERVAL": "2"}
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=heartbeat_env,
    )
    try:
        _read_ready_line(process)
        # Wait for a heartbeat (interval is 2s, give 5s)
        heartbeat = _read_until(process, kind="heartbeat", timeout=5.0)
        assert heartbeat["kind"] == "heartbeat"
        assert heartbeat["payload"]["backend"] == "python-sidecar"
        assert heartbeat["payload"]["protocol_version"]
    finally:
        process.kill()


def test_sidecar_lists_excel_sheets(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sheets.xlsx"
    workbook = Workbook()
    workbook.active.title = "Primary"
    workbook.create_sheet("Secondary")
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_2",
                    "kind": "command",
                    "request_id": "req_2",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "list_sheets",
                        "body": {"path": str(workbook_path)},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert result["kind"] == "result"
        assert result["payload"]["sheets"] == ["Primary", "Secondary"]
    finally:
        process.kill()


def test_sidecar_inspects_input_headers_and_preview(tmp_path: Path) -> None:
    workbook_path = tmp_path / "inspect.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Part Number", "Failure Mode", "Effect"])
    sheet.append(["R1", "Open", "Loss of output"])
    sheet.append(["C2", "Short", "Bias collapse"])
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_3",
                    "kind": "command",
                    "request_id": "req_3",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "inspect_input",
                        "body": {"path": str(workbook_path), "sheet": "Data"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert result["kind"] == "result"
        assert result["payload"]["sheet"] == "Data"
        assert result["payload"]["header_row"] == 1
        assert result["payload"]["columns"] == ["Part Number", "Failure Mode", "Effect"]
        assert result["payload"]["row_count"] == 2
        assert result["payload"]["preview_rows"][0]["Part Number"] == "R1"
    finally:
        process.kill()


def test_sidecar_analyzes_template_metadata(tmp_path: Path) -> None:
    workbook_path = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    sheet.freeze_panes = "B2"
    sheet.append(["FMEA ID", "Failure Mode", "End Effect"])
    sheet.append(["ID-1", "Open", "Subsystem down"])
    sheet.merge_cells("A1:B1")
    sheet.protection.sheet = True
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_4",
                    "kind": "command",
                    "request_id": "req_4",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "analyze_template",
                        "body": {"path": str(workbook_path), "sheet": "Template"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert result["kind"] == "result"
        assert result["payload"]["sheet"] == "Template"
        assert result["payload"]["freeze_panes"] == "B2"
        assert result["payload"]["protected_sheet"] is True
        assert result["payload"]["merged_range_count"] == 1
        assert "FMEA ID" in result["payload"]["columns"][0]
    finally:
        process.kill()


def test_sidecar_inspect_input_reports_row_cap_metadata(tmp_path: Path) -> None:
    workbook_path = tmp_path / "row_cap.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    sheet.append(["Part Number"])
    for index in range(20_005):
        sheet.append([f"PN-{index:05d}"])
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_row_cap",
                    "kind": "command",
                    "request_id": "req_row_cap",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "inspect_input",
                        "body": {"path": str(workbook_path), "sheet": "Data"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert result["payload"]["row_count"] == 20_000
        assert result["payload"]["rows_scanned"] == 20_000
        assert result["payload"]["header_rows_scanned"] == 1
        assert result["payload"]["row_cap_applied"] is True
        assert result["payload"]["column_cap_applied"] is False
        assert result["payload"]["header_search_cap_applied"] is False
    finally:
        process.kill()


def test_sidecar_inspect_input_reports_column_cap_metadata(tmp_path: Path) -> None:
    workbook_path = tmp_path / "column_cap.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wide"
    headers = [f"Column {index}" for index in range(1, 106)]
    sheet.append(headers)
    sheet.append([f"value-{index}" for index in range(1, 106)])
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_col_cap",
                    "kind": "command",
                    "request_id": "req_col_cap",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "inspect_input",
                        "body": {"path": str(workbook_path), "sheet": "Wide"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert len(result["payload"]["columns"]) == 100
        assert result["payload"]["columns_scanned"] == 100
        assert result["payload"]["header_rows_scanned"] == 1
        assert result["payload"]["column_cap_applied"] is True
        assert result["payload"]["row_cap_applied"] is False
        assert result["payload"]["header_search_cap_applied"] is False
    finally:
        process.kill()


def test_sidecar_inspect_input_caps_sparse_sheet_by_physical_rows_scanned(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sparse_row_cap.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet.append(["Part Number"])
    for index in range(20_005):
        if index < 3 or index == 20_004:
            sheet.append([f"PN-{index:05d}"])
        else:
            sheet.append([""])
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_sparse_row_cap",
                    "kind": "command",
                    "request_id": "req_sparse_row_cap",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "inspect_input",
                        "body": {"path": str(workbook_path), "sheet": "Sparse"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="result")
        assert result["payload"]["rows_scanned"] == 20_000
        assert result["payload"]["row_cap_applied"] is True
        assert result["payload"]["row_count"] == 3
        assert result["payload"]["preview_rows"] == [
            {"Part Number": "PN-00000"},
            {"Part Number": "PN-00001"},
            {"Part Number": "PN-00002"},
        ]
    finally:
        process.kill()


def test_sidecar_inspect_input_fails_when_header_search_cap_is_hit(tmp_path: Path) -> None:
    workbook_path = tmp_path / "header_cap.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LateHeader"
    for _ in range(1_001):
        sheet.append([" "])
    sheet.append(["Part Number", "Failure Mode"])
    workbook.save(workbook_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        process.stdin.write(
            json.dumps(
                {
                    "protocol_version": "0.1.0",
                    "id": "cmd_header_cap",
                    "kind": "command",
                    "request_id": "req_header_cap",
                    "run_id": None,
                    "timestamp": "2026-04-03T00:00:00Z",
                    "payload": {
                        "command": "inspect_input",
                        "body": {"path": str(workbook_path), "sheet": "LateHeader"},
                    },
                }
            )
            + "\n"
        )
        process.stdin.flush()

        result = _read_until(process, kind="error")
        assert "Could not locate a non-empty header row within the first 1000 scanned rows" in result["payload"]["message"]
    finally:
        process.kill()


def test_sidecar_validates_phase4_standard_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_validate", "validate_run", _build_phase4_run_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
        assert result["payload"]["reason_code"] == "ok"
        assert result["payload"]["validations"][0]["title"] == "Backend validation passed"
    finally:
        process.kill()


def test_sidecar_rejects_unmigrated_output_strategy(tmp_path: Path) -> None:
    body = _build_phase4_run_body(tmp_path)
    body["outputStrategyId"] = "existing_workbook_best_effort"

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_validate_unsupported", "validate_run", body)
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is False
        assert result["payload"]["reason_code"] == "unsupported_execution_path"
    finally:
        process.kill()


def test_sidecar_executes_phase4_standard_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_execute", "execute_run", _build_phase4_run_body(tmp_path))
        assert ack["kind"] == "ack"
        assert ack["payload"]["workflow_id"] == "piece_part_generate"
        run_id = ack["payload"]["run_id"]

        status = _read_until(process, run_id=run_id, kind="status")
        assert status["payload"]["status"] in {"starting", "running"}

        progress = _read_until(process, run_id=run_id, kind="progress")
        assert progress["payload"]["percent"] >= 1

        log = _read_until(process, run_id=run_id, kind="log")
        assert log["payload"]["line"]

        result = _read_until(process, run_id=run_id, kind="result")
        assert result["payload"]["status"] == "success"
        assert result["payload"]["row_count"] > 0
        assert result["payload"]["output_file"]
        assert result["payload"]["log_lines"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()

        workbook = load_workbook(output_path)
        try:
            assert "FMEA" in workbook.sheetnames
        finally:
            workbook.close()
    finally:
        process.kill()


def test_sidecar_rejects_second_execute_while_run_is_active(tmp_path: Path) -> None:
    first_body = _build_phase4_run_body(tmp_path / "first", group_rows=2000)
    second_body = _build_phase4_run_body(tmp_path / "second", group_rows=2000)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        first_ack = _send_command(process, "req_execute_busy_1", "execute_run", first_body)
        assert first_ack["kind"] == "ack"

        second = _send_command(process, "req_execute_busy_2", "execute_run", second_body)
        assert second["kind"] == "error"
        assert "already active" in second["payload"]["message"]
    finally:
        process.kill()


def test_sidecar_cancels_active_run(tmp_path: Path) -> None:
    body = _build_phase4_run_body(tmp_path / "cancel", group_rows=2000)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_execute_cancel", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        cancel = _send_command(process, "req_cancel", "cancel_run", {"run_id": run_id})
        assert cancel["kind"] == "result"
        assert cancel["payload"]["status"] == "cancelling"

        terminal = _read_until(process, run_id=run_id, timeout=30.0)
        while terminal["kind"] not in {"cancelled", "result", "backend_error"}:
            terminal = _read_until(process, run_id=run_id, timeout=30.0)

        assert terminal["kind"] == "cancelled"
        assert "cancel" in terminal["payload"]["message"].lower()
    finally:
        process.kill()


def test_sidecar_validates_fill_gaps_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_validate_fg", "validate_run", _build_fill_gaps_run_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
        assert result["payload"]["reason_code"] == "ok"
    finally:
        process.kill()


def test_sidecar_executes_fill_gaps_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_execute_fg", "execute_run", _build_fill_gaps_run_body(tmp_path))
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert result["payload"]["row_count"] > 0
        assert result["payload"]["output_file"]
        assert "FillGaps" in result["payload"]["output_file"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()
    finally:
        process.kill()


def _build_template_preserve_run_body(tmp_path: Path) -> dict:
    """Build a run body that uses the template-preserved write strategy."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    grouping_path = tmp_path / "grouping.xlsx"
    bom_path = tmp_path / "bom.xlsx"
    failure_modes_path = tmp_path / "failure_modes.xlsx"
    target_path = tmp_path / "target_fmea.xlsx"

    pd.DataFrame([
        {
            "Component Group": "CPU-000",
            "Reference Designator": "R200",
            "Function Description": "Processor support",
            "Schematic Page": "12",
        },
    ]).to_excel(grouping_path, index=False)

    pd.DataFrame([
        {
            "Reference Designator": "R200",
            "Part Number": "PN-0000",
            "Description": "Resistor",
            "BAE HDA Commodity I": "Resistor",
            "BAE HDA Commodity II": "Chip",
            "Part Usage": "1",
        },
    ]).to_excel(bom_path, index=False)

    pd.DataFrame(
        {
            "FMD-2016 Commodity Type 1": ["Resistor", "Resistor"],
            "FMD-2016 Commodity Type 2": ["Chip", "Chip"],
            "Failure Mode": ["Open", "Drift"],
            "Failure Mode Ratio": [0.5, 0.5],
        }
    ).to_excel(failure_modes_path, index=False)

    # Create a simple target FMEA workbook to merge into
    wb = Workbook()
    ws = wb.active
    ws.title = "FMEA"
    ws.append([
        "FMEA-ID", "FMEA Level", "Failure Mode Causes",
        "Function Description", "Failure Mode", "Failure Mode Ratio",
    ])
    ws.append(["CPU-000", "Circuit Block", "R200", "Processor support", "", ""])
    wb.save(target_path)

    def input_state(role: str, label: str, path: Path) -> dict:
        return {
            "role": role,
            "label": label,
            "path": str(path),
            "selectedSheet": "Sheet1" if role != "targetWorkbook" else "FMEA",
            "source": "desktop-bridge",
            "isResolvingSheets": False,
            "isAnalyzing": False,
            "resolutionError": None,
            "sheets": [{"id": "sheet1", "label": "Sheet1"}]
            if role != "targetWorkbook"
            else [{"id": "fmea", "label": "FMEA"}],
        }

    return {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "existing_workbook_preserve_formatting",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            input_state("grouping", "Grouping workbook", grouping_path),
            input_state("bom", "BOM workbook", bom_path),
            input_state("failureModes", "Failure modes workbook", failure_modes_path),
            input_state("targetWorkbook", "Target workbook", target_path),
        ],
        "mappings": [],
    }


def test_sidecar_validates_template_preserve_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process, "req_validate_tp", "validate_run",
            _build_template_preserve_run_body(tmp_path),
        )
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
        assert result["payload"]["reason_code"] == "ok"
    finally:
        process.kill()


def test_sidecar_executes_template_preserve_run(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(
            process, "req_execute_tp", "execute_run",
            _build_template_preserve_run_body(tmp_path),
        )
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert result["payload"]["row_count"] > 0
        assert result["payload"]["output_file"]
        assert "_Merged_" in result["payload"]["output_file"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()

        workbook = load_workbook(output_path)
        try:
            assert "FMEA" in workbook.sheetnames
            assert "Template_Merge_Summary" in workbook.sheetnames
        finally:
            workbook.close()
    finally:
        process.kill()


# =========================================================================
# Regression tests for known Flet-era failure patterns
# =========================================================================


def test_sidecar_execute_emits_backend_error_on_missing_columns(tmp_path: Path) -> None:
    """Regression: missing required columns should emit backend_error, not crash the sidecar."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    grouping_path = tmp_path / "grouping.xlsx"
    bom_path = tmp_path / "bom.xlsx"
    failure_modes_path = tmp_path / "failure_modes.xlsx"

    # Grouping file missing required "Component Group" column
    pd.DataFrame({"Wrong Column": ["A"], "Reference Designator": ["R1"]}).to_excel(grouping_path, index=False)
    pd.DataFrame({
        "Reference Designator": ["R1"], "Part Number": ["PN-001"],
        "Description": ["Resistor"], "BAE HDA Commodity I": ["Resistor"],
        "BAE HDA Commodity II": ["Chip"], "Part Usage": ["1"],
    }).to_excel(bom_path, index=False)
    pd.DataFrame({
        "FMD-2016 Commodity Type 1": ["Resistor"], "FMD-2016 Commodity Type 2": ["Chip"],
        "Failure Mode": ["Open"], "Failure Mode Ratio": [1.0],
    }).to_excel(failure_modes_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            input_state("grouping", "Grouping", grouping_path),
            input_state("bom", "BOM", bom_path),
            input_state("failureModes", "FM", failure_modes_path),
        ],
        "mappings": [],
    }

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_badcol", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        # Should get a backend_error terminal event, not a hang or crash
        terminal = _read_until(process, run_id=run_id, timeout=15.0)
        while terminal["kind"] not in {"backend_error", "result", "cancelled"}:
            terminal = _read_until(process, run_id=run_id, timeout=15.0)

        assert terminal["kind"] == "backend_error"
        assert terminal["payload"]["message"]  # non-empty error message
        # Phase 5: backend_error must carry an error code and a traceback so
        # the UI can show the user something actionable.
        assert "code" in terminal["payload"], "backend_error missing 'code' field"
        assert isinstance(terminal["payload"]["code"], str)
        assert terminal["payload"]["code"]  # non-empty
        assert "traceback" in terminal["payload"], "backend_error missing 'traceback' field"
        assert isinstance(terminal["payload"]["traceback"], str)
        # The traceback should contain at least the conventional header.
        assert "Traceback" in terminal["payload"]["traceback"]
    finally:
        process.kill()


def test_sidecar_remains_responsive_after_failed_run(tmp_path: Path) -> None:
    """Regression: sidecar must accept new commands after a run fails."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    bad_path = tmp_path / "nonexistent.xlsx"

    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            {"role": "grouping", "label": "G", "path": str(bad_path),
             "selectedSheet": "Sheet1", "source": "desktop-bridge",
             "isResolvingSheets": False, "isAnalyzing": False,
             "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}]},
            {"role": "bom", "label": "B", "path": str(bad_path),
             "selectedSheet": "Sheet1", "source": "desktop-bridge",
             "isResolvingSheets": False, "isAnalyzing": False,
             "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}]},
            {"role": "failureModes", "label": "FM", "path": str(bad_path),
             "selectedSheet": "Sheet1", "source": "desktop-bridge",
             "isResolvingSheets": False, "isAnalyzing": False,
             "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}]},
        ],
        "mappings": [],
    }

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)

        # Execute a run that will fail (files don't exist)
        ack = _send_command(process, "req_fail_run", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        terminal = _read_until(process, run_id=run_id, timeout=15.0)
        while terminal["kind"] not in {"backend_error", "result", "cancelled"}:
            terminal = _read_until(process, run_id=run_id, timeout=15.0)
        assert terminal["kind"] == "backend_error"

        # Sidecar should still respond to new commands after the failed run
        health = _send_command(process, "req_health_after_fail", "health_check", {})
        assert health["kind"] == "result"
        assert health["payload"]["status"] == "ok"
    finally:
        process.kill()


def test_sidecar_validate_rejects_missing_required_files(tmp_path: Path) -> None:
    """Regression: validation must catch missing files before execute_run."""
    body = {
        "workflowId": "piece_part_generate",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "options": {"failureModesStandard": "FMD-2016"},
        "inputs": [
            {"role": "grouping", "label": "Grouping", "path": "",
             "selectedSheet": "", "source": "desktop-bridge",
             "isResolvingSheets": False, "isAnalyzing": False,
             "resolutionError": None, "sheets": []},
        ],
        "mappings": [],
    }

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_validate_missing", "validate_run", body)
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is False
        assert result["payload"]["reason_code"] == "missing_files"
    finally:
        process.kill()


def test_sidecar_cancel_of_nonexistent_run_returns_error() -> None:
    """Regression: cancelling with a bogus run_id must return an error, not crash."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_cancel_bogus", "cancel_run", {"run_id": "run_does_not_exist"})
        assert result["kind"] == "error"
        assert "No active run" in result["payload"]["message"]
    finally:
        process.kill()


# =========================================================================
# BOM Compare integration tests
# =========================================================================


def _build_bom_compare_group_body(tmp_path: Path) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    grouping_path = tmp_path / "grouping.xlsx"
    bom_path = tmp_path / "bom.xlsx"

    pd.DataFrame([
        {"Component Group": "CPU-000", "Reference Designator": "R200, R201", "Function Description": "CPU Support"},
    ]).to_excel(grouping_path, index=False)

    pd.DataFrame([
        {"Reference Designator": "R200", "Part Number": "PN-001", "Description": "Resistor 10K"},
        {"Reference Designator": "R201", "Part Number": "PN-002", "Description": "Resistor 4.7K"},
        {"Reference Designator": "R202", "Part Number": "PN-003", "Description": "Resistor 1K"},
    ]).to_excel(bom_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "bom_compare_group",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            input_state("grouping", "Grouping workbook", grouping_path),
            input_state("bom", "BOM workbook", bom_path),
        ],
        "mappings": [
            {"canonical": "grouping_group_col", "mappedTo": "Component Group", "status": "mapped"},
            {"canonical": "grouping_refdes_col", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "bom_refdes_col", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "bom_desc_col", "mappedTo": "Description", "status": "mapped"},
        ],
        "options": {},
    }


def _build_bom_compare_custom_body(tmp_path: Path) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    bom_a_path = tmp_path / "bom_a.xlsx"
    bom_b_path = tmp_path / "bom_b.xlsx"

    pd.DataFrame([
        {"Reference Designator": "R1", "Part Number": "PN-001"},
        {"Reference Designator": "R2", "Part Number": "PN-002"},
        {"Reference Designator": "R3", "Part Number": "PN-003"},
    ]).to_excel(bom_a_path, index=False)

    pd.DataFrame([
        {"Reference Designator": "R1", "Part Number": "PN-001"},
        {"Reference Designator": "R2", "Part Number": "PN-999"},
        {"Reference Designator": "R4", "Part Number": "PN-004"},
    ]).to_excel(bom_b_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "bom_compare_custom",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            input_state("bomA", "File 1", bom_a_path),
            input_state("bomB", "File 2", bom_b_path),
        ],
        "mappings": [
            {"canonical": "refdes_col_a", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "refdes_col_b", "mappedTo": "Reference Designator", "status": "mapped"},
        ],
        "options": {"key_mode": "refdes_list"},
    }


def test_sidecar_validates_bom_compare_group(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_bcg", "validate_run", _build_bom_compare_group_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
    finally:
        process.kill()


def test_sidecar_executes_bom_compare_group(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_bcg", "execute_run", _build_bom_compare_group_body(tmp_path))
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert result["payload"]["output_file"]
        assert "BomCompare_Group" in result["payload"]["output_file"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()

        # Bug 1 regression guard: the grouping RefDes (R200, R201) both exist
        # in the BOM, so nothing should be reported missing. Before the fix,
        # an omitted dnp_regex compiled to a match-everything pattern that
        # dropped the entire BOM as DNP and flagged every grouping RefDes as
        # "missing in BOM". no_match_count is exposed on the result payload,
        # so this stays a pure-payload assertion (no workbook parsing).
        assert result["payload"]["no_match_count"] == 0
    finally:
        process.kill()


def test_sidecar_validates_bom_compare_custom(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_bcc", "validate_run", _build_bom_compare_custom_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
    finally:
        process.kill()


def test_sidecar_executes_bom_compare_custom(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_bcc", "execute_run", _build_bom_compare_custom_body(tmp_path))
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert result["payload"]["output_file"]
        assert "BomCompare_Custom" in result["payload"]["output_file"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()

        # Custom compare should detect R3 only in A, R4 only in B
        assert result["payload"]["no_match_count"] >= 2
    finally:
        process.kill()


def _build_extraction_compare_body(tmp_path: Path) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    rev_a_path = tmp_path / "extract_rev_a.xlsx"
    rev_b_path = tmp_path / "extract_rev_b.xlsx"

    pd.DataFrame([
        {"Group": "DIG-076 (Verified)", "Failure Mode Causes": "U7, U9", "Component Count": 2, "Pages": "3"},
        {"Group": "DIG-081 (Verified)", "Failure Mode Causes": "C1", "Component Count": 1, "Pages": "4"},
    ]).to_excel(rev_a_path, index=False)

    pd.DataFrame([
        {"Group": "DIG-076 (Verified)", "Failure Mode Causes": "U7", "Component Count": 1, "Pages": "3"},
        {"Group": "DIG-081 (Verified)", "Failure Mode Causes": "U9, R5", "Component Count": 2, "Pages": "4"},
    ]).to_excel(rev_b_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "extraction_compare",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            input_state("extractionA", "Extraction A (older)", rev_a_path),
            input_state("extractionB", "Extraction B (newer)", rev_b_path),
        ],
        "mappings": [],
        "options": {},
    }


def test_sidecar_validates_extraction_compare(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_exc", "validate_run", _build_extraction_compare_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
    finally:
        process.kill()


def test_sidecar_executes_extraction_compare(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_exc", "execute_run", _build_extraction_compare_body(tmp_path))
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert "ExtractionCompare" in result["payload"]["output_file"]
        assert Path(result["payload"]["output_file"]).exists()

        # R5 appeared, C1 disappeared → no_match_count = 2; U9 moved groups
        # → warning_count = 1.
        assert result["payload"]["no_match_count"] == 2
        assert result["payload"]["warning_count"] == 1
    finally:
        process.kill()


def _prefix_env(tmp_path: Path) -> dict:
    """Sidecar env with HOME isolated so prefix tests never touch the real
    ~/.refdes_extractor_config.json."""
    return {
        **SIDECAR_ENV,
        "HOME": str(tmp_path),
        "USERPROFILE": str(tmp_path),
    }


def test_sidecar_reads_and_writes_refdes_prefixes(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, env=_prefix_env(tmp_path),
    )
    try:
        _read_ready_line(process)

        # Fresh home: defaults present, no custom prefixes yet.
        result = _send_command(process, "req_px_read1", "read_refdes_prefixes", {})
        assert result["kind"] == "result"
        assert "U" in result["payload"]["defaults"]
        assert result["payload"]["custom"] == []

        # Write two custom prefixes (lowercase + duplicate + default are cleaned).
        result = _send_command(
            process, "req_px_write", "write_refdes_prefixes",
            {"prefixes": ["ps", "PS", "XU", "U"]},
        )
        assert result["kind"] == "result"
        assert result["payload"]["custom"] == ["PS", "XU"]
        assert result["payload"]["restart_required"] is True

        # Round-trip: the file persists and read returns the cleaned list.
        config_path = tmp_path / ".refdes_extractor_config.json"
        assert config_path.exists()
        assert json.loads(config_path.read_text(encoding="utf-8"))["ref_prefixes"] == ["PS", "XU"]

        result = _send_command(process, "req_px_read2", "read_refdes_prefixes", {})
        assert result["payload"]["custom"] == ["PS", "XU"]
    finally:
        process.kill()


def test_sidecar_read_refdes_prefixes_warns_on_corrupt_config(tmp_path: Path) -> None:
    # Batch 6 #4a: a CORRUPT config (invalid JSON) must not be silently
    # swallowed as "no custom prefixes". The read returns the defaults plus a
    # ``warning`` field so the UI can tell the user their saved list could not
    # be read before a save overwrites it. (An ABSENT file stays warning-free.)
    config_path = tmp_path / ".refdes_extractor_config.json"
    config_path.write_text("{ this is not valid json ", encoding="utf-8")

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, env=_prefix_env(tmp_path),
    )
    try:
        _read_ready_line(process)

        result = _send_command(process, "req_px_corrupt", "read_refdes_prefixes", {})
        assert result["kind"] == "result"
        # Defaults still render; no custom chips from the corrupt file.
        assert "U" in result["payload"]["defaults"]
        assert result["payload"]["custom"] == []
        # The warning is present and mentions the invalid-JSON cause.
        assert "warning" in result["payload"]
        assert "invalid JSON" in result["payload"]["warning"]
    finally:
        process.kill()


def test_sidecar_write_refdes_prefixes_rejects_invalid_and_preserves_keys(tmp_path: Path) -> None:
    # Pre-existing config with OTHER keys the writer must preserve.
    config_path = tmp_path / ".refdes_extractor_config.json"
    config_path.write_text(json.dumps({"other_setting": 42, "ref_prefixes": ["OLD"]}), encoding="utf-8")

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, env=_prefix_env(tmp_path),
    )
    try:
        _read_ready_line(process)

        # Digits are not a prefix — the write is rejected wholesale.
        result = _send_command(
            process, "req_px_bad", "write_refdes_prefixes", {"prefixes": ["P2S"]},
        )
        assert result["kind"] == "error"
        assert "Invalid prefix" in result["payload"]["message"]
        # Rejected write leaves the file untouched.
        assert json.loads(config_path.read_text(encoding="utf-8"))["ref_prefixes"] == ["OLD"]

        # A valid write replaces ref_prefixes but preserves other keys.
        result = _send_command(
            process, "req_px_good", "write_refdes_prefixes", {"prefixes": ["XU"]},
        )
        assert result["kind"] == "result"
        data = json.loads(config_path.read_text(encoding="utf-8"))
        assert data["ref_prefixes"] == ["XU"]
        assert data["other_setting"] == 42
    finally:
        process.kill()


def test_sidecar_write_refdes_prefixes_grandfathers_legacy_tokens(tmp_path: Path) -> None:
    # The extraction loaders accept ANY non-empty string, so a hand-edited
    # legacy prefix (e.g. 6 letters) works today. It must be grandfathered
    # past the pattern check — otherwise the editor is permanently locked
    # for that user (adversarial-review finding). NEW tokens still validate.
    config_path = tmp_path / ".refdes_extractor_config.json"
    config_path.write_text(json.dumps({"ref_prefixes": ["CONNXX"]}), encoding="utf-8")

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, env=_prefix_env(tmp_path),
    )
    try:
        _read_ready_line(process)

        # Echoing the legacy token back alongside a new valid one succeeds.
        result = _send_command(
            process, "req_px_gf", "write_refdes_prefixes",
            {"prefixes": ["CONNXX", "PS"]},
        )
        assert result["kind"] == "result"
        assert result["payload"]["custom"] == ["CONNXX", "PS"]
        data = json.loads(config_path.read_text(encoding="utf-8"))
        assert data["ref_prefixes"] == ["CONNXX", "PS"]

        # A NEW non-conforming token is still rejected wholesale.
        result = _send_command(
            process, "req_px_gf_bad", "write_refdes_prefixes",
            {"prefixes": ["CONNXX", "TOOLONGX"]},
        )
        assert result["kind"] == "error"
        assert "Invalid prefix" in result["payload"]["message"]
        assert json.loads(config_path.read_text(encoding="utf-8"))["ref_prefixes"] == [
            "CONNXX", "PS",
        ]
    finally:
        process.kill()


# =========================================================================
# Failure Rate integration tests
# =========================================================================


def _build_failure_rate_run_body(tmp_path: Path) -> dict:
    tmp_path.mkdir(parents=True, exist_ok=True)
    pred_path = tmp_path / "prediction.xlsx"
    fmea_path = tmp_path / "fmea.xlsx"

    pd.DataFrame([
        {"Reference Designator": "R1", "Failure Rate": 0.00001},
        {"Reference Designator": "C2", "Failure Rate": 0.000005},
    ]).to_excel(pred_path, index=False)

    pd.DataFrame([
        {"Failure Mode Causes": "R1", "Failure Mode Ratio": 0.6, "Part Usage": 1.0},
        {"Failure Mode Causes": "R1", "Failure Mode Ratio": 0.4, "Part Usage": 1.0},
        {"Failure Mode Causes": "C2", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
    ]).to_excel(fmea_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "failure_rate_link",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            input_state("prediction", "Prediction workbook", pred_path),
            input_state("fmea", "FMEA workbook", fmea_path),
        ],
        "mappings": [
            {"canonical": "pred_ref", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "pred_fr", "mappedTo": "Failure Rate", "status": "mapped"},
            {"canonical": "fmea_cause", "mappedTo": "Failure Mode Causes", "status": "mapped"},
            {"canonical": "fmea_ratio", "mappedTo": "Failure Mode Ratio", "status": "mapped"},
            {"canonical": "fmea_usage", "mappedTo": "Part Usage", "status": "mapped"},
        ],
        "options": {"unit_mode": "per_hour", "validate_fmr": False},
    }


def test_sidecar_validates_failure_rate_link(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_fr", "validate_run", _build_failure_rate_run_body(tmp_path))
        assert result["kind"] == "result"
        assert result["payload"]["ok"] is True
    finally:
        process.kill()


def test_sidecar_executes_failure_rate_link(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_fr", "execute_run", _build_failure_rate_run_body(tmp_path))
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        result = _read_until(process, run_id=run_id, kind="result", timeout=30.0)
        assert result["payload"]["status"] == "success"
        assert result["payload"]["row_count"] == 3
        assert result["payload"]["output_file"]
        assert "FailureRate" in result["payload"]["output_file"]

        output_path = Path(result["payload"]["output_file"])
        assert output_path.exists()
    finally:
        process.kill()


# =========================================================================
# RefDes Extractor integration tests
# =========================================================================


def _build_refdes_extract_run_body(tmp_path: Path) -> dict:
    """Build a minimal refdes_extract run body with a 1-page PDF and BOM."""
    fitz = __import__("pytest").importorskip("fitz")

    tmp_path.mkdir(parents=True, exist_ok=True)
    pdf_path = tmp_path / "schematic.pdf"
    bom_path = tmp_path / "bom.xlsx"

    # Create a minimal 1-page PDF with a FreeText annotation named "CPU-001"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.add_freetext_annot(fitz.Rect(50, 50, 300, 200), "CPU-001", fontsize=12)
    doc.save(str(pdf_path))
    doc.close()

    # Create a minimal BOM Excel
    pd.DataFrame({"Reference Designator": ["R1", "R2"]}).to_excel(bom_path, index=False)

    return {
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            {
                "role": "pdf",
                "label": "Schematic PDF",
                "path": str(pdf_path),
                "selectedSheet": "",
                "source": "desktop-bridge",
                "isResolvingSheets": False,
                "isAnalyzing": False,
                "resolutionError": None,
                "sheets": [],
            },
            {
                "role": "bom",
                "label": "BOM workbook",
                "path": str(bom_path),
                "selectedSheet": "Sheet1",
                "source": "desktop-bridge",
                "isResolvingSheets": False,
                "isAnalyzing": False,
                "resolutionError": None,
                "sheets": [{"id": "s1", "label": "Sheet1"}],
            },
        ],
        "mappings": [],
        "options": {"extraction_mode": "functional", "backend_mode": "auto"},
    }


def test_sidecar_validates_refdes_extract(tmp_path: Path) -> None:
    body = _build_refdes_extract_run_body(tmp_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_rd", "validate_run", body)
        assert result["payload"]["ok"] is True
    finally:
        process.kill()


def test_sidecar_executes_refdes_extract(tmp_path: Path) -> None:
    body = _build_refdes_extract_run_body(tmp_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_rd", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        # Wait for terminal event (result or backend_error)
        terminal = _read_until(process, run_id=run_id, timeout=30.0)
        while terminal["kind"] not in {"result", "backend_error", "cancelled"}:
            terminal = _read_until(process, run_id=run_id, timeout=30.0)

        assert terminal["kind"] == "result", (
            f"Expected 'result' but got '{terminal['kind']}': {terminal.get('payload', {}).get('message', '')}"
        )
        assert terminal["payload"]["status"] == "success"
        assert terminal["payload"]["output_file"]

        output_path = Path(terminal["payload"]["output_file"])
        assert output_path.exists()
    finally:
        process.kill()


def test_sidecar_refdes_extract_emits_bom_coverage_sheets(tmp_path: Path) -> None:
    """A BOM whose parts aren't on the schematic must surface a reverse-diff.

    End-to-end guard for the coverage wiring: the BOM (R1, R2) shares no RefDes
    with the PDF's lone CPU-001 group annotation, so both land in 'BOM Not
    Grouped' as Not Extracted and the three coverage sheets are written.
    """
    from openpyxl import load_workbook

    body = _build_refdes_extract_run_body(tmp_path)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_rd_cov", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        terminal = _read_until(process, run_id=run_id, timeout=30.0)
        while terminal["kind"] not in {"result", "backend_error", "cancelled"}:
            terminal = _read_until(process, run_id=run_id, timeout=30.0)

        assert terminal["kind"] == "result"
        output_path = Path(terminal["payload"]["output_file"])
        assert output_path.exists()

        wb = load_workbook(output_path)
        try:
            assert {"Coverage Summary", "BOM Not Grouped", "Extracted Not In BOM"} <= set(wb.sheetnames)
            bng = wb["BOM Not Grouped"]
            rows = {
                bng.cell(row=r, column=1).value: bng.cell(row=r, column=4).value
                for r in range(2, bng.max_row + 1)
            }
            assert "R1" in rows and "R2" in rows
            assert rows["R1"] == "Not Extracted"
        finally:
            wb.close()

        notes = " ".join(terminal["payload"].get("notes", []))
        assert "BOM coverage" in notes
    finally:
        process.kill()


def test_sidecar_rejects_refdes_missing_pdf(tmp_path: Path) -> None:
    __import__("pytest").importorskip("fitz")

    body = {
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            {
                "role": "pdf",
                "label": "Schematic PDF",
                "path": "",
                "selectedSheet": "",
                "source": "desktop-bridge",
                "isResolvingSheets": False,
                "isAnalyzing": False,
                "resolutionError": None,
                "sheets": [],
            },
        ],
        "mappings": [],
        "options": {"extraction_mode": "functional", "backend_mode": "auto"},
    }

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_val_rd_nopdf", "validate_run", body)
        assert result["payload"]["ok"] is False
        assert result["payload"]["reason_code"] == "missing_files"
    finally:
        process.kill()


# =========================================================================
# Cancel-latch terminal selection tests
# =========================================================================


def _run_scripted_cancel_scenario(
    tmp_path: Path,
    *,
    script: str,
    gate_log_line: str,
    extra_body: dict | None = None,
) -> tuple[list[dict], dict, dict]:
    release_gate = tmp_path / "release.gate"
    body = {
        "workflowId": "piece_part_generate",
        "releaseGate": str(release_gate),
        **(extra_body or {}),
    }
    process, reader = _start_scripted_sidecar(script)
    messages: list[dict] = []
    try:
        _write_command(process, "req_scripted_execute", "execute_run", body)
        batch, ack = _read_messages_until(
            reader,
            lambda message: message.get("request_id") == "req_scripted_execute",
        )
        messages.extend(batch)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        batch, _gate_log = _read_messages_until(
            reader,
            lambda message: (
                message.get("run_id") == run_id
                and message.get("kind") == "log"
                and message.get("payload", {}).get("line") == gate_log_line
            ),
        )
        messages.extend(batch)

        _write_command(
            process,
            "req_scripted_cancel",
            "cancel_run",
            {"run_id": run_id},
        )
        batch, cancel_response = _read_messages_until(
            reader,
            lambda message: message.get("request_id") == "req_scripted_cancel",
        )
        messages.extend(batch)
        assert cancel_response["kind"] == "result"
        assert cancel_response["payload"]["status"] == "cancelling"

        release_gate.touch()
        batch, terminal = _read_messages_until(
            reader,
            lambda message: (
                message.get("run_id") == run_id
                and message.get("kind")
                in {"result", "cancelled", "backend_error"}
            ),
        )
        messages.extend(batch)

        run_events = [
            message for message in messages if message.get("run_id") == run_id
        ]
        return run_events, terminal, cancel_response
    finally:
        if process.poll() is None:
            process.kill()
        process.wait(timeout=5)


def test_sidecar_latched_cancel_supersedes_background_error(
    tmp_path: Path,
) -> None:
    script = """
        import time
        from pathlib import Path

        import sidecar_main as sidecar

        sidecar.route_validate = lambda body: {"ok": True}

        def controlled_execute(body, **callbacks):
            callbacks["log_callback"]("TEST: background error gate reached")
            gate = Path(body["releaseGate"])
            while not gate.exists():
                time.sleep(0.005)
            raise RuntimeError("synthetic failure after accepted cancel")

        sidecar.route_execute = controlled_execute
        raise SystemExit(sidecar.main())
    """

    events, terminal, _cancel = _run_scripted_cancel_scenario(
        tmp_path,
        script=script,
        gate_log_line="TEST: background error gate reached",
    )

    assert terminal["kind"] == "cancelled"
    assert not any(event["kind"] == "backend_error" for event in events)
    status_values = [
        event["payload"]["status"]
        for event in events
        if event["kind"] == "status"
    ]
    assert status_values[-1] == "cancelled"

    diagnostics = [
        event
        for event in events
        if event["kind"] == "log"
        and event["payload"]["line"].startswith("Cancellation superseded error:")
    ]
    assert len(diagnostics) == 1
    assert diagnostics[0]["payload"]["level"] == "info"
    assert "RuntimeError" in diagnostics[0]["payload"]["line"]
    assert "synthetic failure after accepted cancel" in diagnostics[0]["payload"]["line"]
    assert any(
        event["kind"] == "log"
        and event["payload"]["level"] == "info"
        and "Traceback" in event["payload"]["line"]
        for event in events
    )
    assert events.index(diagnostics[0]) < events.index(terminal)


def test_sidecar_late_cancel_after_finalize_keeps_success_and_logs_explanation(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "committed-report.xlsx"
    script = """
        import time
        from pathlib import Path

        import sidecar_main as sidecar

        sidecar.route_validate = lambda body: {"ok": True}

        def controlled_execute(body, **callbacks):
            output = Path(body["outputPath"])
            output.write_text("committed", encoding="utf-8")
            callbacks["log_callback"]("TEST: output finalized")
            gate = Path(body["releaseGate"])
            while not gate.exists():
                time.sleep(0.005)
            return {
                "status": "success",
                "output_file": str(output),
                "mode": "desktop-bridge",
            }

        sidecar.route_execute = controlled_execute
        raise SystemExit(sidecar.main())
    """

    events, terminal, _cancel = _run_scripted_cancel_scenario(
        tmp_path,
        script=script,
        gate_log_line="TEST: output finalized",
        extra_body={"outputPath": str(output_path)},
    )

    late_cancel_line = (
        "Cancel arrived after the output was finalized; the run completed and "
        "the report was written."
    )
    assert terminal["kind"] == "result"
    assert terminal["payload"]["status"] == "success"
    assert output_path.exists()
    assert not any(
        event["kind"] in {"cancelled", "backend_error"} for event in events
    )

    late_logs = [
        event
        for event in events
        if event["kind"] == "log"
        and event["payload"]["line"] == late_cancel_line
    ]
    assert len(late_logs) == 1
    assert late_logs[0]["payload"]["level"] == "info"
    assert events.index(late_logs[0]) < events.index(terminal)


# =========================================================================
# BOM Compare and RefDes cancellation tests
#
# These regression tests guard against the _CancelBridge wiring bug where
# the bridge's CancellationToken and stop_event were independent
# threading.Events: cancelling the token would NOT set the stop_event the
# inner helpers checked. The bridges now share a single underlying Event
# (see ``backend/python/bom_compare/runtime.py::_CancelBridge`` and the
# RefDes equivalent), so cancel propagates end-to-end. The unit test in
# ``backend/tests/test_cancel_bridge.py`` covers the wiring directly; the
# tests below cover the full sidecar→ActiveRun→bridge→inner-loop path.
# =========================================================================


def _build_large_bom_compare_group_body(tmp_path: Path, *, group_rows: int = 4000) -> dict:
    """Build a BOM compare body large enough to take >1s so we can cancel it."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    grouping_path = tmp_path / "grouping.xlsx"
    bom_path = tmp_path / "bom.xlsx"

    # group_rows grouping rows, each with refdes range R{n}-R{n}, so the
    # exploder has work to do. The BOM mirrors the same refdes set so the
    # comparison has to scan and match every entry.
    grouping_df = pd.DataFrame([
        {
            "Component Group": f"GRP-{i:05d}",
            "Reference Designator": f"R{i:05d}",
            "Function Description": "Synthetic group for cancel test",
        }
        for i in range(group_rows)
    ])
    grouping_df.to_excel(grouping_path, index=False)

    bom_df = pd.DataFrame([
        {
            "Reference Designator": f"R{i:05d}",
            "Part Number": f"PN-{i:05d}",
            "Description": f"Synthetic resistor {i}",
        }
        for i in range(group_rows)
    ])
    bom_df.to_excel(bom_path, index=False)

    def input_state(role, label, path):
        return {
            "role": role, "label": label, "path": str(path), "selectedSheet": "Sheet1",
            "source": "desktop-bridge", "isResolvingSheets": False, "isAnalyzing": False,
            "resolutionError": None, "sheets": [{"id": "s1", "label": "Sheet1"}],
        }

    return {
        "workflowId": "bom_compare_group",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            input_state("grouping", "Grouping workbook", grouping_path),
            input_state("bom", "BOM workbook", bom_path),
        ],
        "mappings": [
            {"canonical": "grouping_group_col", "mappedTo": "Component Group", "status": "mapped"},
            {"canonical": "grouping_refdes_col", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "bom_refdes_col", "mappedTo": "Reference Designator", "status": "mapped"},
            {"canonical": "bom_desc_col", "mappedTo": "Description", "status": "mapped"},
        ],
        "options": {},
    }


def test_sidecar_cancels_active_bom_compare_run(tmp_path: Path) -> None:
    """Cancellation must reach the BOM Compare ``stop_event``.

    Regression for the _CancelBridge wiring bug. Before the fix, the
    sidecar would emit a ``cancelling`` status but the BOM compare engine
    would continue to a ``success`` terminal because its ``stop_event``
    was a separate threading.Event from the CancellationToken's event.
    """
    body = _build_large_bom_compare_group_body(tmp_path / "bom_cancel", group_rows=4000)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_bcg_cancel", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        # Send cancel as soon as we have the ack. With 4000 rows the run
        # should still be in progress.
        cancel = _send_command(process, "req_cancel_bcg", "cancel_run", {"run_id": run_id})
        assert cancel["kind"] == "result"
        assert cancel["payload"]["status"] == "cancelling"

        # Drain run events until we hit a terminal envelope.
        terminal = _read_until(process, run_id=run_id, timeout=30.0)
        while terminal["kind"] not in {"cancelled", "result", "backend_error"}:
            terminal = _read_until(process, run_id=run_id, timeout=30.0)

        assert terminal["kind"] == "cancelled", (
            f"BOM Compare run reached terminal '{terminal['kind']}' instead of "
            f"'cancelled' after a cancel_run was accepted. This is the exact "
            f"failure mode of the _CancelBridge wiring bug."
        )
        assert "cancel" in terminal["payload"]["message"].lower()
    finally:
        process.kill()


def _build_multi_annot_refdes_body(tmp_path: Path, *, annot_count: int = 200) -> dict:
    """Build a refdes_extract body with many annotations to give time to cancel."""
    fitz = __import__("pytest").importorskip("fitz")

    tmp_path.mkdir(parents=True, exist_ok=True)
    pdf_path = tmp_path / "schematic.pdf"
    bom_path = tmp_path / "bom.xlsx"

    # Multi-page PDF with many FreeText annotations across pages so the
    # extraction loop has to traverse a lot of geometry before finishing.
    doc = fitz.open()
    pages_needed = max(1, annot_count // 20)
    for page_idx in range(pages_needed):
        page = doc.new_page(width=612, height=792)
        per_page = annot_count // pages_needed
        for i in range(per_page):
            row = i % 10
            col = i // 10
            x0 = 50 + col * 60
            y0 = 50 + row * 30
            page.add_freetext_annot(
                fitz.Rect(x0, y0, x0 + 55, y0 + 25),
                f"R{page_idx * 100 + i:04d}",
                fontsize=10,
            )
    doc.save(str(pdf_path))
    doc.close()

    bom_df = pd.DataFrame({"Reference Designator": [f"R{i:04d}" for i in range(annot_count)]})
    bom_df.to_excel(bom_path, index=False)

    return {
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "enrichments": {"functional": False, "piecePart": False},
        "inputs": [
            {
                "role": "pdf",
                "label": "Schematic PDF",
                "path": str(pdf_path),
                "selectedSheet": "",
                "source": "desktop-bridge",
                "isResolvingSheets": False,
                "isAnalyzing": False,
                "resolutionError": None,
                "sheets": [],
            },
            {
                "role": "bom",
                "label": "BOM workbook",
                "path": str(bom_path),
                "selectedSheet": "Sheet1",
                "source": "desktop-bridge",
                "isResolvingSheets": False,
                "isAnalyzing": False,
                "resolutionError": None,
                "sheets": [{"id": "s1", "label": "Sheet1"}],
            },
        ],
        "mappings": [],
        "options": {"extraction_mode": "functional", "backend_mode": "auto"},
    }


def test_sidecar_cancels_active_refdes_run(tmp_path: Path) -> None:
    """Cancellation must reach the RefDes extraction ``stop_event``.

    Regression for the _CancelBridge wiring bug. The RefDes runtime uses
    the same dual-event bridge pattern as BOM Compare and was previously
    only partially honoring cancellation (the explicit ``bridge.cancel.check()``
    after annotation extraction would fire, but the ``stop_event``-driven
    inner helpers would not).
    """
    body = _build_multi_annot_refdes_body(tmp_path / "refdes_cancel", annot_count=400)

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        ack = _send_command(process, "req_exec_rd_cancel", "execute_run", body)
        assert ack["kind"] == "ack"
        run_id = ack["payload"]["run_id"]

        cancel = _send_command(process, "req_cancel_rd", "cancel_run", {"run_id": run_id})
        assert cancel["kind"] == "result"
        assert cancel["payload"]["status"] == "cancelling"

        terminal = _read_until(process, run_id=run_id, timeout=60.0)
        while terminal["kind"] not in {"cancelled", "result", "backend_error"}:
            terminal = _read_until(process, run_id=run_id, timeout=60.0)

        assert terminal["kind"] == "cancelled", (
            f"RefDes run reached terminal '{terminal['kind']}' instead of "
            f"'cancelled' after a cancel_run was accepted. This is the exact "
            f"failure mode of the _CancelBridge wiring bug."
        )
        assert "cancel" in terminal["payload"]["message"].lower()
    finally:
        process.kill()


# =========================================================================
# Sidecar exception envelope tests
#
# These regression tests cover the failure modes where ordinary user
# errors (moved/corrupt files, malformed payloads) used to crash the
# entire sidecar process. Every command branch must now return an
# ``error`` envelope on exception and remain alive for follow-up traffic.
# =========================================================================


def test_sidecar_list_sheets_returns_error_for_missing_file(tmp_path: Path) -> None:
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        missing_path = str(tmp_path / "does_not_exist.xlsx")
        result = _send_command(process, "req_ls_missing", "list_sheets", {"path": missing_path})
        assert result["kind"] == "error"
        assert result["payload"]["message"]

        # Sidecar must still be alive and responsive after the error.
        health = _send_command(process, "req_health_after_ls", "health_check", {})
        assert health["kind"] == "result"
        assert health["payload"]["status"] == "ok"
    finally:
        process.kill()


def test_sidecar_list_sheets_returns_error_for_corrupt_file(tmp_path: Path) -> None:
    corrupt_path = tmp_path / "garbage.xlsx"
    corrupt_path.write_bytes(b"this is not a real xlsx file at all")

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(process, "req_ls_corrupt", "list_sheets", {"path": str(corrupt_path)})
        assert result["kind"] == "error"
        assert result["payload"]["message"]

        # Sidecar must still be alive after the error.
        health = _send_command(process, "req_health_after_corrupt", "health_check", {})
        assert health["kind"] == "result"
    finally:
        process.kill()


def test_sidecar_execute_run_returns_error_when_validate_raises() -> None:
    """If route_validate raises (e.g. malformed body), execute_run must
    emit an ``error`` envelope before any ack and the sidecar must stay alive."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        # A body with no inputs/mappings/options will likely raise inside
        # one of the runtime validators rather than return a clean ok=False.
        body = {"workflowId": "piece_part_generate"}
        result = _send_command(process, "req_exec_bad_body", "execute_run", body)
        # The result is either a normal validation rejection (kind=error,
        # validation rejection path) or an exception envelope (kind=error,
        # exception path). Either way it must be an error and the sidecar
        # must remain alive.
        assert result["kind"] == "error"
        assert result["payload"]["message"]

        health = _send_command(process, "req_health_after_exec", "health_check", {})
        assert health["kind"] == "result"
        assert health["payload"]["status"] == "ok"
    finally:
        process.kill()


def test_sidecar_survives_malformed_envelope() -> None:
    """A malformed JSON line on stdin must not kill the sidecar."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        # Send a malformed line directly (not via _send_command which
        # builds a proper envelope).
        process.stdin.write("this is not json at all\n")
        process.stdin.flush()
        # Drain the error envelope from the malformed line.
        # Use the shared reader and look for any error envelope.
        reader = _get_reader(process)
        # Wait briefly for the error envelope.
        import time as _time
        deadline = _time.monotonic() + 5.0
        saw_error = False
        while _time.monotonic() < deadline and not saw_error:
            try:
                line = reader._queue.get(timeout=0.5).strip()
            except Exception:
                continue
            if not line:
                continue
            msg = json.loads(line)
            if msg.get("kind") == "error" and "Malformed" in msg["payload"].get("message", ""):
                saw_error = True
                break
        assert saw_error, "Expected an 'error' envelope for the malformed line"

        # Sidecar must still respond to a normal command afterward.
        health = _send_command(process, "req_health_after_malformed", "health_check", {})
        assert health["kind"] == "result"
        assert health["payload"]["status"] == "ok"
    finally:
        process.kill()


# ---------------------------------------------------------------------------
# Phase 4a: output_preview on validate_run
# ---------------------------------------------------------------------------

def test_sidecar_validate_attaches_output_preview_for_fmea(tmp_path: Path) -> None:
    """On successful FMEA validation, validate_run includes an output_preview
    sampling RefDes / Part Number / Description from the BOM."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process, "req_val_preview_fmea", "validate_run",
            _build_phase4_run_body(tmp_path, group_rows=3),
        )
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is True
        preview = payload.get("output_preview")
        assert preview is not None, "Expected output_preview on successful validation"
        assert set(preview["columns"]) >= {"RefDes", "Part Number", "Description"}
        assert len(preview["rows"]) == 3
        assert preview["truncated"] is False
        assert preview["total_estimated"] == 3
        # First row matches the fixture's R200 / PN-0000 / Resistor shape.
        ref_idx = preview["columns"].index("RefDes")
        assert preview["rows"][0][ref_idx] == "R200"
    finally:
        process.kill()


def test_sidecar_validate_output_preview_truncated_when_over_cap(
    tmp_path: Path,
) -> None:
    """When the source has more rows than PREVIEW_ROW_CAP (20), truncated=True
    and rows is capped, but total_estimated reports the full count."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        body = _build_phase4_run_body(tmp_path, group_rows=25)
        result = _send_command(process, "req_val_preview_truncate", "validate_run", body)
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is True
        preview = payload.get("output_preview")
        assert preview is not None
        assert len(preview["rows"]) == 20
        assert preview["truncated"] is True
        assert preview["total_estimated"] == 25
    finally:
        process.kill()


def test_sidecar_validate_omits_preview_when_validation_fails(
    tmp_path: Path,
) -> None:
    """A failed validation must NOT carry an output_preview — the preview
    is only a garnish on successful runs."""
    body = _build_phase4_run_body(tmp_path)
    body["outputStrategyId"] = "existing_workbook_best_effort"  # unsupported

    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process, "req_val_preview_fail", "validate_run", body,
        )
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is False
        assert "output_preview" not in payload
    finally:
        process.kill()


def test_sidecar_validate_attaches_output_preview_for_bom_compare(
    tmp_path: Path,
) -> None:
    """BOM Compare attaches a preview from the BOM workbook."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process, "req_val_preview_bc", "validate_run",
            _build_bom_compare_group_body(tmp_path),
        )
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is True
        preview = payload.get("output_preview")
        assert preview is not None
        assert "RefDes" in preview["columns"]
        assert len(preview["rows"]) >= 1
    finally:
        process.kill()


def test_sidecar_validate_attaches_output_preview_for_failure_rate(
    tmp_path: Path,
) -> None:
    """Failure Rate attaches a preview from the prediction workbook."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process, "req_val_preview_fr", "validate_run",
            _build_failure_rate_run_body(tmp_path),
        )
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is True
        preview = payload.get("output_preview")
        # The fixture's prediction workbook uses "Reference Designator" and
        # "Failure Rate" headers, both of which match the synonym lists in
        # failure_rate.runtime._build_failure_rate_output_preview. The preview
        # is therefore NOT best-effort-omitted here — it must be present and
        # well-formed, with the mapped display columns and both source rows.
        assert preview is not None, "Failure Rate preview should be attached for this fixture"
        assert isinstance(preview["columns"], list)
        assert "RefDes" in preview["columns"]
        assert "Failure Rate" in preview["columns"]
        assert isinstance(preview["rows"], list)
        # Prediction fixture has exactly two rows (R1, C2); none truncated.
        assert len(preview["rows"]) == 2
        assert isinstance(preview["truncated"], bool)
        assert preview["truncated"] is False
    finally:
        process.kill()


def test_sidecar_validate_attaches_output_preview_for_refdes(
    tmp_path: Path,
) -> None:
    """RefDes Extractor attaches a preview from the optional BOM workbook."""
    process = subprocess.Popen(
        [sys.executable, str(SIDECAR)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        env=SIDECAR_ENV,
    )
    try:
        _read_ready_line(process)
        result = _send_command(
            process,
            "req_val_preview_rd",
            "validate_run",
            _build_refdes_extract_run_body(tmp_path),
        )
        assert result["kind"] == "result"
        payload = result["payload"]
        assert payload["ok"] is True
        preview = payload.get("output_preview")
        assert preview is not None
        assert "RefDes" in preview["columns"]
        assert len(preview["rows"]) == 2
        assert preview["truncated"] is False
    finally:
        process.kill()
