"""Tests for BOM Compare core comparison logic.

Two concerns are covered:

1. RefDes range expansion (``expand_refdes_range`` / ``split_refdes_list``).
   By project convention hyphens denote PINS (U200-1), never ranges, so the
   group-vs-BOM ``analyze`` orchestrator intentionally does NOT expand ranges
   (see ``test_analyze_does_not_expand_ranges_by_default``). These tests pin the
   *opt-in* expansion helper directly: "R200-R205" -> R200..R205, while
   instance/pin notation (U3-7, J1-4) is preserved.

2. The ``analyze`` orchestrator's set math on a tiny discrete fixture:
   exact matches produce no missing/extra rows, a grouping-only RefDes shows
   up as "Missing in BOM", and a BOM-only RefDes shows up as "BOM Not in
   Groups".
"""

from __future__ import annotations

import re

import pandas as pd
import pytest

from common.exceptions import ValidationError
from common.refdes_utils import expand_refdes_range, split_refdes_list
from bom_compare.bom_compare_logic import ColumnMapping, AnalyzeOptions, compare_two_boms
from bom_compare.group_analysis import analyze, explode_bom, explode_grouping


# ---------------------------------------------------------------------------
# RefDes range expansion (the opt-in helper)
# ---------------------------------------------------------------------------

def test_expand_refdes_range_expands_passive_range() -> None:
    """'R200-R205' expands to the six individual resistors (expand=True)."""
    assert expand_refdes_range("R200-R205", expand=True) == [
        "R200",
        "R201",
        "R202",
        "R203",
        "R204",
        "R205",
    ]


def test_expand_refdes_range_default_does_not_expand() -> None:
    """Default (expand=False) keeps the range token intact — canonicalized."""
    assert expand_refdes_range("R200-R205") == ["R200-R205"]
    assert expand_refdes_range("r200-r205") == ["R200-R205"]  # uppercased


def test_expand_refdes_range_preserves_instance_and_pin_notation() -> None:
    """Instance (U3-7) and connector-pin (J1-4) notation is NOT expanded."""
    assert expand_refdes_range("U3-7", expand=True) == ["U3-7"]
    assert expand_refdes_range("J1-4", expand=True) == ["J1-4"]


def test_expand_refdes_range_preserves_zero_padding() -> None:
    """Leading-zero width of the start token is preserved across the range."""
    assert expand_refdes_range("R08-R10", expand=True) == ["R08", "R09", "R10"]


def test_expand_refdes_range_reversed_bounds_not_expanded() -> None:
    """A descending range (R5-R1) is invalid and returned as-is."""
    assert expand_refdes_range("R5-R1", expand=True) == ["R5-R1"]


def test_split_refdes_list_expands_only_when_opted_in() -> None:
    """split_refdes_list passes through expand_ranges to per-token expansion."""
    # Default: range stays whole, comma still splits the list.
    assert split_refdes_list("R200-R205, C1") == ["R200-R205", "C1"]
    # Opt-in: the range explodes, the standalone token survives.
    assert split_refdes_list("R200-R205, C1", expand_ranges=True) == [
        "R200",
        "R201",
        "R202",
        "R203",
        "R204",
        "R205",
        "C1",
    ]


# ---------------------------------------------------------------------------
# analyze() orchestrator set math on a discrete fixture
# ---------------------------------------------------------------------------

def _run_analyze(grouping_refs, bom_refs):
    """Run analyze() with minimal grouping + BOM frames and no extra checks."""
    group_df = pd.DataFrame(
        {"Group": [f"G{i}" for i in range(len(grouping_refs))], "RefDes": grouping_refs}
    )
    bom_df = pd.DataFrame(
        {"RefDes": bom_refs, "Description": ["" for _ in bom_refs]}
    )
    mapping = ColumnMapping(
        grouping_group_col="Group",
        grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes",
        bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False,
        run_duplicate_checks=False,
        check_part_usage=False,
        check_fmr=False,
    )
    return analyze(
        group_df,
        bom_df,
        mapping,
        options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )


def test_analyze_file_labels_flow_into_summary_duplicates_and_part_usage() -> None:
    """User-typed display names (2026-07-16 request) label every group-mode
    file-identity site: the Summary file rows carry "<name> — <basename>",
    the Duplicates Source column and its BOM description carry the names,
    and the Part Usage Source column swaps its constant "BOM" for the
    user's BOM name."""
    group_df = pd.DataFrame(
        {
            # R1 appears in two groups -> a Grouping-side duplicate row.
            "Group": ["G1", "G2"],
            "RefDes": ["R1", "R1"],
        }
    )
    bom_df = pd.DataFrame(
        {
            # R1 twice -> a BOM-side duplicate row; R2 usage mismatch (usage
            # 2 with a single instance) -> a Part Usage warning row.
            "RefDes": ["R1", "R1", "R2"],
            "Description": ["", "", ""],
            "Part Usage": ["1", "1", "2"],
        }
    )
    mapping = ColumnMapping(
        grouping_group_col="Group",
        grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes",
        bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False,
        run_duplicate_checks=True,
        check_part_usage=True,
        check_fmr=False,
    )

    results = analyze(
        group_df,
        bom_df,
        mapping,
        options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
        file_labels=("CPU Grouping File", "New Digital BOM"),
    )

    summary = dict(zip(results.summary["Item"], results.summary["Value"]))
    assert summary["Grouping File"] == "CPU Grouping File — grouping.xlsx"
    assert summary["BOM File"] == "New Digital BOM — bom.xlsx"

    sources = set(results.duplicates["Source"])
    assert sources == {"CPU Grouping File", "New Digital BOM"}
    bom_dup = results.duplicates[results.duplicates["Source"] == "New Digital BOM"].iloc[0]
    assert bom_dup["Description"] == "Appears 2 times in New Digital BOM"

    assert not results.part_usage_warnings.empty
    assert set(results.part_usage_warnings["Source"]) == {"New Digital BOM"}


def test_analyze_without_file_labels_keeps_generic_sources() -> None:
    """No display names -> byte-identical legacy labels: basenames in the
    Summary, "Grouping"/"BOM" sources, "Appears N times in BOM"."""
    group_df = pd.DataFrame({"Group": ["G1", "G2"], "RefDes": ["R1", "R1"]})
    bom_df = pd.DataFrame({"RefDes": ["R1", "R1"], "Description": ["", ""]})
    mapping = ColumnMapping(
        grouping_group_col="Group",
        grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes",
        bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False,
        run_duplicate_checks=True,
        check_part_usage=False,
        check_fmr=False,
    )

    results = analyze(
        group_df, bom_df, mapping, options, file_paths=("grouping.xlsx", "bom.xlsx")
    )

    summary = dict(zip(results.summary["Item"], results.summary["Value"]))
    assert summary["Grouping File"] == "grouping.xlsx"
    assert summary["BOM File"] == "bom.xlsx"
    assert set(results.duplicates["Source"]) == {"Grouping", "BOM"}
    bom_desc = results.duplicates[results.duplicates["Source"] == "BOM"].iloc[0]["Description"]
    assert bom_desc == "Appears 2 times in BOM"


def test_analyze_exact_overlap_has_no_missing_or_extra() -> None:
    """Identical RefDes sets produce zero missing and zero extra rows."""
    results = _run_analyze(["R1", "R2", "C3"], ["R1", "R2", "C3"])

    assert results.missing_in_bom.empty
    assert results.bom_not_in_groups.empty


def test_analyze_flags_grouping_only_refdes_as_missing() -> None:
    """A RefDes present only in grouping is reported 'Missing in BOM'."""
    results = _run_analyze(["R1", "R2", "R3"], ["R1", "R2"])

    missing_tokens = set(results.missing_in_bom["Token"])
    assert missing_tokens == {"R3"}
    # And it is NOT in the extra list.
    assert results.bom_not_in_groups.empty


def test_analyze_flags_bom_only_refdes_as_extra() -> None:
    """A RefDes present only in the BOM is reported 'BOM Not in Groups'."""
    results = _run_analyze(["R1"], ["R1", "C9"])

    extra_bases = set(results.bom_not_in_groups["Base"])
    assert extra_bases == {"C9"}
    assert results.missing_in_bom.empty


def test_analyze_reports_both_directions_simultaneously() -> None:
    """Disjoint extras on each side are reported independently and correctly."""
    results = _run_analyze(["R1", "R2", "Q5"], ["R1", "R2", "U8"])

    assert set(results.missing_in_bom["Token"]) == {"Q5"}
    assert set(results.bom_not_in_groups["Base"]) == {"U8"}

    # Summary counters mirror the table sizes.
    summary = dict(zip(results.summary["Item"], results.summary["Value"]))
    assert summary["Missing in BOM"] == 1
    assert summary["BOM Not in Groups"] == 1


def test_analyze_does_not_expand_ranges_by_default() -> None:
    """INTENDED behavior: the group analyzer never range-expands hyphenated tokens.

    Per the project's RefDes convention, hyphens denote PINS (U200-1 = pin 1 of
    component U200), never numeric ranges — authors do not write 'R1-R3' to mean
    R1/R2/R3 (an 'R1-R3' token would be its own component). So ``analyze`` calls
    ``split_refdes_list`` WITHOUT ``expand_ranges`` and reduces each token to its
    base RefDes. This pins the no-expansion contract: a synthetic 'R1-R3' token
    reduces to base 'R1' and is NOT exploded into R1/R2/R3. Range expansion stays
    opt-in (and unused by this tool).
    """
    results = _run_analyze(["R1-R3"], ["R1", "R2", "R3"])

    # 'R1-R3' is treated as a single token (base 'R1'), not expanded.
    assert results.missing_in_bom.empty
    assert set(results.bom_not_in_groups["Base"]) == {"R2", "R3"}


def test_analyze_matches_hyphenated_pin_to_base_component() -> None:
    """A hyphenated PIN designator reduces to its base component for matching.

    Real-world case (hyphens are pins): a grouping cell 'U200-1' (pin 1 of U200)
    is treated as component U200, so it matches a BOM that lists 'U200'. No
    missing, no extra.
    """
    results = _run_analyze(["U200-1"], ["U200"])

    assert results.missing_in_bom.empty
    assert results.bom_not_in_groups.empty


# ---------------------------------------------------------------------------
# Phase 1 regression: the group default run is byte-stable with loose OFF.
#
# The frontend "base match" checkbox now defaults to FALSE and maps to
# loose_base_match. A default run (loose off) must NOT fuzzy-prefix-match
# bases: a grouping base 'U20' is NOT covered by a BOM base 'U200'.
# ---------------------------------------------------------------------------

def test_analyze_default_loose_off_no_prefix_match() -> None:
    """Default (loose_base_match=False): a prefix-only base is NOT covered.

    Grouping has U20; BOM has U200. With loose OFF these are distinct bases,
    so U20 is missing-in-BOM and U200 is extra. This pins the behavior a
    default desktop run must preserve (loose matching is opt-in).
    """
    results = _run_analyze(["U20"], ["U200"])

    assert set(results.missing_in_bom["Token"]) == {"U20"}
    assert set(results.bom_not_in_groups["Base"]) == {"U200"}


def test_analyze_loose_on_prefix_base_is_covered() -> None:
    """loose_base_match=True: a BOM base that starts with the grouping base
    (and whose remainder is not all digits) covers it.

    Grouping base 'U20A' is covered by BOM base 'U20AB' (remainder 'B' is not
    all-digits), so nothing is missing and nothing is extra.
    """
    group_df = pd.DataFrame({"Group": ["G0"], "RefDes": ["U20A"]})
    bom_df = pd.DataFrame({"RefDes": ["U20AB"], "Description": [""]})
    mapping = ColumnMapping(
        grouping_group_col="Group",
        grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes",
        bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False,
        run_duplicate_checks=False,
        check_part_usage=False,
        check_fmr=False,
        loose_base_match=True,
    )
    results = analyze(
        group_df, bom_df, mapping, options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )

    assert results.missing_in_bom.empty
    assert results.bom_not_in_groups.empty


# ---------------------------------------------------------------------------
# Phase 2: custom (BOM-vs-BOM) compare must honor exact_match, ignore_dnp,
# loose_base_match, and check_fmr with the SAME user-facing semantics as the
# group path.
# ---------------------------------------------------------------------------

def _custom_df(rows):
    return pd.DataFrame(rows)


def test_custom_base_match_default_reduces_pins_to_base() -> None:
    """Default custom compare (exact_match=False) reduces tokens to base
    RefDes, mirroring the group path. 'U200-1' (pin 1) in File 1 matches
    'U200' in File 2: nothing only-in-A, nothing only-in-B.
    """
    bom_a = _custom_df([{"RefDes": "U200-1"}])
    bom_b = _custom_df([{"RefDes": "U200"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
    )

    assert result.only_in_a == []
    assert result.only_in_b == []
    assert result.in_both_count == 1


def test_custom_exact_match_keeps_full_token() -> None:
    """exact_match=True compares full canonical tokens: 'U200-1' does NOT
    match 'U200', so each side reports a unique entry.
    """
    bom_a = _custom_df([{"RefDes": "U200-1"}])
    bom_b = _custom_df([{"RefDes": "U200"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        exact_match=True,
        check_part_usage=False,
    )

    assert [r["RefDes"] for r in result.only_in_a] == ["U200-1"]
    assert [r["RefDes"] for r in result.only_in_b] == ["U200"]
    assert result.in_both_count == 0


def test_custom_ignore_dnp_skips_dnp_rows() -> None:
    """ignore_dnp=True drops rows whose RefDes/Description marks them DNP.

    File 1 lists C1 (live) and C2 (marked DNP in its description). With DNP
    filtering on, C2 is removed before comparison, so File 2 (only C1) has
    nothing extra and File 1 has nothing only-in-A.
    """
    bom_a = _custom_df([
        {"RefDes": "C1", "Description": "Cap 10uF"},
        {"RefDes": "C2", "Description": "Cap DNP"},
    ])
    bom_b = _custom_df([{"RefDes": "C1", "Description": "Cap 10uF"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        ignore_dnp=True, desc_col_a="Description", desc_col_b="Description",
        check_part_usage=False,
    )

    assert result.only_in_a == []
    assert result.only_in_b == []
    assert result.in_both_count == 1


def test_custom_ignore_dnp_false_keeps_dnp_rows() -> None:
    """ignore_dnp=False keeps the DNP-marked row, so C2 is reported only-in-A."""
    bom_a = _custom_df([
        {"RefDes": "C1", "Description": "Cap 10uF"},
        {"RefDes": "C2", "Description": "Cap DNP"},
    ])
    bom_b = _custom_df([{"RefDes": "C1", "Description": "Cap 10uF"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        ignore_dnp=False, desc_col_a="Description", desc_col_b="Description",
        check_part_usage=False,
    )

    assert [r["RefDes"] for r in result.only_in_a] == ["C2"]


def test_custom_loose_base_match_prefix_covered() -> None:
    """loose_base_match=True suppresses an only-in-A base whose prefix matches
    an only-in-B base (and vice-versa), mirroring the group path.
    """
    bom_a = _custom_df([{"RefDes": "U20A"}])
    bom_b = _custom_df([{"RefDes": "U20AB"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        loose_base_match=True,
        check_part_usage=False,
    )

    assert result.only_in_a == []
    assert result.only_in_b == []


def test_custom_check_fmr_flags_bad_sum() -> None:
    """check_fmr=True validates per-RefDes FMR sums on each file. A RefDes
    whose failure-mode ratios do not sum to 1.0 is reported in fmr_warnings.
    """
    bom_a = _custom_df([
        {"RefDes": "U1", "Ratio": 0.5},
        {"RefDes": "U1", "Ratio": 0.2},  # sums to 0.7 -> bad
        {"RefDes": "U2", "Ratio": 1.0},  # ok
    ])
    bom_b = _custom_df([{"RefDes": "U1", "Ratio": 1.0}, {"RefDes": "U2", "Ratio": 1.0}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_fmr=True,
        check_part_usage=False,
    )

    bad = {w["RefDes"] for w in result.fmr_warnings}
    assert "U1" in bad
    assert "U2" not in bad


def test_custom_check_fmr_off_by_default() -> None:
    """When check_fmr is not requested, fmr_warnings stays empty even if a
    Ratio column exists with non-summing ratios.
    """
    bom_a = _custom_df([{"RefDes": "U1", "Ratio": 0.5}])
    bom_b = _custom_df([{"RefDes": "U1", "Ratio": 0.5}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
    )

    assert result.fmr_warnings == []


# ---------------------------------------------------------------------------
# Tier-1 fix #9: loose_base_match residual-digit guard. 'R1' must NOT be
# treated as covering 'R12' (residual '2' is a digit = a different component).
# Three of the four loose-match directions previously lacked this guard.
# ---------------------------------------------------------------------------

def test_analyze_loose_residual_digit_guard_extra_direction() -> None:
    """Group path: with loose on, grouping base 'R1' must not swallow BOM base
    'R12'. R12 stays flagged as extra; R1 stays flagged as missing.
    """
    group_df = pd.DataFrame({"Group": ["G0"], "RefDes": ["R1"]})
    bom_df = pd.DataFrame({"RefDes": ["R12"], "Description": [""]})
    mapping = ColumnMapping(
        grouping_group_col="Group", grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes", bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False, run_duplicate_checks=False,
        check_part_usage=False, check_fmr=False, loose_base_match=True,
    )
    results = analyze(
        group_df, bom_df, mapping, options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    assert set(results.bom_not_in_groups["Base"]) == {"R12"}
    assert set(results.missing_in_bom["Token"]) == {"R1"}


def test_custom_loose_residual_digit_guard() -> None:
    """Custom path: 'R1' and 'R12' are different components; loose match must
    NOT suppress either (residual '2' is a digit continuation).
    """
    bom_a = _custom_df([{"RefDes": "R1"}])
    bom_b = _custom_df([{"RefDes": "R12"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        loose_base_match=True, check_part_usage=False,
    )

    assert [r["RefDes"] for r in result.only_in_a] == ["R1"]
    assert [r["RefDes"] for r in result.only_in_b] == ["R12"]


def test_analyze_loose_nondigit_residual_is_covered() -> None:
    """Positive side of the residual-digit guard: a NON-digit residual is still
    covered under loose match — grouping base 'CPU' is covered by BOM base 'CPUA'
    (residual 'A'), unlike the digit-residual R1/R12 case which stays flagged.
    """
    group_df = pd.DataFrame({"Group": ["G"], "RefDes": ["CPU"]})
    bom_df = pd.DataFrame({"RefDes": ["CPUA"], "Description": [""]})
    mapping = ColumnMapping(
        grouping_group_col="Group", grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes", bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False, run_duplicate_checks=False,
        check_part_usage=False, check_fmr=False, loose_base_match=True,
    )
    results = analyze(
        group_df, bom_df, mapping, options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    assert results.missing_in_bom.empty


# ---------------------------------------------------------------------------
# Tier-1 fix #6: a NaN/non-numeric ratio cell must be FLAGGED, not silently
# summed. Old behavior: float(NaN) does not raise, NaN poisons the sum, and
# abs(NaN - 1.0) > tol is False, so the inconsistent RefDes silently passed.
# ---------------------------------------------------------------------------

def test_custom_check_fmr_text_ratio_flagged_blank_skipped() -> None:
    """Custom path: a blank (NaN) continuation row whose RefDes still sums to 1.0
    is NOT flagged (no false positive); a non-numeric TEXT ratio IS flagged. The
    underlying #6 fix (a NaN no longer poisons the sum into a silent pass) still
    holds because the blank is excluded from the sum, not summed as NaN.
    """
    bom_a = _custom_df([
        {"RefDes": "U1", "Ratio": float("nan")},  # blank continuation row
        {"RefDes": "U1", "Ratio": 0.6},
        {"RefDes": "U1", "Ratio": 0.4},           # U1 real rows sum to 1.0
        {"RefDes": "U2", "Ratio": "bad"},         # non-numeric text -> data error
        {"RefDes": "U2", "Ratio": 0.5},
    ])
    bom_b = _custom_df([{"RefDes": "U9", "Ratio": 1.0}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_fmr=True, check_part_usage=False,
    )

    statuses = {w["RefDes"]: w["Status"] for w in result.fmr_warnings}
    assert "U1" not in statuses                    # blank skipped, sums to 1.0
    assert "U2" in statuses
    assert "non-numeric" in statuses["U2"].lower()  # text flagged


def test_analyze_check_fmr_text_ratio_flagged_blank_skipped_group_path() -> None:
    """Group path mirrors the custom path: blank ratio skipped, text flagged."""
    group_df = pd.DataFrame({
        "Group": ["G", "G", "G", "G", "G"],
        "RefDes": ["U1", "U1", "U1", "U2", "U2"],
        "Ratio": [float("nan"), 0.6, 0.4, "bad", 0.5],
    })
    bom_df = pd.DataFrame({"RefDes": ["U1", "U2"], "Description": ["", ""]})
    mapping = ColumnMapping(
        grouping_group_col="Group", grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes", bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False, run_duplicate_checks=False,
        check_part_usage=False, check_fmr=True,
    )
    results = analyze(
        group_df, bom_df, mapping, options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    flagged = dict(zip(results.fmr_warnings["RefDes"], results.fmr_warnings["Status"]))
    assert "U1" not in flagged                      # blank skipped, sums to 1.0
    assert "U2" in flagged
    assert "non-numeric" in flagged["U2"].lower()


def test_analyze_check_fmr_canonicalizes_refdes_key_group_path() -> None:
    """Tier-4 pd-checkfmr-key: a PDF-pasted RefDes carrying an invisible char
    (zero-width space) must group with its clean twin, not split the FMR sum
    into a false-positive 'FMR != 1.0'."""
    group_df = pd.DataFrame({
        "Group": ["G", "G"],
        "RefDes": ["U1", "U1​"],   # second cell carries a zero-width space
        "Ratio": [0.5, 0.5],
    })
    bom_df = pd.DataFrame({"RefDes": ["U1"], "Description": [""]})
    mapping = ColumnMapping(
        grouping_group_col="Group", grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes", bom_desc_col="Description",
    )
    options = AnalyzeOptions(
        run_warning_checks=False, run_duplicate_checks=False,
        check_part_usage=False, check_fmr=True,
    )
    results = analyze(
        group_df, bom_df, mapping, options,
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    # Both rows are the same physical U1 (0.5 + 0.5 = 1.0) -> no FMR warning.
    assert list(results.fmr_warnings["RefDes"]) == []


# ---------------------------------------------------------------------------
# Tier-1 #5: custom (BOM-vs-BOM) per-column VALUE diff. With ``compare_columns``
# supplied, a RefDes present in BOTH files whose mapped column values differ is
# reported in ``result.differences``. Without ``compare_columns`` the custom
# compare reports RefDes membership only and emits NO value diffs. These pin the
# backend contract the Tier-1 #5 wiring depends on.
# ---------------------------------------------------------------------------

def test_compare_two_boms_flags_a_changed_column_value() -> None:
    """A RefDes in both files whose Part Number differs is reported in
    ``differences`` with the column label and both values.
    """
    bom_a = _custom_df([
        {"RefDes": "R1", "Part Number": "PN-10K"},
        {"RefDes": "C2", "Part Number": "PN-1UF"},
    ])
    bom_b = _custom_df([
        {"RefDes": "R1", "Part Number": "PN-4K7"},
        {"RefDes": "C2", "Part Number": "PN-1UF"},
    ])

    result = compare_two_boms(
        bom_a, bom_b,
        refdes_col_a="RefDes", refdes_col_b="RefDes",
        compare_columns=[("Part Number", "Part Number", "Text (ignore case)")],
        check_part_usage=False,
    )

    assert len(result.differences) == 1
    diff = result.differences[0]
    assert diff["RefDes"] == "R1"
    assert diff["Column"] == "Part Number"
    assert diff["Value_A"] == "PN-10K"
    assert diff["Value_B"] == "PN-4K7"


def test_compare_two_boms_without_compare_columns_reports_no_value_diffs() -> None:
    """Without ``compare_columns`` the custom compare matches on RefDes
    membership only — even a column that differs produces zero value diffs.
    """
    bom_a = _custom_df([{"RefDes": "R1", "Part Number": "PN-10K"}])
    bom_b = _custom_df([{"RefDes": "R1", "Part Number": "PN-4K7"}])

    result = compare_two_boms(
        bom_a, bom_b,
        refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
    )

    assert result.differences == []


def test_compare_two_boms_numeric_rule_ignores_text_formatting() -> None:
    """The Numeric rule compares values numerically, so '1' and '1.0' are equal
    and produce no diff (a text rule would flag the formatting difference).
    """
    bom_a = _custom_df([{"RefDes": "R1", "Qty": "1"}])
    bom_b = _custom_df([{"RefDes": "R1", "Qty": "1.0"}])

    result = compare_two_boms(
        bom_a, bom_b,
        refdes_col_a="RefDes", refdes_col_b="RefDes",
        compare_columns=[("Qty", "Qty", "Numeric")],
        check_part_usage=False,
    )

    assert result.differences == []


def test_compare_two_boms_skips_pair_with_column_absent_from_a_file() -> None:
    """A compare pair whose column is missing from one file must NOT yield a
    spurious 'value vs empty' diff for every matched RefDes — the pair is
    dropped because the column cannot be compared (e.g. a stale pair after the
    user re-inspected a file with a different schema).
    """
    bom_a = _custom_df([{"RefDes": "R1", "Part Number": "PN-10K"}])
    bom_b = _custom_df([{"RefDes": "R1", "Description": "RES 10K"}])  # no Part Number

    result = compare_two_boms(
        bom_a, bom_b,
        refdes_col_a="RefDes", refdes_col_b="RefDes",
        compare_columns=[("Part Number", "Part Number", "Text (ignore case)")],
        check_part_usage=False,
    )

    assert result.differences == []


# ----- Batch 5 (2026-07 stability sweep) -------------------------------------


def test_custom_fmr_sheet_writes_user_facing_status(tmp_path) -> None:
    """The custom-path Failure Mode Ratio Errors sheet must translate the
    internal "FMR != 1.0" status the same way the group path does — the raw
    token is an unexpanded abbreviation in the delivered report."""
    from openpyxl import load_workbook

    from bom_compare.excel_export import write_bom_compare_excel

    bom_a = _custom_df([
        {"RefDes": "U2", "Ratio": 0.5},
        {"RefDes": "U2", "Ratio": 0.6},  # sums to 1.1 -> FMR mismatch
    ])
    bom_b = _custom_df([{"RefDes": "U9", "Ratio": 1.0}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_fmr=True, check_part_usage=False,
    )
    assert result.fmr_warnings, "fixture should produce an FMR warning"

    out = tmp_path / "custom_fmr.xlsx"
    write_bom_compare_excel(result, str(out), "First BOM", "Second BOM")

    wb = load_workbook(out)
    try:
        ws = wb["Failure Mode Ratio Errors"]
        headers = [c.value for c in ws[1]]
        status_idx = headers.index("Status") + 1
        statuses = [
            ws.cell(row=r, column=status_idx).value
            for r in range(2, ws.max_row + 1)
        ]
    finally:
        wb.close()
    assert any(
        s == "Failure Mode Ratio does not equal 1.0" for s in statuses
    ), statuses
    assert not any(s == "FMR != 1.0" for s in statuses), statuses


# ---------------------------------------------------------------------------
# Batch 6 #1: empty / wrong-sheet inputs must raise a project ValidationError
# with a friendly, UI-labelled message (no raw ValueError, no "DataFrame"
# jargon). The sheet the user picked simply had no data rows.
# ---------------------------------------------------------------------------

def test_explode_grouping_empty_df_raises_friendly_validation_error() -> None:
    with pytest.raises(ValidationError) as excinfo:
        explode_grouping(pd.DataFrame(), "Reference Designator", None)
    msg = str(excinfo.value)
    assert "Grouping file" in msg
    assert "no data rows" in msg
    assert "DataFrame" not in msg  # jargon dropped


def test_explode_bom_empty_df_raises_friendly_validation_error() -> None:
    with pytest.raises(ValidationError) as excinfo:
        explode_bom(
            pd.DataFrame(), "Reference Designator", None, False,
            re.compile(r"\bDNP\b", re.I),
        )
    msg = str(excinfo.value)
    assert "BOM file" in msg
    assert "no data rows" in msg
    assert "DataFrame" not in msg


def test_compare_two_boms_empty_uses_ui_file_labels() -> None:
    good = pd.DataFrame([{"Reference Designator": "R1"}])

    with pytest.raises(ValidationError) as excinfo_a:
        compare_two_boms(pd.DataFrame(), good, "Reference Designator", "Reference Designator")
    msg_a = str(excinfo_a.value)
    assert "File 1" in msg_a
    assert "no data rows" in msg_a
    assert "DataFrame" not in msg_a

    with pytest.raises(ValidationError) as excinfo_b:
        compare_two_boms(good, pd.DataFrame(), "Reference Designator", "Reference Designator")
    msg_b = str(excinfo_b.value)
    assert "File 2" in msg_b
    assert "no data rows" in msg_b
    assert "DataFrame" not in msg_b


def test_group_report_preserves_user_text_in_passthrough_columns(tmp_path) -> None:
    """Final-day audit F2 (sibling of the Failure Rate corruption fix): the
    group-path report expander must run ONLY on tool-authored diagnostic
    columns (Reason/Status). User pass-through text like a BOM Description
    of "Main CB panel" (circuit breaker) must ship verbatim."""
    from openpyxl import load_workbook

    from bom_compare.bom_compare_logic import AnalyzeResults
    from bom_compare.group_analysis import write_excel_report

    results = AnalyzeResults(
        summary=pd.DataFrame([{"Item": "Total groups", "Value": 1}]),
        description_warnings=pd.DataFrame(
            [
                {
                    "Base": "CB1",
                    "Expected_Approx": 2,
                    "Observed_Grouped_Tokens": "CB1-1, CB1-2",
                    "Description": "Main CB panel",
                }
            ]
        ),
        fmr_warnings=pd.DataFrame(
            [{"RefDes": "U2", "Sum": 1.1, "Status": "FMR != 1.0"}]
        ),
    )

    out = tmp_path / "group_report.xlsx"
    write_excel_report(results, str(out))

    wb = load_workbook(out)
    try:
        warn_ws = wb["Warnings"]
        warn_headers = [c.value for c in warn_ws[1]]
        desc_idx = warn_headers.index("Description") + 1
        description = warn_ws.cell(row=2, column=desc_idx).value

        fmr_ws = wb["Failure Mode Ratio Errors"]
        fmr_headers = [c.value for c in fmr_ws[1]]
        status_idx = fmr_headers.index("Status") + 1
        status = fmr_ws.cell(row=2, column=status_idx).value
    finally:
        wb.close()

    # User text is untouched...
    assert description == "Main CB panel", description
    # ...while the tool-authored Status column is still expanded.
    assert status == "Failure Mode Ratio does not equal 1.0", status


def test_detect_column_tolerates_non_string_headers() -> None:
    """Final-day audit F1: openpyxl preserves numeric header cells as
    int/float (a column headed 2024, or the user picking a data row as the
    header row). detect_column must match by string coercion instead of
    dying on int.lower() with a cryptic AttributeError."""
    from common.utils import detect_column

    columns = [2024, "Reference Designator", 3.5]
    assert detect_column(columns, ["Reference Designator"]) == "Reference Designator"
    assert detect_column(columns, ["2024"]) == 2024
    assert detect_column(columns, ["missing"]) is None
    assert (
        detect_column(columns, ["designator"], substring_match=True)
        == "Reference Designator"
    )


# ---------------------------------------------------------------------------
# UX round-2 #4: content-based FMEA detection on the custom path. The FMEA-
# aware checks (scope warnings, duplicate identity, usage CB-skip) used to be
# gated purely on the FILENAME containing fmea/fmeca/piece-part — renaming an
# FMEA export silently dropped every FMEA-aware behavior. A validated FMEA
# Level column (name + value sampling) now triggers the scan too. Content-only
# detection requires EXPLICIT circuit-block/piece-part text in the level
# column — blank-row inference alone is not evidence — so a plain BOM can
# never misfire into FMEA mode.
# ---------------------------------------------------------------------------

def test_custom_fmea_detected_by_content_without_fmea_filename() -> None:
    """A file with a validated FMEA Level column is treated as an FMEA even
    when its filename gives no hint (e.g. 'analysis_revB.xlsx'): the scope
    scan runs and a CB-only token is flagged."""
    bom_a = _custom_df([
        {"RefDes": "U1", "FMEA Level": "Circuit Block"},
        {"RefDes": "U2", "FMEA Level": "Circuit Block"},
        {"RefDes": "U2", "FMEA Level": "Piece-Part"},
    ])
    bom_b = _custom_df([{"RefDes": "U1"}, {"RefDes": "U2"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
        source_name_a="analysis_revB.xlsx",
        source_name_b="parts_list.xlsx",
    )

    cb_only = [
        w for w in result.scope_warnings
        if w["Source"] == "BOM A" and w["ReasonCode"] == "SCOPE_CB_ONLY"
    ]
    assert any(w["RefDes"] == "U1" for w in cb_only), result.scope_warnings
    # U2 appears in both scopes — must not be flagged.
    assert not any(w["RefDes"] == "U2" for w in result.scope_warnings)


def test_custom_plain_bom_never_misfires_into_fmea_mode() -> None:
    """A plain BOM (no FMEA Level column) with grouped multi-token RefDes
    cells must NOT be treated as an FMEA — the loose row-inference path
    ('two tokens in a cell looks like a circuit block') is only reachable
    behind a filename or validated-level-column signal."""
    bom_a = _custom_df([
        {"RefDes": "C1, C2", "Description": "Caps"},
        {"RefDes": "R1", "Description": "Resistor"},
    ])
    bom_b = _custom_df([{"RefDes": "C1"}, {"RefDes": "C2"}, {"RefDes": "R1"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
        source_name_a="bom_revA.xlsx",
        source_name_b="bom_revB.xlsx",
    )

    assert result.scope_warnings == []


def test_custom_level_column_without_explicit_rows_not_fmea() -> None:
    """Content-only detection demands explicit CB/PP text. A 'Record Type'
    column whose values merely contain 'cb' as a substring ('PCB Assembly')
    opens the detection gate but classifies zero explicit rows — the file
    must NOT enter FMEA mode, even though its blank rows would be inferred
    as piece-part by the level-column classifier."""
    bom_a = _custom_df([
        {"RefDes": "U1", "Record Type": "PCB Assembly"},
        {"RefDes": "U2", "Record Type": None},
        {"RefDes": "U3", "Record Type": None},
    ])
    bom_b = _custom_df([{"RefDes": "U1"}, {"RefDes": "U2"}, {"RefDes": "U3"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
        source_name_a="assembly_export.xlsx",
        source_name_b="parts.xlsx",
    )

    assert result.scope_warnings == []


def test_custom_fmea_filename_fallback_unchanged() -> None:
    """The legacy filename trigger still works without a level column: an
    FMEA-named file whose rows carry CB/PP keywords in a plain text column
    goes through the row-scan fallback exactly as before."""
    bom_a = _custom_df([
        {"RefDes": "U1", "Notes": "Circuit Block"},
        {"RefDes": "U2", "Notes": "Circuit Block"},
        {"RefDes": "U2", "Notes": "Piece-Part"},
    ])
    bom_b = _custom_df([{"RefDes": "U1"}, {"RefDes": "U2"}])

    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=False,
        source_name_a="board_FMEA.xlsx",
        source_name_b="parts_list.xlsx",
    )

    cb_only = [
        w for w in result.scope_warnings
        if w["Source"] == "BOM A" and w["ReasonCode"] == "SCOPE_CB_ONLY"
    ]
    assert any(w["RefDes"] == "U1" for w in cb_only), result.scope_warnings


# ---------------------------------------------------------------------------
# UX round-2 #3: the custom-path report uses the SAME human sheet-naming
# scheme as the group path (spaces, Title Case) — no more Only_In_* /
# Part_Usage / Failure_Mode_Ratio underscore names in one tool's output.
# ---------------------------------------------------------------------------

def test_custom_report_sheet_names_use_group_path_scheme(tmp_path) -> None:
    from openpyxl import load_workbook

    from bom_compare.bom_compare_logic import BomCompareResult
    from bom_compare.excel_export import write_bom_compare_excel

    result = BomCompareResult(
        only_in_a=[{"RefDes": "U1"}],
        only_in_b=[{"RefDes": "U9"}],
        part_usage_warnings=[{
            "Source": "File 1", "RefDes": "U1", "Base": "U1", "Usage": 0.5,
            "Expected": 1.0, "Count": 1, "ReasonCode": "PU_MISMATCH",
            "Reason": "Usage mismatch",
        }],
        fmr_warnings=[{
            "Source": "File 1", "RefDes": "U2", "Sum": 1.1,
            "Status": "FMR != 1.0",
        }],
        scope_warnings=[{
            "Source": "BOM A", "RefDes": "U1", "ReasonCode": "SCOPE_CB_ONLY",
            "Scope Status": "Only found in Circuit Block rows",
            "Cross-File": "Only in BOM A", "Details": "",
            "InCircuitBlock": "Yes", "InPiecePart": "No", "ScopeColumn": "RefDes",
        }],
    )

    out = tmp_path / "custom_names.xlsx"
    write_bom_compare_excel(result, str(out), "First BOM", "Second BOM")

    wb = load_workbook(out)
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()

    assert names == [
        "Summary",
        "Only In First BOM",
        "Only In Second BOM",
        "Part Usage",
        "Failure Mode Ratio Errors",
        "Scope Warnings",
    ], names
    # Lockstep guard: the custom path must never reintroduce underscore names.
    assert not any("_" in n for n in names), names


# ---------------------------------------------------------------------------
# 2026-07-20 cross-file checks. Domain facts these encode:
#   1. FMR columns only exist in FMEA-generated files (header "Failure Mode
#      Ratio"), never in grouping files — the group-path FMR check must follow
#      the column to whichever file has it instead of only ever reading the
#      grouping file.
#   2. Dash suffixes are PINS of one base component, and Part Usage 1/N means
#      the base has exactly N pin/instance tokens in BOTH files (suffixes need
#      not be sequential: usage 1/4 <=> U60-1, U60-2, U60-3, U60-100).
# ---------------------------------------------------------------------------

def _cross_mapping() -> ColumnMapping:
    return ColumnMapping(
        grouping_group_col="Group", grouping_refdes_col="RefDes",
        bom_refdes_col="RefDes", bom_desc_col="Description",
    )


def _cross_options(**overrides) -> AnalyzeOptions:
    base = dict(
        run_warning_checks=False, run_duplicate_checks=False,
        check_part_usage=False, check_fmr=False,
    )
    base.update(overrides)
    return AnalyzeOptions(**base)


def _bom_usage_df(refs, usages):
    return pd.DataFrame({
        "RefDes": refs,
        "Description": ["" for _ in refs],
        "Part Usage": usages,
    })


def test_analyze_check_fmr_detects_ratio_column_in_bom_file() -> None:
    """FMRs live in FMEA-generated files, which load in the BOM slot — the
    check must follow the Failure Mode Ratio column there instead of
    reporting 'column not found' against the (always ratio-less) grouping."""
    group_df = pd.DataFrame({"Group": ["G"], "RefDes": ["U1"]})
    bom_df = pd.DataFrame({
        "RefDes": ["U1", "U1"],
        "Description": ["", ""],
        "Failure Mode Ratio": [0.6, 0.3],  # sums to 0.9 -> must be flagged
    })
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_fmr=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    rows = results.fmr_warnings.to_dict("records")
    assert any(r["RefDes"] == "U1" and r["Status"] == "FMR != 1.0" for r in rows), rows
    assert all("not found" not in str(r["Status"]).lower() for r in rows), rows
    flagged = [r for r in rows if r["RefDes"] == "U1"]
    assert flagged[0]["Source"] == "BOM", flagged


def test_analyze_check_fmr_neither_file_reports_plain_language_skip() -> None:
    """With no FMR column in either file the report says so in plain language
    (naming the FMEA-generator origin), not the bare 'Ratio column not found'."""
    group_df = pd.DataFrame({"Group": ["G"], "RefDes": ["U1"]})
    bom_df = pd.DataFrame({"RefDes": ["U1"], "Description": [""]})
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_fmr=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    rows = results.fmr_warnings.to_dict("records")
    assert len(rows) == 1, rows
    assert "FMEA-generated" in str(rows[0]["Status"]), rows


def test_analyze_check_fmr_ignores_generic_percentage_column() -> None:
    """Now that BOTH files are scanned, detection uses the strict FMR synonym
    list — a generic 'Percentage' column (common on plain BOMs, unrelated to
    failure modes) must not trigger the check."""
    group_df = pd.DataFrame({
        "Group": ["G"], "RefDes": ["U1"], "Percentage": [0.2],
    })
    bom_df = pd.DataFrame({"RefDes": ["U1"], "Description": [""]})
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_fmr=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    rows = results.fmr_warnings.to_dict("records")
    assert not any(r["Status"] == "FMR != 1.0" for r in rows), rows
    assert len(rows) == 1 and "FMEA-generated" in str(rows[0]["Status"]), rows


def test_analyze_fmr_source_uses_display_labels() -> None:
    """The FMR sheet's new Source column carries the user's display name for
    the file the bad sum came from (mirrors the Part Usage sheet)."""
    group_df = pd.DataFrame({"Group": ["G"], "RefDes": ["U1"]})
    bom_df = pd.DataFrame({
        "RefDes": ["U1"], "Description": [""], "Failure Mode Ratio": [0.4],
    })
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_fmr=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
        file_labels=("Groups", "Main FMEA"),
    )
    rows = results.fmr_warnings.to_dict("records")
    assert rows and rows[0]["Source"] == "Main FMEA", rows


def test_analyze_cross_usage_counts_agree_no_warnings() -> None:
    """Usage 1/4 with four pin tokens in the BOM AND four in the grouping
    (non-sequential suffixes) is fully consistent — no warnings."""
    group_df = pd.DataFrame({
        "Group": ["G"], "RefDes": ["U60-1, U60-2, U60-3, U60-100"],
    })
    bom_df = _bom_usage_df(
        ["U60-1", "U60-2", "U60-3", "U60-100"], [0.25, 0.25, 0.25, 0.25],
    )
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_part_usage=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    assert results.part_usage_warnings.empty, results.part_usage_warnings


def test_analyze_cross_usage_flags_grouping_count_mismatch() -> None:
    """Usage 1/4 matches the BOM's four pin tokens but the grouping only lists
    three -> exactly one cross-file warning naming the grouping side."""
    group_df = pd.DataFrame({
        "Group": ["G"], "RefDes": ["U60-1, U60-2, U60-3"],
    })
    bom_df = _bom_usage_df(
        ["U60-1", "U60-2", "U60-3", "U60-100"], [0.25, 0.25, 0.25, 0.25],
    )
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_part_usage=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    rows = results.part_usage_warnings.to_dict("records")
    assert [r["ReasonCode"] for r in rows] == ["PU_COUNT_MATCHES_THIS_FILE_ONLY"], rows
    assert rows[0]["Base"] == "U60", rows
    assert "Grouping" in rows[0]["Reason"], rows
    assert "3" in rows[0]["Reason"] and "4" in rows[0]["Reason"], rows


def test_analyze_cross_usage_flags_bom_count_mismatch() -> None:
    """One U60 in the BOM with usage 1/4 while the grouping lists four pins:
    the within-file mismatch still fires, and the cross-file check adds WHICH
    side the usage value actually agrees with."""
    group_df = pd.DataFrame({
        "Group": ["G"], "RefDes": ["U60-1, U60-2, U60-3, U60-100"],
    })
    bom_df = _bom_usage_df(["U60"], [0.25])
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_part_usage=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    codes = set(results.part_usage_warnings["ReasonCode"])
    assert "PU_EXPECTED_MISMATCH_BASIC" in codes, codes
    assert "PU_COUNT_MATCHES_OTHER_FILE_ONLY" in codes, codes


def test_analyze_cross_usage_conflict_matches_neither() -> None:
    """Usage 1/4, two pins in the BOM, three in the grouping — the value
    agrees with neither file's count."""
    group_df = pd.DataFrame({
        "Group": ["G"], "RefDes": ["U60-1, U60-2, U60-3"],
    })
    bom_df = _bom_usage_df(["U60-1", "U60-2"], [0.25, 0.25])
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_part_usage=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    codes = set(results.part_usage_warnings["ReasonCode"])
    assert "PU_CROSS_COUNT_CONFLICT" in codes, codes


def test_analyze_cross_usage_skips_bases_absent_from_grouping() -> None:
    """A base entirely absent from the grouping file is a membership finding
    (BOM Not Grouped sheet) — the usage cross-check must not double-report."""
    group_df = pd.DataFrame({"Group": ["G"], "RefDes": ["R1"]})
    bom_df = _bom_usage_df(["R1", "U60"], [1.0, 1.0])
    results = analyze(
        group_df, bom_df, _cross_mapping(), _cross_options(check_part_usage=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    assert results.part_usage_warnings.empty, results.part_usage_warnings
    assert not results.bom_not_in_groups.empty  # membership finding remains


def test_analyze_cross_usage_respects_prov_removal() -> None:
    """The grouping-side instance count uses the SAME filtered token set as the
    membership checks: treat_prov_as_covered=False removes PROV-group tokens
    before counting."""
    group_df = pd.DataFrame({
        "Group": ["G1", "PROV"],
        "RefDes": ["U60-1, U60-2, U60-3", "U60-100"],
    })
    bom_df = _bom_usage_df(
        ["U60-1", "U60-2", "U60-3", "U60-100"], [0.25, 0.25, 0.25, 0.25],
    )
    covered = analyze(
        group_df, bom_df, _cross_mapping(),
        _cross_options(check_part_usage=True, treat_prov_as_covered=True),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    assert covered.part_usage_warnings.empty, covered.part_usage_warnings

    excluded = analyze(
        group_df, bom_df, _cross_mapping(),
        _cross_options(check_part_usage=True, treat_prov_as_covered=False),
        file_paths=("grouping.xlsx", "bom.xlsx"),
    )
    codes = list(excluded.part_usage_warnings["ReasonCode"])
    assert codes == ["PU_COUNT_MATCHES_THIS_FILE_ONLY"], codes


def test_compare_two_boms_cross_usage_between_files() -> None:
    """Custom path: each file's usage is cross-checked against the OTHER
    file's instance count. Here each file is internally consistent but they
    disagree with each other (4 pins vs 3 pins), so BOTH get flagged."""
    third = 1.0 / 3.0
    bom_a = _custom_df([
        {"RefDes": r, "Part Usage": 0.25}
        for r in ("U60-1", "U60-2", "U60-3", "U60-100")
    ])
    bom_b = _custom_df([
        {"RefDes": r, "Part Usage": third}
        for r in ("U60-1", "U60-2", "U60-3")
    ])
    result = compare_two_boms(
        bom_a, bom_b, refdes_col_a="RefDes", refdes_col_b="RefDes",
        check_part_usage=True,
    )
    cross = [
        w for w in result.part_usage_warnings
        if w["ReasonCode"] == "PU_COUNT_MATCHES_THIS_FILE_ONLY"
    ]
    assert {w["Source"] for w in cross} == {"File 1", "File 2"}, result.part_usage_warnings


def test_cross_usage_reason_codes_have_user_facing_labels() -> None:
    """Lockstep: every emitted cross-file ReasonCode has a plain-language
    label so the report never ships a raw underscore code."""
    from common.user_facing_labels import REASON_CODE_LABELS

    for code in (
        "PU_COUNT_MATCHES_THIS_FILE_ONLY",
        "PU_COUNT_MATCHES_OTHER_FILE_ONLY",
        "PU_CROSS_COUNT_CONFLICT",
    ):
        assert code in REASON_CODE_LABELS, code
