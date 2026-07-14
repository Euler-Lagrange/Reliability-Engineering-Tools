"""Test infrastructure: ensure backend/python is on sys.path so direct imports
of FMEA / common / runtime modules work for in-process unit tests.

The existing test_sidecar_main.py spawns the sidecar as a subprocess and
therefore doesn't need this — but Phase D's test_fmea_phase_d.py imports
FMEAProcessor directly for focused unit tests of the new inheritance,
BOM Additions, and functional_to_piecepart logic.
"""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
import sys
from pathlib import Path
from typing import Any, Callable, Iterator, Mapping

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

_BACKEND_PYTHON = Path(__file__).resolve().parents[1] / "python"
if str(_BACKEND_PYTHON) not in sys.path:
    sys.path.insert(0, str(_BACKEND_PYTHON))


@dataclass(frozen=True)
class PreserveTargetWorkbook:
    """Round-trip helpers and stable coordinates for a preserve target."""

    path: Path
    header_row: int
    columns: Mapping[str, int]
    rows: Mapping[str, int]
    merged_range: str

    @contextmanager
    def open(self, output_path: Path | str | None = None) -> Iterator[Any]:
        workbook = load_workbook(output_path or self.path, data_only=False)
        try:
            yield workbook
        finally:
            workbook.close()

    def assert_cells(
        self,
        output_path: Path | str,
        expected: Mapping[tuple[str, str], Any],
    ) -> None:
        """Round-trip *output_path* and assert exact cell values/formulas."""
        with self.open(output_path) as workbook:
            actual = {
                (sheet_name, coordinate): workbook[sheet_name][coordinate].value
                for sheet_name, coordinate in expected
            }
        assert actual == dict(expected)


@pytest.fixture
def preserve_target_factory(tmp_path: Path) -> Callable[..., PreserveTargetWorkbook]:
    """Build the shared, user-authored preserve-formatting target workbook.

    The shape is deliberately richer than the legacy preserve fixture: its
    header starts on row 3, it has two real function groups and piece-part
    rows, and it carries user content that later Wave-4 tests must prove is
    not lost or misattributed.
    """

    def build(
        filename: str = "preserve_target.xlsx",
        *,
        configure: Callable[[Any, PreserveTargetWorkbook], None] | None = None,
    ) -> PreserveTargetWorkbook:
        path = tmp_path / filename
        header_row = 3
        headers = [
            "FMEA-ID",
            "FMEA Level",
            "Failure Mode Causes",
            "Component Part Number",
            "Component Part Description",
            "Failure Mode",
            "Failure Mode Ratio",
            "Local Effect",
            "Next Higher Effect",
            "End Effect",
            "Part Usage",
            "Engineering Notes",
        ]
        columns = {header: index for index, header in enumerate(headers, start=1)}
        rows = {
            "duplicate_group": 4,
            "duplicate_a": 5,
            "duplicate_b": 6,
            "unique_group": 8,
            "unique_piece_part": 9,
            "merged_note": 12,
        }
        fixture = PreserveTargetWorkbook(
            path=path,
            header_row=header_row,
            columns=columns,
            rows=rows,
            merged_range="H12:J12",
        )

        workbook = Workbook()
        fmea = workbook.active
        fmea.title = "FMEA"
        fmea["A1"] = "Legacy qualification FMEA - user maintained"
        fmea["A2"] = "Do not discard hand-authored effects or engineering notes"
        for column, header in enumerate(headers, start=1):
            cell = fmea.cell(row=header_row, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True)

        fmea.append(
            [
                "CPU-001-A",
                "Circuit Block",
                "U1",
                "",
                "Processor control",
                "",
                "",
                "",
                "",
                "",
                "",
                "User-owned CPU group note",
            ]
        )
        fmea.append(
            [
                "CPU-001-U1-A",
                "Piece-Part",
                "U1",
                "PN-U1-A",
                "Processor channel A",
                "OPEN",
                0.5,
                "Hand-authored local effect A",
                '=H5&" -> system"',
                "Hand-authored end effect A",
                "=1/2",
                "Keep with PN-U1-A",
            ]
        )
        fmea.append(
            [
                "CPU-001-U1-B",
                "Piece-Part",
                "U1",
                "PN-U1-B",
                "Processor channel B",
                "OPEN",
                0.5,
                "Hand-authored local effect B",
                '=H6&" -> system"',
                "Hand-authored end effect B",
                "=1/2",
                "Keep with PN-U1-B",
            ]
        )
        fmea.append(
            [
                "",
                "Separator",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "User-authored separator between function groups",
            ]
        )
        fmea.append(
            [
                "PWR-002-A",
                "Circuit Block",
                "R1",
                "",
                "Power conditioning",
                "",
                "",
                "",
                "",
                "",
                "",
                "User-owned power group note",
            ]
        )
        fmea.append(
            [
                "PWR-002-R1-A",
                "Piece-Part",
                "R1",
                "PN-R1",
                "Sense resistor",
                "OPEN",
                1.0,
                "Hand-authored local effect R1",
                '=H9&" -> shutdown"',
                "Hand-authored end effect R1",
                1,
                "Calibration-critical note",
            ]
        )
        fmea["B10"] = "Separator"
        fmea["B11"] = "Separator"
        fmea["B12"] = "Note"
        fmea.merge_cells(fixture.merged_range)
        fmea["H12"] = "User approval block below the insertion region"
        fmea.row_dimensions[12].height = 31.5

        validation = workbook.create_sheet("Validation_Warnings")
        validation["A1"] = "User-authored validation register"
        validation["A2"] = "Waiver 17 remains open"
        validation["B2"] = "Owner: Reliability"

        calculations = workbook.create_sheet("Calculations")
        calculations["A1"] = "User calculation sheet"
        calculations["A2"] = 21
        calculations["B2"] = "=A2*2"

        if configure is not None:
            configure(workbook, fixture)

        workbook.save(path)
        workbook.close()
        return fixture

    return build
