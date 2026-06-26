"""Pure-logic tests for the RefDes BOM-coverage reverse-diff.

The RefDes Extractor already labels every *extracted* component
``(Verified)``/``(Unverified)`` against the BOM. ``coverage_report`` adds the
two missing directions an engineer needs when reconciling a board:

    * BOM Not Grouped     — in the BOM, but not pulled into a real group
                            (further split: never extracted vs extracted but
                            left UNGROUPED / PROVISIONAL).
    * Extracted Not In BOM — pulled off the schematic, absent from the BOM.

Coverage is *component level*: pins/instances collapse to their base RefDes
(``U200-1`` and ``U200-2`` both cover BOM ``U200``; ``J1-4`` covers ``J1``),
via ``get_usage_base_refdes`` — the same family of base-normalization the rest
of the suite uses. These tests assert the exact buckets and counts so the
classification can never silently drift.
"""

from __future__ import annotations

from openpyxl import Workbook

from refdes_extractor.coverage_report import build_coverage, write_coverage_sheets


def _row(group: str, refdes_csv: str = "", pages: str = "", *, is_gap: bool = False) -> dict:
    """Build a result row matching the extractor's real output schema."""
    row = {
        "group": group,
        "failure mode causes": refdes_csv,
        "component count": len([t for t in refdes_csv.split(",") if t.strip()]),
        "pages": pages,
    }
    if is_gap:
        row["_is_gap"] = True
    return row


# --------------------------------------------------------------------------
# BOM Not Grouped direction
# --------------------------------------------------------------------------

def test_bom_components_absent_from_groups_are_not_extracted():
    coverage = build_coverage(
        bom_set={"R1", "R2", "R3"},
        results=[_row("GRP1 (Verified)", "R1", "2")],
    )
    statuses = {r["refdes"]: r["status"] for r in coverage.bom_not_grouped}
    assert statuses == {"R2": "Not Extracted", "R3": "Not Extracted"}
    assert coverage.summary["not_extracted_count"] == 2
    assert coverage.summary["grouped_from_bom"] == 1


def test_bom_component_only_in_provisional_is_extracted_provisional():
    coverage = build_coverage(
        bom_set={"U5"},
        results=[_row("PROVISIONAL", "U5", "7")],
    )
    assert [r["status"] for r in coverage.bom_not_grouped] == ["Extracted-Provisional"]
    assert coverage.summary["extracted_provisional_count"] == 1
    assert coverage.summary["not_extracted_count"] == 0


def test_bom_component_only_ungrouped_is_extracted_ungrouped():
    # UNGROUPED rows are emitted only by the legacy fallback backend (the default
    # NextGen backend discards non-annotation words); this pins the classification
    # for that legacy-shaped input.
    coverage = build_coverage(
        bom_set={"U7"},
        results=[_row("UNGROUPED (IN BOM)", "U7", "3")],
    )
    assert [r["status"] for r in coverage.bom_not_grouped] == ["Extracted-Ungrouped"]
    assert coverage.summary["extracted_ungrouped_count"] == 1


def test_grouped_component_is_not_reported_as_missing():
    coverage = build_coverage(
        bom_set={"R1"},
        results=[_row("GRP1 (Verified)", "R1", "2")],
    )
    assert coverage.bom_not_grouped == []
    assert coverage.summary["bom_not_grouped_count"] == 0


# --------------------------------------------------------------------------
# Extracted Not In BOM direction
# --------------------------------------------------------------------------

def test_extracted_component_absent_from_bom_is_listed():
    coverage = build_coverage(
        bom_set={"R1"},
        results=[
            _row("GRP1 (Verified)", "R1", "2"),
            _row("GRP2 (Unverified)", "C9", "4"),
        ],
    )
    extra = {r["refdes"]: r for r in coverage.extracted_not_in_bom}
    assert set(extra) == {"C9"}
    assert extra["C9"]["group"] == "GRP2 (Unverified)"
    assert extra["C9"]["pages"] == "4"
    assert coverage.summary["extracted_not_in_bom_count"] == 1


def test_verified_components_never_appear_as_extracted_not_in_bom():
    coverage = build_coverage(
        bom_set={"R1", "R2"},
        results=[_row("GRP1 (Verified)", "R1, R2", "2")],
    )
    assert coverage.extracted_not_in_bom == []


# --------------------------------------------------------------------------
# Normalization: coverage is component-level
# --------------------------------------------------------------------------

def test_instance_and_pin_suffixes_collapse_to_base():
    # BOM lists the base parts; schematic carries instance + pin suffixes.
    coverage = build_coverage(
        bom_set={"U200", "J1"},
        results=[_row("G1 (Verified)", "U200-1, U200-2, J1-4", "5")],
    )
    assert coverage.bom_not_grouped == []  # both bases are covered
    assert coverage.extracted_not_in_bom == []


def test_matching_is_case_insensitive():
    coverage = build_coverage(
        bom_set={"R1"},
        results=[_row("G1 (Verified)", "r1", "2")],
    )
    assert coverage.bom_not_grouped == []


# --------------------------------------------------------------------------
# Gap rows are never extraction evidence
# --------------------------------------------------------------------------

def test_gap_rows_are_ignored():
    coverage = build_coverage(
        bom_set={"R1"},
        results=[
            _row("CPU-003 (GROUP NOT DETECTED)", "", "", is_gap=True),
        ],
    )
    # The gap is not extraction, so R1 is genuinely Not Extracted...
    assert [r["status"] for r in coverage.bom_not_grouped] == ["Not Extracted"]
    # ...and the gap placeholder must not leak into the over-extraction list.
    assert coverage.extracted_not_in_bom == []


# --------------------------------------------------------------------------
# Enrichment: Part Number + Description + pages from the BOM
# --------------------------------------------------------------------------

def test_not_extracted_rows_are_enriched_from_bom_metadata():
    coverage = build_coverage(
        bom_set={"R5"},
        results=[],
        bom_meta={"R5": {"part_number": "CRCW0603", "description": "RES 10K 1%"}},
        bom_page_map={"R5": {5, 3}},
    )
    row = coverage.bom_not_grouped[0]
    assert row["refdes"] == "R5"
    assert row["part_number"] == "CRCW0603"
    assert row["description"] == "RES 10K 1%"
    assert row["pages"] == "3, 5"  # sorted, comma-joined


def test_missing_metadata_yields_blank_columns():
    coverage = build_coverage(bom_set={"R5"}, results=[])
    row = coverage.bom_not_grouped[0]
    assert row["part_number"] == ""
    assert row["description"] == ""
    assert row["pages"] == ""


# --------------------------------------------------------------------------
# Summary + gate
# --------------------------------------------------------------------------

def test_summary_counts_for_mixed_board():
    coverage = build_coverage(
        bom_set={"R1", "R2", "U3", "U4"},
        results=[
            _row("GRP1 (Verified)", "R1", "1"),       # covered
            _row("PROVISIONAL", "U3", "2"),            # extracted-provisional
            _row("GRP9 (Unverified)", "Q99", "9"),     # extracted, not in BOM
        ],
    )
    s = coverage.summary
    assert s["bom_count"] == 4
    assert s["grouped_from_bom"] == 1                 # R1
    assert s["bom_not_grouped_count"] == 3            # R2, U3, U4
    assert s["not_extracted_count"] == 2              # R2, U4
    assert s["extracted_provisional_count"] == 1      # U3
    assert s["extracted_not_in_bom_count"] == 1       # Q99


def test_empty_bom_reports_nothing_missing():
    coverage = build_coverage(
        bom_set=set(),
        results=[_row("GRP1 (Unverified)", "R1", "1")],
    )
    assert coverage.bom_not_grouped == []
    assert coverage.summary["bom_count"] == 0


# --------------------------------------------------------------------------
# Workbook writer
# --------------------------------------------------------------------------

def test_collapsed_bom_rows_merge_pages_and_metadata_deterministically():
    # A BOM that lists multiple instance/pin rows sharing a usage-base must
    # collapse to ONE component-level row — deterministically, without dropping
    # the other rows' pages or metadata (review finding: non-deterministic
    # set-iteration winner + lost pages/meta).
    coverage = build_coverage(
        bom_set={"U200-1", "U200-2"},
        results=[],
        bom_meta={
            "U200-1": {"part_number": "PN-A", "description": ""},
            "U200-2": {"part_number": "PN-A", "description": "MCU"},
        },
        bom_page_map={"U200-1": {3}, "U200-2": {7}},
    )
    assert len(coverage.bom_not_grouped) == 1
    row = coverage.bom_not_grouped[0]
    assert row["refdes"] == "U200-1"          # deterministic (sorted)
    assert row["pages"] == "3, 7"             # unioned across collapsed rows
    assert row["part_number"] == "PN-A"
    assert row["description"] == "MCU"         # first non-blank across the group
    assert coverage.summary["bom_count"] == 1


def test_nan_cells_do_not_create_phantom_components():
    # A NaN cell in a result row must not stringify to a 'NAN' phantom RefDes
    # (project NaN-guard gotcha).
    coverage = build_coverage(
        bom_set={"R1"},
        results=[
            {"group": "G (Unverified)", "failure mode causes": float("nan"), "pages": float("nan")},
        ],
    )
    assert coverage.extracted_not_in_bom == []
    assert [r["refdes"] for r in coverage.bom_not_grouped] == ["R1"]


def test_mixed_type_pages_do_not_raise():
    # _format_pages must tolerate a mixed int/str page set rather than crash
    # the whole run (defensive: coverage is on the extraction critical path).
    coverage = build_coverage(
        bom_set={"R9"},
        results=[],
        bom_page_map={"R9": {1, "5", 2}},
    )
    assert coverage.bom_not_grouped[0]["pages"]  # non-empty, no exception


def test_unparented_pins_do_not_become_a_phantom_component():
    # Piece-part mode emits unparented pins as the literal token PIN-{text}.
    # These are NOT components: get_usage_base_refdes('PIN-3') == 'PIN', so every
    # distinct pin would collapse into one bogus 'PIN' over-extraction row and
    # inflate the count. A RefDes always has a digit; 'PIN' has none.
    coverage = build_coverage(
        bom_set={"U200"},
        results=[
            _row("U200 (Verified)", "U200", "4"),
            _row("U200 (Unverified)", "PIN-3, PIN-99, PIN-A7", "4"),
        ],
    )
    assert coverage.extracted_not_in_bom == []
    assert coverage.summary["extracted_not_in_bom_count"] == 0


def test_over_extracted_token_across_rows_unions_pages_and_prefers_real_group():
    # A not-in-BOM token seen in several result rows must union its pages and
    # show a real group (not PROVISIONAL), mirroring the BOM-side merge — not
    # silently keep only the first row (setdefault data loss).
    coverage = build_coverage(
        bom_set={"R1"},
        results=[
            _row("PROVISIONAL", "C9", "9"),
            _row("POWER (Unverified)", "C9", "2, 5"),
        ],
    )
    rows = coverage.extracted_not_in_bom
    assert len(rows) == 1
    assert rows[0]["refdes"] == "C9"
    assert rows[0]["group"] == "POWER (Unverified)"   # real group preferred
    assert rows[0]["pages"] == "2, 5, 9"              # unioned, numeric sort


def test_write_coverage_sheets_adds_three_sheets_with_rows():
    coverage = build_coverage(
        bom_set={"R1", "R2"},
        results=[
            _row("GRP1 (Verified)", "R1", "1"),
            _row("GRP2 (Unverified)", "C9", "4"),
        ],
        bom_meta={"R2": {"part_number": "PN2", "description": "RES"}},
    )
    wb = Workbook()
    wb.active.title = "RefDes Extraction"
    write_coverage_sheets(wb, coverage)

    assert "Coverage Summary" in wb.sheetnames
    assert "BOM Not Grouped" in wb.sheetnames
    assert "Extracted Not In BOM" in wb.sheetnames

    bng = wb["BOM Not Grouped"]
    header = [c.value for c in bng[1]]
    assert header == ["RefDes", "Part Number", "Description", "Status", "Pages"]
    # one data row (R2), plus header
    assert bng.max_row == 2
    assert bng.cell(row=2, column=1).value == "R2"

    eni = wb["Extracted Not In BOM"]
    assert [c.value for c in eni[1]] == ["RefDes", "Group", "Pages"]
    assert eni.cell(row=2, column=1).value == "C9"
