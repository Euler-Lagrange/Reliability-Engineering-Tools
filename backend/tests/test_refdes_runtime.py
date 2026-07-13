"""Tests for RefDes Extractor runtime helpers (backend/python/refdes_extractor/runtime.py)."""

from __future__ import annotations

from pathlib import Path

import pytest

from common.cancellation import CancellationError
from common.exceptions import ValidationError
from refdes_extractor import runtime as refdes_runtime
from refdes_extractor.runtime import (
    _results_dataframe,
    _validate_options,
    execute_run_request,
    validate_run_request,
)


def test_dig4xx_combined_silent_loss_regression(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Wave R5: the test that would have caught the DIG-4xx incident. One
    run through the REAL pipeline exercises every silent-loss fix at once:

    * R1 — a valid RefDes outside every group lands in UNGROUPED (IN BOM);
    * R2 — a page whose annotation extraction times out becomes a result
      note + warning_count, and later pages still contribute;
    * R3 — a group label whose /Contents is empty is recovered from its
      appearance text and its group extracts normally;
    * R4 — a long numbering hole surfaces as ONE range summary row and the
      per-family gap note names the family.
    """
    import time

    fitz = pytest.importorskip("fitz")
    from openpyxl import Workbook, load_workbook

    # --- fixture PDF: 3 pages. Self-labeled FreeText groups use the
    # annotation's own rect, so member words must sit inside those rects. ---
    doc = fitz.open()
    page1 = doc.new_page(width=612, height=792)
    page1.add_freetext_annot(fitz.Rect(50, 40, 300, 200), "DIG-001", fontsize=10)
    page1.insert_text((100, 150), "R55", fontsize=10)  # inside DIG-001 rect
    ann5 = page1.add_freetext_annot(fitz.Rect(320, 40, 560, 200), "DIG-005", fontsize=10)
    page1.insert_text((400, 120), "U12", fontsize=10)  # inside DIG-005 rect
    page1.insert_text((100, 700), "C77", fontsize=10)  # outside all groups, in BOM
    doc.xref_set_key(ann5.xref, "Contents", "()")  # R3 trigger: appearance-only label

    page2 = doc.new_page(width=612, height=792)  # the page that times out (R2)
    page2.add_freetext_annot(fitz.Rect(50, 40, 300, 200), "DIG-010", fontsize=10)

    page3 = doc.new_page(width=612, height=792)
    page3.add_freetext_annot(fitz.Rect(50, 40, 300, 200), "DIG-020", fontsize=10)

    pdf_path = tmp_path / "dig4xx.pdf"
    doc.save(str(pdf_path))
    doc.close()

    # --- fixture BOM ---
    bom_path = tmp_path / "bom.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Reference Designator"])
    for refdes in ("R55", "U12", "C77"):
        ws.append([refdes])
    wb.save(bom_path)

    # --- deterministic page-2 timeout: real annots(), slowed on page index 1 ---
    real_annots = fitz.Page.annots

    def slow_annots(self, *args, **kwargs):
        if self.number == 1:
            time.sleep(0.6)
        return real_annots(self, *args, **kwargs)

    monkeypatch.setattr(fitz.Page, "annots", slow_annots)

    result = execute_run_request(
        {
            "workflowId": "refdes_extract",
            "outputStrategyId": "new_workbook_standard",
            "outputDirectory": str(tmp_path),
            "inputs": [
                {"role": "pdf", "path": str(pdf_path)},
                {"role": "bom", "path": str(bom_path), "selectedSheet": "BOM"},
            ],
            "options": {"annotation_page_timeout_seconds": 0.2},
        }
    )

    # R2: the timed-out page is a first-class warning.
    assert result["warning_count"] >= 1
    assert any("timed out on page 2" in note for note in result["notes"])

    # R4: the per-family gap note names the DIG family.
    assert any("DIG" in note and "not detected" in note.lower() for note in result["notes"])

    # Row-level assertions from the written workbook.
    out_files = list(tmp_path.glob("RefDesExtract_*.xlsx"))
    assert len(out_files) == 1
    out_wb = load_workbook(out_files[0])
    try:
        sheet = out_wb.worksheets[0]
        headers = [c.value for c in sheet[1]]
        group_idx = headers.index("Group")
        causes_idx = headers.index("Failure Mode Causes")
        rows = {
            str(r[group_idx].value): str(r[causes_idx].value or "")
            for r in sheet.iter_rows(min_row=2)
        }
    finally:
        out_wb.close()

    # R3: the appearance-only DIG-005 label was recovered and extracted.
    assert "U12" in rows.get("DIG-005 (Verified)", "")
    # R1: the out-of-group C77 surfaces in the UNGROUPED bucket.
    assert "C77" in rows.get("UNGROUPED (IN BOM)", "")
    # R4: the long hole between DIG-005 and DIG-020 is one summary row.
    assert any(
        "RANGE NOT DETECTED" in name and "consecutive" in name for name in rows
    )
    # R2 side effect: page 2's DIG-010 vanished (that IS the incident) — but
    # page 3's DIG-020 still extracted, proving the run continued.
    assert not any(name.startswith("DIG-010") for name in rows)
    assert any(name.startswith("DIG-020") for name in rows)


def test_bom_collision_orphans_do_not_count_as_dropped_pins(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Wave R6: 'bom-collision' orphan records flag pins that were KEPT — the
    'N pins dropped before output' note must exclude them, and they get their
    own kept-and-flagged note instead."""
    fitz = pytest.importorskip("fitz")
    from refdes_test import refdes_test_logic

    monkeypatch.setattr(
        refdes_test_logic, "extract_annotations_from_doc", lambda _doc, **_kw: []
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "detect_groups_with_fallback",
        lambda _doc, _annotations, **_kwargs: ([], False, {}),
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "extract_with_geometry_analysis_detailed",
        lambda **_kwargs: (
            [],
            {
                "backend_used": "test",
                "token_diagnostics": {},
                "orphan_pins": [
                    {"page": 1, "group": "G", "pin_text": "U7",
                     "disposition": "bom-collision", "detail": "kept"},
                    {"page": 1, "group": "G", "pin_text": "38",
                     "disposition": "excluded", "detail": "dropped"},
                ],
            },
        ),
    )
    monkeypatch.setattr(refdes_runtime, "verify_excel_readable", lambda _path: True)

    pdf_path = tmp_path / "schematic.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    result = execute_run_request(
        {
            "workflowId": "refdes_extract",
            "outputStrategyId": "new_workbook_standard",
            "outputDirectory": str(tmp_path),
            "inputs": [{"role": "pdf", "path": str(pdf_path)}],
            "options": {},
        }
    )

    assert any(n.startswith("1 pin dropped before output") for n in result["notes"])
    assert any("match a BOM RefDes (kept, flagged)" in n for n in result["notes"])


def test_annotation_timeout_option_plumbed_and_surfaces_warning(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Wave R2: annotation_page_timeout_seconds is a validated RefDesConfig
    option, the runtime forwards it into annotation extraction, and a
    timed-out page becomes a result note + warning_count increment instead
    of a log-only whisper (the DIG-4xx incident shipped with zero visible
    warning surface)."""
    fitz = pytest.importorskip("fitz")
    from refdes_test import refdes_test_logic

    captured: dict = {}

    def fake_extract_annotations(
        doc,
        stop_event=None,
        log_func=None,
        page_timeout=None,
        max_timeouts=3,
        timed_out_pages=None,
    ):
        captured["page_timeout"] = page_timeout
        if timed_out_pages is not None:
            timed_out_pages.append(2)
        return []

    monkeypatch.setattr(
        refdes_test_logic, "extract_annotations_from_doc", fake_extract_annotations
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "detect_groups_with_fallback",
        lambda _doc, _annotations, **_kwargs: ([], False, {}),
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "extract_with_geometry_analysis_detailed",
        lambda **_kwargs: ([], {"backend_used": "test", "token_diagnostics": {}}),
    )
    monkeypatch.setattr(refdes_runtime, "verify_excel_readable", lambda _path: True)

    pdf_path = tmp_path / "schematic.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    result = execute_run_request(
        {
            "workflowId": "refdes_extract",
            "outputStrategyId": "new_workbook_standard",
            "outputDirectory": str(tmp_path),
            "inputs": [{"role": "pdf", "path": str(pdf_path)}],
            "options": {"annotation_page_timeout_seconds": 45},
        }
    )

    assert captured["page_timeout"] == 45
    assert result["warning_count"] >= 1
    assert any("timed out on page 2" in note for note in result["notes"])
    # The option is a first-class validated field: bad values fail validation.
    bad = _validate_options({"annotation_page_timeout_seconds": 0})
    assert any(name == "annotation_page_timeout_seconds" for name, _ in bad)


def test_adaptive_orphan_threshold_zero_is_valid() -> None:
    # Adversarial-review fix: 0 is a legitimate threshold ("trigger full geometry
    # even with zero orphans"), the frontend control allows a minimum of 0, and
    # it is the integer analogue of adaptive_orphan_ratio=0.0. Validating it as
    # >= 1 would hard-block a UI-permitted value.
    assert _validate_options({"adaptive_orphan_threshold": 0}) == []
    bad = _validate_options({"adaptive_orphan_threshold": -1})
    assert any(name == "adaptive_orphan_threshold" for name, _ in bad)


def test_every_config_field_is_validated() -> None:
    # Guard against a future RefDesConfig field being added without a matching
    # validation category (which would leave it silently unvalidated, mirroring
    # from_options' silent drop).
    from dataclasses import fields

    from refdes_extractor.runtime import (
        RefDesConfig,
        _BOOL_OPTIONS,
        _ENUM_OPTIONS,
        _NON_NEGATIVE_INT_OPTIONS,
        _POSITIVE_INT_OPTIONS,
        _POSITIVE_NUMBER_OPTIONS,
        _RATIO_OPTION,
    )

    validated = (
        set(_BOOL_OPTIONS)
        | set(_ENUM_OPTIONS)
        | set(_POSITIVE_INT_OPTIONS)
        | set(_NON_NEGATIVE_INT_OPTIONS)
        | set(_POSITIVE_NUMBER_OPTIONS)
        | {_RATIO_OPTION}
    )
    assert validated == {f.name for f in fields(RefDesConfig)}


def test_results_dataframe_drops_internal_columns_after_annotation() -> None:
    # #4: internal markers ("_is_gap" from detect_sequence_gaps, "_row_style"
    # from annotate_results) must not leak into the user-facing sheet — and
    # the display headers are Title Case with the Validation Notes column.
    from refdes_extractor.validation_notes import annotate_results

    rows = [
        {"group": "U200", "failure mode causes": "Open", "component count": 3, "pages": "1"},
        {
            "group": "U201 (GROUP NOT DETECTED)",
            "failure mode causes": "",
            "component count": 0,
            "pages": "",
            "_is_gap": True,  # internal styling marker from detect_sequence_gaps
        },
    ]
    df = _results_dataframe(annotate_results(rows, bom_provided=True))
    assert "_is_gap" not in df.columns
    assert "_row_style" not in df.columns
    assert list(df.columns) == [
        "Group",
        "Failure Mode Causes",
        "Component Count",
        "Pages",
        "Validation Notes",
    ]
    assert len(df) == 2  # rows preserved; only the internal columns dropped


def test_results_dataframe_renames_to_display_headers() -> None:
    rows = [
        {"group": "U200", "failure mode causes": "Open", "component count": 3, "pages": "1"},
    ]
    df = _results_dataframe(rows)
    assert list(df.columns) == ["Group", "Failure Mode Causes", "Component Count", "Pages"]
    assert len(df) == 1


# ---------------------------------------------------------------------------
# #5 Stage 2 — engine-option validation
# ---------------------------------------------------------------------------

# The 16 user-settable engine parameters at their shipped defaults (mirrors the
# frontend RefDesOptions defaults, which mirror RefDesConfig). Every value is
# valid, so this doubles as the "all-defaults pass clean" fixture.
_DEFAULT_OPTIONS = {
    "extraction_mode": "functional",
    "backend_mode": "auto",
    "geometry_analysis_enabled": True,
    "adaptive_geometry_enabled": True,
    "geometry_batch_size": 10,
    "max_pin_label_length": 4,
    "prov_distance": 15.0,
    "geometry_subprocess_enabled": False,
    "geometry_batch_timeout_seconds": 240.0,
    "geometry_batch_checkpoint_enabled": True,
    "pin_assignment_threshold": 50.0,
    "refdes_search_radius": 100.0,
    "adaptive_orphan_threshold": 5,
    "adaptive_orphan_ratio": 0.3,
    "adaptive_max_pages": 10,
    "pinlist_prefers_annotation_mode": True,
}


def _offenders(errors) -> list[str]:
    return sorted(name for name, _ in errors)


def test_validate_options_all_defaults_pass_clean() -> None:
    # (c) valid values + all-defaults produce no errors.
    assert _validate_options(dict(_DEFAULT_OPTIONS)) == []


def test_validate_options_empty_and_non_dict_pass_clean() -> None:
    assert _validate_options({}) == []
    assert _validate_options(None) == []
    assert _validate_options("nonsense") == []


def test_validate_options_valid_non_default_values_pass_clean() -> None:
    # (c) valid non-default values across every category are accepted.
    assert _validate_options({
        "extraction_mode": "piece_part",
        "backend_mode": "legacy",
        "geometry_subprocess_enabled": True,
        "geometry_batch_size": 25,
        "pin_assignment_threshold": 75.0,
        "adaptive_orphan_ratio": 0.0,
    }) == []


def test_validate_options_rejects_wrong_typed_number() -> None:
    # (a) a wrong-typed value (non-numeric string) is flagged.
    assert _offenders(_validate_options({"geometry_batch_size": "abc"})) == ["geometry_batch_size"]
    assert _offenders(_validate_options({"prov_distance": "wide"})) == ["prov_distance"]


def test_validate_options_rejects_out_of_range_ratio() -> None:
    # (b) an out-of-range ratio is flagged in both directions.
    assert _offenders(_validate_options({"adaptive_orphan_ratio": 5})) == ["adaptive_orphan_ratio"]
    assert _offenders(_validate_options({"adaptive_orphan_ratio": -0.1})) == ["adaptive_orphan_ratio"]


def test_validate_options_ratio_boundaries_are_valid() -> None:
    for value in (0, 0.0, 1, 1.0, 0.3):
        assert _validate_options({"adaptive_orphan_ratio": value}) == [], value


def test_validate_options_positive_int_rejects_zero_negative_fractional() -> None:
    assert _offenders(_validate_options({"geometry_batch_size": 0})) == ["geometry_batch_size"]
    assert _offenders(_validate_options({"adaptive_max_pages": -1})) == ["adaptive_max_pages"]
    assert _offenders(_validate_options({"max_pin_label_length": 2.5})) == ["max_pin_label_length"]
    assert _offenders(_validate_options({"adaptive_orphan_threshold": "x"})) == ["adaptive_orphan_threshold"]


def test_validate_options_positive_int_accepts_whole_valued_float() -> None:
    # A NumberField that emits 10.0 for an integer control must still pass.
    assert _validate_options({"geometry_batch_size": 10.0}) == []


def test_validate_options_positive_number_rejects_zero_and_negative() -> None:
    assert _offenders(_validate_options({"prov_distance": 0})) == ["prov_distance"]
    assert _offenders(_validate_options({"pin_assignment_threshold": -5})) == ["pin_assignment_threshold"]
    assert _offenders(_validate_options({"geometry_batch_timeout_seconds": 0.0})) == [
        "geometry_batch_timeout_seconds"
    ]


def test_validate_options_number_rejects_non_finite() -> None:
    # NaN / inf must not slip past a ">0" comparison.
    assert _offenders(_validate_options({"refdes_search_radius": float("inf")})) == ["refdes_search_radius"]
    assert _offenders(_validate_options({"prov_distance": float("nan")})) == ["prov_distance"]


def test_validate_options_bool_flag_rejects_non_bool() -> None:
    # bool is an int subclass — an unguarded numeric check would accept 1/0.
    assert _offenders(_validate_options({"geometry_analysis_enabled": 1})) == ["geometry_analysis_enabled"]
    assert _offenders(_validate_options({"pinlist_prefers_annotation_mode": "yes"})) == [
        "pinlist_prefers_annotation_mode"
    ]
    assert _validate_options({"geometry_subprocess_enabled": True}) == []
    assert _validate_options({"geometry_subprocess_enabled": False}) == []


def test_validate_options_number_field_rejects_bool() -> None:
    # A boolean must not masquerade as a numeric value.
    assert _offenders(_validate_options({"geometry_batch_size": True})) == ["geometry_batch_size"]


def test_validate_options_enum_rejects_unknown_value() -> None:
    assert _offenders(_validate_options({"extraction_mode": "bogus"})) == ["extraction_mode"]
    assert _offenders(_validate_options({"backend_mode": "turbo"})) == ["backend_mode"]
    assert _validate_options({"extraction_mode": "piece_part", "backend_mode": "nextgen"}) == []


def test_validate_options_ignores_unknown_keys() -> None:
    # Mirrors RefDesConfig.from_options' silent drop of unknown keys.
    assert _validate_options({"totally_unknown_key": "abc", "another": 5}) == []


def test_validate_options_reports_every_offender() -> None:
    errors = _validate_options({
        "geometry_batch_size": "abc",
        "adaptive_orphan_ratio": 5,
        "extraction_mode": "bogus",
    })
    assert _offenders(errors) == ["adaptive_orphan_ratio", "extraction_mode", "geometry_batch_size"]


# --- Integration: validate_run_request / execute_run_request wiring -----------

def _base_body(options=None) -> dict:
    return {
        "workflowId": "refdes_extract",
        "inputs": [{"role": "pdf", "path": "C:/schematic.pdf"}],
        "options": {} if options is None else options,
    }


def test_validate_run_request_accepts_default_options() -> None:
    pytest.importorskip("fitz")
    resp = validate_run_request(_base_body(dict(_DEFAULT_OPTIONS)))
    assert resp["ok"] is True
    assert resp["reason_code"] == "ok"


def test_validate_run_request_blocks_wrong_typed_option() -> None:
    # (a) wrong-typed value is a BLOCKING validation error, with a per-option entry.
    pytest.importorskip("fitz")
    resp = validate_run_request(_base_body({"geometry_batch_size": "abc"}))
    assert resp["ok"] is False
    assert resp["reason_code"] == "invalid_option"
    assert "geometry_batch_size" in resp["toast_text"]
    assert any(
        m["area"] == "Options" and "geometry_batch_size" in m["title"]
        for m in resp["validations"]
    )


def test_validate_run_request_blocks_out_of_range_ratio() -> None:
    # (b) out-of-range value is a BLOCKING validation error.
    pytest.importorskip("fitz")
    resp = validate_run_request(_base_body({"adaptive_orphan_ratio": 5}))
    assert resp["ok"] is False
    assert resp["reason_code"] == "invalid_option"
    assert "adaptive_orphan_ratio" in resp["toast_text"]


def test_validate_run_request_missing_file_precedes_option_check() -> None:
    # A missing required PDF must win over a bad option (setup blockers first).
    pytest.importorskip("fitz")
    body = {"workflowId": "refdes_extract", "inputs": [], "options": {"geometry_batch_size": "abc"}}
    resp = validate_run_request(body)
    assert resp["ok"] is False
    assert resp["reason_code"] == "missing_files"


def test_execute_run_request_raises_on_invalid_option() -> None:
    # A bad option fails visibly at execute time too, before any engine work.
    pytest.importorskip("fitz")
    with pytest.raises(ValidationError):
        execute_run_request(_base_body({"adaptive_orphan_ratio": 5}))


@pytest.mark.parametrize(
    ("bom_outcome", "expected_note"),
    [
        ("loader_error", "could not be loaded"),
        ("empty_result", "0 recognizable RefDes"),
    ],
)
def test_bom_cross_check_soft_failure_counts_as_a_warning(
    bom_outcome: str,
    expected_note: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fitz = pytest.importorskip("fitz")
    from refdes_extractor import bom_loader, extraction_engine
    from refdes_extractor.bom_loader import BomLoadResult
    from refdes_test import refdes_test_logic

    pdf_path = tmp_path / "schematic.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()
    bom_path = tmp_path / "bom.xlsx"
    bom_path.touch()

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_annotations_from_doc",
        lambda _doc, **_kwargs: [],
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "detect_groups_with_fallback",
        lambda _doc, _annotations, **_kwargs: ([], False, {}),
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "extract_with_geometry_analysis_detailed",
        lambda **_kwargs: ([], {"backend_used": "test", "token_diagnostics": {}}),
    )
    monkeypatch.setattr(extraction_engine, "cleanup_words_extraction_threads", lambda: 0)
    monkeypatch.setattr(refdes_runtime, "verify_excel_readable", lambda _path: True)

    if bom_outcome == "loader_error":
        def fail_bom_load(*_args, **_kwargs):
            raise RuntimeError("test BOM read failure")

        monkeypatch.setattr(bom_loader, "load_bom_data", fail_bom_load)
    else:
        monkeypatch.setattr(
            bom_loader,
            "load_bom_data",
            lambda *_args, **_kwargs: BomLoadResult(refdes=set(), page_map={}),
        )

    result = execute_run_request({
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "outputDirectory": str(tmp_path),
        "inputs": [
            {"role": "pdf", "path": str(pdf_path)},
            {"role": "bom", "path": str(bom_path)},
        ],
        "options": {},
    })

    assert result["status"] == "success"
    assert result["warning_count"] == 1
    assert result["no_match_count"] == 0
    assert "BOM cross-check FAILED" in result["title"]
    assert any(
        "BOM CROSS-CHECK FAILED" in note and expected_note in note
        for note in result["notes"]
    )


def test_refdes_cancel_checks_after_engine_and_before_promote(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fitz = pytest.importorskip("fitz")
    from refdes_test import refdes_test_logic

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_annotations_from_doc",
        lambda _doc, **_kwargs: [],
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "detect_groups_with_fallback",
        lambda _doc, _annotations, **_kwargs: ([], False, {}),
    )

    def build_body(case_dir: Path) -> dict:
        case_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = case_dir / "schematic.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(pdf_path)
        doc.close()
        return {
            "workflowId": "refdes_extract",
            "outputStrategyId": "new_workbook_standard",
            "outputDirectory": str(case_dir),
            "inputs": [{"role": "pdf", "path": str(pdf_path)}],
            "options": {},
        }

    post_engine_dir = tmp_path / "post-engine"

    def cancel_before_engine_return(**kwargs):
        kwargs["stop_event"].set()
        return [], {"backend_used": "test", "token_diagnostics": {}}

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_with_geometry_analysis_detailed",
        cancel_before_engine_return,
    )
    monkeypatch.setattr(refdes_runtime, "verify_excel_readable", lambda _path: True)

    with pytest.raises(CancellationError):
        execute_run_request(build_body(post_engine_dir))

    assert not list(post_engine_dir.glob("RefDesExtract_*.xlsx"))
    assert not list(post_engine_dir.glob(".*.part.xlsx"))

    pre_promote_dir = tmp_path / "pre-promote"
    captured = {}

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_with_geometry_analysis_detailed",
        lambda **_kwargs: (
            [],
            {"backend_used": "test", "token_diagnostics": {}},
        ),
    )

    def capture_processor(processor) -> None:
        captured["processor"] = processor

    def cancel_after_verify(temp_path: Path) -> bool:
        assert Path(temp_path).exists()
        captured["processor"].cancel.cancel()
        return True

    monkeypatch.setattr(refdes_runtime, "verify_excel_readable", cancel_after_verify)

    with pytest.raises(CancellationError):
        execute_run_request(
            build_body(pre_promote_dir),
            processor_ready_callback=capture_processor,
        )

    assert not list(pre_promote_dir.glob("RefDesExtract_*.xlsx"))
    assert not list(pre_promote_dir.glob(".*.part.xlsx"))


def test_refdes_runtime_passes_stop_event_to_group_detection(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fitz = pytest.importorskip("fitz")
    from refdes_test import refdes_test_logic

    pdf_path = tmp_path / "schematic.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()
    body = {
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "outputDirectory": str(tmp_path),
        "inputs": [{"role": "pdf", "path": str(pdf_path)}],
        "options": {},
    }
    captured = {}

    def capture_processor(processor) -> None:
        captured["processor"] = processor

    def set_cancel_at_group_detection(
        stage: str,
        _message: str,
        _percent: int,
        _current,
        _total,
    ) -> None:
        if stage == "Detecting groups":
            captured["processor"].stop_event.set()

    def detect_with_required_stop_event(
        _doc,
        _annotations,
        *,
        log_func,
        stop_event,
    ):
        assert log_func is not None
        assert stop_event is captured["processor"].stop_event
        assert stop_event.is_set()
        raise CancellationError("Cancelled during group detection")

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_annotations_from_doc",
        lambda _doc, **_kwargs: [],
    )
    monkeypatch.setattr(
        refdes_test_logic,
        "detect_groups_with_fallback",
        detect_with_required_stop_event,
    )

    with pytest.raises(CancellationError, match="group detection"):
        execute_run_request(
            body,
            processor_ready_callback=capture_processor,
            progress_callback=set_cancel_at_group_detection,
        )


def _run_refdes_doc_cleanup_case(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    *,
    remaining_threads: int,
) -> tuple[object, list[str]]:
    fitz = pytest.importorskip("fitz")
    from refdes_extractor import extraction_engine
    from refdes_test import refdes_test_logic

    class _FakeDocument:
        def __init__(self) -> None:
            self.close_calls = 0

        def __len__(self) -> int:
            return 1

        def close(self) -> None:
            self.close_calls += 1

    pdf_path = tmp_path / "schematic.pdf"
    pdf_path.touch()
    body = {
        "workflowId": "refdes_extract",
        "outputStrategyId": "new_workbook_standard",
        "outputDirectory": str(tmp_path),
        "inputs": [{"role": "pdf", "path": str(pdf_path)}],
        "options": {},
    }
    fake_doc = _FakeDocument()
    logs: list[str] = []

    monkeypatch.setattr(fitz, "open", lambda _path: fake_doc)
    monkeypatch.setattr(
        extraction_engine,
        "cleanup_words_extraction_threads",
        lambda: remaining_threads,
    )

    def stop_after_open(*_args, **_kwargs):
        raise RuntimeError("stop after PDF open")

    monkeypatch.setattr(
        refdes_test_logic,
        "extract_annotations_from_doc",
        stop_after_open,
    )

    with pytest.raises(RuntimeError, match="stop after PDF open"):
        execute_run_request(body, log_callback=logs.append)

    return fake_doc, logs


def test_refdes_zombie_extraction_thread_skips_pdf_close_and_warns(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake_doc, logs = _run_refdes_doc_cleanup_case(
        tmp_path,
        monkeypatch,
        remaining_threads=1,
    )
    warning = (
        "1 extraction thread(s) still running; leaving the PDF handle open to "
        "avoid a native crash."
    )

    assert fake_doc.close_calls == 0
    assert logs.count(warning) == 1


def test_refdes_zero_extraction_threads_closes_pdf(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake_doc, logs = _run_refdes_doc_cleanup_case(
        tmp_path,
        monkeypatch,
        remaining_threads=0,
    )

    assert fake_doc.close_calls == 1
    assert not any("leaving the PDF handle open" in line for line in logs)


# ----- Batch 5 (2026-07 stability sweep) -------------------------------------


def test_group_verification_summary_requires_components() -> None:
    """The verified/unverified metrics must reflect BOM verification content,
    not the "(Verified)" row-label suffix (which every regular group emits
    even when its Verified row is empty)."""
    from refdes_extractor.runtime import _summarize_group_verification

    results = [
        # Group A: Verified row EMPTY, all components unverified.
        {"group": "A (Verified)", "component count": 0},
        {"group": "A (Unverified)", "component count": 3},
        # Group B: genuinely verified components.
        {"group": "B (Verified)", "component count": 2},
        {"group": "B (Unverified)", "component count": 0},
        # Gap placeholder rows are skipped entirely.
        {"group": "C", "component count": 0, "_is_gap": True},
        {"group": "GROUP NOT DETECTED: D", "component count": 0},
    ]
    total, verified, unverified = _summarize_group_verification(results)
    assert total == 2, (total, verified, unverified)
    assert verified == 1, (total, verified, unverified)
    assert unverified == 1, (total, verified, unverified)


def test_config_manager_carries_out_folder_for_checkpoints(tmp_path) -> None:
    """The geometry-batch checkpoint writer early-returns without an
    out_folder; the runtime must pass the resolved output directory through
    to_config_manager or the checkpoint control is inert."""
    from refdes_extractor.runtime import RefDesConfig

    cm = RefDesConfig().to_config_manager(out_folder=str(tmp_path))
    assert cm.get("out_folder", "") == str(tmp_path)

    # Default stays empty (no checkpoint writes without a destination).
    cm_default = RefDesConfig().to_config_manager()
    assert cm_default.get("out_folder", "") == ""


def test_sidecar_main_installs_multiprocessing_freeze_support() -> None:
    """The sidecar ships as a PyInstaller --onefile exe and the geometry
    subprocess option uses mp spawn: without freeze_support() at the top of
    main(), frozen children re-run the sidecar loop instead of the worker."""
    import inspect

    import sidecar_main

    source = inspect.getsource(sidecar_main.main)
    assert "freeze_support()" in source, (
        "sidecar_main.main() must call multiprocessing.freeze_support() "
        "before anything else"
    )
