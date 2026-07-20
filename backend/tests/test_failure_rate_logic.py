"""Focused numeric tests for the Failure Rate linker math.

The entire purpose of the Failure Rate tool is the per-mode failure-rate
arithmetic:

    Mode_FR = Part_FR * Usage * Corrected_Ratio

and the unit-mode scaling that converts the prediction file's reported
failure rate into per-hour space before the multiply.  These tests build
minimal in-memory DataFrames, drive ``FMEALinkerLogic.process`` directly,
and assert the EXACT computed numbers — so they fail loudly if the linker
math or the unit conversion ever regresses.
"""

from __future__ import annotations

import math

import pandas as pd
import pytest

from failure_rate.failure_rate_logic import (
    FMEALinkerLogic,
    normalize_refdes_for_lookup,
)


# Column map shared by most tests. Keys mirror what the runtime adapter
# hands ``process`` (see failure_rate/runtime.py).
COL_MAP = {
    "pred_ref": "Reference Designator",
    "pred_fr": "Failure Rate",
    "fmea_cause": "Failure Mode Causes",
    "fmea_ratio": "Failure Mode Ratio",
    "fmea_usage": "Part Usage",
    "fmea_func": "",
}


def _make_logic(pred_df: pd.DataFrame, fmea_df: pd.DataFrame) -> FMEALinkerLogic:
    """Build a linker with DataFrames injected directly (no Excel IO)."""
    logic = FMEALinkerLogic(log_func=lambda _msg: None)
    logic.prediction_df = pred_df
    logic.fmea_df = fmea_df
    return logic


def test_mode_fr_is_part_fr_times_usage_times_ratio() -> None:
    """Mode_FR for a linked row equals Part_FR * Usage * Corrected_Ratio."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R1", "Failure Rate": 0.00001},
            {"Reference Designator": "C2", "Failure Rate": 0.000004},
        ]
    )
    fmea = pd.DataFrame(
        [
            # R1: two modes splitting the failure rate 60/40, usage 1.0
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 0.6, "Part Usage": 1.0},
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 0.4, "Part Usage": 1.0},
            # C2: single mode, ratio 0.5, usage 2.0 -> exercises usage multiply
            {"Failure Mode Causes": "C2", "Failure Mode Ratio": 0.5, "Part Usage": 2.0},
        ]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    # Part_FR is the (unconverted, per_hour) lookup value per row.
    assert list(result["Part_FR"]) == pytest.approx([0.00001, 0.00001, 0.000004])
    assert list(result["Corrected_Ratio"]) == pytest.approx([0.6, 0.4, 0.5])

    # Mode_FR = Part_FR * Usage * Ratio, computed independently here.
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.00001 * 1.0 * 0.6)  # 6e-6
    assert result["Mode_FR"].iloc[1] == pytest.approx(0.00001 * 1.0 * 0.4)  # 4e-6
    assert result["Mode_FR"].iloc[2] == pytest.approx(0.000004 * 2.0 * 0.5)  # 4e-6

    # The two R1 modes must sum back to the full part failure rate (usage 1.0).
    assert result["Mode_FR"].iloc[0] + result["Mode_FR"].iloc[1] == pytest.approx(0.00001)


@pytest.mark.parametrize(
    "unit_mode, scale",
    [
        ("per_hour", 1.0),
        ("per_million_hours", 1e-6),
        ("per_billion_hours", 1e-9),
    ],
)
def test_unit_mode_scales_part_fr_and_mode_fr(unit_mode: str, scale: float) -> None:
    """Each supported unit mode divides the raw prediction value before linking.

    The prediction file reports ``500`` for U10. Interpreted per_hour that's
    500/hr; per_million_hours that's 500e-6/hr; per_billion_hours 500e-9/hr.
    With ratio 1.0 and usage 1.0, Mode_FR must equal the scaled Part_FR.
    """
    raw_fr = 500.0
    pred = pd.DataFrame([{"Reference Designator": "U10", "Failure Rate": raw_fr}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "U10", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )
    col_map = dict(COL_MAP, unit_mode=unit_mode)

    result = _make_logic(pred, fmea).process(col_map)

    expected = raw_fr * scale
    assert result["Part_FR"].iloc[0] == pytest.approx(expected)
    # ratio=1.0, usage=1.0 -> Mode_FR collapses to the scaled Part_FR.
    assert result["Mode_FR"].iloc[0] == pytest.approx(expected)


def test_unit_modes_produce_distinct_scaled_numbers() -> None:
    """Cross-check: the three unit modes must yield 1 : 1e-6 : 1e-9 ratios.

    This guards against a regression where the conversion branch silently
    falls through to per_hour for the scaled modes.
    """
    raw_fr = 730.0
    pred = pd.DataFrame([{"Reference Designator": "U1", "Failure Rate": raw_fr}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "U1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )

    def mode_fr_for(unit: str) -> float:
        result = _make_logic(
            pred.copy(), fmea.copy()
        ).process(dict(COL_MAP, unit_mode=unit))
        return float(result["Mode_FR"].iloc[0])

    per_hour = mode_fr_for("per_hour")
    per_million = mode_fr_for("per_million_hours")
    per_billion = mode_fr_for("per_billion_hours")

    assert per_hour == pytest.approx(730.0)
    assert per_million == pytest.approx(730.0e-6)
    assert per_billion == pytest.approx(730.0e-9)
    # The three are genuinely different (no silent fall-through to per_hour).
    assert per_hour != pytest.approx(per_million)
    assert per_million != pytest.approx(per_billion)


def test_invalid_unit_mode_defaults_to_per_hour() -> None:
    """An unrecognized unit_mode falls back to per_hour (no scaling)."""
    pred = pd.DataFrame([{"Reference Designator": "R5", "Failure Rate": 12.0}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "R5", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )

    result = _make_logic(pred, fmea).process(dict(COL_MAP, unit_mode="per_zorp"))

    # No conversion applied -> Part_FR is the raw value.
    assert result["Part_FR"].iloc[0] == pytest.approx(12.0)
    assert result["Mode_FR"].iloc[0] == pytest.approx(12.0)


def test_unmatched_refdes_yields_zero_part_fr_and_note() -> None:
    """A FMEA cause with no prediction match gets Part_FR=0 and Mode_FR=0."""
    pred = pd.DataFrame([{"Reference Designator": "R1", "Failure Rate": 0.001}])
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
            {"Failure Mode Causes": "Q99", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
        ]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    # Matched row.
    assert result["Part_FR"].iloc[0] == pytest.approx(0.001)
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.001)
    # Unmatched row: zeroed and annotated.
    assert result["Part_FR"].iloc[1] == 0.0
    assert result["Mode_FR"].iloc[1] == 0.0
    assert "RefDes not in Prediction" in result["Validation_Notes"].iloc[1]


def test_zero_padded_refdes_links_across_sources() -> None:
    """Prediction 'U007' must link to FMEA cause 'U7' (leading zeros stripped)."""
    pred = pd.DataFrame([{"Reference Designator": "U007", "Failure Rate": 0.02}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "U7", "Failure Mode Ratio": 0.25, "Part Usage": 1.0}]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Extracted_RefDes"].iloc[0] == "U7"
    assert result["Part_FR"].iloc[0] == pytest.approx(0.02)
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.02 * 0.25)  # 0.005


@pytest.mark.parametrize(
    "prediction_rows",
    [
        [
            {"Reference Designator": "R01", "Failure Rate": 0.001},
            {"Reference Designator": "R1", "Failure Rate": 0.9},
        ],
        [
            {"Reference Designator": "R1", "Failure Rate": 0.9},
            {"Reference Designator": "R01", "Failure Rate": 0.001},
        ],
    ],
    ids=["low-rate-first", "high-rate-first"],
)
def test_conflicting_normalized_prediction_duplicates_choose_max_and_warn(
    prediction_rows: list[dict[str, object]],
) -> None:
    """R01/R1 conflicts are order-independent, conservative, and visible."""
    pred = pd.DataFrame(prediction_rows)
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Part_FR"].iloc[0] == pytest.approx(0.9)
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.9)
    note = result["Validation_Notes"].iloc[0]
    assert "R01=0.001" in note
    assert "R1=0.9" in note
    assert "maximum" in note.lower()
    assert "conservative" in note.lower()


def test_conflicting_normalized_prediction_duplicates_are_order_independent() -> None:
    """Swapping conflicting prediction rows must not change any output cell."""
    prediction_rows = [
        {"Reference Designator": "R01", "Failure Rate": 0.001},
        {"Reference Designator": "R1", "Failure Rate": 0.9},
    ]
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )

    forward = _make_logic(pd.DataFrame(prediction_rows), fmea.copy()).process(COL_MAP)
    reverse = _make_logic(
        pd.DataFrame(list(reversed(prediction_rows))), fmea.copy()
    ).process(COL_MAP)

    pd.testing.assert_frame_equal(forward, reverse)


def test_identical_normalized_prediction_duplicates_are_clean() -> None:
    """Equivalent R01/R1 rates deduplicate without an output warning."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R01", "Failure Rate": 0.25},
            {"Reference Designator": "R1", "Failure Rate": 0.25},
        ]
    )
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Part_FR"].iloc[0] == pytest.approx(0.25)
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.25)
    assert result["Validation_Notes"].iloc[0] == ""


def test_invalid_usage_and_ratio_default_to_one() -> None:
    """Out-of-range usage/ratio default to 1.0 and are annotated.

    Usage of 5_000_000 exceeds the 1e6 cap; ratio of 50 exceeds the 0..10
    band. Both default to 1.0, so Mode_FR collapses to Part_FR.
    """
    pred = pd.DataFrame([{"Reference Designator": "C1", "Failure Rate": 0.003}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "C1", "Failure Mode Ratio": 50.0, "Part Usage": 5_000_000.0}]
    )

    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Corrected_Ratio"].iloc[0] == pytest.approx(1.0)
    assert result["Mode_FR"].iloc[0] == pytest.approx(0.003)  # 0.003 * 1.0 * 1.0
    notes = result["Validation_Notes"].iloc[0]
    assert "Invalid Usage" in notes
    assert "Invalid Ratio" in notes


def test_function_fr_sums_mode_fr_within_function_group() -> None:
    """When a function column is mapped, Function_FR sums Mode_FR per function."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R1", "Failure Rate": 0.001},
            {"Reference Designator": "R2", "Failure Rate": 0.002},
            {"Reference Designator": "R3", "Failure Rate": 0.004},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0, "Function": "F_A"},
            {"Failure Mode Causes": "R2", "Failure Mode Ratio": 1.0, "Part Usage": 1.0, "Function": "F_A"},
            {"Failure Mode Causes": "R3", "Failure Mode Ratio": 1.0, "Part Usage": 1.0, "Function": "F_B"},
        ]
    )
    col_map = dict(COL_MAP, fmea_func="Function")

    result = _make_logic(pred, fmea).process(col_map)

    # F_A rows: 0.001 + 0.002 = 0.003 for both R1 and R2 rows.
    f_a = result[result["Function"] == "F_A"]["Function_FR"]
    assert list(f_a) == pytest.approx([0.003, 0.003])
    # F_B row: just R3 = 0.004.
    f_b = result[result["Function"] == "F_B"]["Function_FR"]
    assert list(f_b) == pytest.approx([0.004])


def test_normalize_refdes_strips_leading_zeros_and_handles_nan() -> None:
    """Lookup normalization: zero-strip, uppercase, suffix-preserve, NaN-safe."""
    assert normalize_refdes_for_lookup("u01") == "U1"
    assert normalize_refdes_for_lookup("R001") == "R1"
    assert normalize_refdes_for_lookup("U007A") == "U7A"
    assert normalize_refdes_for_lookup("R10") == "R10"  # no spurious stripping
    assert normalize_refdes_for_lookup(float("nan")) == ""
    assert normalize_refdes_for_lookup(None) == ""


def test_circuit_block_rollup_sums_children_by_fmea_id() -> None:
    """A circuit-block row rolls up to the SUM of its piece-part children, tied
    by FMEA-ID prefix, and is excluded from the function total (no double-count).

    Worked example: block CPU-001 owns R201 (two modes 0.6/0.4) and C5.
        lambda(R201)=1e-6, lambda(C5)=4e-7
        roll-up = lambda(R201) + lambda(C5) = 1.4e-6
    The block row's failure rate becomes the roll-up; Function_FR is that same
    roll-up broadcast to every row (NOT block + parts = 2x).
    """
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "C5", "Failure Rate": 4e-7},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA-ID": "CPU-001", "FMEA Level": "Circuit Block",
             "Failure Mode Causes": "R201, C5", "Failure Mode Ratio": "",
             "Part Usage": "", "Function": "CPU power"},
            {"FMEA-ID": "CPU-001-R201-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": 0.6,
             "Part Usage": 1.0, "Function": "CPU power"},
            {"FMEA-ID": "CPU-001-R201-B", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": 0.4,
             "Part Usage": 1.0, "Function": "CPU power"},
            {"FMEA-ID": "CPU-001-C5-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "C5", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "CPU power"},
        ]
    )
    result = _make_logic(pred, fmea).process(dict(COL_MAP, fmea_func="Function"))

    # Piece-part Mode_FR unchanged by the roll-up.
    assert result["Mode_FR"].iloc[1] == pytest.approx(1e-6 * 0.6)
    assert result["Mode_FR"].iloc[2] == pytest.approx(1e-6 * 0.4)
    assert result["Mode_FR"].iloc[3] == pytest.approx(4e-7 * 1.0)

    rollup = 1e-6 + 4e-7  # 1.4e-6
    # Block row failure rate IS the roll-up (no more phantom first-RefDes rate).
    assert result["Mode_FR"].iloc[0] == pytest.approx(rollup)
    assert result["Part_FR"].iloc[0] == pytest.approx(rollup)
    # Function_FR is the roll-up broadcast to every row, NOT double-counted.
    assert list(result["Function_FR"]) == pytest.approx([rollup] * 4)
    assert "Circuit-block roll-up" in result["Validation_Notes"].iloc[0]


def test_circuit_block_rollup_falls_back_to_listed_refdes_without_id() -> None:
    """With no FMEA-ID column, the block's listed RefDes are matched to children."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "C5", "Failure Rate": 4e-7},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA Level": "Circuit Block", "Failure Mode Causes": "R201, C5",
             "Failure Mode Ratio": "", "Part Usage": ""},
            {"FMEA Level": "Piece-Part", "Failure Mode Causes": "R201",
             "Failure Mode Ratio": 0.6, "Part Usage": 1.0},
            {"FMEA Level": "Piece-Part", "Failure Mode Causes": "R201",
             "Failure Mode Ratio": 0.4, "Part Usage": 1.0},
            {"FMEA Level": "Piece-Part", "Failure Mode Causes": "C5",
             "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    rollup = 1e-6 + 4e-7
    assert result["Mode_FR"].iloc[0] == pytest.approx(rollup)
    assert "Circuit-block roll-up" in result["Validation_Notes"].iloc[0]


def test_plain_piece_part_file_keeps_legacy_function_rollup() -> None:
    """No block rows -> Function_FR still sums Mode_FR by Function Description."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R1", "Failure Rate": 0.001},
            {"Reference Designator": "R2", "Failure Rate": 0.002},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "F_A"},
            {"Failure Mode Causes": "R2", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "F_A"},
        ]
    )
    result = _make_logic(pred, fmea).process(dict(COL_MAP, fmea_func="Function"))
    # Legacy behaviour: both rows carry the function total 0.003.
    assert list(result["Function_FR"]) == pytest.approx([0.003, 0.003])


def test_unparseable_prediction_fr_is_flagged_not_silent() -> None:
    """A non-numeric Prediction FR is coerced to 0.0 but FLAGGED, not silent."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R1", "Failure Rate": "TBD"},
            {"Reference Designator": "R2", "Failure Rate": 0.002},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
            {"Failure Mode Causes": "R2", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    # R1's FR was unparseable -> 0.0 but explicitly flagged.
    assert result["Part_FR"].iloc[0] == 0.0
    assert result["Mode_FR"].iloc[0] == 0.0
    assert "unparseable" in result["Validation_Notes"].iloc[0].lower()
    # R2 is a clean numeric link with no false flag.
    assert result["Mode_FR"].iloc[1] == pytest.approx(0.002)
    assert "unparseable" not in result["Validation_Notes"].iloc[1].lower()


def test_block_only_file_leaf_block_keeps_rate_not_zeroed() -> None:
    """Regression guard: a block-only FMEA (Circuit Block rows, NO piece-part
    children) must keep real rates. A multi-RefDes leaf block = sum of its listed
    components' lambda; a single-RefDes leaf block keeps its own computed rate.
    The roll-up must NOT wipe childless blocks to 0.
    """
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "C5", "Failure Rate": 4e-7},
            {"Reference Designator": "U9", "Failure Rate": 2e-6},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA Level": "Circuit Block", "Failure Mode Causes": "R201, C5",
             "Failure Mode Ratio": "", "Part Usage": ""},   # multi-RefDes leaf
            {"FMEA Level": "Circuit Block", "Failure Mode Causes": "U9",
             "Failure Mode Ratio": 1.0, "Part Usage": 1.0},  # single-RefDes leaf
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    # Multi-RefDes leaf -> sum of listed lambdas (NOT zero, NOT phantom-first-only).
    assert result["Mode_FR"].iloc[0] == pytest.approx(1e-6 + 4e-7)
    # Single-RefDes leaf -> keeps its computed rate, never zeroed.
    assert result["Mode_FR"].iloc[1] == pytest.approx(2e-6)


def test_infinite_prediction_fr_is_flagged_and_zeroed() -> None:
    """A prediction FR of 'inf' must be flagged and zeroed, never propagated."""
    pred = pd.DataFrame(
        [
            {"Reference Designator": "U1", "Failure Rate": "inf"},
            {"Reference Designator": "U2", "Failure Rate": 0.002},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "U1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
            {"Failure Mode Causes": "U2", "Failure Mode Ratio": 1.0, "Part Usage": 1.0},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Part_FR"].iloc[0] == 0.0
    assert result["Mode_FR"].iloc[0] == 0.0
    assert "unparseable" in result["Validation_Notes"].iloc[0].lower()
    assert math.isfinite(result["Mode_FR"].iloc[1])
    assert result["Mode_FR"].iloc[1] == pytest.approx(0.002)


def test_function_fr_consistent_when_block_and_ungrouped_pp_share_function() -> None:
    """Function_FR is one per-function total for ALL rows of a function, even
    when the function mixes a block+children with an unrelated ungrouped
    piece-part — no two different 'function totals' in a single function.
    """
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "R900", "Failure Rate": 5e-6},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA-ID": "CPU-001", "FMEA Level": "Circuit Block",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": "",
             "Part Usage": "", "Function": "f"},
            {"FMEA-ID": "CPU-001-R201-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "f"},
            {"FMEA-ID": "PP-R900", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R900", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "f"},
        ]
    )
    result = _make_logic(pred, fmea).process(dict(COL_MAP, fmea_func="Function"))

    # function total = child R201 (1e-6) + ungrouped R900 (5e-6) = 6e-6, all rows.
    assert list(result["Function_FR"]) == pytest.approx([6e-6, 6e-6, 6e-6])
    # the block's OWN rate is just its child = 1e-6.
    assert result["Mode_FR"].iloc[0] == pytest.approx(1e-6)


def test_fmr_validation_groups_by_instance_not_base() -> None:
    """FMR (ratio) validation sums to 1.0 PER instance/pin designator (U200-1),
    NOT per base component.

    Domain rule: each instance carries its own open/short failure modes whose
    ratios sum to 1.0; component-level roll-up is handled by the Part Usage
    column, not by collapsing instances to a base RefDes. Here U200-1's ratios
    sum to 1.0 (valid) and U200-2's to 0.8 (broken). Instance grouping flags
    ONLY U200-2. If validation were (wrongly) grouped by base U200, the combined
    sum would be 1.8 and U200-1's rows would be falsely flagged too -- this test
    guards against that regression.
    """
    pred = pd.DataFrame([{"Reference Designator": "U200", "Failure Rate": 0.001}])
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "U200-1", "Failure Mode Ratio": 0.6, "Part Usage": 0.5},
            {"Failure Mode Causes": "U200-1", "Failure Mode Ratio": 0.4, "Part Usage": 0.5},
            {"Failure Mode Causes": "U200-2", "Failure Mode Ratio": 0.5, "Part Usage": 0.5},
            {"Failure Mode Causes": "U200-2", "Failure Mode Ratio": 0.3, "Part Usage": 0.5},
        ]
    )

    result = _make_logic(pred, fmea).process(COL_MAP, check_fmr=True)
    notes = result["Validation_Notes"].astype(str)

    # Grouping is by the instance designator extracted from the cause text.
    assert set(result["Validation_RefDes"]) == {"U200-1", "U200-2"}
    # U200-1 ratios sum to 1.0 -> no FMR warning.
    assert all("FMR Sum" not in n for n in notes[result["Validation_RefDes"] == "U200-1"])
    # U200-2 ratios sum to 0.8 -> FMR warning on its rows.
    assert all("FMR Sum" in n for n in notes[result["Validation_RefDes"] == "U200-2"])


# ----- Tier-1 #1 (Side B): genuine-gap Part Usage -> blank Mode_FR -------------
#
# Decision: a genuinely missing Part Usage (NaN/blank, INCLUDING a "=1/N"
# formula cell with no cached value -> pd.to_numeric NaN) on a row that HAS a
# real failure rate (part_fr > 0) must NOT be silently defaulted to usage 1.0.
# Defaulting overstated the rate. Instead Mode_FR becomes NaN (blank cell) and
# the row is flagged. A part_fr == 0 (unmatched RefDes) keeps Mode_FR 0 — a
# meaningful "not in Prediction", not a gap. The circuit-block roll-up must SKIP
# blank children rather than propagating NaN into the whole block.


def test_blank_usage_with_real_fr_yields_nan_mode_fr_and_flag() -> None:
    """A linked row (part_fr > 0) whose Part Usage is blank/NaN must produce a
    NaN Mode_FR (blank cell) + an "Invalid Usage" flag — NOT part_fr * 1.0 * ratio.
    """
    pred = pd.DataFrame([{"Reference Designator": "R1", "Failure Rate": 0.001}])
    fmea = pd.DataFrame(
        [
            # blank Part Usage, real FR -> genuine gap
            {"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": ""},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    # part_fr links fine.
    assert result["Part_FR"].iloc[0] == pytest.approx(0.001)
    # Mode_FR is NaN (blank), NOT 0.001 * 1.0 * 1.0.
    assert pd.isna(result["Mode_FR"].iloc[0]), result["Mode_FR"].iloc[0]
    # The gap is flagged — the note reflects the blanked outcome (not a default).
    assert "left blank" in result["Validation_Notes"].iloc[0].lower()


def test_formula_usage_cell_with_no_cached_value_is_treated_as_gap() -> None:
    """A "=1/250" formula string that openpyxl/pandas reads with no cached
    numeric value (pd.to_numeric -> NaN) must be treated as a genuine gap:
    Mode_FR blank, NOT part_fr * 1.0 * ratio (which overstated the rate).
    """
    pred = pd.DataFrame([{"Reference Designator": "U7", "Failure Rate": 0.5}])
    fmea = pd.DataFrame(
        [
            {"Failure Mode Causes": "U7", "Failure Mode Ratio": 1.0, "Part Usage": "=1/250"},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Part_FR"].iloc[0] == pytest.approx(0.5)
    # Must be blank, not 0.5 (which is what the old usage=1.0 default produced).
    assert pd.isna(result["Mode_FR"].iloc[0]), result["Mode_FR"].iloc[0]
    assert "left blank" in result["Validation_Notes"].iloc[0].lower()


def test_unmatched_refdes_keeps_zero_mode_fr_not_blank() -> None:
    """A RefDes absent from the Prediction (part_fr == 0) must keep Mode_FR 0 —
    a meaningful "not in Prediction", never turned into a blank — even when its
    Part Usage is also blank.
    """
    pred = pd.DataFrame([{"Reference Designator": "R1", "Failure Rate": 0.001}])
    fmea = pd.DataFrame(
        [
            # R99 is not in the prediction -> part_fr 0; usage also blank.
            {"Failure Mode Causes": "R99", "Failure Mode Ratio": 1.0, "Part Usage": ""},
        ]
    )
    result = _make_logic(pred, fmea).process(COL_MAP)

    assert result["Part_FR"].iloc[0] == 0.0
    # Mode_FR stays 0 (not NaN) — the part has no rate, this is not a usage gap.
    assert result["Mode_FR"].iloc[0] == 0.0
    assert not pd.isna(result["Mode_FR"].iloc[0])


def test_block_rollup_skips_blank_child_and_flags_block() -> None:
    """Circuit-block roll-up must SKIP a blank-usage child rather than
    propagating its NaN Mode_FR into the whole block. The block FR must equal
    the sum of the PRESENT children, and the block row must be flagged that a
    child usage was unknown.

    Block CPU-001 owns R201 (real FR, usage 1.0 -> 1e-6) and C5 (real FR, BLANK
    usage -> NaN Mode_FR). The block FR must be 1e-6 (R201 only), not NaN.
    """
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "C5", "Failure Rate": 4e-7},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA-ID": "CPU-001", "FMEA Level": "Circuit Block",
             "Failure Mode Causes": "R201, C5", "Failure Mode Ratio": "",
             "Part Usage": "", "Function": "CPU power"},
            {"FMEA-ID": "CPU-001-R201-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": 1.0,
             "Part Usage": 1.0, "Function": "CPU power"},
            # C5 child: real FR but BLANK usage -> NaN Mode_FR.
            {"FMEA-ID": "CPU-001-C5-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "C5", "Failure Mode Ratio": 1.0,
             "Part Usage": "", "Function": "CPU power"},
        ]
    )
    result = _make_logic(pred, fmea).process(dict(COL_MAP, fmea_func="Function"))

    # The R201 child is a clean 1e-6.
    assert result["Mode_FR"].iloc[1] == pytest.approx(1e-6)
    # The C5 child is a genuine gap -> NaN.
    assert pd.isna(result["Mode_FR"].iloc[2]), result["Mode_FR"].iloc[2]
    # The block roll-up SKIPS the blank child: block FR == present child (1e-6),
    # NOT NaN.
    assert not pd.isna(result["Mode_FR"].iloc[0]), result["Mode_FR"].iloc[0]
    assert result["Mode_FR"].iloc[0] == pytest.approx(1e-6)
    # The block row is flagged that a child usage was unknown.
    block_notes = str(result["Validation_Notes"].iloc[0]).lower()
    assert "unknown" in block_notes, result["Validation_Notes"].iloc[0]


def test_block_rollup_all_gap_children_blanks_block_fr() -> None:
    """When EVERY child of a circuit block has a usage gap (all NaN Mode_FR),
    the block FR is entirely unknown and must be left BLANK (NaN), not a
    definitive 0.0 that reads as 'this block contributes nothing'.
    """
    pred = pd.DataFrame(
        [
            {"Reference Designator": "R201", "Failure Rate": 1e-6},
            {"Reference Designator": "C5", "Failure Rate": 4e-7},
        ]
    )
    fmea = pd.DataFrame(
        [
            {"FMEA-ID": "CPU-001", "FMEA Level": "Circuit Block",
             "Failure Mode Causes": "R201, C5", "Failure Mode Ratio": "",
             "Part Usage": "", "Function": "CPU power"},
            # BOTH children: real FR but BLANK usage -> NaN Mode_FR.
            {"FMEA-ID": "CPU-001-R201-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "R201", "Failure Mode Ratio": 1.0,
             "Part Usage": "", "Function": "CPU power"},
            {"FMEA-ID": "CPU-001-C5-A", "FMEA Level": "Piece-Part",
             "Failure Mode Causes": "C5", "Failure Mode Ratio": 1.0,
             "Part Usage": "", "Function": "CPU power"},
        ]
    )
    result = _make_logic(pred, fmea).process(dict(COL_MAP, fmea_func="Function"))

    assert pd.isna(result["Mode_FR"].iloc[1])
    assert pd.isna(result["Mode_FR"].iloc[2])
    # Block FR is entirely unknown -> BLANK (NaN), not 0.0.
    assert pd.isna(result["Mode_FR"].iloc[0]), result["Mode_FR"].iloc[0]
    assert "unknown" in str(result["Validation_Notes"].iloc[0]).lower()


# ----- Batch 5 (2026-07 stability sweep): output integrity ------------------


def test_save_results_preserves_user_text_in_passthrough_columns(tmp_path) -> None:
    """The abbreviation expander must run ONLY on the tool-written notes
    column. Running it over every column rewrote legitimate user text:
    "Main CB panel" (circuit breaker) shipped as "Main Circuit Block panel"."""
    from openpyxl import load_workbook

    pred = pd.DataFrame([{"Reference Designator": "R1", "Failure Rate": 0.00001}])
    fmea = pd.DataFrame(
        [
            {
                "Failure Mode Causes": "R1",
                "Failure Mode Ratio": 1.0,
                "Part Usage": 1.0,
                "Component Description": "Main CB panel",
            }
        ]
    )
    logic = _make_logic(pred, fmea)
    logic.process(COL_MAP)

    out = tmp_path / "fr_out.xlsx"
    logic.save_results(str(out))

    wb = load_workbook(out)
    try:
        ws = wb["Main"]
        headers = [c.value for c in ws[1]]
        desc_idx = headers.index("Component Description") + 1
        assert ws.cell(row=2, column=desc_idx).value == "Main CB panel"
    finally:
        wb.close()


def test_save_results_drops_internal_validation_refdes_column(tmp_path) -> None:
    """Validation_RefDes is an internal FMR-groupby key — it must not ship
    in the delivered workbook (it renders as a fully blank column on
    default runs)."""
    from openpyxl import load_workbook

    pred = pd.DataFrame([{"Reference Designator": "R1", "Failure Rate": 0.00001}])
    fmea = pd.DataFrame(
        [{"Failure Mode Causes": "R1", "Failure Mode Ratio": 1.0, "Part Usage": 1.0}]
    )
    logic = _make_logic(pred, fmea)
    logic.process(COL_MAP)

    out = tmp_path / "fr_cols.xlsx"
    logic.save_results(str(out))

    wb = load_workbook(out)
    try:
        headers = [c.value for c in wb["Main"][1]]
    finally:
        wb.close()
    assert "Validation_RefDes" not in headers, headers


def test_rollup_only_notes_are_informational() -> None:
    """Pure roll-up annotations are informational, not warnings — they must
    not inflate warning_count or draw amber row styling. Mixed notes that
    also carry a real warning stay warnings."""
    from failure_rate.failure_rate_logic import note_is_informational_only

    assert note_is_informational_only(
        "Circuit-block roll-up of 3 piece-part row(s); "
    ) is True
    assert note_is_informational_only(
        "Block roll-up of 2 listed component(s); "
    ) is True
    # Mixed: roll-up + genuine warning -> still a warning.
    assert note_is_informational_only(
        "Circuit-block roll-up of 3 piece-part row(s); "
        "1 child usage(s) unknown — block FR may be understated; "
    ) is False
    assert note_is_informational_only("RefDes not in Prediction") is False
    assert note_is_informational_only("") is False
