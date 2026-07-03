"""Tests for the RefDes extraction Validation Notes + styled workbook output.

Feature: cross-group duplicate extractions get ordinal reason notes and a row
highlight; populated Unverified rows get an explicit not-in-BOM note; gap
placeholder rows get an explanation. The workbook write path applies the
suite's modern styling (Aptos Narrow, frozen header, header fill, banding).
"""

from __future__ import annotations

from openpyxl import Workbook, load_workbook

from refdes_extractor.validation_notes import annotate_results
from refdes_extractor.runtime import (
    _write_component_detail_sheet,
    _write_extraction_sheet,
    _write_orphan_pins_sheet,
)


def _row(group: str, causes: str, pages: str = "", **extra) -> dict:
    row = {
        "group": group,
        "failure mode causes": causes,
        "component count": len([t for t in causes.split(", ") if t]),
        "pages": pages,
    }
    row.update(extra)
    return row


# ---------------------------------------------------------------------------
# annotate_results
# ---------------------------------------------------------------------------

def test_second_extraction_gets_ordinal_note_and_warning_style() -> None:
    rows = annotate_results(
        [
            _row("DIG-076 (Verified)", "U7-38, U9", pages="3"),
            _row("DIG-081 (Verified)", "U7-38", pages="7"),
        ],
        bom_provided=True,
    )
    # First occurrence (page 3) stays clean.
    assert rows[0]["validation notes"] == ""
    assert rows[0]["_row_style"] == "default"
    # Second occurrence names the token, the ordinal, and where it was first seen.
    assert rows[1]["validation notes"] == (
        "2nd extraction of U7-38 — first seen in DIG-076 (p. 3)."
    )
    assert rows[1]["_row_style"] == "warning"


def test_ordinals_follow_page_order_not_row_order() -> None:
    # The row emitted first sits on a LATER page — it must be the "2nd"
    # extraction, because ordinals follow schematic reading order.
    rows = annotate_results(
        [
            _row("DIG-081 (Verified)", "U7", pages="9"),
            _row("DIG-076 (Verified)", "U7", pages="2"),
        ],
        bom_provided=True,
    )
    assert "2nd extraction of U7 — first seen in DIG-076 (p. 2)." in rows[0]["validation notes"]
    assert rows[1]["validation notes"] == ""


def test_third_extraction_counts_up_and_ungrouped_is_an_occurrence_site() -> None:
    rows = annotate_results(
        [
            _row("DIG-076 (Verified)", "U7", pages="1"),
            _row("DIG-081 (Verified)", "U7", pages="2"),
            _row("UNGROUPED (IN BOM)", "U7", pages="5"),
        ],
        bom_provided=True,
    )
    assert "2nd extraction of U7" in rows[1]["validation notes"]
    assert "3rd extraction of U7" in rows[2]["validation notes"]


def test_verified_unverified_split_of_one_group_does_not_self_flag() -> None:
    # The V/U rows of a single group are two halves of ONE extraction —
    # different tokens by construction, so nothing may flag.
    rows = annotate_results(
        [
            _row("DIG-076 (Verified)", "U7, U9", pages="3"),
            _row("DIG-076 (Unverified)", "U12", pages="3"),
        ],
        bom_provided=True,
    )
    assert "extraction of" not in rows[0]["validation notes"]
    assert "extraction of" not in rows[1]["validation notes"]


def test_unverified_rows_get_not_in_bom_note_only_when_bom_provided() -> None:
    base = [
        _row("DIG-076 (Verified)", "U7", pages="3"),
        _row("DIG-076 (Unverified)", "U12, U13", pages="3"),
    ]
    with_bom = annotate_results(base, bom_provided=True)
    assert with_bom[1]["validation notes"] == "Not found in BOM: U12, U13."
    assert with_bom[1]["_row_style"] == "unverified"

    without_bom = annotate_results(base, bom_provided=False)
    assert without_bom[1]["validation notes"] == ""
    assert without_bom[1]["_row_style"] == "default"


def test_gap_rows_get_sequence_gap_note_and_gap_style() -> None:
    rows = annotate_results(
        [
            _row("CPU-003 (GROUP NOT DETECTED)", "", _is_gap=True),
        ],
        bom_provided=True,
    )
    assert rows[0]["validation notes"] == (
        "Expected group missing from schematic (sequence gap)."
    )
    assert rows[0]["_row_style"] == "gap"
    # Gap rows never participate in duplicate counting.
    assert "extraction of" not in rows[0]["validation notes"]


def test_multiple_notes_join_as_sentences() -> None:
    rows = annotate_results(
        [
            _row("DIG-076 (Unverified)", "U99", pages="2"),
            _row("DIG-081 (Unverified)", "U99", pages="6"),
        ],
        bom_provided=True,
    )
    assert rows[1]["validation notes"] == (
        "2nd extraction of U99 — first seen in DIG-076 (p. 2). "
        "Not found in BOM: U99."
    )
    # Duplicate styling outranks the unverified grey.
    assert rows[1]["_row_style"] == "warning"


def test_ambiguous_tokens_get_note_and_warning_style() -> None:
    rows = annotate_results(
        [_row("DIG-076 (Verified)", "U7-38, U9", pages="3")],
        bom_provided=True,
        ambiguous_tokens={"U7-38": 3},
    )
    assert rows[0]["validation notes"] == (
        "U7-38 assigned ambiguously (3 candidates) — see Component Detail."
    )
    assert rows[0]["_row_style"] == "warning"


def test_annotate_without_ambiguous_param_is_unchanged() -> None:
    rows = annotate_results(
        [_row("DIG-076 (Verified)", "U7-38", pages="3")],
        bom_provided=True,
    )
    assert rows[0]["validation notes"] == ""
    assert rows[0]["_row_style"] == "default"


# ---------------------------------------------------------------------------
# diagnostics sheets (read-back)
# ---------------------------------------------------------------------------

def test_component_detail_sheet_rows_sorted_and_flagged(tmp_path) -> None:
    details = {
        "token_diagnostics": {
            "U7-38": {"group": "DIG-076", "pages": [4, 9], "confidence": 0.82,
                      "source": "geometry", "candidates": 3},
            "R1": {"group": "CPU-001", "pages": [2]},
        }
    }
    wb = Workbook()
    _write_component_detail_sheet(wb, details)
    path = tmp_path / "detail.xlsx"
    wb.save(path)

    ws = load_workbook(path)["Component Detail"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
    assert headers == ["Component", "Group", "Pages", "Confidence", "Source", "Flags"]
    # Sorted by group: CPU-001 first.
    assert ws.cell(row=2, column=1).value == "R1"
    assert ws.cell(row=3, column=1).value == "U7-38"
    assert ws.cell(row=3, column=3).value == "4, 9"
    assert ws.cell(row=3, column=6).value == "ambiguous (3 candidates)"
    assert ws.cell(row=1, column=1).font.name == "Aptos Narrow"
    assert ws.freeze_panes == "A2"


def test_orphan_pins_sheet_and_skip_when_empty(tmp_path) -> None:
    details = {
        "orphan_pins": [
            {"page": 7, "group": "DIG-076", "pin_text": "38",
             "disposition": "excluded", "detail": "No qualified pinlist cluster matched this pin."},
        ]
    }
    wb = Workbook()
    _write_orphan_pins_sheet(wb, details)
    path = tmp_path / "orphans.xlsx"
    wb.save(path)

    ws = load_workbook(path)["Orphan Pins"]
    assert [ws.cell(row=1, column=c).value for c in range(1, 6)] == [
        "Pin", "Page", "Group", "Disposition", "Detail",
    ]
    assert ws.cell(row=2, column=1).value == "38"
    assert ws.cell(row=2, column=4).value == "excluded"

    # Empty diagnostics: neither sheet is created (legacy-fallback runs).
    wb_empty = Workbook()
    _write_component_detail_sheet(wb_empty, {})
    _write_orphan_pins_sheet(wb_empty, None)
    assert "Component Detail" not in wb_empty.sheetnames
    assert "Orphan Pins" not in wb_empty.sheetnames


# ---------------------------------------------------------------------------
# styled workbook output (read-back)
# ---------------------------------------------------------------------------

def test_extraction_sheet_is_styled_modern(tmp_path) -> None:
    results = annotate_results(
        [
            _row("DIG-076 (Verified)", "U7-38", pages="3"),
            _row("DIG-081 (Verified)", "U7-38", pages="7"),
        ],
        bom_provided=True,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "RefDes Extraction"
    _write_extraction_sheet(ws, results)
    path = tmp_path / "out.xlsx"
    wb.save(path)
    wb.close()

    ws = load_workbook(path)["RefDes Extraction"]
    # Header: Title Case, Aptos Narrow bold on the steel-blue fill, frozen.
    headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
    assert headers == [
        "Group",
        "Failure Mode Causes",
        "Component Count",
        "Pages",
        "Validation Notes",
    ]
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.name == "Aptos Narrow"
    assert header_cell.font.bold is True
    assert header_cell.fill.start_color.rgb.endswith("2E5A88")
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None

    # Data font follows the presets; the duplicate row carries the amber fill.
    assert ws.cell(row=2, column=1).font.name == "Aptos Narrow"
    dup_cell = ws.cell(row=3, column=1)
    assert dup_cell.fill.start_color.rgb.endswith("FFEB9C")
    # The internal _row_style marker must not appear as a column.
    assert ws.max_column == 5
